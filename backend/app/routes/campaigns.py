import csv
import base64
import difflib
import io
import json
import logging
import re
import asyncio
import hashlib
import threading
import unicodedata
import time
from typing import Any
from datetime import datetime, timedelta
from pathlib import Path
from urllib import error as urllib_error
from urllib import request as urllib_request
from uuid import uuid4

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response as FastAPIResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app import config
from app.auth import get_current_user
from app.db import SessionLocal, get_db
from app.email_service import render_campaign_email, render_campaign_subject
from app.models import (
    Campaign,
    CampaignAttachment,
    CampaignItem,
    CampaignMarketEstimateOffer,
    CampaignRecipient,
    CampaignStatus,
    DeliveryStatus,
    EmailJob,
    EmailJobStatus,
    Group,
    GroupMember,
    InboundInvoice,
    InvoiceExtractionItem,
    InvoiceItemOfferAnalysisCache,
    InvoiceItemCharacteristicMatch,
    OfferAnalysisResult,
    OpenAIUsageLog,
    ExtractionStatus,
    TechSheet,
    TechSheetLink,
    Product,
    Response as ResponseModel,
    ResponseItem,
    SuppressedEmail,
    Supplier,
    SupplierType,
    User,
)
from app.schemas import (
    CampaignCreate,
    CampaignAttachmentOut,
    CampaignItemOut,
    CampaignItemsPayload,
    CampaignOut,
    CampaignRecipientOut,
    CampaignRecipientPayload,
    CampaignStats,
    CampaignUpdate,
    ImportedEmailFileOut,
    ImportedEmailFileExtractPayload,
    ImportedEmailFileExtractResult,
    ImportedEmailFileReassignPayload,
    ImportedEmailFileReassignResult,
    InboundInvoiceOut,
    OfferSummaryEntry,
    ParseCharacteristicsFromTextIn,
    ParseCharacteristicsFromTextOut,
    OfferEntry,
    ProductOffers,
    ProductOffersPage,
    SimpleProductOfferEntry,
    SimpleProductOffers,
    SimpleProductOffersPage,
    ResponseOut,
    TechSheetLinkPayload,
)
from app.settings_service import get_settings, get_user_embedding_thresholds
from app.token_service import create_token, token_expiry
from app.routes.invoice_extractions import (
    _build_campaign_requested_index,
    _flatten_requested_components,
    _get_or_compute_offer_analysis,
    _pick_requested_reference,
    _pick_requested_references_batch,
    _strip_name_noise_for_match,
    extract_invoice as extract_imported_invoice,
)
from app.routes.tech_sheets import link_tech_sheets
from app.services.campaign_attachment_extraction import (
    extract_attachment_products_preview,
    learn_anchor_linking_from_products,
)
from app.services.tech_sheet_extraction import extract_tech_sheet, parse_characteristics
from app.services.openai_embeddings import (
    cosine_similarity,
    enqueue_embeddings,
    fetch_embeddings,
    get_cached_embeddings,
)
from app.services.invoice_extraction import (
    _collect_campaign_requested_characteristics,
    _extract_offered_characteristics,
    _select_requested_characteristics_for_product,
)
from app.services.results_cache import (
    build_results_cache_signature,
    get_results_cache,
    invalidate_results_cache,
    set_results_cache,
)
from app.services.openai_payload_cache import (
    get_openai_payload_cache,
    set_openai_payload_cache,
)
from app.services.openai_chat import call_chat_completion_json_with_retry
from app.services.openai_observability import record_openai_api_call

logger = logging.getLogger("app")

router = APIRouter(prefix="/campaigns", tags=["campaigns"], dependencies=[Depends(get_current_user)])

_CAMPAIGN_PROCESSING_STANDARD = "standard"
_CAMPAIGN_PROCESSING_STRICT = "strict"
_STANDARD_MIN_GPT_MATCH = 0.70
_RESULTS_MIN_GPT_MATCH = _STANDARD_MIN_GPT_MATCH
_SIMPLE_PRODUCT_ENABLED = False
_SIMPLE_PRODUCT_LOGIC_VERSION = 1
_SIMPLE_PRODUCT_CACHE_TYPE = "simple_product_page"
_SIMPLE_PRODUCT_PRODUCT_CACHE_TYPE = "simple_product_product"
_SIMPLE_PRODUCT_MODEL = "gpt-5.1"
_SIMPLE_PRODUCT_TOP_CANDIDATES = 5
_SIMPLE_PRODUCT_TOP_OFFERS_PER_INVOICE = 2
_SIMPLE_PRODUCT_GPT_BATCH_SIZE = 2
_QTY_SOURCE_MANUAL = "manual"
_QTY_SOURCE_CARRIED_FORWARD = "carried_forward"
_QTY_SOURCE_ATTACHMENT_AUTO = "attachment_auto"
_QTY_SOURCE_ATTACHMENT_AUTO_GPT = "attachment_auto_gpt"


def _normalize_campaign_processing_mode(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    if text in {"strict", "strict_mode", "mode_strict"}:
        return _CAMPAIGN_PROCESSING_STRICT
    if text in {"standard", "std", "normal", "classic", "non_strict", "non-strict"}:
        return _CAMPAIGN_PROCESSING_STANDARD
    return None
_QTY_COMPLEX_RE = re.compile(
    r"\b(pack|boite|boîte|bte|lot|x|×|ml|cl|l|mg|g|kg|flacon|ampoule|comprime|capsule|sachet)\b",
    re.IGNORECASE,
)


def _cost_group_for_operation(operation: str) -> tuple[str, str]:
    op = str(operation or "").strip().lower()
    if op.startswith("attachment_"):
        return "attachments", "Extraction pièces jointes"
    if op.startswith("invoice_"):
        return "invoices", "Extraction + matching factures"
    if op.startswith("tech_sheet_") or op.startswith("tech_similarity"):
        return "tech_sheets", "Fiches techniques + conformité"
    if op.startswith("campaign_qty_") or op.startswith("price_equivalence"):
        return "pricing", "Quantités / conversions / prix"
    if op.startswith("embeddings_"):
        return "embeddings", "Embeddings"
    if op.startswith("cache_"):
        return "cache", "Cache OpenAI"
    return "other", "Autres"


def _safe_meta_int(meta: dict[str, Any], key: str) -> int | None:
    value = meta.get(key)
    try:
        if value is None:
            return None
        return int(value)
    except Exception:
        return None


def _campaign_openai_cost_breakdown(
    db: Session,
    *,
    campaign: Campaign,
) -> dict[str, Any]:
    tenant_id = int(campaign.tenant_id)
    campaign_id = int(campaign.id)
    now = datetime.utcnow()
    since = campaign.created_at or (now - timedelta(days=30))

    # Strict attribution by campaign_id metadata only.
    window_rows = (
        db.query(OpenAIUsageLog)
        .filter(OpenAIUsageLog.created_at >= since, OpenAIUsageLog.created_at <= now)
        .order_by(OpenAIUsageLog.created_at.desc())
        .all()
    )

    attributed_rows: list[OpenAIUsageLog] = []
    for row in window_rows:
        meta = row.metadata_json if isinstance(row.metadata_json, dict) else {}
        row_campaign_id = _safe_meta_int(meta, "campaign_id")
        if row_campaign_id is not None and row_campaign_id == campaign_id:
            attributed_rows.append(row)

    scope = "strict_campaign_id"
    rows_for_summary = attributed_rows

    summary = {
        "scope": scope,
        "since_utc": since.isoformat(timespec="seconds"),
        "until_utc": now.isoformat(timespec="seconds"),
        "window_calls": int(len(window_rows)),
        "attributed_calls": int(len(attributed_rows)),
        "calls": 0,
        "success_calls": 0,
        "failed_calls": 0,
        "prompt_tokens": 0,
        "completion_tokens": 0,
        "total_tokens": 0,
        "total_usd": 0.0,
    }
    by_group: dict[str, dict[str, Any]] = {}
    by_operation: dict[str, dict[str, Any]] = {}
    by_model: dict[str, dict[str, Any]] = {}

    for row in rows_for_summary:
        operation = str(row.operation or "unknown").strip() or "unknown"
        model = str(row.model or "unknown").strip() or "unknown"
        cost = float(row.estimated_cost_usd or 0.0)
        prompt_tokens = int(row.prompt_tokens or 0)
        completion_tokens = int(row.completion_tokens or 0)
        total_tokens = int(row.total_tokens or 0)

        summary["calls"] += 1
        if bool(row.success):
            summary["success_calls"] += 1
        else:
            summary["failed_calls"] += 1
        summary["prompt_tokens"] += prompt_tokens
        summary["completion_tokens"] += completion_tokens
        summary["total_tokens"] += total_tokens
        summary["total_usd"] += cost

        group_key, group_label = _cost_group_for_operation(operation)
        group_row = by_group.setdefault(
            group_key,
            {
                "key": group_key,
                "label": group_label,
                "calls": 0,
                "success_calls": 0,
                "failed_calls": 0,
                "prompt_tokens": 0,
                "completion_tokens": 0,
                "total_tokens": 0,
                "total_usd": 0.0,
            },
        )
        group_row["calls"] = int(group_row["calls"]) + 1
        if bool(row.success):
            group_row["success_calls"] = int(group_row["success_calls"]) + 1
        else:
            group_row["failed_calls"] = int(group_row["failed_calls"]) + 1
        group_row["prompt_tokens"] = int(group_row["prompt_tokens"]) + prompt_tokens
        group_row["completion_tokens"] = int(group_row["completion_tokens"]) + completion_tokens
        group_row["total_tokens"] = int(group_row["total_tokens"]) + total_tokens
        group_row["total_usd"] = float(group_row["total_usd"]) + cost

        op_row = by_operation.setdefault(
            operation,
            {
                "operation": operation,
                "calls": 0,
                "success_calls": 0,
                "failed_calls": 0,
                "prompt_tokens": 0,
                "completion_tokens": 0,
                "total_tokens": 0,
                "total_usd": 0.0,
            },
        )
        op_row["calls"] = int(op_row["calls"]) + 1
        if bool(row.success):
            op_row["success_calls"] = int(op_row["success_calls"]) + 1
        else:
            op_row["failed_calls"] = int(op_row["failed_calls"]) + 1
        op_row["prompt_tokens"] = int(op_row["prompt_tokens"]) + prompt_tokens
        op_row["completion_tokens"] = int(op_row["completion_tokens"]) + completion_tokens
        op_row["total_tokens"] = int(op_row["total_tokens"]) + total_tokens
        op_row["total_usd"] = float(op_row["total_usd"]) + cost

        model_row = by_model.setdefault(
            model,
            {
                "model": model,
                "calls": 0,
                "success_calls": 0,
                "failed_calls": 0,
                "prompt_tokens": 0,
                "completion_tokens": 0,
                "total_tokens": 0,
                "total_usd": 0.0,
            },
        )
        model_row["calls"] = int(model_row["calls"]) + 1
        if bool(row.success):
            model_row["success_calls"] = int(model_row["success_calls"]) + 1
        else:
            model_row["failed_calls"] = int(model_row["failed_calls"]) + 1
        model_row["prompt_tokens"] = int(model_row["prompt_tokens"]) + prompt_tokens
        model_row["completion_tokens"] = int(model_row["completion_tokens"]) + completion_tokens
        model_row["total_tokens"] = int(model_row["total_tokens"]) + total_tokens
        model_row["total_usd"] = float(model_row["total_usd"]) + cost

    summary["total_usd"] = round(float(summary["total_usd"]), 6)

    def _sorted_rows(rows: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for row in rows.values():
            normalized = dict(row)
            normalized["total_usd"] = round(float(normalized.get("total_usd") or 0.0), 6)
            out.append(normalized)
        out.sort(
            key=lambda r: (float(r.get("total_usd") or 0.0), int(r.get("calls") or 0)),
            reverse=True,
        )
        return out

    return {
        "summary": summary,
        "by_group": _sorted_rows(by_group),
        "by_operation": _sorted_rows(by_operation),
        "by_model": _sorted_rows(by_model),
    }


_AUTO_CAMPAIGN_TITLE_RE = re.compile(r"^\s*camp(?:agne)?[\s\-_]*\d+\s*$", re.IGNORECASE)


def _normalized_campaign_title(campaign_id: int, raw_title: str | None) -> str:
    title = (raw_title or "").strip()
    if not title:
        return f"CAMP {campaign_id}"
    if _AUTO_CAMPAIGN_TITLE_RE.match(title):
        return f"CAMP {campaign_id}"
    return title


def _clean_display_text(value: str | None) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\u00A0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return text
    # Repair common mojibake when UTF-8 bytes were decoded as latin1/cp1252.
    if any(h in text for h in ("Ã", "Â", "â", "Ð", "Ñ")):
        for enc in ("latin1", "cp1252"):
            try:
                repaired = text.encode(enc).decode("utf-8")
            except Exception:
                continue
            if repaired and repaired.count("\uFFFD") <= text.count("\uFFFD"):
                text = re.sub(r"\s+", " ", repaired).strip()
                break
    return text


def _get_campaign(db: Session, campaign_id: int, current_user: User) -> Campaign:
    query = db.query(Campaign).filter(Campaign.id == campaign_id)
    if not current_user.is_admin:
        query = query.filter(
            Campaign.tenant_id == current_user.tenant_id,
            Campaign.owner_user_id == current_user.id,
        )
    campaign = query.first()
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    return campaign


def _campaign_effective_processing_mode(
    db: Session,
    campaign: Campaign,
) -> str:
    explicit = _normalize_campaign_processing_mode(getattr(campaign, "processing_mode", None))
    if explicit:
        return explicit
    has_strict_preview = (
        db.query(CampaignAttachment.id)
        .filter(
            CampaignAttachment.campaign_id == campaign.id,
            CampaignAttachment.tenant_id == campaign.tenant_id,
            CampaignAttachment.strict_extraction_preview.isnot(None),
        )
        .limit(1)
        .first()
        is not None
    )
    return _CAMPAIGN_PROCESSING_STRICT if has_strict_preview else _CAMPAIGN_PROCESSING_STANDARD


def _decorate_campaign_processing_mode(db: Session, campaign: Campaign) -> Campaign:
    campaign.processing_mode = _campaign_effective_processing_mode(db, campaign)
    return campaign


def _safe_attachment_name(name: str) -> str:
    base = Path(name).name
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._")
    return cleaned or "attachment.pdf"


def _normalize_file_token(value: str | None) -> str:
    text = _clean_display_text(value) or (value or "")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def _resolve_existing_file(path: Path, filename: str | None = None) -> Path | None:
    if path.exists() and path.is_file():
        return path

    names: list[str] = []
    for raw in (filename, path.name):
        if not raw:
            continue
        cleaned = _clean_display_text(raw) or raw
        for variant in (raw, cleaned):
            variant = str(variant).strip()
            if variant and variant not in names:
                names.append(variant)

    search_dirs = [
        config.BASE_DIR,
        config.BASE_DIR / "Autoclave",
        config.BASE_DIR / "legacy_files",
        config.CAMPAIGN_ATTACHMENTS_DIR,
    ]

    for root in search_dirs:
        if not root.exists():
            continue
        for name in names:
            candidate = root / name
            if candidate.exists() and candidate.is_file():
                return candidate

    target_tokens = {_normalize_file_token(name) for name in names if name}
    target_tokens.discard("")
    if not target_tokens:
        return None

    for root in search_dirs:
        if not root.exists():
            continue
        try:
            for pattern in ("*.pdf", "*.xlsx", "*.xlsm"):
                for candidate in root.rglob(pattern):
                    token = _normalize_file_token(candidate.name)
                    if token in target_tokens:
                        return candidate
        except Exception:
            continue
    return None


def _run_attachment_reextract_job(
    *,
    campaign_id: int,
    user_id: int,
    force: bool,
    expected_default: int | None,
    expected_map: dict[str, Any],
    strict_expected_count: bool,
    strict_count_mode: str,
    document_type_hint_default: str | None,
    layout_hint_default: str | None,
    document_type_hints_map: dict[str, Any],
    layout_hints_map: dict[str, Any],
    require_specs_for_acceptance: bool,
    requested_attachment_ids: set[int],
) -> None:
    def _parse_expected_count(raw: object) -> int | None:
        try:
            value = int(raw)  # type: ignore[arg-type]
        except Exception:
            return None
        return value if value > 0 else None

    def _parse_doc_type_hint(raw: object) -> str | None:
        value = str(raw or "").strip().lower()
        if value in {
            "product_list",
            "requirement_sheet",
            "catalog_list",
            "accessories_appendix",
            "strict_produit",
            "mixed",
        }:
            return value
        return None

    def _parse_layout_hint(raw: object) -> str | None:
        value = str(raw or "").strip().lower()
        if value in {"table", "paragraph", "mixed"}:
            return value
        return None

    def _preview_products_count(preview: dict[str, Any] | None) -> int:
        if not isinstance(preview, dict):
            return 0
        products = preview.get("products")
        if isinstance(products, list):
            return len([p for p in products if isinstance(p, dict)])
        items = preview.get("items")
        if isinstance(items, list):
            return len(items)
        return 0

    def _preview_unassigned_count(preview: dict[str, Any] | None) -> int:
        if not isinstance(preview, dict):
            return 0
        rows = preview.get("unassigned")
        return len(rows) if isinstance(rows, list) else 0

    def _preview_placeholder_count(preview: dict[str, Any] | None) -> int:
        if not isinstance(preview, dict):
            return 0
        try:
            return int(preview.get("expected_count_placeholders") or 0)
        except Exception:
            return 0

    def _preview_avg_quality(preview: dict[str, Any] | None) -> float:
        if not isinstance(preview, dict):
            return 0.0
        products = preview.get("products")
        if not isinstance(products, list):
            return 0.0
        scores: list[float] = []
        for row in products:
            if not isinstance(row, dict):
                continue
            try:
                score = float(row.get("quality_score") or 0.0)
            except Exception:
                score = 0.0
            if score > 0:
                scores.append(score)
        if not scores:
            return 0.0
        return max(0.0, min(1.0, sum(scores) / len(scores)))

    def _preview_rank(preview: dict[str, Any] | None, expected: int | None) -> float:
        count = _preview_products_count(preview)
        placeholders = _preview_placeholder_count(preview)
        real_count = max(0, count - placeholders)
        if expected and expected > 0:
            coverage = min(1.0, real_count / float(expected))
        else:
            coverage = min(1.0, real_count / 20.0)
        quality = _preview_avg_quality(preview)
        unassigned = _preview_unassigned_count(preview)
        penalty = min(1.0, unassigned / float(max(1, real_count + 1)))
        return (coverage * 0.65) + (quality * 0.25) + ((1.0 - penalty) * 0.10)

    def _needs_fallback_preview(preview: dict[str, Any] | None, expected: int | None) -> bool:
        count = _preview_products_count(preview)
        placeholders = _preview_placeholder_count(preview)
        real_count = max(0, count - placeholders)
        if real_count <= 0:
            return True
        if placeholders > 0:
            return True
        if expected and expected > 0:
            min_target = max(1, int(round(expected * 0.60)))
            if real_count < min_target:
                return True
        unassigned = _preview_unassigned_count(preview)
        if unassigned > max(12, real_count * 2):
            return True
        return False

    db = SessionLocal()
    try:
        current_user = db.query(User).filter(User.id == user_id).first()
        if not current_user:
            logger.warning("Attachment OCR async job aborted: user %s not found", user_id)
            return
        campaign = _get_campaign(db, campaign_id, current_user)
        explicit_processing_mode = _normalize_campaign_processing_mode(campaign.processing_mode)
        strict_hint_requested = bool(document_type_hint_default == "strict_produit") or any(
            _parse_doc_type_hint(v) == "strict_produit" for v in document_type_hints_map.values()
        )
        campaign_processing_mode = explicit_processing_mode
        if campaign_processing_mode is None and strict_hint_requested:
            campaign_processing_mode = _CAMPAIGN_PROCESSING_STRICT
            campaign.processing_mode = _CAMPAIGN_PROCESSING_STRICT
            db.flush()
        if campaign_processing_mode is None:
            campaign_processing_mode = _campaign_effective_processing_mode(db, campaign)

        rows = (
            db.query(CampaignAttachment)
            .filter(
                CampaignAttachment.campaign_id == campaign.id,
                CampaignAttachment.tenant_id == campaign.tenant_id,
            )
            .order_by(CampaignAttachment.created_at.desc())
            .all()
        )
        pending_rows = [
            row
            for row in rows
            if row.extracted_at is None or row.extraction_preview is None
        ]
        target_rows = rows if force else pending_rows
        if requested_attachment_ids:
            target_rows = [row for row in target_rows if int(row.id) in requested_attachment_ids]
        if not target_rows:
            return

        thresholds = get_user_embedding_thresholds(db, current_user.tenant_id, current_user.id)
        for row in target_rows:
            try:
                expected_for_row = expected_default
                document_type_hint_for_row = document_type_hint_default
                layout_hint_for_row = layout_hint_default
                for key in (
                    row.id,
                    str(row.id),
                    row.file_name,
                    Path(row.file_name).name if row.file_name else "",
                    Path(row.file_path).name if row.file_path else "",
                ):
                    if not key:
                        continue
                    if key in expected_map:
                        parsed_expected = _parse_expected_count(expected_map.get(key))
                        if parsed_expected is not None:
                            expected_for_row = parsed_expected
                            break
                for key in (
                    row.id,
                    str(row.id),
                    row.file_name,
                    Path(row.file_name).name if row.file_name else "",
                    Path(row.file_path).name if row.file_path else "",
                ):
                    if not key:
                        continue
                    if key in document_type_hints_map:
                        parsed_doc_hint = _parse_doc_type_hint(document_type_hints_map.get(key))
                        if parsed_doc_hint is not None:
                            document_type_hint_for_row = parsed_doc_hint
                            break
                for key in (
                    row.id,
                    str(row.id),
                    row.file_name,
                    Path(row.file_name).name if row.file_name else "",
                    Path(row.file_path).name if row.file_path else "",
                ):
                    if not key:
                        continue
                    if key in layout_hints_map:
                        parsed_layout_hint = _parse_layout_hint(layout_hints_map.get(key))
                        if parsed_layout_hint is not None:
                            layout_hint_for_row = parsed_layout_hint
                            break

                preview = extract_attachment_products_preview(
                    row.file_path,
                    thresholds.extraction_threshold,
                    expected_for_row,
                    strict_expected_count if expected_for_row is not None else False,
                    document_type_hint_for_row,
                    layout_hint_for_row,
                    strict_count_mode,
                    require_specs_for_acceptance,
                )
                strict_preview = None
                if isinstance(preview, dict) and "strict_payload" in preview:
                    strict_preview = preview.pop("strict_payload")

                if _needs_fallback_preview(preview, expected_for_row):
                    candidates: list[tuple[str | None, str | None]] = [
                        ("mixed", "mixed"),
                        ("catalog_list", "mixed"),
                        ("catalog_list", "table"),
                        ("requirement_sheet", "mixed"),
                        ("accessories_appendix", "mixed"),
                    ]
                    seen_profiles: set[tuple[str | None, str | None]] = {
                        (document_type_hint_for_row, layout_hint_for_row)
                    }
                    best_preview = preview
                    best_rank = _preview_rank(preview, expected_for_row)
                    for cand_doc_hint, cand_layout_hint in candidates:
                        profile = (cand_doc_hint, cand_layout_hint)
                        if profile in seen_profiles:
                            continue
                        seen_profiles.add(profile)
                        try:
                            candidate_preview = extract_attachment_products_preview(
                                row.file_path,
                                thresholds.extraction_threshold,
                                expected_for_row,
                                strict_expected_count if expected_for_row is not None else False,
                                cand_doc_hint,
                                cand_layout_hint,
                                strict_count_mode,
                                require_specs_for_acceptance,
                            )
                        except Exception:
                            continue
                        candidate_rank = _preview_rank(candidate_preview, expected_for_row)
                        if candidate_rank > best_rank + 0.03:
                            best_preview = candidate_preview
                            best_rank = candidate_rank
                    preview = best_preview

                row.parse_engine = str(preview.get("engine") or "document_ai_ocr")
                row.parse_score = float(preview.get("avg_regex_score") or 0.0)
                if campaign_processing_mode == _CAMPAIGN_PROCESSING_STRICT and strict_preview is not None:
                    row.strict_extraction_preview = strict_preview
                elif campaign_processing_mode != _CAMPAIGN_PROCESSING_STRICT:
                    row.strict_extraction_preview = None
                row.extraction_preview = preview
                row.extracted_at = datetime.utcnow()
                db.flush()
            except Exception as exc:
                logger.warning("Attachment async re-extraction failed for %s: %s", row.file_path, exc)
                row.parse_engine = "error"
                row.parse_score = 0.0
                db.flush()

        auto_qty_updates = _autofill_campaign_item_quantities_from_attachments(
            db,
            tenant_id=campaign.tenant_id,
            campaign_id=campaign.id,
        )
        if auto_qty_updates > 0:
            logger.info(
                "Auto-filled requested_qty from attachment extraction: tenant=%s campaign=%s updated=%s",
                campaign.tenant_id,
                campaign.id,
                auto_qty_updates,
            )
        db.commit()
    except Exception as exc:
        logger.exception("Attachment OCR async job failed: campaign=%s err=%s", campaign_id, exc)
        try:
            db.rollback()
        except Exception:
            pass
    finally:
        db.close()


def _build_file_response(path: Path, filename: str | None = None) -> FileResponse:
    resolved = _resolve_existing_file(path, filename)
    if resolved is None:
        raise HTTPException(status_code=404, detail="File not found")
    suffix = resolved.suffix.lower()
    if suffix == ".pdf":
        media_type = "application/pdf"
    elif suffix == ".xlsx":
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif suffix == ".xlsm":
        media_type = "application/vnd.ms-excel.sheet.macroenabled.12"
    else:
        media_type = "application/octet-stream"
    return FileResponse(
        path=str(resolved),
        media_type=media_type,
        filename=(filename or resolved.name),
    )


def _normalize_requested_characteristics_payload(raw_kv: dict | None) -> dict:
    out: dict[str, list[dict[str, str]]] = {}
    if not isinstance(raw_kv, dict):
        return out
    for raw_key, raw_values in raw_kv.items():
        key = str(raw_key or "").strip()
        if not key:
            continue
        values = raw_values if isinstance(raw_values, list) else [raw_values]
        bucket: list[dict[str, str]] = []
        seen: set[str] = set()
        for raw_value in values:
            value = " ".join(str(raw_value or "").split()).strip()
            if not value:
                continue
            token = value.lower()
            if token in seen:
                continue
            seen.add(token)
            bucket.append({"value": value})
        if bucket:
            out[key] = bucket
    return out


def _normalize_product_name_for_qty_match(value: str | None) -> str:
    text = _clean_display_text(value) or ""
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^a-z0-9]+", " ", text.lower())
    return " ".join(text.split())


def _to_positive_float(value: Any) -> float | None:
    try:
        num = float(value)
    except Exception:
        return None
    if num <= 0:
        return None
    return num


_QTY_EXPLICIT_SIGNAL_RE = re.compile(
    r"\b(qte|quantite|quantité|qty|quantity)\s*[:=]?\s*\d"
    r"|\b\d+(?:[.,]\d+)?\s*(boite|boites|bte|packs?|lots?|flacons?|ampoules?|sachets?|pieces?|pcs?|unites?|u)\b"
    r"|\b(?:x|×|\*)\s*\d+\b",
    re.IGNORECASE,
)
_QTY_TECH_HEAVY_HINT_RE = re.compile(
    r"\b(caracteristiques techniques|caracteristique|tension|puissance|hz|volt|voltage|ipx|temperature|humidite|classe|altitude|pression|frequence)\b",
    re.IGNORECASE,
)
_QTY_NUM_UNIT_RE = re.compile(
    r"\b\d+(?:[.,]\d+)?\s*(v|w|kw|hz|mm|cm|m|kg|g|mg|l|ml|bar|a|°c|c)\b",
    re.IGNORECASE,
)


def _has_explicit_qty_signal(name: str | None, description: str | None) -> bool:
    text = f"{name or ''} || {description or ''}"
    norm = _price_eq_norm_text(text)
    if not norm:
        return False
    return _QTY_EXPLICIT_SIGNAL_RE.search(norm) is not None


def _looks_like_tech_spec_block(name: str | None, description: str | None) -> bool:
    text = f"{name or ''} || {description or ''}"
    norm = _price_eq_norm_text(text)
    if not norm:
        return False
    if _QTY_TECH_HEAVY_HINT_RE.search(norm) is None:
        return False
    unit_hits = len(list(_QTY_NUM_UNIT_RE.finditer(norm)))
    return unit_hits >= 3


def _is_suspicious_default_qty(
    *,
    qty: float | None,
    name: str | None,
    description: str | None,
) -> bool:
    if qty is None:
        return False
    # Common false-positive: model injects quantity=1 for technical paragraphs.
    if abs(float(qty) - 1.0) > 1e-9:
        return False
    if _has_explicit_qty_signal(name, description):
        return False
    return _looks_like_tech_spec_block(name, description)


_PRICE_EQ_UNIT_ALIASES = {
    "l": ("volume", "ml", 1000.0),
    "litre": ("volume", "ml", 1000.0),
    "litres": ("volume", "ml", 1000.0),
    "liter": ("volume", "ml", 1000.0),
    "liters": ("volume", "ml", 1000.0),
    "ml": ("volume", "ml", 1.0),
    "cl": ("volume", "ml", 10.0),
    "dl": ("volume", "ml", 100.0),
    "kg": ("mass", "mg", 1000000.0),
    "g": ("mass", "mg", 1000.0),
    "mg": ("mass", "mg", 1.0),
    "mcg": ("mass", "mg", 0.001),
    "ug": ("mass", "mg", 0.001),
    "µg": ("mass", "mg", 0.001),
    "m": ("length", "mm", 1000.0),
    "cm": ("length", "mm", 10.0),
    "mm": ("length", "mm", 1.0),
    "ui": ("count", "ct", 1.0),
    "iu": ("count", "ct", 1.0),
    "u": ("count", "ct", 1.0),
    "unite": ("count", "ct", 1.0),
    "unites": ("count", "ct", 1.0),
    "unit": ("count", "ct", 1.0),
    "units": ("count", "ct", 1.0),
    "cp": ("count", "ct", 1.0),
    "cpr": ("count", "ct", 1.0),
    "comprime": ("count", "ct", 1.0),
    "comprimes": ("count", "ct", 1.0),
    "capsule": ("count", "ct", 1.0),
    "capsules": ("count", "ct", 1.0),
    "gelule": ("count", "ct", 1.0),
    "gelules": ("count", "ct", 1.0),
    "ampoule": ("count", "ct", 1.0),
    "ampoules": ("count", "ct", 1.0),
    "flacon": ("count", "ct", 1.0),
    "flacons": ("count", "ct", 1.0),
    "sachet": ("count", "ct", 1.0),
    "sachets": ("count", "ct", 1.0),
    "boite": ("count", "ct", 1.0),
    "boites": ("count", "ct", 1.0),
    "bte": ("count", "ct", 1.0),
    "pack": ("count", "ct", 1.0),
    "lot": ("count", "ct", 1.0),
}
_PRICE_EQ_PACK_WORD_RE = re.compile(
    r"\b(pack|boite|boite|bte|lot|blister|sachet|flacon|ampoule|carton|kit)\b",
    re.IGNORECASE,
)
_PRICE_EQ_MULTIPLIER_RE = re.compile(
    r"(?<!\d)(\d+(?:[.,]\d+)?)\s*(?:x|×|\*)\s*(\d+(?:[.,]\d+)?)(?:\s*([a-zA-Zµ]+))?\b"
)
_PRICE_EQ_VALUE_UNIT_RE = re.compile(
    r"(?<![A-Za-z0-9])(\d+(?:[.,]\d+)?)\s*([a-zA-Zµ]+)\b"
)
_PRICE_EQ_PACK_COUNT_RE = re.compile(
    r"\b(?:boite|boite|bte|pack|lot|sachet|flacon|ampoule|blister|carton|kit)\s*(?:de)?\s*(\d+(?:[.,]\d+)?)\b",
    re.IGNORECASE,
)
_PRICE_EQ_X_COUNT_RE = re.compile(r"(?:x|×|\*)\s*(\d+(?:[.,]\d+)?)\b", re.IGNORECASE)



_PRICE_EQ_MACHINE_HINTS = (
    "autoclave",
    "thermosoudeuse",
    "thermo soudeuse",
    "soudeuse",
    "compresseur",
    "fauteuil",
    "station",
    "radiographie",
    "radio dentaire",
    "sterilisateur de paillasse",
    "armoire de sterilisation",
    "table d examen",
    "table d'examen",
    "paravent",
    "nebuliseur",
    "debitmetre",
    "otoscope",
    "laryngoscope",
    "negatoscope",
)
_PRICE_EQ_CONSUMABLE_HINTS = (
    "reactif",
    "reagent",
    "consommable",
    "flacon",
    "ampoule",
    "tube",
    "seringue",
    "aiguille",
    "gant",
    "compresse",
    "bandelette",
    "capsule",
    "gelule",
    "comprime",
    "sachet",
    "boite",
    "kit",
)
_PRICE_EQ_PHARMA_UNIT_HINT_RE = re.compile(
    r"\b(mg|g|kg|mcg|ug|ui|iu|ml|cl|dl|l)\b",
    re.IGNORECASE,
)


def _price_eq_keyword_hits(text: str, keywords: tuple[str, ...]) -> int:
    if not text:
        return 0
    hits = 0
    for kw in keywords:
        token = kw.strip().lower()
        if not token:
            continue
        if " " in token:
            if token in text:
                hits += 1
            continue
        if re.search(rf"\b{re.escape(token)}\b", text):
            hits += 1
    return hits


def _infer_price_eq_context(requested_text: str | None, offered_text: str | None) -> str:
    req = _price_eq_norm_text(requested_text)
    off = _price_eq_norm_text(offered_text)
    combined = f"{req} | {off}".strip(" |")
    if not combined:
        return "unknown"
    machine_hits = _price_eq_keyword_hits(combined, _PRICE_EQ_MACHINE_HINTS)
    consumable_hits = _price_eq_keyword_hits(combined, _PRICE_EQ_CONSUMABLE_HINTS)
    has_pharma_units = _PRICE_EQ_PHARMA_UNIT_HINT_RE.search(combined) is not None

    if machine_hits >= 1 and consumable_hits == 0:
        return "machine"
    if machine_hits >= 2 and consumable_hits <= 1:
        return "machine"
    if consumable_hits >= 2 and machine_hits == 0:
        return "consommable"
    if consumable_hits >= 1 and machine_hits == 0 and has_pharma_units:
        return "consommable"
    if machine_hits == 0 and consumable_hits == 0:
        return "unknown"
    return "mixed"


def _is_price_eq_conversion_allowed(
    *,
    context: str,
    req_measure: dict[str, Any],
    off_measure: dict[str, Any],
    score: float,
    factor: float,
) -> bool:
    if factor <= 0:
        return False
    if abs(factor - 1.0) <= 1e-9:
        return True
    req_dim = str(req_measure.get("dimension") or "").strip().lower()
    off_dim = str(off_measure.get("dimension") or "").strip().lower()
    if not req_dim or req_dim != off_dim:
        return False
    if context == "machine":
        # Machine/equipment: keep direct PU only, no auto conversion.
        return False
    if context in {"unknown", "mixed"}:
        # Unknown nature: only very strong count conversions.
        if req_dim != "count":
            return False
        req_kind = str(req_measure.get("kind") or "").strip().lower()
        off_kind = str(off_measure.get("kind") or "").strip().lower()
        # Packaging evidence (boite/pack/lot counts) can be accepted with a
        # slightly lower threshold in mixed context.
        min_score = 0.93
        if req_kind == "pack_count" and off_kind == "pack_count":
            min_score = 0.88
        elif "pack_count" in {req_kind, off_kind} and "explicit_unit" in {req_kind, off_kind}:
            min_score = 0.90
        if score < min_score:
            return False
        if factor < 0.25 or factor > 4.0:
            return False
        return True
    # context == "consommable"
    if req_dim == "length":
        return score >= 0.94 and 0.5 <= factor <= 2.0
    if req_dim == "count":
        if score < 0.82:
            return False
        return 0.25 <= factor <= 8.0
    if req_dim in {"volume", "mass"}:
        if factor < 0.1 or factor > 10.0:
            return False
        if score < 0.82 and (factor < 0.5 or factor > 2.0):
            return False
        return True
    return False
def _price_eq_norm_text(value: str | None) -> str:
    text = _clean_display_text(value) or str(value or "")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def _price_eq_to_float(raw: str | None) -> float | None:
    if raw is None:
        return None
    cleaned = str(raw).strip().replace(",", ".")
    try:
        value = float(cleaned)
    except Exception:
        return None
    if value <= 0:
        return None
    return value


def _price_eq_collect_text(raw: Any) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, str):
        text = _clean_display_text(raw) or raw.strip()
        return [text] if text else []
    if isinstance(raw, (int, float, bool)):
        return [str(raw)]
    out: list[str] = []
    if isinstance(raw, (list, tuple, set)):
        for entry in raw:
            out.extend(_price_eq_collect_text(entry))
        return out
    if isinstance(raw, dict):
        preferred = ("text", "value", "item", "name", "details")
        seeded = False
        for key in preferred:
            if key in raw:
                seeded = True
                out.extend(_price_eq_collect_text(raw.get(key)))
        if not seeded:
            for value in raw.values():
                out.extend(_price_eq_collect_text(value))
        return out
    return []


def _build_price_eq_requested_text(
    *,
    product_name: str | None,
    requested_description: str | None,
    requested_payload: dict[str, Any] | None,
) -> str:
    chunks: list[str] = []
    if product_name:
        chunks.append(str(product_name))
    if requested_description:
        chunks.append(str(requested_description))
    if isinstance(requested_payload, dict):
        for key in ("characteristics", "requirements", "accessories", "description"):
            values = _price_eq_collect_text(requested_payload.get(key))
            if values:
                chunks.extend(values[:8])
    cleaned: list[str] = []
    seen: set[str] = set()
    for row in chunks:
        text = _price_eq_norm_text(row)
        if not text or text in seen:
            continue
        seen.add(text)
        cleaned.append(text)
    return " | ".join(cleaned[:20])


def _build_price_eq_offered_text(
    *,
    corrected_name: str | None,
    extracted_name: str | None,
    extracted_description: str | None,
) -> str:
    chunks = [corrected_name, extracted_name, extracted_description]
    cleaned: list[str] = []
    seen: set[str] = set()
    for row in chunks:
        text = _price_eq_norm_text(row)
        if not text or text in seen:
            continue
        seen.add(text)
        cleaned.append(text)
    return " | ".join(cleaned[:12])


def _extract_price_eq_measures(text: str | None) -> list[dict[str, Any]]:
    norm = _price_eq_norm_text(text)
    if not norm:
        return []
    out: dict[tuple[str, float, str], dict[str, Any]] = {}

    def _store(
        *,
        dimension: str,
        value_base: float,
        unit_base: str,
        evidence: str,
        kind: str,
        confidence: float,
    ) -> None:
        if value_base <= 0:
            return
        key = (dimension, round(float(value_base), 6), kind)
        prev = out.get(key)
        payload = {
            "dimension": dimension,
            "value": round(float(value_base), 6),
            "unit": unit_base,
            "evidence": evidence.strip(),
            "kind": kind,
            "confidence": max(0.0, min(1.0, float(confidence))),
        }
        if prev is None or payload["confidence"] > float(prev.get("confidence") or 0.0):
            out[key] = payload

    for match in _PRICE_EQ_MULTIPLIER_RE.finditer(norm):
        a = _price_eq_to_float(match.group(1))
        b = _price_eq_to_float(match.group(2))
        if a is None or b is None:
            continue
        raw_unit = (match.group(3) or "").strip().lower()
        unit_payload = _PRICE_EQ_UNIT_ALIASES.get(raw_unit)
        evidence = match.group(0)
        if unit_payload is not None:
            dim, base_unit, factor = unit_payload
            _store(
                dimension=dim,
                value_base=(a * b * factor),
                unit_base=base_unit,
                evidence=evidence,
                kind="multiplier_unit",
                confidence=0.94,
            )
            continue
        window_left = max(0, match.start() - 20)
        window_right = min(len(norm), match.end() + 20)
        has_pack_word = _PRICE_EQ_PACK_WORD_RE.search(norm[window_left:window_right]) is not None
        _store(
            dimension="count",
            value_base=(a * b),
            unit_base="ct",
            evidence=evidence,
            kind="multiplier_count",
            confidence=0.86 if has_pack_word else 0.76,
        )

    for match in _PRICE_EQ_VALUE_UNIT_RE.finditer(norm):
        value = _price_eq_to_float(match.group(1))
        if value is None:
            continue
        raw_unit = (match.group(2) or "").strip().lower()
        payload = _PRICE_EQ_UNIT_ALIASES.get(raw_unit)
        if payload is None:
            continue
        dim, base_unit, factor = payload
        kind = "explicit_unit"
        confidence = 0.95 if dim in {"volume", "mass", "length"} else 0.86
        _store(
            dimension=dim,
            value_base=(value * factor),
            unit_base=base_unit,
            evidence=match.group(0),
            kind=kind,
            confidence=confidence,
        )

    for match in _PRICE_EQ_PACK_COUNT_RE.finditer(norm):
        value = _price_eq_to_float(match.group(1))
        if value is None:
            continue
        _store(
            dimension="count",
            value_base=value,
            unit_base="ct",
            evidence=match.group(0),
            kind="pack_count",
            confidence=0.88,
        )

    for match in _PRICE_EQ_X_COUNT_RE.finditer(norm):
        value = _price_eq_to_float(match.group(1))
        if value is None:
            continue
        _store(
            dimension="count",
            value_base=value,
            unit_base="ct",
            evidence=match.group(0),
            kind="x_count",
            confidence=0.72,
        )

    return sorted(
        out.values(),
        key=lambda row: (
            float(row.get("confidence") or 0.0),
            float(row.get("value") or 0.0),
        ),
        reverse=True,
    )


def _compute_price_equivalence(
    *,
    unit_price: float | None,
    requested_text: str | None,
    offered_text: str | None,
) -> dict[str, Any]:
    raw_price = _to_positive_float(unit_price)
    if raw_price is None:
        return {
            "equivalent_unit_price": None,
            "equivalence_factor": None,
            "equivalence_basis": None,
            "equivalence_confidence": None,
        }
    req_measures = _extract_price_eq_measures(requested_text)
    off_measures = _extract_price_eq_measures(offered_text)
    if not req_measures or not off_measures:
        return {
            "equivalent_unit_price": None,
            "equivalence_factor": None,
            "equivalence_basis": None,
            "equivalence_confidence": None,
        }

    best_pair: tuple[dict[str, Any], dict[str, Any], float, float] | None = None
    for req in req_measures:
        req_dim = str(req.get("dimension") or "")
        req_val = _to_positive_float(req.get("value"))
        req_conf = float(req.get("confidence") or 0.0)
        if not req_dim or req_val is None:
            continue
        for off in off_measures:
            off_dim = str(off.get("dimension") or "")
            off_val = _to_positive_float(off.get("value"))
            off_conf = float(off.get("confidence") or 0.0)
            if req_dim != off_dim or off_val is None:
                continue
            factor = req_val / off_val
            if factor <= 0:
                continue
            score = min(req_conf, off_conf)
            if str(req.get("kind") or "") == str(off.get("kind") or ""):
                score += 0.03
            if "x_count" in {str(req.get("kind") or ""), str(off.get("kind") or "")}:
                score -= 0.05
            score = max(0.0, min(1.0, score))

            # Plausibility guards against accidental matches from OCR noise.
            if score < 0.75 and (factor < 0.20 or factor > 5.0):
                continue
            if score < 0.90 and (factor < 0.02 or factor > 50.0):
                continue

            if best_pair is None or score > best_pair[2]:
                best_pair = (req, off, score, factor)

    if best_pair is None:
        return {
            "equivalent_unit_price": None,
            "equivalence_factor": None,
            "equivalence_basis": None,
            "equivalence_confidence": None,
        }

    req_best, off_best, score, factor = best_pair
    context = _infer_price_eq_context(requested_text, offered_text)
    if not _is_price_eq_conversion_allowed(
        context=context,
        req_measure=req_best,
        off_measure=off_best,
        score=float(score),
        factor=float(factor),
    ):
        return {
            "equivalent_unit_price": None,
            "equivalence_factor": None,
            "equivalence_basis": None,
            "equivalence_confidence": None,
        }
    eq_price = raw_price * factor
    if eq_price <= 0:
        return {
            "equivalent_unit_price": None,
            "equivalence_factor": None,
            "equivalence_basis": None,
            "equivalence_confidence": None,
        }
    basis = (
        f"{req_best.get('evidence') or '?'} -> {off_best.get('evidence') or '?'}"
    )
    return {
        "equivalent_unit_price": round(float(eq_price), 6),
        "equivalence_factor": round(float(factor), 6),
        "equivalence_basis": str(basis)[:220],
        "equivalence_confidence": round(float(score), 4),
    }


_PRICE_EQ_GPT_MIN_CONFIDENCE = 0.70
_PRICE_EQ_GPT_MAX_FACTOR = 100.0
_PRICE_EQ_GPT_MIN_FACTOR = 0.01
_PRICE_EQ_GPT_TIMEOUT_SECONDS = 45
_PRICE_EQ_GPT_MAX_TEXT_CHARS = 1800


def _price_eq_judge_cache_key(
    *,
    requested_text: str,
    offered_text: str,
    context: str,
    model: str,
) -> str:
    payload = {
        "version": "price_eq_gpt_v1",
        "requested_text": _price_eq_norm_text(requested_text)[:_PRICE_EQ_GPT_MAX_TEXT_CHARS],
        "offered_text": _price_eq_norm_text(offered_text)[:_PRICE_EQ_GPT_MAX_TEXT_CHARS],
        "context": str(context or "unknown").strip().lower(),
        "model": str(model or "").strip(),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _gpt_decide_price_equivalence(
    *,
    requested_text: str | None,
    offered_text: str | None,
    context: str,
) -> dict[str, Any] | None:
    if not config.OPENAI_API_KEY:
        return None
    req = _price_eq_norm_text(requested_text)
    off = _price_eq_norm_text(offered_text)
    if not req or not off:
        return None
    model = (
        str(getattr(config, "OPENAI_PRODUCT_MATCH_MODEL", "") or "").strip()
        or str(getattr(config, "OPENAI_ATTACHMENT_MODEL", "") or "").strip()
        or "gpt-5.1"
    )
    fallback_model = str(
        getattr(config, "OPENAI_PRODUCT_MATCH_FALLBACK_MODEL", "") or ""
    ).strip()
    model_candidates: list[str] = [model]
    if fallback_model and fallback_model != model:
        model_candidates.append(fallback_model)
    cache_key = _price_eq_judge_cache_key(
        requested_text=req,
        offered_text=off,
        context=context,
        model=model,
    )
    for model_name in model_candidates:
        cached = get_openai_payload_cache(
            cache_type="price_eq_gpt_judge",
            model=model_name,
            cache_key=cache_key,
        )
        if isinstance(cached, dict):
            return cached

    timeout_s = max(
        20,
        int(getattr(config, "OPENAI_PRODUCT_MATCH_TIMEOUT_SECONDS", _PRICE_EQ_GPT_TIMEOUT_SECONDS) or _PRICE_EQ_GPT_TIMEOUT_SECONDS),
    )
    prompt_req = req[:_PRICE_EQ_GPT_MAX_TEXT_CHARS]
    prompt_off = off[:_PRICE_EQ_GPT_MAX_TEXT_CHARS]
    system_prompt = (
        "Tu es un expert achat biomedical. "
        "Decide uniquement si une conversion quantitative est valable entre produit demande et produit offert. "
        "Retourne strictement un JSON. "
        "convertible='yes' seulement si l'unite/conditionnement est vraiment comparable. "
        "Pour machines/mobilier/instrument durable: en general 'no'. "
        "Si doute: 'uncertain'. "
        "Ne fabrique pas de facteur."
    )
    user_payload = {
        "requested_text": prompt_req,
        "offered_text": prompt_off,
        "context_hint": str(context or "unknown"),
    }
    payload = {
        "model": model,
        "temperature": 0,
        "seed": 42,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
        ],
        "response_format": {
            "type": "json_schema",
            "json_schema": {
                "name": "price_equivalence_judge",
                "strict": True,
                "schema": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "convertible": {
                            "type": "string",
                            "enum": ["yes", "no", "uncertain"],
                        },
                        "factor_requested_per_offered": {
                            "anyOf": [
                                {"type": "number"},
                                {"type": "null"},
                            ]
                        },
                        "confidence": {
                            "type": "number",
                            "minimum": 0,
                            "maximum": 1,
                        },
                        "basis_type": {"type": "string"},
                        "reason": {"type": "string"},
                    },
                    "required": [
                        "convertible",
                        "factor_requested_per_offered",
                        "confidence",
                        "basis_type",
                        "reason",
                    ],
                },
            },
        },
    }
    parsed, model_used = call_chat_completion_json_with_retry(
        operation="price_equivalence_gpt_judge",
        payload=payload,
        primary_model=model,
        fallback_model=fallback_model,
        timeout_s=timeout_s,
        retry_attempts=max(
            1,
            int(getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_ATTEMPTS", 3) or 3),
        ),
        retry_backoff_s=max(
            0.2,
            float(
                getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_BACKOFF_SECONDS", 1.5)
                or 1.5
            ),
        ),
        metadata={"context": context},
    )
    if not isinstance(parsed, dict):
        return None

    convertible = str(parsed.get("convertible") or "").strip().lower()
    if convertible not in {"yes", "no", "uncertain"}:
        convertible = "uncertain"
    confidence = _to_positive_float(parsed.get("confidence"))
    if confidence is None:
        confidence = 0.0
    confidence = max(0.0, min(1.0, float(confidence)))
    factor = _to_positive_float(parsed.get("factor_requested_per_offered"))
    if factor is not None:
        factor = float(factor)
    result = {
        "convertible": convertible,
        "factor_requested_per_offered": factor,
        "confidence": confidence,
        "basis_type": str(parsed.get("basis_type") or "").strip().lower()[:40],
        "reason": str(parsed.get("reason") or "").strip()[:240],
    }
    set_openai_payload_cache(
        cache_type="price_eq_gpt_judge",
        model=(model_used or model),
        cache_key=cache_key,
        payload=result,
    )
    return result

def _compute_price_equivalence_effective(
    *,
    unit_price: float | None,
    requested_text: str | None,
    offered_text: str | None,
    use_ai_prompt: bool,
) -> dict[str, Any]:
    legacy = _compute_price_equivalence(
        unit_price=unit_price,
        requested_text=requested_text,
        offered_text=offered_text,
    )
    if not use_ai_prompt:
        return legacy

    raw_price = _to_positive_float(unit_price)
    if raw_price is None:
        return legacy

    context = _infer_price_eq_context(requested_text, offered_text)
    decision = _gpt_decide_price_equivalence(
        requested_text=requested_text,
        offered_text=offered_text,
        context=context,
    )
    if not isinstance(decision, dict):
        return legacy

    convertible = str(decision.get("convertible") or "").strip().lower()
    confidence = float(decision.get("confidence") or 0.0)
    factor = _to_positive_float(decision.get("factor_requested_per_offered"))

    if (
        convertible == "yes"
        and factor is not None
        and _PRICE_EQ_GPT_MIN_FACTOR <= float(factor) <= _PRICE_EQ_GPT_MAX_FACTOR
        and confidence >= _PRICE_EQ_GPT_MIN_CONFIDENCE
    ):
        eq_price = raw_price * float(factor)
        if eq_price > 0:
            basis_type = str(decision.get("basis_type") or "none").strip().lower() or "none"
            reason = str(decision.get("reason") or "").strip()
            basis = f"gpt_{basis_type}"
            if reason:
                basis = f"{basis}:{reason}"
            return {
                "equivalent_unit_price": round(float(eq_price), 6),
                "equivalence_factor": round(float(factor), 6),
                "equivalence_basis": basis[:220],
                "equivalence_confidence": round(float(confidence), 4),
            }
    if convertible == "no" and confidence >= _PRICE_EQ_GPT_MIN_CONFIDENCE:
        return {
            "equivalent_unit_price": None,
            "equivalence_factor": None,
            "equivalence_basis": "gpt_no_conversion",
            "equivalence_confidence": round(float(confidence), 4),
        }

    return legacy


def _offer_comparable_price(raw: dict[str, Any]) -> float:
    eq = _to_positive_float(raw.get("equivalent_unit_price"))
    if eq is not None:
        return float(eq)
    return float("inf")


def _apply_strict_compensation_price_equivalence(
    *,
    eq_payload: dict[str, Any] | None,
    match_payload: dict[str, Any] | None,
    unit_price: float | None,
) -> dict[str, Any]:
    eq = dict(eq_payload or {})
    if _to_positive_float(eq.get("equivalent_unit_price")) is not None:
        return eq

    raw_price = _to_positive_float(unit_price)
    if raw_price is None:
        return eq

    match = match_payload if isinstance(match_payload, dict) else {}
    compensable = _as_optional_bool(match.get("match_compensable"))
    mode = str(match.get("match_compensation_mode") or "").strip().lower()
    factor = _to_positive_float(match.get("match_compensation_factor"))
    if compensable is not True or factor is None or factor <= 0:
        return eq

    if mode in {"non_compensable", "not_compensated", ""}:
        return eq

    eq_price = raw_price * factor
    if eq_price <= 0:
        return eq

    basis = str(match.get("match_compensation_basis") or "").strip() or None
    if basis:
        basis = f"strict_compensation:{basis}"
    else:
        basis = f"strict_compensation:{mode}"

    confidence = _to_optional_float(match.get("match_compensation_confidence"))
    if confidence is None:
        confidence = 0.74

    eq["equivalent_unit_price"] = round(float(eq_price), 6)
    eq["equivalence_factor"] = round(float(factor), 6)
    eq["equivalence_basis"] = basis[:220]
    eq["equivalence_confidence"] = round(float(confidence), 4)
    return eq


def _as_optional_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    raw = str(value or "").strip().lower()
    if raw in {"true", "1", "yes"}:
        return True
    if raw in {"false", "0", "no"}:
        return False
    return None


def _to_optional_float(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value or "").strip()
    if not raw:
        return None
    raw = raw.replace(" ", "").replace(",", ".")
    try:
        return float(raw)
    except Exception:
        return None


_MEASURE_UNIT_TOKEN_RE = re.compile(
    r"\b\d+(?:[.,]\d+)?\s*([a-zµ]+)\b",
    re.IGNORECASE,
)
_MEASURE_UNIT_CANONICAL: dict[str, str] = {
    "mm": "mm",
    "cm": "cm",
    "m": "m",
    "mg": "mg",
    "g": "g",
    "gr": "g",
    "gram": "g",
    "grams": "g",
    "gramme": "g",
    "grammes": "g",
    "kg": "kg",
    "ml": "ml",
    "cl": "cl",
    "l": "l",
    "litre": "l",
    "litres": "l",
    "liter": "l",
    "liters": "l",
}
_MEASURE_UNIT_FAMILY: dict[str, str] = {
    "mm": "length",
    "cm": "length",
    "m": "length",
    "mg": "mass",
    "g": "mass",
    "kg": "mass",
    "ml": "volume",
    "cl": "volume",
    "l": "volume",
}
_LENGTH_UNITS = {"mm", "cm", "m"}
_MEASURE_UNIT_BASE_FACTOR: dict[str, tuple[str, str, float]] = {
    "mm": ("length", "cm", 0.1),
    "cm": ("length", "cm", 1.0),
    "m": ("length", "cm", 100.0),
    "mg": ("mass", "g", 0.001),
    "g": ("mass", "g", 1.0),
    "kg": ("mass", "g", 1000.0),
    "ml": ("volume", "ml", 1.0),
    "cl": ("volume", "ml", 10.0),
    "l": ("volume", "ml", 1000.0),
}
_PACKAGING_HINT_RE = re.compile(
    r"\b(bt|bte|boite|boites|box|pack|paquet|carton|sachet|b)\b",
    re.IGNORECASE,
)


def _flatten_text_values(raw: Any) -> list[str]:
    out: list[str] = []

    def _walk(value: Any) -> None:
        if value is None:
            return
        if isinstance(value, str):
            text = " ".join(value.split()).strip()
            if text:
                out.append(text)
            return
        if isinstance(value, list):
            for row in value:
                _walk(row)
            return
        if isinstance(value, dict):
            for row in value.values():
                _walk(row)
            return
        text = " ".join(str(value).split()).strip()
        if text:
            out.append(text)

    _walk(raw)
    seen: set[str] = set()
    deduped: list[str] = []
    for row in out:
        token = row.lower()
        if token in seen:
            continue
        seen.add(token)
        deduped.append(row)
    return deduped


def _canonical_measure_unit_token(value: str | None) -> str | None:
    raw = str(value or "").strip().lower()
    if not raw:
        return None
    return _MEASURE_UNIT_CANONICAL.get(raw)


def _extract_measure_units_from_text(text: str | None) -> list[str]:
    value = _clean_display_text(text) or ""
    if not value:
        return []
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    out: list[str] = []
    seen: set[str] = set()
    for m in _MEASURE_UNIT_TOKEN_RE.finditer(value):
        unit = _canonical_measure_unit_token(m.group(1))
        if not unit or unit in seen:
            continue
        seen.add(unit)
        out.append(unit)
    if not out:
        # Allow bare unit value when field already isolates unit.
        unit = _canonical_measure_unit_token(value)
        if unit:
            out.append(unit)
    return out


def _extract_measure_values_from_text(text: str | None) -> list[dict[str, Any]]:
    value = _clean_display_text(text) or ""
    if not value:
        return []
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    out: list[dict[str, Any]] = []
    seen: set[tuple[str, float]] = set()
    for m in _MEASURE_UNIT_TOKEN_RE.finditer(value):
        raw_unit = m.group(1)
        unit = _canonical_measure_unit_token(raw_unit)
        if not unit:
            continue
        base_cfg = _MEASURE_UNIT_BASE_FACTOR.get(unit)
        if base_cfg is None:
            continue
        raw_num = str(m.group(0) or "")
        raw_num_match = re.search(r"\d+(?:[.,]\d+)?", raw_num)
        if not raw_num_match:
            continue
        try:
            numeric = float(raw_num_match.group(0).replace(",", "."))
        except Exception:
            continue
        family, base_unit, factor = base_cfg
        base_value = numeric * factor
        token = (family, round(base_value, 9))
        if token in seen:
            continue
        seen.add(token)
        out.append(
            {
                "unit": unit,
                "family": family,
                "base_unit": base_unit,
                "value": numeric,
                "base_value": base_value,
            }
        )
    return out


_DIMENSION_EXPR_RE = re.compile(
    r"(\d+(?:[.,]\d+)?(?:\s*(?:mm|cm|m))?)\s*(?:x|\*)\s*"
    r"(\d+(?:[.,]\d+)?(?:\s*(?:mm|cm|m))?)"
    r"(?:\s*(?:x|\*)\s*(\d+(?:[.,]\d+)?(?:\s*(?:mm|cm|m))?))?",
    re.IGNORECASE,
)


def _parse_dim_part(part: str | None) -> tuple[float | None, str | None]:
    raw = str(part or "").strip().lower()
    if not raw:
        return None, None
    num_match = re.search(r"\d+(?:[.,]\d+)?", raw)
    if not num_match:
        return None, None
    try:
        value = float(num_match.group(0).replace(",", "."))
    except Exception:
        value = None
    unit_match = re.search(r"(mm|cm|m)\b", raw)
    unit = _canonical_measure_unit_token(unit_match.group(1)) if unit_match else None
    if unit not in _LENGTH_UNITS:
        unit = None
    return value, unit


def _extract_dimension_vectors_from_text(text: str | None) -> list[list[dict[str, Any]]]:
    value = _clean_display_text(text) or ""
    if not value:
        return []
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    out: list[list[dict[str, Any]]] = []
    for m in _DIMENSION_EXPR_RE.finditer(value):
        parts = [m.group(1), m.group(2), m.group(3)]
        parsed: list[dict[str, Any]] = []
        explicit_unit = None
        for p in parts:
            if not p:
                continue
            num, unit = _parse_dim_part(p)
            if num is None:
                continue
            if unit and explicit_unit is None:
                explicit_unit = unit
            parsed.append({"value": num, "unit": unit})
        if not parsed:
            continue
        if explicit_unit:
            for row in parsed:
                if row.get("unit") is None:
                    row["unit"] = explicit_unit
        out.append(parsed)
    return out


def _first_dimension_vector_from_values(values: list[str]) -> list[dict[str, Any]] | None:
    for raw in values:
        vectors = _extract_dimension_vectors_from_text(raw)
        if vectors:
            return vectors[0]
    return None


def _to_base_length(value: float, unit: str) -> float | None:
    cfg = _MEASURE_UNIT_BASE_FACTOR.get(unit)
    if cfg is None:
        return None
    family, _base_unit, factor = cfg
    if family != "length":
        return None
    return float(value) * float(factor)


def _dimension_vectors_compatible(
    offered_vec: list[dict[str, Any]] | None,
    requested_vec: list[dict[str, Any]] | None,
) -> tuple[bool | None, str]:
    if not offered_vec or not requested_vec:
        return None, "dimension_incomplete"
    if len(offered_vec) != len(requested_vec):
        return False, "dimension_count_mismatch"

    off_norm: list[float] = []
    req_norm: list[float] = []
    for i in range(len(offered_vec)):
        off_row = offered_vec[i]
        req_row = requested_vec[i]
        off_val = _to_optional_float(off_row.get("value"))
        req_val = _to_optional_float(req_row.get("value"))
        off_unit = _canonical_measure_unit_token(off_row.get("unit"))
        req_unit = _canonical_measure_unit_token(req_row.get("unit"))
        if off_val is None or req_val is None:
            return None, "dimension_incomplete"

        # Infer missing side unit from the opposite side for normalized comparison.
        if off_unit is None and req_unit in _LENGTH_UNITS:
            off_unit = req_unit
        if req_unit is None and off_unit in _LENGTH_UNITS:
            req_unit = off_unit

        # If both sides have no explicit length unit (ex: "5x5" vs "7.5*7.5"),
        # compare raw dimension values directly as implicit same-unit dimensions.
        if off_unit is None and req_unit is None:
            off_norm.append(float(off_val))
            req_norm.append(float(req_val))
            continue

        if off_unit not in _LENGTH_UNITS or req_unit not in _LENGTH_UNITS:
            return None, "dimension_incomplete"

        off_base = _to_base_length(off_val, off_unit)
        req_base = _to_base_length(req_val, req_unit)
        if off_base is None or req_base is None:
            return None, "dimension_incomplete"
        off_norm.append(off_base)
        req_norm.append(req_base)

    off_norm_sorted = sorted(off_norm)
    req_norm_sorted = sorted(req_norm)
    for left, right in zip(off_norm_sorted, req_norm_sorted):
        if not _values_close(left, right):
            return False, "dimension_mismatch"
    return True, "dimension_normalized_exact"


def _infer_length_unit_from_dimension_vectors(
    offered_vec: list[dict[str, Any]] | None,
    requested_vec: list[dict[str, Any]] | None,
    *,
    compatible: bool | None,
) -> str | None:
    if compatible is not True:
        return None
    if not offered_vec or not requested_vec:
        return None

    explicit_units: list[str] = []
    values: list[float] = []

    for row in [*offered_vec, *requested_vec]:
        unit = _canonical_measure_unit_token(row.get("unit"))
        if unit in _LENGTH_UNITS:
            explicit_units.append(unit)
        value = _to_optional_float(row.get("value"))
        if value is not None:
            values.append(abs(float(value)))

    if explicit_units:
        uniq = sorted(set(explicit_units))
        if len(uniq) == 1:
            return uniq[0]
        # Conflicting explicit length units (example: mixed m/cm) are not
        # safe for a single inferred unit label.
        return None

    # Soft fallback for common medical dimension patterns without explicit unit.
    # Typical "5x5", "7.5x7.5", "10x10" are treated as cm.
    if values and all(0.5 <= v <= 50.0 for v in values):
        return "cm"
    return None


def _first_measure_unit_from_values(values: list[str]) -> str | None:
    for raw in values:
        units = _extract_measure_units_from_text(raw)
        if units:
            return units[0]
    return None


def _first_measure_value_from_values(values: list[str]) -> dict[str, Any] | None:
    for raw in values:
        entries = _extract_measure_values_from_text(raw)
        if entries:
            return entries[0]
    return None


def _extract_packaging_counts_from_text(text: str | None) -> list[int]:
    value = _clean_display_text(text) or ""
    if not value:
        return []
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    out: list[int] = []

    for m in re.finditer(
        r"\b(?:bt|bte|boite|boites|box|pack|paquet|carton|sachet|b)\s*(?:de)?\s*[/\-]?\s*(\d{1,5})\b",
        value,
    ):
        try:
            n = int(m.group(1))
        except Exception:
            continue
        if n > 0:
            out.append(n)

    for m in re.finditer(
        r"\b(\d{1,4})\s*(?:boite|boites|box|pack|packs|paquet|paquets|carton|cartons|sachet|sachets)\s*(?:de)?\s*(\d{1,5})\b",
        value,
    ):
        try:
            qty = int(m.group(1))
            size = int(m.group(2))
        except Exception:
            continue
        if qty > 0 and size > 0:
            out.append(qty * size)

    for m in re.finditer(r"\b(\d{1,4})\s*(?:x|\*)\s*(\d{1,5})\b", value):
        start, end = m.span()
        ctx = value[max(0, start - 18) : min(len(value), end + 18)]
        if _PACKAGING_HINT_RE.search(ctx) is None:
            continue
        try:
            left = int(m.group(1))
            right = int(m.group(2))
        except Exception:
            continue
        total = left * right
        if total > 0:
            out.append(total)

    deduped: list[int] = []
    seen: set[int] = set()
    for n in out:
        if n in seen:
            continue
        seen.add(n)
        deduped.append(n)
    return deduped


def _first_packaging_count_from_values(values: list[str]) -> int | None:
    for raw in values:
        counts = _extract_packaging_counts_from_text(raw)
        if counts:
            return counts[0]
    return None


def _first_measure_value_for_family_from_values(
    values: list[str],
    family: str,
) -> dict[str, Any] | None:
    target = str(family or "").strip().lower()
    if not target:
        return None
    for raw in values:
        entries = _extract_measure_values_from_text(raw)
        for row in entries:
            if str(row.get("family") or "").strip().lower() == target:
                return row
    return None


def _extract_commercial_form_from_values(values: list[str]) -> str | None:
    if not values:
        return None
    token_map = {
        "rouleau": "rouleau",
        "flacon": "flacon",
        "boite": "boite",
        "box": "boite",
        "pack": "pack",
        "sachet": "sachet",
        "piece": "piece",
        "pcs": "piece",
    }
    for raw in values:
        value = _clean_display_text(raw) or ""
        if not value:
            continue
        value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
        value = value.lower()
        for token, canonical in token_map.items():
            if re.search(rf"\b{re.escape(token)}\b", value):
                return canonical
    return None


def _strict_numeric_status(
    offered: dict[str, Any] | None,
    requested: dict[str, Any] | None,
    *,
    convertible_family: bool,
) -> str:
    if not offered and not requested:
        return "not_applicable"
    if not offered or not requested:
        return "missing_one_side"
    offered_base = _to_optional_float(offered.get("base_value"))
    requested_base = _to_optional_float(requested.get("base_value"))
    if offered_base is None or requested_base is None:
        return "missing_one_side"
    if _values_close(float(offered_base), float(requested_base)):
        return "same"
    if convertible_family:
        return "convertible"
    return "different"


def _strict_packaging_status(
    offered_packaging: int | None,
    requested_packaging: int | None,
) -> str:
    if offered_packaging is None and requested_packaging is None:
        return "not_applicable"
    if offered_packaging is None or requested_packaging is None:
        return "missing_one_side"
    if int(offered_packaging) == int(requested_packaging):
        return "same"
    return "convertible"


def _strict_dimension_status(
    offered_dimensions: list[dict[str, Any]] | None,
    requested_dimensions: list[dict[str, Any]] | None,
) -> str:
    if not offered_dimensions and not requested_dimensions:
        return "not_applicable"
    if not offered_dimensions or not requested_dimensions:
        return "missing_one_side"
    compatible, _basis = _dimension_vectors_compatible(offered_dimensions, requested_dimensions)
    if compatible is True:
        return "same"
    if compatible is False:
        return "different"
    return "missing_one_side"


def _strict_commercial_form_status(
    offered_form: str | None,
    requested_form: str | None,
) -> str:
    if not offered_form and not requested_form:
        return "not_applicable"
    if not offered_form or not requested_form:
        return "unknown"
    if offered_form == requested_form:
        return "same"
    return "different"


def _dedupe_alerts(values: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for raw in values:
        token = str(raw or "").strip()
        if not token or token in seen:
            continue
        seen.add(token)
        out.append(token)
    return out


def _build_strict_relation_payload(
    *,
    strict_payload: dict[str, Any] | None,
    unit_payload: dict[str, Any] | None,
    requested_product_name: str | None,
    requested_description: str | None,
    requested_characteristics: Any,
    requested_line_text: str | None,
    strict_match_score: float | None,
    strict_match_threshold: float | None,
    strict_match_decision: str | None,
    match_badges: list[str] | None,
) -> dict[str, Any]:
    strict = strict_payload if isinstance(strict_payload, dict) else {}
    unit_eval = unit_payload if isinstance(unit_payload, dict) else {}
    requested_values = [
        str(_clean_display_text(requested_line_text) or requested_line_text or ""),
        str(requested_product_name or ""),
        str(requested_description or ""),
        *_flatten_text_values(requested_characteristics),
    ]
    offered_values = [
        str(strict.get("name") or ""),
        str(strict.get("unit") or ""),
        str(strict.get("dimensions") or ""),
        str(strict.get("size") or ""),
        str(strict.get("weight") or ""),
        str(strict.get("packaging") or ""),
    ]

    requested_weight = _first_measure_value_for_family_from_values(requested_values, "mass")
    offered_weight = _first_measure_value_for_family_from_values(offered_values, "mass")
    requested_volume = _first_measure_value_for_family_from_values(requested_values, "volume")
    offered_volume = _first_measure_value_for_family_from_values(offered_values, "volume")
    requested_dimensions = _first_dimension_vector_from_values(requested_values)
    offered_dimensions = _first_dimension_vector_from_values(offered_values)
    requested_packaging = _first_packaging_count_from_values(requested_values)
    offered_packaging = _first_packaging_count_from_values(offered_values)
    requested_form = _extract_commercial_form_from_values(requested_values)
    offered_form = _extract_commercial_form_from_values(offered_values)

    spec_comparison = {
        "weight_status": _strict_numeric_status(
            offered=offered_weight,
            requested=requested_weight,
            convertible_family=False,
        ),
        "volume_status": _strict_numeric_status(
            offered=offered_volume,
            requested=requested_volume,
            convertible_family=True,
        ),
        "dimension_status": _strict_dimension_status(offered_dimensions, requested_dimensions),
        "packaging_status": _strict_packaging_status(offered_packaging, requested_packaging),
        "commercial_form_status": _strict_commercial_form_status(offered_form, requested_form),
    }

    score = _to_optional_float(strict_match_score)
    threshold = _to_optional_float(strict_match_threshold)
    if threshold is None:
        threshold = 0.0
    decision = str(strict_match_decision or "").strip().lower()
    family_match = bool(decision == "matched" and score is not None and float(score) >= float(threshold))
    family_reason = "same_product_family" if family_match else "different_or_uncertain_family"

    comp_mode = str(unit_eval.get("match_compensation_mode") or "").strip().lower()
    comp_factor = _to_optional_float(unit_eval.get("match_compensation_factor"))
    comp_basis = str(unit_eval.get("match_compensation_basis") or "").strip() or None
    compensable = _as_optional_bool(unit_eval.get("match_compensable"))
    conversion_allowed = bool(
        compensable is True and comp_mode not in {"", "unknown", "non_compensable", "not_compensated"}
    )
    conversion_mode = comp_mode or "none"
    conversion_factor = comp_factor
    conversion_basis = comp_basis
    # Conservative fallback requested by business: keep price baseline when strict
    # match is accepted but no conversion factor can be proven.
    if conversion_factor is None and family_match and conversion_mode in {"unknown", "none"}:
        conversion_factor = 1.0
        if not conversion_basis:
            conversion_basis = "fallback_identity_unknown"

    blocking_diff = (
        spec_comparison["weight_status"] == "different"
        or spec_comparison["dimension_status"] == "different"
        or spec_comparison["commercial_form_status"] == "different"
        or compensable is False
        or comp_mode == "non_compensable"
    )
    has_convertible = (
        spec_comparison["volume_status"] == "convertible"
        or spec_comparison["packaging_status"] == "convertible"
        or conversion_allowed
    )
    has_uncertain = any(
        spec_comparison[key] in {"missing_one_side", "unknown"}
        for key in (
            "weight_status",
            "volume_status",
            "dimension_status",
            "packaging_status",
            "commercial_form_status",
        )
    ) or conversion_mode in {"unknown", "none"}
    all_same_or_na = all(
        spec_comparison[key] in {"same", "not_applicable"}
        for key in (
            "weight_status",
            "volume_status",
            "dimension_status",
            "packaging_status",
            "commercial_form_status",
        )
    )

    if not family_match:
        final_decision = "NO_MATCH"
        decision_reason = "family_mismatch"
    elif blocking_diff:
        final_decision = "MATCH_SAME_PRODUCT_DIFFERENT_SPEC"
        decision_reason = "blocking_spec_difference"
    elif all_same_or_na:
        final_decision = "EXACT_MATCH"
        decision_reason = "all_specs_same_or_not_applicable"
    elif has_convertible:
        final_decision = "MATCH_WITH_CONVERSION"
        decision_reason = "convertible_spec_detected"
    elif has_uncertain:
        final_decision = "MATCH_UNCERTAIN"
        decision_reason = "missing_or_unknown_spec_info"
    else:
        final_decision = "MATCH_SAME_PRODUCT_DIFFERENT_SPEC"
        decision_reason = "non_convertible_spec_difference"

    alerts: list[str] = []
    if final_decision == "EXACT_MATCH":
        alerts.append("EXACT")
    elif final_decision == "MATCH_WITH_CONVERSION":
        alerts.append("CONVERSION")
    elif final_decision == "MATCH_SAME_PRODUCT_DIFFERENT_SPEC":
        alerts.append("SPEC_DIFF")
    elif final_decision == "MATCH_UNCERTAIN":
        alerts.append("UNCERTAIN_SPEC")
    elif final_decision == "NO_MATCH":
        alerts.append("NO_MATCH")

    badge_map = {
        "TVA_GLOBALE": "GLOBAL_VAT",
        "QTY_CORRIGEE": "QTY_FIXED",
        "PRIX_CORRIGE": "PRICE_FIXED",
        "SANS_UNITE": "NO_UNIT",
        "UNITE_DEVINEE": "UNIT_INFERRED",
    }
    for badge in (match_badges or []):
        mapped = badge_map.get(str(badge or "").strip())
        if mapped:
            alerts.append(mapped)

    alert_colors = {
        "EXACT": "green",
        "CONVERSION": "green",
        "SPEC_DIFF": "red",
        "UNCERTAIN_SPEC": "gray",
        "NO_MATCH": "red",
        "GLOBAL_VAT": "blue",
        "QTY_FIXED": "amber",
        "PRICE_FIXED": "amber",
        "NO_UNIT": "gray",
        "UNIT_INFERRED": "yellow",
    }
    alerts_dedup = _dedupe_alerts(alerts)

    return {
        "strict_family_match": family_match,
        "strict_family_reason": family_reason,
        "strict_spec_comparison": spec_comparison,
        "strict_conversion_allowed": conversion_allowed,
        "strict_conversion_mode": conversion_mode,
        "strict_conversion_factor": round(float(conversion_factor), 6) if conversion_factor is not None else None,
        "strict_conversion_basis": conversion_basis,
        "strict_final_decision": final_decision,
        "strict_decision_reason": decision_reason,
        "strict_alerts": alerts_dedup,
        "strict_alert_colors": {
            key: alert_colors.get(key, "gray")
            for key in alerts_dedup
        },
    }


def _units_are_compatible(offered_unit: str | None, requested_unit: str | None) -> bool | None:
    if not offered_unit or not requested_unit:
        return None
    left = _MEASURE_UNIT_FAMILY.get(offered_unit)
    right = _MEASURE_UNIT_FAMILY.get(requested_unit)
    if not left or not right:
        return None
    return left == right


def _values_close(a: float, b: float) -> bool:
    return abs(float(a) - float(b)) <= max(1e-6, 1e-6 * max(abs(float(a)), abs(float(b))))


def _evaluate_unit_compensation(
    *,
    offered_measure: dict[str, Any] | None,
    requested_measure: dict[str, Any] | None,
    offered_packaging: int | None,
    requested_packaging: int | None,
    dimension_compatible: bool | None = None,
    dimension_basis: str | None = None,
    inferred_phase: bool,
) -> dict[str, Any]:
    mode = "unknown"
    compatible: bool | None = None
    compensable: bool | None = None
    factor: float | None = None
    basis = None
    confidence: float | None = None

    if dimension_compatible is True:
        mode = "equivalent_exact"
        compatible = True
        compensable = True
        factor = 1.0
        basis = dimension_basis or "dimension_normalized_exact"
        confidence = 1.0 if not inferred_phase else 0.92
    elif dimension_compatible is False:
        mode = "non_compensable"
        compatible = False
        compensable = False
        factor = None
        basis = dimension_basis or "dimension_mismatch"
        confidence = 1.0
    elif isinstance(offered_measure, dict) and isinstance(requested_measure, dict):
        of_family = str(offered_measure.get("family") or "")
        rq_family = str(requested_measure.get("family") or "")
        of_base = _to_optional_float(offered_measure.get("base_value"))
        rq_base = _to_optional_float(requested_measure.get("base_value"))
        if of_family and rq_family and of_family == rq_family:
            compatible = True
            if of_base is not None and rq_base is not None and of_base > 0 and rq_base > 0:
                if _values_close(of_base, rq_base):
                    mode = "equivalent_exact"
                    compensable = True
                    factor = 1.0
                    basis = f"exact_{of_family}"
                    confidence = 1.0
                elif of_family == "volume":
                    mode = "compensable_liquid"
                    compensable = True
                    factor = rq_base / of_base if of_base else None
                    basis = "liquid_conversion"
                    confidence = 0.98 if not inferred_phase else 0.90
                else:
                    mode = "non_compensable"
                    compensable = False
                    factor = None
                    basis = f"{of_family}_not_compensable_unless_equal"
                    confidence = 1.0 if not inferred_phase else 0.85
            else:
                if of_family == "volume":
                    mode = "compensable_liquid"
                    compensable = True
                    factor = None
                    basis = "liquid_family_inferred"
                    confidence = 0.70
                else:
                    mode = "non_compensable"
                    compensable = False
                    factor = None
                    basis = f"{of_family}_family_only_non_compensable"
                    confidence = 0.70
        elif of_family and rq_family:
            mode = "non_compensable"
            compatible = False
            compensable = False
            basis = "different_measure_families"
            confidence = 1.0
    elif offered_packaging or requested_packaging:
        # Business rule: missing packaging count defaults to 1 unit.
        # Example: requested "boite 300" vs offered without packaging -> factor 300.
        offered_pack = int(offered_packaging) if (offered_packaging and offered_packaging > 0) else 1
        requested_pack = int(requested_packaging) if (requested_packaging and requested_packaging > 0) else 1
        compatible = True
        if offered_pack == requested_pack:
            mode = "equivalent_exact"
            compensable = True
            factor = 1.0
            basis = "packaging_equal" if (offered_packaging and requested_packaging) else "packaging_assumed_one_equal"
            confidence = 1.0 if (offered_packaging and requested_packaging) else 0.85
        else:
            mode = "compensable_packaging"
            compensable = True
            factor = float(requested_pack) / float(offered_pack)
            basis = "packaging_ratio" if (offered_packaging and requested_packaging) else "packaging_ratio_assumed_one"
            confidence = 0.95 if (offered_packaging and requested_packaging) else 0.80

    return {
        "match_compensation_mode": mode,
        "match_unit_compatible": compatible,
        "match_compensable": compensable,
        "match_compensation_factor": round(float(factor), 6) if factor is not None else None,
        "match_compensation_basis": basis,
        "match_compensation_confidence": round(float(confidence), 4) if confidence is not None else None,
    }


def _build_match_unit_payload(
    *,
    strict_payload: dict[str, Any] | None,
    requested_product_name: str | None,
    requested_description: str | None,
    requested_characteristics: Any,
    requested_line_text: str | None = None,
) -> dict[str, Any]:
    requested_line = _clean_display_text(requested_line_text) or str(requested_line_text or "")
    strict = strict_payload if isinstance(strict_payload, dict) else {}
    requested_values = [
        str(requested_line or ""),
        str(requested_product_name or ""),
        str(requested_description or ""),
        *_flatten_text_values(requested_characteristics),
    ]
    offered_values = [
        str(strict.get("unit") or ""),
        str(strict.get("name") or ""),
        str(strict.get("dimensions") or ""),
        str(strict.get("size") or ""),
        str(strict.get("weight") or ""),
        str(strict.get("packaging") or ""),
    ]

    unit_offered = _first_measure_unit_from_values(offered_values)
    unit_requested = _first_measure_unit_from_values(requested_values)
    units_present_both = bool(unit_offered and unit_requested)
    offered_measure = _first_measure_value_from_values(offered_values)
    requested_measure = _first_measure_value_from_values(requested_values)
    offered_infer_values = [
        str(strict.get("name") or ""),
        str(strict.get("dimensions") or ""),
        str(strict.get("size") or ""),
        str(strict.get("weight") or ""),
        str(strict.get("packaging") or ""),
    ]
    requested_infer_values = requested_values
    inferred_offered_from_text = _first_measure_unit_from_values(offered_infer_values)
    inferred_requested_from_text = _first_measure_unit_from_values(requested_infer_values)
    offered_dimensions = _first_dimension_vector_from_values(offered_values)
    requested_dimensions = _first_dimension_vector_from_values(requested_values)
    dim_compatible, dim_basis = _dimension_vectors_compatible(offered_dimensions, requested_dimensions)
    offered_packaging = _first_packaging_count_from_values(offered_values)
    requested_packaging = _first_packaging_count_from_values(requested_values)

    explicit_comp_eval = {
        "match_compensable": None,
        "match_compensation_mode": "unknown",
        "match_compensation_factor": None,
        "match_compensation_basis": None,
        "match_compensation_confidence": None,
        "match_unit_compatible": None,
    }
    explicit_compatible: bool | None = None
    if units_present_both:
        explicit_comp_eval = _evaluate_unit_compensation(
            offered_measure=offered_measure,
            requested_measure=requested_measure,
            offered_packaging=offered_packaging,
            requested_packaging=requested_packaging,
            dimension_compatible=dim_compatible,
            dimension_basis=dim_basis,
            inferred_phase=False,
        )
        explicit_compatible = _as_optional_bool(explicit_comp_eval.get("match_unit_compatible"))

    inferred_offered: str | None = None
    inferred_requested: str | None = None
    inferred_comp_eval = {
        "match_compensable": None,
        "match_compensation_mode": "unknown",
        "match_compensation_factor": None,
        "match_compensation_basis": None,
        "match_compensation_confidence": None,
        "match_unit_compatible": None,
    }
    inferred_compatible: bool | None = None
    compare_mode = "missing"
    compatible: bool | None = None
    comp_eval = {
        "match_compensable": None,
        "match_compensation_mode": "unknown",
        "match_compensation_factor": None,
        "match_compensation_basis": None,
        "match_compensation_confidence": None,
    }

    inferred_offered = (
        _canonical_measure_unit_token(str(strict.get("unit_inferred") or "").strip())
        or inferred_offered_from_text
        or unit_offered
    )
    inferred_requested = inferred_requested_from_text or unit_requested
    inferred_from_dimensions = _infer_length_unit_from_dimension_vectors(
        offered_dimensions,
        requested_dimensions,
        compatible=dim_compatible,
    )
    if not inferred_offered and inferred_from_dimensions:
        inferred_offered = inferred_from_dimensions
    if not inferred_requested and inferred_from_dimensions:
        inferred_requested = inferred_from_dimensions
    inferred_offered_measure = offered_measure
    inferred_requested_measure = requested_measure
    if inferred_offered and (not inferred_offered_measure):
        fam = _MEASURE_UNIT_FAMILY.get(inferred_offered)
        if fam:
            inferred_offered_measure = {
                "unit": inferred_offered,
                "family": fam,
                "base_value": None,
            }
    if inferred_requested and (not inferred_requested_measure):
        fam = _MEASURE_UNIT_FAMILY.get(inferred_requested)
        if fam:
            inferred_requested_measure = {
                "unit": inferred_requested,
                "family": fam,
                "base_value": None,
            }
    if (
        inferred_offered_measure
        or inferred_requested_measure
        or offered_packaging
        or requested_packaging
        or (dim_compatible is not None)
    ):
        inferred_comp_eval = _evaluate_unit_compensation(
            offered_measure=inferred_offered_measure,
            requested_measure=inferred_requested_measure,
            offered_packaging=offered_packaging,
            requested_packaging=requested_packaging,
            dimension_compatible=dim_compatible,
            dimension_basis=dim_basis,
            inferred_phase=True,
        )
        inferred_compatible = _as_optional_bool(inferred_comp_eval.get("match_unit_compatible"))

    if units_present_both:
        compare_mode = "present_both"
        comp_eval = explicit_comp_eval
        compatible = explicit_compatible
        # Explicit units remain the primary truth when both sides are present.
        # Keep inferred phase only as diagnostic context for badge rendering.
        if (
            explicit_compatible is False
            and inferred_compatible is True
            and inferred_comp_eval.get("match_compensation_mode") != "unknown"
        ):
            compare_mode = "present_both_inferred_rescue"
    else:
        if inferred_comp_eval.get("match_compensation_mode") != "unknown":
            compare_mode = "inferred_both"
            comp_eval = inferred_comp_eval
            compatible = inferred_compatible
        else:
            compare_mode = "inferred_incomplete"
            compatible = None

    return {
        "match_units_present_both": units_present_both,
        "match_unit_offered": unit_offered,
        "match_unit_requested": unit_requested,
        "match_unit_inferred_offered": inferred_offered if not units_present_both else None,
        "match_unit_inferred_requested": inferred_requested if not units_present_both else None,
        "match_unit_compare_mode": compare_mode,
        "match_unit_compatible": compatible,
        "match_unit_missing": not units_present_both,
        "match_unit_inferred": inferred_offered if not units_present_both else None,
        "match_unit_explicit_compatible": explicit_compatible,
        "match_unit_inferred_compatible": inferred_compatible,
        **comp_eval,
    }


def _has_measure_unit_in_strict_payload(strict_payload: dict[str, Any] | None) -> bool:
    strict = strict_payload if isinstance(strict_payload, dict) else {}
    probe = " ".join(
        [
            str(strict.get("name") or ""),
            str(strict.get("dimensions") or ""),
            str(strict.get("size") or ""),
            str(strict.get("weight") or ""),
        ]
    )
    return _MEASURE_UNIT_TOKEN_RE.search(probe) is not None


def _build_match_badges(
    *,
    eq_payload: dict[str, Any] | None,
    strict_payload: dict[str, Any] | None,
    unit_payload: dict[str, Any] | None = None,
) -> dict[str, Any]:
    strict = strict_payload if isinstance(strict_payload, dict) else {}
    eq = eq_payload if isinstance(eq_payload, dict) else {}
    unit_eval = unit_payload if isinstance(unit_payload, dict) else {}
    if not strict:
        return {
            "match_badges": [],
            "match_badge_score": None,
            "match_badge_band": None,
            "match_badge_breakdown": {},
        }
    trace = strict.get("numeric_trace") if isinstance(strict.get("numeric_trace"), dict) else {}
    factor = _to_optional_float(unit_eval.get("match_compensation_factor"))
    if factor is None:
        factor = _to_optional_float(eq.get("equivalence_factor"))
    unit_missing = bool(unit_payload.get("match_unit_missing")) if isinstance(unit_payload, dict) else bool(strict.get("unit_missing"))
    unit_inferred = (
        str(unit_payload.get("match_unit_inferred") or "").strip() or None
        if isinstance(unit_payload, dict)
        else (str(strict.get("unit_inferred") or "").strip() or None)
    )
    measure_unit_present = (
        _has_measure_unit_in_strict_payload(strict)
        or bool(str(unit_eval.get("match_unit_offered") or "").strip())
        or bool(str(unit_eval.get("match_unit_inferred_offered") or "").strip())
    )
    compare_mode = str(unit_eval.get("match_unit_compare_mode") or "").strip().lower() or None
    explicit_compatible = _as_optional_bool(unit_eval.get("match_unit_explicit_compatible"))
    inferred_compatible = _as_optional_bool(unit_eval.get("match_unit_inferred_compatible"))
    final_compatible = _as_optional_bool(unit_eval.get("match_unit_compatible"))
    eval_compensable = _as_optional_bool(unit_eval.get("match_compensable"))
    comp_mode = str(unit_eval.get("match_compensation_mode") or "").strip().lower()

    badges: list[str] = []
    qty_final = _to_optional_float(strict.get("quantity"))
    qty_raw = _to_optional_float(trace.get("quantity_raw"))
    if qty_final is not None and qty_raw is not None and abs(qty_final - qty_raw) > 1e-9:
        badges.append("QTY_CORRIGEE")

    price_final = _to_optional_float(strict.get("unit_price"))
    price_raw = _to_optional_float(trace.get("unit_price_raw"))
    if price_final is not None and price_raw is not None and abs(price_final - price_raw) > 1e-9:
        badges.append("PRIX_CORRIGE")

    if bool(trace.get("swap_applied")):
        badges.append("SWAP_QTY_PRIX")

    if str(strict.get("vat_source") or "").strip().lower() == "calculated_global":
        badges.append("TVA_GLOBALE")

    if factor is not None and abs(factor - 1.0) > 1e-9:
        badges.append("COEF_CONVERSION")

    # Match-level compensability badges (strict comparison only).
    if eval_compensable is False:
        badges.append("NON_COMPENSABLE")
    elif eval_compensable is True:
        badges.append("COMPENSABLE")

    # Color-coded strict alerts requested by business.
    # RED: non-compensable strict mismatch (typically dimensions/weight mismatch).
    # ORANGE: explicit/inferred units remain incompatible.
    # YELLOW: explicit mismatch but inferred-compatible fallback exists.
    if eval_compensable is False or comp_mode == "non_compensable":
        badges.append("RED_NON_COMPENSABLE")
    else:
        if explicit_compatible is False and inferred_compatible is True:
            badges.append("YELLOW_UNIT_INFERRED_COMPATIBLE")
        elif explicit_compatible is False:
            badges.append("ORANGE_UNIT_INCOMPATIBLE")
        elif compare_mode == "inferred_both" and final_compatible is False:
            badges.append("ORANGE_UNIT_INCOMPATIBLE")

    # Missing/inferred unit badges.
    # `SANS_UNITE` means no explicit commercial unit AND no measurement unit
    # detected in strict fields (cm/m/l/g...).
    if unit_missing and not measure_unit_present:
        badges.append("SANS_UNITE")
    if measure_unit_present:
        badges.append("UNITE_MESURE_PRESENTE")
    if unit_inferred:
        badges.append("UNITE_DEVINEE")

    if not badges:
        return {
            "match_badges": [],
            "match_badge_score": 1.0,
            "match_badge_band": "high",
            "match_badge_breakdown": {},
        }

    penalties = {
        "QTY_CORRIGEE": 0.12,
        "PRIX_CORRIGE": 0.10,
        "SWAP_QTY_PRIX": 0.20,
        "TVA_GLOBALE": 0.05,
        "COEF_CONVERSION": 0.10,
        "NON_COMPENSABLE": 0.00,
        "COMPENSABLE": 0.00,
        "RED_NON_COMPENSABLE": 0.22,
        "ORANGE_UNIT_INCOMPATIBLE": 0.14,
        "YELLOW_UNIT_INFERRED_COMPATIBLE": 0.07,
        "SANS_UNITE": 0.15,
        "UNITE_MESURE_PRESENTE": 0.00,
        "UNITE_DEVINEE": 0.05,
    }
    colors = {
        "QTY_CORRIGEE": "amber",
        "PRIX_CORRIGE": "amber",
        "SWAP_QTY_PRIX": "orange",
        "TVA_GLOBALE": "blue",
        "COEF_CONVERSION": "purple",
        "NON_COMPENSABLE": "red",
        "COMPENSABLE": "green",
        "RED_NON_COMPENSABLE": "red",
        "ORANGE_UNIT_INCOMPATIBLE": "orange",
        "YELLOW_UNIT_INFERRED_COMPATIBLE": "yellow",
        "SANS_UNITE": "gray",
        "UNITE_MESURE_PRESENTE": "green",
        "UNITE_DEVINEE": "yellow",
    }
    score = 1.0
    breakdown: dict[str, Any] = {}
    for badge in badges:
        score -= penalties.get(badge, 0.0)
        breakdown[badge] = {
            "penalty": penalties.get(badge, 0.0),
            "color": colors.get(badge),
        }
    score = max(0.0, min(1.0, score))
    band = "high" if score >= 0.85 else ("medium" if score >= 0.60 else "low")
    return {
        "match_badges": badges,
        "match_badge_score": round(score, 4),
        "match_badge_band": band,
        "match_badge_breakdown": breakdown,
    }


def _build_match_comparison_payload(
    *,
    eq_payload: dict[str, Any] | None,
    strict_payload: dict[str, Any] | None,
    unit_payload: dict[str, Any] | None = None,
    requested_product_name: str | None = None,
    requested_description: str | None = None,
    requested_characteristics: Any = None,
    requested_line_text: str | None = None,
    strict_match_score: float | None = None,
    strict_match_threshold: float | None = None,
    strict_match_decision: str | None = None,
    strict_mode: bool = False,
) -> dict[str, Any]:
    eq = eq_payload if isinstance(eq_payload, dict) else {}
    strict = strict_payload if isinstance(strict_payload, dict) else {}

    factor = _to_positive_float(eq.get("equivalence_factor"))
    confidence = _to_positive_float(eq.get("equivalence_confidence"))
    basis = str(eq.get("equivalence_basis") or "").strip() or None

    if factor is None:
        mode = "not_compensated"
        compensable = False
    elif abs(float(factor) - 1.0) <= 1e-9:
        mode = "equal"
        compensable = True
    else:
        mode = "converted"
        compensable = True

    out = {
        "match_compensable": compensable,
        "match_compensation_mode": mode,
        "match_compensation_factor": round(float(factor), 6) if factor is not None else None,
        "match_compensation_basis": basis,
        "match_compensation_confidence": round(float(confidence), 4) if confidence is not None else None,
        "match_unit_missing": None,
        "match_unit_inferred": None,
        "match_units_present_both": None,
        "match_unit_offered": None,
        "match_unit_requested": None,
        "match_unit_inferred_offered": None,
        "match_unit_inferred_requested": None,
        "match_unit_compare_mode": None,
        "match_unit_compatible": None,
        "match_badges": None,
        "match_badge_score": None,
        "match_badge_band": None,
        "match_badge_breakdown": None,
        "strict_family_match": None,
        "strict_family_reason": None,
        "strict_spec_comparison": None,
        "strict_conversion_allowed": None,
        "strict_conversion_mode": None,
        "strict_conversion_factor": None,
        "strict_conversion_basis": None,
        "strict_final_decision": None,
        "strict_decision_reason": None,
        "strict_alerts": None,
        "strict_alert_colors": None,
    }
    if strict_mode:
        unit_eval = unit_payload if isinstance(unit_payload, dict) else {}
        unit_compensable = _as_optional_bool(unit_eval.get("match_compensable"))
        unit_mode = str(unit_eval.get("match_compensation_mode") or "").strip() or None
        unit_factor = _to_optional_float(unit_eval.get("match_compensation_factor"))
        unit_basis = str(unit_eval.get("match_compensation_basis") or "").strip() or None
        unit_confidence = _to_optional_float(unit_eval.get("match_compensation_confidence"))

        if unit_compensable is not None:
            out["match_compensable"] = unit_compensable
        if unit_mode:
            out["match_compensation_mode"] = unit_mode
        out["match_compensation_factor"] = (
            round(float(unit_factor), 6) if unit_factor is not None else None
        )
        out["match_compensation_basis"] = unit_basis
        out["match_compensation_confidence"] = (
            round(float(unit_confidence), 4) if unit_confidence is not None else None
        )

        badges_payload = _build_match_badges(
            eq_payload=eq,
            strict_payload=strict,
            unit_payload=unit_eval,
        )
        out["match_unit_missing"] = _as_optional_bool(unit_eval.get("match_unit_missing"))
        out["match_unit_inferred"] = str(unit_eval.get("match_unit_inferred") or "").strip() or None
        out["match_units_present_both"] = _as_optional_bool(unit_eval.get("match_units_present_both"))
        out["match_unit_offered"] = str(unit_eval.get("match_unit_offered") or "").strip() or None
        out["match_unit_requested"] = str(unit_eval.get("match_unit_requested") or "").strip() or None
        out["match_unit_inferred_offered"] = str(unit_eval.get("match_unit_inferred_offered") or "").strip() or None
        out["match_unit_inferred_requested"] = str(unit_eval.get("match_unit_inferred_requested") or "").strip() or None
        out["match_unit_compare_mode"] = str(unit_eval.get("match_unit_compare_mode") or "").strip() or None
        out["match_unit_compatible"] = _as_optional_bool(unit_eval.get("match_unit_compatible"))
        out.update(badges_payload)
        out.update(
            _build_strict_relation_payload(
                strict_payload=strict,
                unit_payload=unit_eval,
                requested_product_name=requested_product_name,
                requested_description=requested_description,
                requested_characteristics=requested_characteristics,
                requested_line_text=requested_line_text,
                strict_match_score=strict_match_score,
                strict_match_threshold=strict_match_threshold,
                strict_match_decision=strict_match_decision,
                match_badges=out.get("match_badges") if isinstance(out.get("match_badges"), list) else [],
            )
        )
    return out


def _collect_attachment_qty_candidates(
    db: Session,
    *,
    tenant_id: int,
    campaign_id: int,
) -> list[dict[str, Any]]:
    rows = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.campaign_id == campaign_id,
            CampaignAttachment.tenant_id == tenant_id,
            CampaignAttachment.extraction_preview.isnot(None),
        )
        .all()
    )
    out: list[dict[str, Any]] = []
    seen: set[tuple[str, float]] = set()
    for row in rows:
        preview = row.extraction_preview if isinstance(row.extraction_preview, dict) else {}
        products = preview.get("products") if isinstance(preview.get("products"), list) else []
        for product in products:
            if not isinstance(product, dict):
                continue
            raw_name = product.get("name") or product.get("product_name")
            display_name = _clean_display_text(str(raw_name or "")) or str(raw_name or "")
            norm_name = _normalize_product_name_for_qty_match(str(raw_name or ""))
            qty = _to_positive_float(product.get("quantity"))
            description = _clean_display_text(str(product.get("description") or "")) or None
            if not norm_name or qty is None:
                continue
            if _is_suspicious_default_qty(
                qty=qty,
                name=display_name,
                description=description,
            ):
                continue
            token = (norm_name, round(qty, 6))
            if token in seen:
                continue
            seen.add(token)
            out.append(
                {
                    "name": display_name,
                    "name_norm": norm_name,
                    "quantity": qty,
                    "description": description,
                    "attachment_file": str(row.file_name or ""),
                }
            )
    return out


def _collect_attachment_source_candidates(
    db: Session,
    *,
    tenant_id: int,
    campaign_id: int,
) -> list[dict[str, Any]]:
    rows = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.campaign_id == campaign_id,
            CampaignAttachment.tenant_id == tenant_id,
            CampaignAttachment.extraction_preview.isnot(None),
        )
        .all()
    )
    out: list[dict[str, Any]] = []
    seen: set[tuple[int, str]] = set()
    for row in rows:
        preview = row.extraction_preview if isinstance(row.extraction_preview, dict) else {}
        products = preview.get("products") if isinstance(preview.get("products"), list) else []
        for product in products:
            if not isinstance(product, dict):
                continue
            raw_name = product.get("name") or product.get("product_name")
            display_name = _clean_display_text(str(raw_name or "")) or str(raw_name or "")
            norm_name = _normalize_product_name_for_qty_match(str(raw_name or ""))
            if not norm_name:
                continue
            token = (int(row.id), norm_name)
            if token in seen:
                continue
            seen.add(token)
            out.append(
                {
                    "name": display_name,
                    "name_norm": norm_name,
                    "attachment_id": int(row.id),
                    "attachment_file": str(row.file_name or ""),
                    "attachment_path": str(row.file_path or ""),
                }
            )
    return out


def _rank_attachment_source_candidates(
    product_name: str | None,
    candidates: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    target = _normalize_product_name_for_qty_match(product_name)
    if not target or not candidates:
        return []

    target_tokens = set(target.split())
    ranked: list[dict[str, Any]] = []
    for candidate in candidates:
        cand_name = str(candidate.get("name_norm") or "").strip()
        if not cand_name:
            continue
        method = "fuzzy"
        score = 0.0
        if cand_name == target:
            method = "exact"
            score = 1.0
        elif len(target) >= 10 and (target in cand_name or cand_name in target):
            method = "substring"
            seq = difflib.SequenceMatcher(None, target, cand_name).ratio()
            score = max(0.93, seq)
        else:
            seq = difflib.SequenceMatcher(None, target, cand_name).ratio()
            cand_tokens = set(cand_name.split())
            overlap = (len(target_tokens & cand_tokens) / max(1, len(target_tokens))) if target_tokens else 0.0
            score = (0.75 * seq) + (0.25 * overlap)
            if score < 0.70:
                continue
        ranked.append(
            {
                "score": round(float(score), 6),
                "method": method,
                "name": str(candidate.get("name") or ""),
                "name_norm": cand_name,
                "attachment_id": int(candidate.get("attachment_id") or 0),
                "attachment_file": str(candidate.get("attachment_file") or ""),
                "attachment_path": str(candidate.get("attachment_path") or ""),
            }
        )
    ranked.sort(key=lambda row: float(row.get("score") or 0.0), reverse=True)
    return ranked


def _resolve_attachment_source_candidate(
    product_name: str | None,
    candidates: list[dict[str, Any]],
) -> dict[str, Any] | None:
    ranked = _rank_attachment_source_candidates(product_name, candidates)
    if not ranked:
        return None
    best = ranked[0]
    best_score = float(best.get("score") or 0.0)
    if best_score < 0.84:
        return None

    # Guard ambiguous mapping across different attachments.
    if len(ranked) > 1:
        second = ranked[1]
        second_score = float(second.get("score") or 0.0)
        best_attachment_id = int(best.get("attachment_id") or 0)
        second_attachment_id = int(second.get("attachment_id") or 0)
        if (
            best_attachment_id > 0
            and second_attachment_id > 0
            and best_attachment_id != second_attachment_id
            and second_score >= 0.84
            and abs(best_score - second_score) <= 0.02
        ):
            return None
    return best


def _rank_qty_candidates(
    product_name: str | None,
    candidates: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    target = _normalize_product_name_for_qty_match(product_name)
    if not target or not candidates:
        return []

    target_tokens = set(target.split())
    ranked: list[dict[str, Any]] = []
    for candidate in candidates:
        cand_name = str(candidate.get("name_norm") or "").strip()
        qty = _to_positive_float(candidate.get("quantity"))
        if not cand_name or qty is None:
            continue
        method = "fuzzy"
        score = 0.0
        if cand_name == target:
            method = "exact"
            score = 1.0
        elif len(target) >= 10 and (target in cand_name or cand_name in target):
            method = "substring"
            seq = difflib.SequenceMatcher(None, target, cand_name).ratio()
            score = max(0.93, seq)
        else:
            seq = difflib.SequenceMatcher(None, target, cand_name).ratio()
            cand_tokens = set(cand_name.split())
            overlap = (len(target_tokens & cand_tokens) / max(1, len(target_tokens))) if target_tokens else 0.0
            score = (0.75 * seq) + (0.25 * overlap)
            if score < 0.70:
                continue
        ranked.append(
            {
                "score": round(float(score), 6),
                "method": method,
                "quantity": qty,
                "name": str(candidate.get("name") or ""),
                "name_norm": cand_name,
                "description": candidate.get("description"),
            }
        )
    ranked.sort(key=lambda row: float(row.get("score") or 0.0), reverse=True)
    return ranked


def _infer_qty_from_ranked_candidates(
    ranked: list[dict[str, Any]],
) -> tuple[float, float, str] | None:
    if not ranked:
        return None
    best = ranked[0]
    best_score = float(best.get("score") or 0.0)
    best_qty = _to_positive_float(best.get("quantity"))
    best_method = str(best.get("method") or "fuzzy")
    if best_qty is None or best_score < 0.90:
        return None

    # Avoid ambiguous auto-fill when top candidates are too close but imply
    # different quantities.
    if len(ranked) > 1:
        second = ranked[1]
        second_score = float(second.get("score") or 0.0)
        second_qty = _to_positive_float(second.get("quantity")) or 0.0
        if abs(best_score - second_score) <= 0.02 and abs(best_qty - second_qty) > 1e-6:
            return None

    if best_method == "exact":
        confidence = 0.99
    elif best_method == "substring":
        confidence = min(0.97, max(0.93, best_score))
    else:
        confidence = min(0.96, max(0.90, best_score))
    return best_qty, round(float(confidence), 4), best_method


def _has_complex_qty_pattern(text: str | None) -> bool:
    value = _normalize_product_name_for_qty_match(text)
    if not value:
        return False
    if _QTY_COMPLEX_RE.search(value):
        return True
    return bool(re.search(r"\d+(?:[.,]\d+)?", value))


def _should_use_qty_gpt_fallback(
    product_name: str | None,
    ranked: list[dict[str, Any]],
) -> bool:
    if not bool(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_ENABLED", True)):
        return False
    if not config.OPENAI_API_KEY:
        return False
    if not ranked:
        return True
    best = ranked[0]
    best_score = float(best.get("score") or 0.0)
    if best_score < 0.90:
        return True
    trigger_gap = max(
        0.01,
        min(
            0.20,
            float(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_TRIGGER_GAP", 0.03) or 0.03),
        ),
    )
    if len(ranked) > 1:
        second = ranked[1]
        second_score = float(second.get("score") or 0.0)
        best_qty = _to_positive_float(best.get("quantity")) or 0.0
        second_qty = _to_positive_float(second.get("quantity")) or 0.0
        if abs(best_score - second_score) <= trigger_gap and abs(best_qty - second_qty) > 1e-6:
            return True
    if _has_complex_qty_pattern(product_name) and best_score < 0.99:
        return True
    return False


def _campaign_qty_match_cache_key(
    *,
    product_name: str,
    candidates: list[dict[str, Any]],
) -> str:
    normalized_candidates = []
    for row in candidates:
        name = str(row.get("name") or "").strip()
        qty = _to_positive_float(row.get("quantity"))
        if not name or qty is None:
            continue
        normalized_candidates.append(
            {
                "name": name,
                "quantity": round(float(qty), 6),
            }
        )
    payload = {
        "product_name": " ".join(product_name.split()),
        "candidates": normalized_candidates,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _infer_qty_from_candidates_gpt(
    *,
    product_name: str | None,
    ranked: list[dict[str, Any]],
) -> tuple[float, float] | None:
    product_label = _clean_display_text(product_name) or str(product_name or "")
    product_label = " ".join(product_label.split())[:500]
    if not product_label:
        return None
    if not ranked:
        return None
    if not config.OPENAI_API_KEY:
        return None

    model = str(
        getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_MODEL", "")
        or getattr(config, "OPENAI_PRODUCT_MATCH_MODEL", "")
        or getattr(config, "OPENAI_ATTACHMENT_MODEL", "gpt-5.1")
        or "gpt-5.1"
    ).strip()
    timeout_s = max(
        15,
        int(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_TIMEOUT_SECONDS", 35) or 35),
    )
    top_k = max(3, int(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_TOP_K", 12) or 12))
    min_conf = max(
        0.0,
        min(
            1.0,
            float(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_MIN_CONFIDENCE", 0.60) or 0.60),
        ),
    )
    shortlist = ranked[:top_k]
    candidate_payload: list[dict[str, Any]] = []
    for idx, row in enumerate(shortlist, start=1):
        qty = _to_positive_float(row.get("quantity"))
        name = str(row.get("name") or "").strip()
        if not name or qty is None:
            continue
        candidate_payload.append(
            {
                "id": idx,
                "name": name,
                "quantity": round(float(qty), 6),
                "local_score": round(float(row.get("score") or 0.0), 4),
            }
        )
    if not candidate_payload:
        return None

    cache_key = _campaign_qty_match_cache_key(
        product_name=product_label,
        candidates=candidate_payload,
    )
    fallback_model = str(
        getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_FALLBACK_MODEL", "") or ""
    ).strip()
    model_candidates: list[str] = [model]
    if fallback_model and fallback_model != model:
        model_candidates.append(fallback_model)
    for model_name in model_candidates:
        cached = get_openai_payload_cache(
            cache_type="campaign_qty_match",
            model=model_name,
            cache_key=cache_key,
        )
        if not isinstance(cached, dict):
            continue
        try:
            best_id = int(cached.get("best_id") or 0)
            confidence = float(cached.get("confidence") or 0.0)
            if best_id <= 0 or confidence < min_conf:
                return None
            for row in candidate_payload:
                if int(row.get("id") or 0) == best_id:
                    qty = _to_positive_float(row.get("quantity"))
                    if qty is not None:
                        return qty, round(float(confidence), 4)
            return None
        except Exception:
            continue

    system_prompt = (
        "Tu fais du rattachement quantite produit. "
        "Choisis le seul candidat qui correspond exactement au produit demande. "
        "N'invente pas de quantite. Utilise strictement la quantite du candidat choisi. "
        "Si aucun candidat ne correspond, retourne best_id=0 avec confidence faible."
    )
    user_prompt = json.dumps(
        {
            "requested_product": product_label,
            "candidates": candidate_payload,
        },
        ensure_ascii=False,
    )
    payload = {
        "model": model,
        "temperature": 0,
        "seed": 42,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "response_format": {
            "type": "json_schema",
            "json_schema": {
                "name": "campaign_qty_match",
                "strict": True,
                "schema": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "best_id": {"type": "integer", "minimum": 0},
                        "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                        "reason": {"type": "string"},
                    },
                    "required": ["best_id", "confidence", "reason"],
                },
            },
        },
    }
    parsed, model_used = call_chat_completion_json_with_retry(
        operation="campaign_qty_match_gpt",
        payload=payload,
        primary_model=model,
        fallback_model=fallback_model,
        timeout_s=timeout_s,
        retry_attempts=max(
            1,
            int(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_RETRY_ATTEMPTS", 3) or 3),
        ),
        retry_backoff_s=max(
            0.2,
            float(
                getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_RETRY_BACKOFF_SECONDS", 1.0)
                or 1.0
            ),
        ),
        metadata={"candidate_count": len(candidate_payload)},
    )
    if not isinstance(parsed, dict):
        return None
    set_openai_payload_cache(
        cache_type="campaign_qty_match",
        model=(model_used or model),
        cache_key=cache_key,
        payload=parsed,
    )
    try:
        best_id = int(parsed.get("best_id") or 0)
        confidence = float(parsed.get("confidence") or 0.0)
    except Exception:
        return None
    if best_id <= 0 or confidence < min_conf:
        return None
    for row in candidate_payload:
        if int(row.get("id") or 0) != best_id:
            continue
        qty = _to_positive_float(row.get("quantity"))
        if qty is None:
            return None
        return qty, round(float(confidence), 4)
    return None


def _resolve_requested_qty_for_item(
    incoming_qty: Any,
    *,
    existing_qty: float | None,
    existing_source: str | None,
    existing_confidence: float | None,
    product_name: str | None,
    qty_candidates: list[dict[str, Any]],
    gpt_budget: dict[str, int] | None = None,
) -> tuple[float | None, str | None, float | None]:
    manual_qty = _to_positive_float(incoming_qty)
    if manual_qty is not None:
        return manual_qty, _QTY_SOURCE_MANUAL, 1.0

    source_norm = _normalize_qty_source(existing_source)
    if existing_qty is not None and source_norm in {_QTY_SOURCE_MANUAL, _QTY_SOURCE_CARRIED_FORWARD, None}:
        # Preserve user-entered and carried values by default.
        source = (existing_source or _QTY_SOURCE_CARRIED_FORWARD).strip()[:32] or _QTY_SOURCE_CARRIED_FORWARD
        confidence = _to_positive_float(existing_confidence)
        if confidence is None:
            confidence = 1.0 if source in {_QTY_SOURCE_MANUAL, _QTY_SOURCE_CARRIED_FORWARD} else 0.90
        return existing_qty, source, round(float(confidence), 4)

    ranked = _rank_qty_candidates(product_name, qty_candidates)
    inferred = _infer_qty_from_ranked_candidates(ranked)
    if existing_qty is not None and source_norm in {_QTY_SOURCE_ATTACHMENT_AUTO, _QTY_SOURCE_ATTACHMENT_AUTO_GPT}:
        # Revalidate auto-filled historical values against current extraction.
        if inferred is not None:
            inferred_qty, inferred_conf, _method = inferred
            if abs(float(inferred_qty) - float(existing_qty)) <= 1e-9:
                confidence = _to_positive_float(existing_confidence)
                if confidence is None:
                    confidence = inferred_conf
                return (
                    existing_qty,
                    source_norm,
                    round(float(max(confidence, inferred_conf)), 4),
                )
        # stale/unsupported auto value -> do not preserve
        existing_qty = None

    if inferred is not None:
        qty, confidence, _method = inferred
        return qty, _QTY_SOURCE_ATTACHMENT_AUTO, round(float(confidence), 4)

    if not _should_use_qty_gpt_fallback(product_name, ranked):
        return None, None, None

    gpt_max_calls = max(
        0,
        int(getattr(config, "OPENAI_CAMPAIGN_QTY_GPT_MAX_CALLS_PER_RUN", 80) or 80),
    )
    if gpt_budget is not None:
        used_calls = int(gpt_budget.get("calls") or 0)
        if used_calls >= gpt_max_calls:
            return None, None, None
        gpt_budget["calls"] = used_calls + 1

    inferred_gpt = _infer_qty_from_candidates_gpt(
        product_name=product_name,
        ranked=ranked,
    )
    if inferred_gpt is None:
        return None, None, None
    qty, confidence = inferred_gpt
    return qty, _QTY_SOURCE_ATTACHMENT_AUTO_GPT, round(float(confidence), 4)


def _refresh_qty_source_on_manual_edit(
    existing_source: str | None,
    old_qty: float | None,
    new_qty: float | None,
) -> tuple[str | None, float | None]:
    if new_qty is None:
        return None, None
    if old_qty is None:
        return _QTY_SOURCE_MANUAL, 1.0
    if abs(float(old_qty) - float(new_qty)) > 1e-9:
        return _QTY_SOURCE_MANUAL, 1.0
    if existing_source:
        return existing_source, 1.0
    return _QTY_SOURCE_MANUAL, 1.0


def _normalize_qty_source(value: str | None) -> str | None:
    raw = str(value or "").strip().lower()
    if raw in {
        _QTY_SOURCE_MANUAL,
        _QTY_SOURCE_CARRIED_FORWARD,
        _QTY_SOURCE_ATTACHMENT_AUTO,
        _QTY_SOURCE_ATTACHMENT_AUTO_GPT,
    }:
        return raw
    return None


def _autofill_campaign_item_quantities_from_attachments(
    db: Session,
    *,
    tenant_id: int,
    campaign_id: int,
) -> int:
    candidates = _collect_attachment_qty_candidates(
        db,
        tenant_id=tenant_id,
        campaign_id=campaign_id,
    )
    if not candidates:
        return 0
    gpt_budget: dict[str, int] = {"calls": 0}

    rows = (
        db.query(CampaignItem, Product)
        .join(Product, CampaignItem.product_id == Product.id)
        .filter(
            CampaignItem.campaign_id == campaign_id,
            CampaignItem.tenant_id == tenant_id,
            Product.tenant_id == tenant_id,
        )
        .all()
    )
    updated = 0
    for item_row, product in rows:
        old_qty = _to_positive_float(item_row.requested_qty)
        old_source = _normalize_qty_source(getattr(item_row, "requested_qty_source", None))
        old_conf = _to_positive_float(getattr(item_row, "requested_qty_confidence", None))
        qty, qty_source, qty_conf = _resolve_requested_qty_for_item(
            None,
            existing_qty=old_qty,
            existing_source=old_source,
            existing_confidence=old_conf,
            product_name=product.name,
            qty_candidates=candidates,
            gpt_budget=gpt_budget,
        )
        new_qty = _to_positive_float(qty)
        same_qty = (
            (old_qty is None and new_qty is None)
            or (
                old_qty is not None
                and new_qty is not None
                and abs(float(old_qty) - float(new_qty)) <= 1e-9
            )
        )
        same_source = (old_source or None) == (qty_source or None)
        old_conf_norm = _to_positive_float(old_conf)
        new_conf_norm = _to_positive_float(qty_conf)
        same_conf = (
            (old_conf_norm is None and new_conf_norm is None)
            or (
                old_conf_norm is not None
                and new_conf_norm is not None
                and abs(float(old_conf_norm) - float(new_conf_norm)) <= 1e-9
            )
        )
        if same_qty and same_source and same_conf:
            continue
        item_row.requested_qty = new_qty
        item_row.requested_qty_source = qty_source
        item_row.requested_qty_confidence = qty_conf
        updated += 1
    return updated


def _ensure_group_ids(
    db: Session,
    group_ids: list[int],
    tenant_id: int,
    owner_user_id: int | None = None,
) -> set[int]:
    if not group_ids:
        return set()
    query = db.query(Group.id).filter(Group.id.in_(group_ids), Group.tenant_id == tenant_id)
    rows = query.all()
    found = {row[0] for row in rows}
    missing = set(group_ids) - found
    if missing:
        raise HTTPException(status_code=404, detail="Group not found")
    return found


def _ensure_supplier_ids(
    db: Session,
    supplier_ids: list[int],
    tenant_id: int,
    owner_user_id: int | None = None,
) -> set[int]:
    if not supplier_ids:
        return set()
    query = db.query(Supplier.id).filter(
        Supplier.id.in_(supplier_ids),
        Supplier.tenant_id == tenant_id,
    )
    rows = query.all()
    found = {row[0] for row in rows}
    missing = set(supplier_ids) - found
    if missing:
        raise HTTPException(status_code=404, detail="Supplier not found")
    return found


def _ensure_product_ids(
    db: Session,
    product_ids: set[int],
    tenant_id: int,
    owner_user_id: int | None = None,
) -> None:
    if not product_ids:
        return
    query = db.query(Product.id).filter(
        Product.id.in_(product_ids),
        Product.tenant_id == tenant_id,
    )
    rows = query.all()
    found = {row[0] for row in rows}
    missing = product_ids - found
    if missing:
        raise HTTPException(status_code=404, detail="Product not found")


def _dedupe_supplier_ids_by_email(
    db: Session,
    supplier_ids: set[int],
    tenant_id: int,
) -> set[int]:
    if not supplier_ids:
        return set()
    rows = (
        db.query(Supplier.id, Supplier.email, Supplier.active)
        .filter(Supplier.id.in_(supplier_ids), Supplier.tenant_id == tenant_id)
        .order_by(Supplier.id.asc())
        .all()
    )
    selected_by_key: dict[str, tuple[int, bool]] = {}
    for supplier_id, email, active in rows:
        key = (email or "").strip().lower()
        if not key:
            key = f"id:{supplier_id}"
        if key not in selected_by_key:
            selected_by_key[key] = (int(supplier_id), bool(active))
            continue
        existing_id, existing_active = selected_by_key[key]
        if (not existing_active) and bool(active):
            selected_by_key[key] = (int(supplier_id), True)
        elif bool(active) == existing_active and int(supplier_id) < existing_id:
            selected_by_key[key] = (int(supplier_id), existing_active)
    return {supplier_id for supplier_id, _ in selected_by_key.values()}


def _offers_by_product(
    campaign_id: int,
    db: Session,
    current_user: User,
    limit: int | None = None,
    conform_only: bool = False,
    use_ai_prompt_for_conversion: bool | None = None,
) -> list[ProductOffers]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    campaign_processing_mode = _campaign_effective_processing_mode(db, campaign)
    if use_ai_prompt_for_conversion is None:
        app_settings = get_settings(db, tenant_id)
        use_ai_prompt_for_conversion = bool(
            getattr(app_settings, "use_ai_prompt_for_conversion", True)
        )
    thresholds = get_user_embedding_thresholds(db, tenant_id, current_user.id)
    embed_confirm_threshold = float(thresholds.confirm_threshold)
    embed_verify_threshold = float(thresholds.verify_threshold)
    char_match_threshold = float(thresholds.characteristics_threshold)
    items = (
        db.query(CampaignItem)
        .filter(CampaignItem.campaign_id == campaign_id, CampaignItem.tenant_id == tenant_id)
        .all()
    )
    campaign_item_by_product_id: dict[int, CampaignItem] = {
        item.product_id: item for item in items if item.product_id is not None
    }
    # Preserve campaign item order but avoid duplicate product cards in Results.
    product_ids: list[int] = []
    seen_product_ids: set[int] = set()
    for item in items:
        pid = item.product_id
        if pid in seen_product_ids:
            continue
        seen_product_ids.add(pid)
        product_ids.append(pid)
    if not product_ids:
        return []

    product_query = db.query(Product).filter(
        Product.id.in_(product_ids),
        Product.tenant_id == tenant_id,
    )
    products = product_query.all()
    products_by_id = {p.id: p for p in products}
    requested_payload_by_product_id: dict[int, dict[str, Any]] = {}
    requested_equivalence_text_by_product_id: dict[int, str] = {}
    for product in products:
        campaign_item = campaign_item_by_product_id.get(product.id)
        requested_description = None
        requested_characteristics = None
        if campaign_item is not None:
            requested_description = campaign_item.requested_description
            requested_characteristics = campaign_item.requested_characteristics
        if requested_characteristics is None:
            requested_characteristics = product.characteristics
        requested_payload_by_product_id[product.id] = _flatten_requested_components(
            requested_description=requested_description,
            requested_characteristics=requested_characteristics,
        )
        requested_equivalence_text_by_product_id[product.id] = _build_price_eq_requested_text(
            product_name=product.name,
            requested_description=requested_description,
            requested_payload=requested_payload_by_product_id.get(product.id),
        )
    attachment_source_candidates = _collect_attachment_source_candidates(
        db,
        tenant_id=tenant_id,
        campaign_id=campaign_id,
    )
    attachment_source_cache: dict[str, dict[str, Any] | None] = {}

    def _resolve_requested_source_attachment(product_name: str | None) -> dict[str, Any] | None:
        key = _normalize_product_name_for_qty_match(product_name)
        if not key:
            return None
        if key not in attachment_source_cache:
            attachment_source_cache[key] = _resolve_attachment_source_candidate(
                product_name,
                attachment_source_candidates,
            )
        return attachment_source_cache.get(key)

    campaign_supplier_ids = {
        row[0]
        for row in db.query(CampaignRecipient.supplier_id)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
        )
        .all()
    }
    supplier_query = db.query(Supplier).filter(
        Supplier.tenant_id == tenant_id,
        Supplier.active.is_(True),
    )
    if campaign_supplier_ids:
        supplier_query = supplier_query.filter(Supplier.id.in_(campaign_supplier_ids))
    suppliers = supplier_query.order_by(Supplier.id.asc()).all()

    suppliers_by_email: dict[str, Supplier] = {}
    suppliers_by_name: dict[str, Supplier] = {}
    for supplier in suppliers:
        if supplier.email:
            email_key = supplier.email.strip().lower()
            # deterministic first-win to avoid random alias when duplicate emails exist
            suppliers_by_email.setdefault(email_key, supplier)
        if supplier.name:
            name_key = (_clean_display_text(supplier.name) or "").lower()
            if name_key:
                suppliers_by_name.setdefault(name_key, supplier)

    def _norm(text: str | None) -> str:
        return " ".join((text or "").strip().lower().split())

    def _hash_norm(text: str) -> str:
        return hashlib.sha256(text.encode("utf-8")).hexdigest()

    def _extract_emb_score_from_notes(notes: str | None) -> float | None:
        if not notes:
            return None
        gm = re.search(r"\bgpt_score=([0-9]+(?:\.[0-9]+)?)", notes)
        if gm:
            try:
                return round(float(gm.group(1)), 4)
            except Exception:
                pass
        m = re.search(r"\bemb_score=([0-9]+(?:\.[0-9]+)?)", notes)
        if not m:
            return None
        try:
            return round(float(m.group(1)), 4)
        except Exception:
            return None

    def _extract_gpt_score_from_notes(notes: str | None) -> float | None:
        if not notes:
            return None
        gm = re.search(r"\bgpt_score=([0-9]+(?:\.[0-9]+)?)", notes)
        if not gm:
            return None
        try:
            return round(float(gm.group(1)), 4)
        except Exception:
            return None

    def _extract_gpt_match_from_notes(notes: str | None) -> str | None:
        if not notes:
            return None
        gm = re.search(r"\bgpt_match=([^\n]+?)(?:\s+gpt_score=|$)", notes)
        if not gm:
            return None
        value = str(gm.group(1) or "").strip()
        return value or None

    def _extract_gpt_matches_from_notes(notes: str | None) -> dict[str, float]:
        out: dict[str, float] = {}
        if not notes:
            return out
        m = re.search(r"\bgpt_matches_b64=([A-Za-z0-9_\-]+=*)", notes)
        if not m:
            return out
        raw_b64 = m.group(1).strip()
        if not raw_b64:
            return out
        try:
            padded = raw_b64 + ("=" * ((4 - len(raw_b64) % 4) % 4))
            raw = base64.urlsafe_b64decode(padded.encode("ascii")).decode("utf-8", errors="ignore")
            parsed = json.loads(raw)
        except Exception:
            return out
        if not isinstance(parsed, list):
            return out
        for row in parsed:
            if not isinstance(row, dict):
                continue
            name = str(row.get("name") or "").strip()
            if not name:
                continue
            try:
                score = float(row.get("score"))
            except Exception:
                continue
            score = max(0.0, min(1.0, score))
            norm_name = _norm(name)
            if not norm_name:
                continue
            prev = out.get(norm_name)
            if prev is None or score > prev:
                out[norm_name] = score
        return out

    def _extract_price_source_from_notes(notes: str | None) -> str | None:
        if not notes:
            return None
        m = re.search(r"\bprice_source=([a-z_]+)\b", notes)
        if not m:
            return None
        return m.group(1)

    _NAME_STOPWORDS = {
        "de",
        "du",
        "des",
        "ou",
        "plus",
        "pour",
        "avec",
        "une",
        "un",
        "et",
        "la",
        "le",
        "les",
        "a",
        "au",
        "aux",
        "en",
        "d",
    }
    _SECONDARY_MATCH_GENERIC_TOKENS = {
        "bande",
        "piece",
        "rouleau",
        "rouleaux",
        "boite",
        "bt",
        "box",
        "paquet",
        "pack",
        "sachet",
        "unite",
        "compresse",
        "compresses",
        "sterile",
        "steriles",
    }
    _TOKEN_SYNONYMS = {
        # FR/EN medical naming variants
        "sterilisateur": "autoclave",
        "sterilizer": "autoclave",
        "sterilisation": "sterilisation",
        "stérilisateur": "autoclave",
        "autoclave": "autoclave",
        "benchtop": "paillasse",
        "tabletop": "paillasse",
        "bench": "paillasse",
        "paillasse": "paillasse",
        # accessory variants
        "trolley": "chariot",
        "cart": "chariot",
        "charriot": "chariot",
        "stand": "support",
        "holder": "support",
        "rack": "panier",
        "tray": "panier",
    }
    _UNIT_ALIASES = {
        "l": "l",
        "litre": "l",
        "litres": "l",
        "liter": "l",
        "liters": "l",
        "ml": "ml",
        "cl": "cl",
        "g": "g",
        "kg": "kg",
        "mg": "mg",
        "m": "m",
        "cm": "cm",
        "mm": "mm",
    }
    _UNIT_TO_BASE = {
        "ml": ("l", 0.001),
        "cl": ("l", 0.01),
        "l": ("l", 1.0),
        "mg": ("g", 0.001),
        "g": ("g", 1.0),
        "kg": ("g", 1000.0),
        "mm": ("m", 0.001),
        "cm": ("m", 0.01),
        "m": ("m", 1.0),
    }
    _ACCESSORY_TOKENS = {
        "chariot",
        "charriot",
        "trolley",
        "sutrolley",
        "support",
        "panier",
        "socle",
        "table",
        "tiroir",
        "roulette",
    }
    _STRICT_COMPLEMENT_MARKERS = ("pour", "avec", "sans", "compatible", "compatibles")
    _STRICT_COMPLEMENT_GENERIC_TOKENS = {
        "accessoire",
        "accessoires",
        "kit",
        "set",
        "modele",
        "model",
        "usage",
        "unique",
        "systeme",
        "dispositif",
    }
    _STRICT_COMPLEMENT_MISSING_PENALTY = 0.12

    def _canonical_token(token: str) -> str:
        t = token.strip().lower()
        if not t:
            return t
        t = _TOKEN_SYNONYMS.get(t, t)
        if t.endswith("s") and len(t) > 4:
            singular = t[:-1]
            t = _TOKEN_SYNONYMS.get(singular, singular)
        return t

    def _extract_unit_tokens(text: str) -> list[str]:
        out: list[str] = []
        if not text:
            return out
        # Accept "1000ml", "1000 ml", "1,5 litre", etc.
        for m in re.finditer(r"(\d+(?:[.,]\d+)?)\s*([a-z]+)\b", text):
            raw_num = m.group(1).replace(",", ".")
            raw_unit = m.group(2).lower()
            unit = _UNIT_ALIASES.get(raw_unit)
            if unit is None:
                continue
            try:
                num = float(raw_num)
            except Exception:
                continue
            base = _UNIT_TO_BASE.get(unit)
            if base is None:
                continue
            base_unit, factor = base
            val = num * factor
            # Canonical numeric-unit token for overlap matching.
            out.append(f"u_{val:.6g}{base_unit}")
        return out

    def _name_tokens(text: str | None) -> list[str]:
        value_raw = _norm(text)
        if not value_raw:
            return []
        # Remove accents so FR variants remain tokenizable (e.g. stérilisateur).
        value = unicodedata.normalize("NFKD", value_raw).encode("ascii", "ignore").decode("ascii")
        unit_tokens = _extract_unit_tokens(value)
        raw_toks = [t for t in re.split(r"[^a-z0-9]+", value) if t]
        norm_toks: list[str] = []
        for tok in raw_toks:
            if tok in _NAME_STOPWORDS:
                continue
            canon = _canonical_token(tok)
            if not canon or canon in _NAME_STOPWORDS:
                continue
            norm_toks.append(canon)
        return norm_toks + unit_tokens

    def _secondary_match_specific_tokens(text: str | None) -> set[str]:
        out: set[str] = set()
        for tok in _name_tokens(text):
            if not tok or tok.startswith("u_"):
                continue
            if tok in _SECONDARY_MATCH_GENERIC_TOKENS:
                continue
            if tok.isdigit():
                continue
            if len(tok) <= 2:
                continue
            out.add(tok)
        return out

    def _normalized_ascii_text(text: str | None) -> str:
        value_raw = _norm(text)
        if not value_raw:
            return ""
        return unicodedata.normalize("NFKD", value_raw).encode("ascii", "ignore").decode("ascii")

    def _extract_structured_complements(text: str | None) -> list[tuple[str, set[str]]]:
        normalized = _normalized_ascii_text(text)
        if not normalized:
            return []
        pattern = re.compile(
            r"\b(pour|avec|sans|compatible|compatibles)\b\s+(.+?)(?=(?:\b(?:pour|avec|sans|compatible|compatibles)\b)|$)"
        )
        complements: list[tuple[str, set[str]]] = []
        for match in pattern.finditer(normalized):
            marker = str(match.group(1) or "").strip().lower()
            raw_tail = str(match.group(2) or "").strip(" ,;:-()[]{}")
            if not raw_tail:
                continue
            tokens = {
                tok
                for tok in _name_tokens(raw_tail)
                if tok
                and not tok.startswith("u_")
                and tok not in _STRICT_COMPLEMENT_GENERIC_TOKENS
                and len(tok) > 2
            }
            if tokens:
                complements.append((marker, tokens))
        return complements

    def _strict_missing_complement_penalty(
        offered_name: str | None,
        requested_name: str | None,
    ) -> tuple[float, list[str]]:
        offered_complements = _extract_structured_complements(offered_name)
        if not offered_complements:
            return 0.0, []
        requested_tokens = {
            tok
            for tok in _name_tokens(requested_name)
            if tok and not tok.startswith("u_")
        }
        requested_complements = _extract_structured_complements(requested_name)
        missing_labels: list[str] = []
        for marker, offered_tokens in offered_complements:
            if not offered_tokens:
                continue
            if offered_tokens & requested_tokens:
                continue
            if any(offered_tokens & req_tokens for _req_marker, req_tokens in requested_complements):
                continue
            label = f"{marker} {' '.join(sorted(offered_tokens))}".strip()
            missing_labels.append(label)
        if not missing_labels:
            return 0.0, []
        return _STRICT_COMPLEMENT_MISSING_PENALTY, missing_labels

    def _semantic_overlap_ratio(
        required_tokens: list[str],
        offered_tokens: list[str],
        token_vectors: dict[str, list[float]],
        *,
        token_threshold: float = 0.65,
    ) -> float:
        req = [t for t in dict.fromkeys(required_tokens) if t]
        off = [t for t in dict.fromkeys(offered_tokens) if t]
        if not req:
            return 0.0
        off_set = set(off)
        matched = 0
        for rt in req:
            # canonical unit token (e.g. u_1l) should match exactly.
            if rt.startswith("u_"):
                if rt in off_set:
                    matched += 1
                continue
            rv = token_vectors.get(rt)
            if rv is None:
                if rt in off_set:
                    matched += 1
                continue
            best = 0.0
            for ot in off:
                if ot.startswith("u_"):
                    continue
                ov = token_vectors.get(ot)
                if ov is None:
                    if ot == rt:
                        best = 1.0
                        break
                    continue
                score = cosine_similarity(rv, ov)
                if score > best:
                    best = score
            if best >= token_threshold:
                matched += 1
        return matched / len(req)

    use_gpt_multi_match = str(
        getattr(config, "OPENAI_PRODUCT_MATCH_MODE", "embedding") or "embedding"
    ).strip().lower() in {"gpt", "gpt_only", "gpt_hybrid", "hybrid_gpt"}
    strict_requested_available = (
        campaign_processing_mode == _CAMPAIGN_PROCESSING_STRICT
        and (
            db.query(CampaignAttachment.id)
            .filter(
                CampaignAttachment.campaign_id == campaign_id,
                CampaignAttachment.tenant_id == tenant_id,
                CampaignAttachment.strict_extraction_preview.isnot(None),
            )
            .limit(1)
            .first()
            is not None
        )
    )
    use_strict_requested_match = bool(strict_requested_available)

    pdf_statuses = [ExtractionStatus.CONFIRMED]
    if not conform_only:
        pdf_statuses.append(ExtractionStatus.TO_VERIFY)
        # In GPT/strict modes, Results can display offers based on GPT confidence
        # even when line status is still PENDING (< verify threshold).
        if use_gpt_multi_match or use_strict_requested_match:
            pdf_statuses.append(ExtractionStatus.PENDING)

    # Pull PDF lines once, then reuse for all requested products.
    pdf_items = (
        db.query(InvoiceExtractionItem, InboundInvoice)
        .join(InboundInvoice, InvoiceExtractionItem.invoice_id == InboundInvoice.id)
        .filter(
            InvoiceExtractionItem.campaign_id == campaign_id,
            InvoiceExtractionItem.tenant_id == tenant_id,
            InvoiceExtractionItem.status.in_(pdf_statuses),
            InboundInvoice.tenant_id == tenant_id,
        )
        .all()
    )
    # Multi-match index: one extracted line can satisfy several requested products.
    product_norms = [(p.id, p.name or "", _norm(p.name)) for p in products if (p.name or "").strip()]
    requested_qty_by_product_norm: dict[str, float] = {}
    for pid, _pname, pnorm in product_norms:
        campaign_item = campaign_item_by_product_id.get(pid)
        requested_qty = _to_positive_float(getattr(campaign_item, "requested_qty", None))
        if pnorm and requested_qty is not None:
            requested_qty_by_product_norm[pnorm] = requested_qty
    product_tokens_by_norm: dict[str, list[str]] = {
        norm_name: _name_tokens(raw_name) for _pid, raw_name, norm_name in product_norms
    }
    requested_rows = _collect_campaign_requested_characteristics(db, tenant_id, campaign_id)
    requested_by_product_norm: dict[str, dict[str, list[str]]] = {}
    for _pid, pname, pnorm in product_norms:
        requested_by_product_norm[pnorm] = _select_requested_characteristics_for_product(
            pname, requested_rows
        )
    strict_requested_index_by_campaign: dict[int, list[dict[str, Any]]] = {}
    strict_requested_candidates: list[dict[str, Any]] = []
    strict_product_norm_by_id: dict[int, str] = {}
    if use_strict_requested_match:
        strict_requested_index_by_campaign = _build_campaign_requested_index(
            db=db,
            tenant_id=tenant_id,
            campaign_ids=[campaign_id],
        )
        strict_requested_candidates = strict_requested_index_by_campaign.get(campaign_id, [])
        for row in strict_requested_candidates:
            try:
                product_id = int(row.get("product_id"))
            except Exception:
                continue
            product_obj = products_by_id.get(product_id)
            product_name = (
                product_obj.name
                if product_obj is not None
                else str(row.get("product_name") or "").strip()
            )
            product_norm = _norm(product_name)
            if product_norm:
                strict_product_norm_by_id[product_id] = product_norm
        use_strict_requested_match = bool(strict_requested_candidates)

    pdf_item_ids = [item_row.id for item_row, _invoice in pdf_items]
    tech_link_by_item_id: dict[int, tuple[TechSheetLink, TechSheet]] = {}
    pair_cache_by_key: dict[tuple[int, str], InvoiceItemCharacteristicMatch] = {}
    if pdf_item_ids:
        links = (
            db.query(TechSheetLink, TechSheet)
            .join(TechSheet, TechSheetLink.tech_sheet_id == TechSheet.id)
            .filter(
                TechSheetLink.tenant_id == tenant_id,
                TechSheetLink.campaign_id == campaign_id,
                TechSheetLink.invoice_item_id.in_(pdf_item_ids),
                TechSheet.tenant_id == tenant_id,
            )
            .all()
        )
        for link, sheet in links:
            tech_link_by_item_id[link.invoice_item_id] = (link, sheet)
        pair_rows = (
            db.query(InvoiceItemCharacteristicMatch)
            .filter(
                InvoiceItemCharacteristicMatch.tenant_id == tenant_id,
                InvoiceItemCharacteristicMatch.campaign_id == campaign_id,
                InvoiceItemCharacteristicMatch.invoice_item_id.in_(pdf_item_ids),
            )
            .all()
        )
        for row in pair_rows:
            pair_cache_by_key[(row.invoice_item_id, _norm(row.requested_product_norm))] = row
    offer_analysis_by_key: dict[tuple[int, str], OfferAnalysisResult] = {}
    if pdf_item_ids and product_norms:
        requested_norms = [norm_name for _pid, _pname, norm_name in product_norms if norm_name]
        analysis_rows = (
            db.query(OfferAnalysisResult)
            .filter(
                OfferAnalysisResult.tenant_id == tenant_id,
                OfferAnalysisResult.campaign_id == campaign_id,
                OfferAnalysisResult.invoice_item_id.in_(pdf_item_ids),
                OfferAnalysisResult.requested_product_norm.in_(requested_norms),
            )
            .all()
        )
        for row in analysis_rows:
            offer_analysis_by_key[(int(row.invoice_item_id), _norm(row.requested_product_norm))] = row

    pdf_multi_match_by_item_id: dict[int, set[str]] = {}
    pair_embed_scores: dict[tuple[int, str], float] = {}
    pair_fuzzy_scores: dict[tuple[int, str], float] = {}
    strict_match_diag_by_item_product: dict[tuple[int, str], dict[str, Any]] = {}
    offered_chars_by_item_id: dict[int, dict[str, list[str]]] = {}
    pair_char_points: dict[tuple[int, str], tuple[int | None, int | None]] = {}
    item_row_by_id: dict[int, InvoiceExtractionItem] = {row.id: row for row, _inv in pdf_items}

    def _line_base_for_matching(it: InvoiceExtractionItem) -> str:
        # Always prioritize corrected target name when available so Results
        # stays aligned with user corrections/rematch decisions.
        return it.corrected_product_name or it.product_name or ""

    def _pair_points_for(item_row: InvoiceExtractionItem, product_norm: str) -> tuple[int | None, int | None]:
        pair_key = (item_row.id, product_norm)
        if pair_key in pair_char_points:
            return pair_char_points[pair_key]
        requested = requested_by_product_norm.get(product_norm) or {}
        cached_pair = pair_cache_by_key.get(pair_key)
        if cached_pair is not None:
            matched_points: int | None = None
            total_points: int | None = None
            try:
                matched_points = (
                    int(cached_pair.matched_points)
                    if cached_pair.matched_points is not None
                    else None
                )
            except Exception:
                matched_points = None
            try:
                total_points = (
                    int(cached_pair.total_points)
                    if cached_pair.total_points is not None
                    else None
                )
            except Exception:
                total_points = None
            # Ignore stale cached rows (legacy 0/0) when requested characteristics now exist.
            cached_stale = bool(requested) and (total_points is None or total_points <= 0)
            if not cached_stale:
                pair_char_points[pair_key] = (matched_points, total_points)
                return pair_char_points[pair_key]

        if not requested:
            pair_char_points[pair_key] = (None, None)
            return pair_char_points[pair_key]

        # Reuse persisted similarity when this line is already matched to the same requested product.
        line_matched_norm = _norm(_line_base_for_matching(item_row))
        if line_matched_norm == product_norm and isinstance(item_row.similarity_points, dict):
            mp = item_row.similarity_points.get("matched_points")
            tp = item_row.similarity_points.get("total_points")
            matched_points: int | None = None
            total_points: int | None = None
            try:
                matched_points = int(mp) if mp is not None else None
            except Exception:
                matched_points = None
            try:
                total_points = int(tp) if tp is not None else None
            except Exception:
                total_points = None
            pair_char_points[pair_key] = (matched_points, total_points)
            return pair_char_points[pair_key]

        offered = offered_chars_by_item_id.get(item_row.id)
        if offered is None:
            offered, _source = _extract_offered_characteristics(db, item_row, tenant_id, campaign_id)
            offered_chars_by_item_id[item_row.id] = offered

        def _value_match(req_val: str, off_val: str) -> bool:
            req_norm = _norm(req_val)
            off_norm = _norm(off_val)
            if not req_norm or not off_norm:
                return False
            if req_norm == off_norm:
                return True
            if req_norm in off_norm or off_norm in req_norm:
                return True
            return difflib.SequenceMatcher(None, req_norm, off_norm).ratio() >= char_match_threshold

        total_points = 0
        matched_points = 0
        for req_key, req_vals in requested.items():
            req_values = [str(v) for v in (req_vals or []) if str(v).strip()]
            if not req_values:
                continue
            total_points += 1
            off_values = [str(v) for v in (offered.get(req_key, []) or []) if str(v).strip()]
            key_matched = False
            for rv in req_values:
                for ov in off_values:
                    if _value_match(rv, ov):
                        key_matched = True
                        break
                if key_matched:
                    break
            if key_matched:
                matched_points += 1

        if total_points == 0:
            matched_points = None
            total_points = None
        pair_char_points[pair_key] = (matched_points, total_points)
        return pair_char_points[pair_key]

    known_product_norms = {norm_name for _pid, _name, norm_name in product_norms}
    if use_strict_requested_match and pdf_items:
        xlsx_batch_results: dict[int, dict[str, Any]] = {}
        xlsx_items = [
            item_row
            for item_row, _invoice in pdf_items
            if str(item_row.source_engine or "").strip().lower() == "xlsx_parser"
        ]
        if xlsx_items:
            xlsx_batch_results = _pick_requested_references_batch(
                items=xlsx_items,
                matched_names_by_item_id={
                    item_row.id: (item_row.corrected_product_name or "").strip() or None
                    for item_row in xlsx_items
                },
                candidates=strict_requested_candidates,
            )
        for item_row, _invoice in pdf_items:
            pdf_multi_match_by_item_id[item_row.id] = set()
            diagnostics: dict[str, Any] = {}
            source_engine = str(item_row.source_engine or "").strip().lower()
            matched_row = None
            if source_engine == "xlsx_parser":
                batch_result = xlsx_batch_results.get(item_row.id)
                if isinstance(batch_result, dict):
                    matched_raw = batch_result.get("matched_row")
                    if isinstance(matched_raw, dict):
                        matched_row = dict(matched_raw)
                    diag_raw = batch_result.get("diagnostics")
                    if isinstance(diag_raw, dict):
                        diagnostics.update(diag_raw)
            if matched_row is None and source_engine != "xlsx_parser":
                matched_row = _pick_requested_reference(
                    item=item_row,
                    matched_name=(item_row.corrected_product_name or "").strip() or None,
                    candidates=strict_requested_candidates,
                    diagnostics=diagnostics,
                )
            if not isinstance(matched_row, dict):
                continue

            product_norm = None
            try:
                matched_product_id = int(matched_row.get("product_id"))
                product_norm = strict_product_norm_by_id.get(matched_product_id)
            except Exception:
                product_norm = None
            if not product_norm:
                product_norm = _norm(str(matched_row.get("product_name") or "").strip())
            if not product_norm or product_norm not in known_product_norms:
                continue

            conf = diagnostics.get("confidence")
            try:
                conf_value = round(float(conf), 4) if conf is not None else None
            except Exception:
                conf_value = None
            complement_penalty, missing_complements = _strict_missing_complement_penalty(
                item_row.corrected_product_name or item_row.product_name or "",
                matched_row.get("product_name") or "",
            )
            if conf_value is not None and complement_penalty > 0:
                conf_value = round(max(0.0, float(conf_value) - float(complement_penalty)), 4)
                diagnostics["confidence"] = conf_value
                diagnostics["penalty_applied"] = True
                diagnostics["penalty_value"] = round(float(complement_penalty), 4)
                diagnostics["penalty_labels"] = list(missing_complements)
                reason_text = str(diagnostics.get("reason") or "").strip()
                complement_note = (
                    "Penalite complement structurant absent cote demande: "
                    + ", ".join(missing_complements)
                    + f" (-{complement_penalty:.2f})."
                )
                diagnostics["reason"] = (
                    f"{reason_text} {complement_note}".strip() if reason_text else complement_note
                )
            if conf_value is not None:
                pair_embed_scores[(item_row.id, product_norm)] = conf_value
            eligible_product_norms: set[str] = {product_norm}
            target_base = _strip_name_noise_for_match(
                item_row.corrected_product_name or item_row.product_name or ""
            )
            target_tokens = set(_name_tokens(target_base))
            target_specific_tokens = _secondary_match_specific_tokens(target_base)
            strict_threshold = None
            try:
                strict_threshold = float(diagnostics.get("threshold"))
            except Exception:
                strict_threshold = None
            if strict_threshold is None:
                strict_threshold = 0.70
            if target_base:
                for candidate in strict_requested_candidates:
                    try:
                        candidate_product_id = int(candidate.get("product_id"))
                    except Exception:
                        candidate_product_id = 0
                    candidate_norm = strict_product_norm_by_id.get(candidate_product_id)
                    if not candidate_norm or candidate_norm == product_norm:
                        continue
                    candidate_name = str(candidate.get("product_name") or "").strip()
                    candidate_base = _strip_name_noise_for_match(candidate_name)
                    if not candidate_base:
                        continue
                    candidate_tokens = set(_name_tokens(candidate_base))
                    candidate_specific_tokens = _secondary_match_specific_tokens(candidate_base)
                    if not (target_specific_tokens & candidate_specific_tokens):
                        continue
                    overlap = (
                        len(target_tokens & candidate_tokens) / max(len(target_tokens | candidate_tokens), 1)
                        if target_tokens or candidate_tokens
                        else 0.0
                    )
                    similarity_score = max(
                        difflib.SequenceMatcher(None, target_base, candidate_base).ratio(),
                        overlap,
                    )
                    if (
                        candidate_base == target_base
                        or candidate_base in target_base
                        or target_base in candidate_base
                    ):
                        eligible_product_norms.add(candidate_norm)
                        similarity_score = max(similarity_score, 0.97)
                    if similarity_score >= strict_threshold:
                        eligible_product_norms.add(candidate_norm)
                        pair_embed_scores[(item_row.id, candidate_norm)] = round(
                            float(similarity_score), 4
                        )
                        strict_match_diag_by_item_product.setdefault(
                            (item_row.id, candidate_norm),
                            {
                                "score": round(float(similarity_score), 4),
                                "threshold": strict_threshold,
                                "decision": "matched",
                                "reason": "display_similarity_threshold",
                                "matched_name": candidate_name or None,
                                "matched_row": dict(candidate),
                            },
                        )
            pdf_multi_match_by_item_id[item_row.id] = eligible_product_norms
            strict_match_diag_by_item_product[(item_row.id, product_norm)] = {
                "score": conf_value,
                "threshold": diagnostics.get("threshold"),
                "decision": diagnostics.get("decision"),
                "reason": diagnostics.get("reason"),
                "penalty_applied": diagnostics.get("penalty_applied"),
                "penalty_value": diagnostics.get("penalty_value"),
                "penalty_labels": diagnostics.get("penalty_labels"),
                "matched_name": str(matched_row.get("product_name") or "").strip() or None,
                "matched_row": dict(matched_row),
            }

    if product_norms and pdf_items and not use_strict_requested_match:
        fuzzy_threshold = 0.80
        fuzzy_priority_threshold = 0.85

        # First pass: exact + fuzzy.
        item_texts: list[str] = []
        item_ids: list[int] = []
        item_text_by_id: dict[int, str] = {}
        item_tokens_by_id: dict[int, list[str]] = {}
        item_required_embed_threshold: dict[int, float] = {}
        for item_row, _invoice in pdf_items:
            extracted_text = _norm(f"{item_row.product_name or ''} {item_row.description or ''}")
            if not extracted_text:
                continue
            matched = pdf_multi_match_by_item_id.setdefault(item_row.id, set())
            corrected_norm = _norm(_line_base_for_matching(item_row))
            for _pid, _pname, product_norm in product_norms:
                if not product_norm:
                    continue
                if corrected_norm and corrected_norm == product_norm:
                    matched.add(product_norm)
                    pair_fuzzy_scores[(item_row.id, product_norm)] = 1.0
                    continue
                fuzzy_score = difflib.SequenceMatcher(None, extracted_text, product_norm).ratio()
                pair_fuzzy_scores[(item_row.id, product_norm)] = round(float(fuzzy_score), 4)
                if fuzzy_score >= fuzzy_threshold:
                    matched.add(product_norm)

            item_texts.append(extracted_text)
            item_ids.append(item_row.id)
            item_text_by_id[item_row.id] = extracted_text
            item_tokens_by_id[item_row.id] = _name_tokens(extracted_text)
            if item_row.status == ExtractionStatus.TO_VERIFY:
                item_required_embed_threshold[item_row.id] = embed_verify_threshold
            else:
                item_required_embed_threshold[item_row.id] = embed_confirm_threshold

        # Second pass: embeddings for all requested products, not only the best match.
        if config.OPENAI_API_KEY and item_texts:
            product_texts = [norm for _, _, norm in product_norms]
            token_texts = sorted(
                {
                    tok
                    for toks in list(product_tokens_by_norm.values()) + list(item_tokens_by_id.values())
                    for tok in toks
                    if tok and not tok.startswith("u_")
                }
            )
            token_text_set = set(token_texts)
            combined = product_texts + item_texts + token_texts
            cached, missing_texts, _ = get_cached_embeddings(combined)
            if missing_texts:
                enqueue_embeddings(missing_texts)
                token_missing = [t for t in missing_texts if t in token_text_set]
                if token_missing:
                    try:
                        fetched = fetch_embeddings(token_missing)
                        for t, vec in zip(token_missing, fetched):
                            cached[_hash_norm(t)] = vec
                    except Exception:
                        pass

            product_embeddings: dict[str, list[float]] = {}
            for product_norm in product_texts:
                emb = cached.get(_hash_norm(product_norm))
                if emb is not None:
                    product_embeddings[product_norm] = emb

            item_embeddings: dict[int, list[float]] = {}
            for item_id, item_text in zip(item_ids, item_texts):
                emb = cached.get(_hash_norm(item_text))
                if emb is not None:
                    item_embeddings[item_id] = emb

            token_embeddings: dict[str, list[float]] = {}
            for tok in token_texts:
                emb = cached.get(_hash_norm(tok))
                if emb is not None:
                    token_embeddings[tok] = emb

            if product_embeddings and item_embeddings:
                for item_id, item_emb in item_embeddings.items():
                    matched = pdf_multi_match_by_item_id.setdefault(item_id, set())
                    required_threshold = item_required_embed_threshold.get(item_id, embed_confirm_threshold)
                    item_tokens = item_tokens_by_id.get(item_id, _name_tokens(item_text_by_id.get(item_id, "")))
                    item_has_accessory = bool(set(item_tokens).intersection(_ACCESSORY_TOKENS))
                    score_by_product: dict[str, float] = {}
                    for product_norm, product_emb in product_embeddings.items():
                        emb_score = cosine_similarity(item_emb, product_emb)
                        score_by_product[product_norm] = emb_score
                    if not score_by_product:
                        continue
                    for product_norm, emb_score in score_by_product.items():
                        fuzzy_pair = pair_fuzzy_scores.get((item_id, product_norm), 0.0)
                        if fuzzy_pair >= fuzzy_priority_threshold:
                            emb_adjusted = float(fuzzy_pair)
                            pair_embed_scores[(item_id, product_norm)] = round(emb_adjusted, 4)
                            if emb_adjusted >= required_threshold:
                                matched.add(product_norm)
                            continue
                        req_tokens = product_tokens_by_norm.get(product_norm, [])
                        # Historical behavior for result ranking:
                        # keep embedding score simple and only add a small
                        # characteristics bonus. Do not apply overlap/accessory penalties.
                        matched_points = 0
                        row_obj = item_row_by_id.get(item_id)
                        if row_obj is not None:
                            mp, _tp = _pair_points_for(row_obj, product_norm)
                            if isinstance(mp, int) and mp > 0:
                                matched_points = mp

                        emb_adjusted = float(emb_score)
                        emb_adjusted += 0.015 * matched_points
                        emb_adjusted = max(0.0, min(1.0, emb_adjusted))

                        pair_embed_scores[(item_id, product_norm)] = round(emb_adjusted, 4)
                        if emb_adjusted >= required_threshold:
                            matched.add(product_norm)

    gpt_min_conf = max(
        _STANDARD_MIN_GPT_MATCH if not use_strict_requested_match else 0.0,
        min(
            1.0,
            float(getattr(config, "OPENAI_PRODUCT_MATCH_MIN_CONFIDENCE", 0.30) or 0.30),
        ),
    )
    if use_gpt_multi_match and pdf_items and not use_strict_requested_match:
        for item_row, _invoice in pdf_items:
            # In GPT mode, reset any previous embedding-derived candidates for
            # this line. Results should be driven only by GPT confidences.
            pdf_multi_match_by_item_id[item_row.id] = set()
            gpt_scores = _extract_gpt_matches_from_notes(item_row.notes)
            if not gpt_scores:
                continue
            offered_qty = _to_positive_float(item_row.corrected_quantity)
            if offered_qty is None:
                offered_qty = _to_positive_float(item_row.quantity)
            filtered_scores = dict(gpt_scores)
            # Quantity-coherence guard: when one candidate has a very strong
            # quantity agreement, avoid leaking the same invoice line to
            # unrelated requested products.
            if offered_qty is not None and offered_qty > 0:
                known_qty_rows: list[tuple[str, float, float, float]] = []
                for pnorm, score in gpt_scores.items():
                    req_qty = requested_qty_by_product_norm.get(pnorm)
                    if req_qty is None or req_qty <= 0:
                        continue
                    ratio = max(offered_qty, req_qty) / max(1e-9, min(offered_qty, req_qty))
                    known_qty_rows.append((pnorm, float(score), req_qty, ratio))
                strong_known: tuple[str, float, float, float] | None = None
                if known_qty_rows:
                    known_qty_rows.sort(key=lambda row: (row[3], -row[1]))
                    best_row = known_qty_rows[0]
                    if best_row[3] <= 1.15 and best_row[1] >= gpt_min_conf:
                        strong_known = best_row
                if strong_known is not None:
                    strong_norm, strong_score, _strong_req_qty, _strong_ratio = strong_known
                    narrowed: dict[str, float] = {}
                    for pnorm, score in gpt_scores.items():
                        req_qty = requested_qty_by_product_norm.get(pnorm)
                        if req_qty is not None and req_qty > 0:
                            ratio = max(offered_qty, req_qty) / max(
                                1e-9, min(offered_qty, req_qty)
                            )
                            # Keep quantity-compatible candidates; allow an
                            # outlier only if semantic confidence is clearly higher.
                            if ratio <= 1.50 or score >= strong_score + 0.20:
                                narrowed[pnorm] = float(score)
                            continue
                        # For candidates without requested quantity, require a
                        # strong semantic lead when a quantity-strong match exists.
                        if score >= strong_score + 0.25 or pnorm == strong_norm:
                            narrowed[pnorm] = float(score)
                    if narrowed:
                        filtered_scores = narrowed
            eligible: set[str] = set()
            for pnorm, score in filtered_scores.items():
                if pnorm not in known_product_norms:
                    continue
                pair_embed_scores[(item_row.id, pnorm)] = round(float(score), 4)
                if score >= gpt_min_conf:
                    eligible.add(pnorm)
            if eligible:
                # In GPT mode, use GPT eligible set as source of truth.
                pdf_multi_match_by_item_id[item_row.id] = eligible

    # Build the PDF-offer candidate index once so Results does not rescan all
    # extracted invoice lines for every requested product card.
    pdf_offer_candidates_by_product_norm: dict[str, list[tuple[InvoiceExtractionItem, InboundInvoice]]] = {}
    if pdf_items:
        for item_row, invoice in pdf_items:
            resolved_unit_price = (
                item_row.corrected_unit_price
                if item_row.corrected_unit_price is not None
                else item_row.unit_price
            )
            if resolved_unit_price is None or resolved_unit_price <= 0:
                continue

            matched_name = _norm(_line_base_for_matching(item_row))
            eligible_products = pdf_multi_match_by_item_id.get(item_row.id, set())
            gpt_matched_norm = None
            gpt_match_rate = _extract_gpt_score_from_notes(item_row.notes)
            gpt_matched_name = _extract_gpt_match_from_notes(item_row.notes)
            if gpt_matched_name:
                gpt_matched_norm = _norm(gpt_matched_name)
            gpt_match_is_eligible = bool(
                gpt_matched_norm
                and gpt_match_rate is not None
                and float(gpt_match_rate) >= _STANDARD_MIN_GPT_MATCH
            )

            candidate_norms: list[str] = []
            if use_strict_requested_match:
                if not eligible_products:
                    continue
                candidate_norms.extend(sorted(eligible_products))
            elif use_gpt_multi_match:
                if not eligible_products:
                    # Fallback to GPT-matched name if available.
                    if gpt_match_is_eligible:
                        candidate_norms.append(gpt_matched_norm)
                    else:
                        continue
                candidate_norms.extend(sorted(eligible_products))
            else:
                if matched_name:
                    candidate_norms.append(matched_name)
                if gpt_match_is_eligible and gpt_matched_norm not in candidate_norms:
                    candidate_norms.append(gpt_matched_norm)
                if not candidate_norms and matched_name:
                    best_norm = None
                    best_score = -1.0
                    for product_norm in known_product_norms:
                        score = difflib.SequenceMatcher(None, matched_name, product_norm).ratio()
                        if score > best_score:
                            best_score = score
                            best_norm = product_norm
                    if best_norm and best_score >= 0.50:
                        candidate_norms.append(best_norm)

            seen_candidate_norms: set[str] = set()
            for product_norm in candidate_norms:
                if (
                    not product_norm
                    or product_norm not in known_product_norms
                    or product_norm in seen_candidate_norms
                ):
                    continue
                seen_candidate_norms.add(product_norm)
                pdf_offer_candidates_by_product_norm.setdefault(product_norm, []).append(
                    (item_row, invoice)
                )

    # Pull supplier form offers once and group them by requested product.
    response_items_by_product_id: dict[int, list[tuple[ResponseItem, ResponseModel, Supplier]]] = {}
    response_rows = (
        db.query(ResponseItem, ResponseModel, Supplier)
        .join(ResponseModel, ResponseItem.response_id == ResponseModel.id)
        .join(Supplier, ResponseModel.supplier_id == Supplier.id)
        .filter(
            ResponseModel.campaign_id == campaign_id,
            ResponseModel.tenant_id == tenant_id,
            ResponseItem.product_id.in_(product_ids),
            ResponseItem.tenant_id == tenant_id,
            Supplier.tenant_id == tenant_id,
        )
        .all()
    )
    for item_row, response, supplier in response_rows:
        response_items_by_product_id.setdefault(item_row.product_id, []).append(
            (item_row, response, supplier)
        )

    results: list[ProductOffers] = []
    for product_id in product_ids:
        product = products_by_id.get(product_id)
        if not product:
            continue
        campaign_item = campaign_item_by_product_id.get(product_id)
        requested_qty = _to_positive_float(
            getattr(campaign_item, "requested_qty", None) if campaign_item is not None else None
        )
        source_candidate = _resolve_requested_source_attachment(product.name)
        requested_source_attachment_id = (
            int(source_candidate.get("attachment_id"))
            if isinstance(source_candidate, dict) and source_candidate.get("attachment_id")
            else None
        )
        requested_source_attachment_file = (
            str(source_candidate.get("attachment_file") or "").strip()
            if isinstance(source_candidate, dict)
            else None
        )
        if requested_source_attachment_file == "":
            requested_source_attachment_file = None
        requested_source_attachment_path = (
            str(source_candidate.get("attachment_path") or "").strip()
            if isinstance(source_candidate, dict)
            else None
        )
        if requested_source_attachment_path == "":
            requested_source_attachment_path = None
        requested_source_attachment_score = (
            float(source_candidate.get("score"))
            if isinstance(source_candidate, dict) and source_candidate.get("score") is not None
            else None
        )
        product_norm = _norm(product.name)
        response_items = response_items_by_product_id.get(product_id, [])
        if conform_only:
            response_items = [row for row in response_items if (row[0].notes or "") == "Conforme"]
        offers = []
        for item_row, response, supplier in response_items:
            direct_price = _to_positive_float(item_row.unit_price)
            form_eq_payload = {
                "equivalent_unit_price": direct_price,
                "equivalence_factor": 1.0 if direct_price is not None else None,
                "equivalence_basis": "direct_form" if direct_price is not None else None,
                "equivalence_confidence": 1.0 if direct_price is not None else None,
            }
            match_payload = _build_match_comparison_payload(
                eq_payload=form_eq_payload,
                strict_payload=None,
                unit_payload=None,
                requested_product_name=product.name,
                requested_description=None,
                requested_characteristics=None,
                requested_line_text=None,
                strict_match_score=None,
                strict_match_threshold=None,
                strict_match_decision=None,
                strict_mode=use_strict_requested_match,
            )
            offers.append(
                {
                    "supplier_id": supplier.id,
                    "supplier_name": _clean_display_text(supplier.name),
                    "supplier_email": supplier.email,
                    "supplier_phone": supplier.phone,
                    "sender_email": supplier.email,
                    "invoice_sender_email": supplier.email,
                    "requested_catalog_name": _clean_display_text(product.name),
                    "requested_strict_name_used": None,
                    "requested_strict_line_used": None,
                    "requested_strict_source": None,
                    "requested_qty": requested_qty,
                    "invoice_item_id": None,
                    "invoice_id": None,
                    "invoice_file_name": None,
                    "invoice_file_path": None,
                    "tech_sheet_id": None,
                    "tech_sheet_file": None,
                    "tech_sheet_file_path": None,
                    "offered_qty": item_row.offered_qty,
                    "offered_qty_raw": item_row.offered_qty,
                    "offered_qty_strict": None,
                    "unit_price": item_row.unit_price,
                    "equivalent_unit_price": form_eq_payload.get("equivalent_unit_price"),
                    "equivalence_factor": form_eq_payload.get("equivalence_factor"),
                    "equivalence_basis": form_eq_payload.get("equivalence_basis"),
                    "equivalence_confidence": form_eq_payload.get("equivalence_confidence"),
                    "currency": item_row.currency,
                    "price_source": "supplier_form",
                    "extracted_name": None,
                    "extracted_description": None,
                    "name_match_rate": 1.0,
                    "matched_characteristics": None,
                    "total_characteristics": None,
                    "characteristics_threshold": None,
                    "characteristics_details": None,
                    "tech_sheet_similarity_score": None,
                    "tech_sheet_similarity_status": None,
                    "tech_sheet_similarity_breakdown": None,
                    "tech_sheet_similarity_display_score": None,
                    "tech_sheet_similarity_display_status": None,
                    "strict_match_score": None,
                    "strict_match_threshold": None,
                    "strict_match_decision": None,
                    "strict_match_reason": None,
                    **match_payload,
                }
            )

        # Include PDF extraction items as offers (multi-match by threshold).
        for item_row, invoice in pdf_offer_candidates_by_product_norm.get(product_norm, []):
            resolved_unit_price = (
                item_row.corrected_unit_price
                if item_row.corrected_unit_price is not None
                else item_row.unit_price
            )
            # Results page should compare real offers; skip lines without PU.
            if resolved_unit_price is None or resolved_unit_price <= 0:
                continue
            matched_name = _norm(_line_base_for_matching(item_row))
            extracted_name_norm = _norm(item_row.product_name or "")
            name_match_rate = None
            if product_norm and extracted_name_norm:
                name_match_rate = round(
                    difflib.SequenceMatcher(None, product_norm, extracted_name_norm).ratio(),
                    4,
                )
            matched_points, total_points = _pair_points_for(item_row, product_norm)
            pair_cache = pair_cache_by_key.get((item_row.id, product_norm))
            pair_threshold = None
            pair_details = None
            if pair_cache is not None:
                try:
                    pair_threshold = (
                        float(pair_cache.threshold)
                        if pair_cache.threshold is not None
                        else None
                    )
                except Exception:
                    pair_threshold = None
                pair_details = pair_cache.points if isinstance(pair_cache.points, list) else None
            embedding_match_rate = pair_embed_scores.get((item_row.id, product_norm))
            fuzzy_pair = pair_fuzzy_scores.get((item_row.id, product_norm))
            if (
                not use_strict_requested_match
                and embedding_match_rate is None
                and isinstance(fuzzy_pair, float)
                and fuzzy_pair >= 0.85
            ):
                embedding_match_rate = round(fuzzy_pair, 4)
            if (
                not use_strict_requested_match
                and embedding_match_rate is None
                and matched_name == product_norm
            ):
                embedding_match_rate = _extract_emb_score_from_notes(item_row.notes)
            gpt_match_rate = _extract_gpt_score_from_notes(item_row.notes)
            gpt_matched_name = _extract_gpt_match_from_notes(item_row.notes)
            strict_diag = strict_match_diag_by_item_product.get((item_row.id, product_norm))
            strict_match_score = strict_diag.get("score") if isinstance(strict_diag, dict) else None
            strict_match_threshold = strict_diag.get("threshold") if isinstance(strict_diag, dict) else None
            strict_match_decision = (
                str(strict_diag.get("decision") or "").strip() or None
                if isinstance(strict_diag, dict)
                else None
            )
            strict_match_reason = (
                str(strict_diag.get("reason") or "").strip() or None
                if isinstance(strict_diag, dict)
                else None
            )
            strict_penalty_applied = (
                bool(strict_diag.get("penalty_applied"))
                if isinstance(strict_diag, dict)
                else None
            )
            strict_penalty_value = None
            if isinstance(strict_diag, dict):
                try:
                    strict_penalty_value = (
                        round(float(strict_diag.get("penalty_value")), 4)
                        if strict_diag.get("penalty_value") is not None
                        else None
                    )
                except Exception:
                    strict_penalty_value = None
            strict_penalty_labels = (
                [
                    str(value).strip()
                    for value in list(strict_diag.get("penalty_labels") or [])
                    if str(value).strip()
                ]
                if isinstance(strict_diag, dict)
                else None
            )
            strict_matched_name = (
                str(strict_diag.get("matched_name") or "").strip() or None
                if isinstance(strict_diag, dict)
                else None
            )
            strict_matched_row = (
                strict_diag.get("matched_row")
                if isinstance(strict_diag, dict) and isinstance(strict_diag.get("matched_row"), dict)
                else None
            )
            if use_strict_requested_match:
                gpt_match_rate = strict_match_score
                if strict_matched_name:
                    gpt_matched_name = strict_matched_name
            elif use_gpt_multi_match:
                pair_gpt = pair_embed_scores.get((item_row.id, product_norm))
                if pair_gpt is not None:
                    if gpt_match_rate is None:
                        gpt_match_rate = round(float(pair_gpt), 4)
                    if not gpt_matched_name:
                        gpt_matched_name = product.name
            supplier_match = None
            if invoice.sender_email:
                supplier_match = suppliers_by_email.get(invoice.sender_email.lower())
            if supplier_match is None and invoice.sender_name:
                sender_name_key = (_clean_display_text(invoice.sender_name) or "").lower()
                if sender_name_key:
                    supplier_match = suppliers_by_name.get(sender_name_key)
            supplier_name = (
                supplier_match.name
                if supplier_match is not None
                else (invoice.sender_name or invoice.sender_email or "PDF")
            )
            supplier_phone = supplier_match.phone if supplier_match is not None else invoice.sender_phone
            linked_sheet = tech_link_by_item_id.get(item_row.id)
            linked_sheet_obj = linked_sheet[1] if linked_sheet is not None else None
            linked_sheet_link = linked_sheet[0] if linked_sheet is not None else None
            requested_payload = requested_payload_by_product_id.get(product_id) or {}
            strict_requested_line_text = None
            if isinstance(strict_matched_row, dict):
                line_chunks: list[str] = []
                name_chunk = str(strict_matched_row.get("product_name") or "").strip()
                if name_chunk:
                    line_chunks.append(name_chunk)
                desc_chunk = str(strict_matched_row.get("requested_description") or "").strip()
                if desc_chunk:
                    line_chunks.append(desc_chunk)
                line_chunks.extend(_flatten_text_values(strict_matched_row.get("requested_characteristics")))
                dedup_chunks: list[str] = []
                seen_chunks: set[str] = set()
                for chunk in line_chunks:
                    token = _norm(chunk)
                    if not token or token in seen_chunks:
                        continue
                    seen_chunks.add(token)
                    dedup_chunks.append(chunk)
                strict_requested_line_text = " | ".join(dedup_chunks).strip() or None
            has_requested_reference = bool(
                (requested_payload.get("characteristics") or [])
                or (requested_payload.get("requirements") or [])
                or (requested_payload.get("accessories") or [])
                or requested_payload.get("description")
            )
            tech_raw_score: float | None = None
            tech_raw_status: str | None = None
            tech_breakdown: dict[str, Any] | None = None
            tech_display_score: float | None = None
            tech_display_status: str | None = None
            if use_strict_requested_match:
                tech_raw_score = None
                tech_raw_status = None
                tech_breakdown = None
                tech_display_score = None
                tech_display_status = None
            elif not has_requested_reference:
                tech_raw_status = "no_requested_reference"
                tech_display_status = "no_requested_reference"
            else:
                materialized = offer_analysis_by_key.get((item_row.id, product_norm))
                if materialized is not None and isinstance(materialized.payload, dict):
                    tech_raw_score = materialized.tech_sheet_similarity_score
                    tech_raw_status = materialized.tech_sheet_similarity_status
                    tech_breakdown = materialized.payload.get("tech_sheet_similarity_breakdown")
                    tech_display_score = materialized.tech_sheet_similarity_display_score
                    tech_display_status = materialized.tech_sheet_similarity_display_status
                elif linked_sheet_obj is None:
                    analysis_payload, _analysis_dirty = _get_or_compute_offer_analysis(
                        db=db,
                        item=item_row,
                        requested_product_name=product.name,
                        requested_product_norm=product_norm,
                        requested_product_id=product.id,
                        requested_payload=requested_payload,
                        linked_sheet_obj=None,
                        link_score=None,
                        gpt_match_score=gpt_match_rate,
                        invoice_item_name_for_sheet=item_row.product_name or _line_base_for_matching(item_row),
                    )
                    tech_raw_score = analysis_payload.get("tech_sheet_similarity_score")
                    tech_raw_status = analysis_payload.get("tech_sheet_similarity_status")
                    tech_breakdown = analysis_payload.get("tech_sheet_similarity_breakdown")
                    tech_display_score = analysis_payload.get("tech_sheet_similarity_display_score")
                    tech_display_status = analysis_payload.get("tech_sheet_similarity_display_status")
                else:
                    analysis_payload, _analysis_dirty = _get_or_compute_offer_analysis(
                        db=db,
                        item=item_row,
                        requested_product_name=product.name,
                        requested_product_norm=product_norm,
                        requested_product_id=product.id,
                        requested_payload=requested_payload,
                        linked_sheet_obj=linked_sheet_obj,
                        link_score=(float(linked_sheet_link.score) if linked_sheet_link is not None else None),
                        gpt_match_score=gpt_match_rate,
                        invoice_item_name_for_sheet=item_row.product_name or _line_base_for_matching(item_row),
                    )
                    tech_raw_score = analysis_payload.get("tech_sheet_similarity_score")
                    tech_raw_status = analysis_payload.get("tech_sheet_similarity_status")
                    tech_breakdown = analysis_payload.get("tech_sheet_similarity_breakdown")
                    tech_display_score = analysis_payload.get("tech_sheet_similarity_display_score")
                    tech_display_status = analysis_payload.get("tech_sheet_similarity_display_status")
            inferred_price_source = _extract_price_source_from_notes(item_row.notes)
            if not inferred_price_source and item_row.corrected_unit_price is not None:
                inferred_price_source = "manual_correction"
            if not inferred_price_source:
                inferred_price_source = (item_row.source_engine or "unknown_price_source").lower()
            requested_eq_text = requested_equivalence_text_by_product_id.get(product_id) or ""
            offered_eq_text = _build_price_eq_offered_text(
                corrected_name=item_row.corrected_product_name,
                extracted_name=item_row.product_name,
                extracted_description=item_row.description,
            )
            eq_payload = _compute_price_equivalence_effective(
                unit_price=resolved_unit_price,
                requested_text=requested_eq_text,
                offered_text=offered_eq_text,
                use_ai_prompt=bool(use_ai_prompt_for_conversion),
            )
            eq_context = _infer_price_eq_context(requested_eq_text, offered_eq_text)
            # Equipment/machines: never apply capacity-based conversion for price
            # equivalence. Keep direct unit price baseline only.
            if eq_context == "machine":
                eq_payload = {
                    "equivalent_unit_price": _to_positive_float(resolved_unit_price),
                    "equivalence_factor": 1.0 if _to_positive_float(resolved_unit_price) is not None else None,
                    "equivalence_basis": "direct_machine",
                    "equivalence_confidence": 0.72,
                }
            # If no proven conversion was found, keep direct comparable price for
            # unknown context only. For consommables/mixed, avoid unsafe fallback.
            elif _to_positive_float(eq_payload.get("equivalent_unit_price")) is None:
                if eq_context == "unknown":
                    eq_payload = {
                        "equivalent_unit_price": _to_positive_float(resolved_unit_price),
                        "equivalence_factor": 1.0 if _to_positive_float(resolved_unit_price) is not None else None,
                        "equivalence_basis": "direct_unknown",
                        "equivalence_confidence": 0.58,
                    }
            strict_payload = (
                item_row.strict_components_payload
                if isinstance(item_row.strict_components_payload, dict)
                else None
            )
            unit_payload = _build_match_unit_payload(
                strict_payload=strict_payload,
                requested_product_name=product.name,
                requested_description=requested_payload.get("description"),
                requested_characteristics=requested_payload,
                requested_line_text=strict_requested_line_text,
            )
            match_payload = _build_match_comparison_payload(
                eq_payload=eq_payload,
                strict_payload=strict_payload,
                unit_payload=unit_payload,
                requested_product_name=product.name,
                requested_description=requested_payload.get("description"),
                requested_characteristics=requested_payload,
                requested_line_text=strict_requested_line_text,
                strict_match_score=strict_match_score,
                strict_match_threshold=strict_match_threshold,
                strict_match_decision=strict_match_decision,
                strict_mode=use_strict_requested_match,
            )
            if use_strict_requested_match:
                eq_payload = _apply_strict_compensation_price_equivalence(
                    eq_payload=eq_payload,
                    match_payload=match_payload,
                    unit_price=resolved_unit_price,
                )
            requested_strict_name_used = None
            requested_strict_line_used = None
            requested_strict_source = None
            if isinstance(strict_matched_row, dict):
                requested_strict_name_used = _clean_display_text(str(strict_matched_row.get("product_name") or ""))
                requested_strict_line_used = _clean_display_text(strict_requested_line_text)
                requested_strict_source = "gpt_attachment_strict"
            offered_qty_raw = _to_positive_float(item_row.corrected_quantity)
            if offered_qty_raw is None:
                offered_qty_raw = _to_positive_float(item_row.quantity)
            offered_qty_strict = (
                _to_positive_float(strict_payload.get("quantity"))
                if isinstance(strict_payload, dict)
                else None
            )
            offered_qty_display = (
                offered_qty_strict
                if use_strict_requested_match and offered_qty_strict is not None
                else offered_qty_raw
            )
            vat_rate = (
                _to_positive_float(strict_payload.get("vat_rate"))
                if isinstance(strict_payload, dict)
                else None
            )
            if vat_rate is None and isinstance(strict_payload, dict):
                vat_rate = _to_positive_float(strict_payload.get("vat_rate_calc"))
            if offered_qty_display is None:
                offered_qty_display = 1.0
            offers.append(
                {
                    "supplier_id": supplier_match.id if supplier_match is not None else 0,
                    "supplier_name": _clean_display_text(supplier_name),
                    "supplier_email": (
                        supplier_match.email
                        if supplier_match is not None and str(supplier_match.email or "").strip()
                        else invoice.sender_email
                    ),
                    "supplier_phone": supplier_phone,
                    "sender_email": invoice.sender_email,
                    "invoice_sender_email": invoice.sender_email,
                    "requested_catalog_name": _clean_display_text(product.name),
                    "requested_strict_name_used": requested_strict_name_used,
                    "requested_strict_line_used": requested_strict_line_used,
                    "requested_strict_source": requested_strict_source,
                    "requested_qty": requested_qty,
                    "invoice_item_id": item_row.id,
                    "invoice_id": item_row.invoice_id,
                    "invoice_file_name": invoice.file_name,
                    "invoice_file_path": invoice.file_path,
                    "tech_sheet_id": linked_sheet_obj.id if linked_sheet_obj is not None else None,
                    "tech_sheet_file": linked_sheet_obj.file_name if linked_sheet_obj is not None else None,
                    "tech_sheet_file_path": linked_sheet_obj.file_path if linked_sheet_obj is not None else None,
                    "offered_qty": offered_qty_display,
                    "offered_qty_raw": offered_qty_raw,
                    "offered_qty_strict": offered_qty_strict,
                    "unit_price": resolved_unit_price,
                    "vat_rate": vat_rate,
                    "equivalent_unit_price": eq_payload.get("equivalent_unit_price"),
                    "equivalence_factor": eq_payload.get("equivalence_factor"),
                    "equivalence_basis": eq_payload.get("equivalence_basis"),
                    "equivalence_confidence": eq_payload.get("equivalence_confidence"),
                    "currency": item_row.corrected_currency or item_row.currency or "DZD",
                    "price_source": inferred_price_source,
                    "extracted_name": _clean_display_text(item_row.product_name),
                    "extracted_description": _clean_display_text(item_row.description),
                    "gpt_matched_name": _clean_display_text(gpt_matched_name),
                    "gpt_match_rate": gpt_match_rate,
                    "embedding_match_rate": embedding_match_rate,
                    "strict_match_score": strict_match_score,
                    "strict_match_threshold": strict_match_threshold,
                    "strict_match_decision": strict_match_decision,
                    "strict_match_reason": strict_match_reason,
                    "strict_penalty_applied": strict_penalty_applied,
                    "strict_penalty_value": strict_penalty_value,
                    "strict_penalty_labels": strict_penalty_labels,
                    "name_match_rate": name_match_rate,
                    "matched_characteristics": matched_points,
                    "total_characteristics": total_points,
                    "characteristics_threshold": pair_threshold,
                    "characteristics_details": pair_details,
                    "tech_sheet_similarity_score": tech_raw_score,
                    "tech_sheet_similarity_status": tech_raw_status,
                    "tech_sheet_similarity_breakdown": tech_breakdown,
                    "tech_sheet_similarity_display_score": tech_display_score,
                    "tech_sheet_similarity_display_status": tech_display_status,
                    **match_payload,
                }
            )
        offers_sorted = sorted(
            offers,
            key=_offer_comparable_price,
        )
        if limit is not None:
            offers_sorted = offers_sorted[:limit]
        best_offer = None
        for candidate in offers_sorted:
            if _offer_comparable_price(candidate) != float("inf"):
                best_offer = candidate
                break
        results.append(
            ProductOffers(
                product_id=product.id,
                product_name=product.name,
                requested_qty=requested_qty,
                requested_source_attachment_id=requested_source_attachment_id,
                requested_source_attachment_file=requested_source_attachment_file,
                requested_source_attachment_path=requested_source_attachment_path,
                requested_source_attachment_score=requested_source_attachment_score,
                offers=offers_sorted,
                best_offer=best_offer,
            )
        )
    return results


def _result_offer_gpt_score(raw: dict[str, Any]) -> float | None:
    value = raw.get("gpt_match_rate")
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except Exception:
            return None
    return None


def _serialize_results_payload(
    rows: list[ProductOffers],
    *,
    useful_only: bool,
    min_match_score: float,
) -> list[dict[str, Any]]:
    payload: list[dict[str, Any]] = []
    for row in rows:
        row_data = row.dict()
        offers_raw = row_data.get("offers") or []
        normalized_offers: list[dict[str, Any]] = []
        for offer_raw in offers_raw:
            if not isinstance(offer_raw, dict):
                continue
            score = _result_offer_gpt_score(offer_raw)
            if useful_only and (score is None or score < min_match_score):
                continue
            normalized_offers.append(dict(offer_raw))
        if useful_only and not normalized_offers:
            continue
        row_data["offers"] = normalized_offers
        row_data["best_offer"] = normalized_offers[0] if normalized_offers else None
        payload.append(row_data)
    return payload


def _summarize_offer_components(raw: dict[str, Any]) -> dict[str, Any]:
    out = dict(raw)
    breakdown = out.get("tech_sheet_similarity_breakdown")
    components = breakdown.get("components") if isinstance(breakdown, dict) else None
    if isinstance(components, dict):
        component_map = {
            "characteristics": "tech_characteristics",
            "requirements": "tech_requirements",
            "accessories": "tech_accessories",
        }
        for source_key, target_prefix in component_map.items():
            component = components.get(source_key)
            if not isinstance(component, dict):
                continue
            matched = component.get("matched")
            total = component.get("total")
            coverage = component.get("coverage")
            if isinstance(matched, (int, float)):
                out[f"{target_prefix}_matched"] = int(matched)
            if isinstance(total, (int, float)):
                out[f"{target_prefix}_total"] = int(total)
            if isinstance(coverage, (int, float)):
                out[f"{target_prefix}_coverage"] = float(coverage)
        description = components.get("description")
        if isinstance(description, dict):
            desc_score = description.get("score")
            if isinstance(desc_score, (int, float)):
                out["tech_description_score"] = float(desc_score)
    return out


def _lighten_offer_payload(raw: dict[str, Any]) -> dict[str, Any]:
    out = _summarize_offer_components(raw)
    out["characteristics_details"] = None
    out["tech_sheet_similarity_breakdown"] = None
    return out


def _lighten_product_payload(raw: dict[str, Any]) -> dict[str, Any]:
    out = dict(raw)
    offers_raw = out.get("offers") or []
    light_offers: list[dict[str, Any]] = []
    for offer_raw in offers_raw:
        if isinstance(offer_raw, dict):
            light_offers.append(_lighten_offer_payload(offer_raw))
    out["offers"] = light_offers
    best_offer = out.get("best_offer")
    if isinstance(best_offer, dict):
        out["best_offer"] = _lighten_offer_payload(best_offer)
    else:
        out["best_offer"] = None
    return out


def _cached_results_payload(
    campaign_id: int,
    db: Session,
    current_user: User,
    *,
    conform_only: bool,
    offer_limit: int | None,
    useful_only: bool,
    min_match_score: float,
) -> list[dict[str, Any]]:
    def _shared_cache_profile() -> str:
        if cache_campaign_mode == _CAMPAIGN_PROCESSING_STRICT:
            return "strict_campaign_global"
        if conform_only:
            return "standard_conform_only"
        if useful_only:
            return "standard_useful_only"
        return "standard_default"

    cache_campaign = (
        db.query(Campaign)
        .filter(
            Campaign.tenant_id == current_user.tenant_id,
            Campaign.id == campaign_id,
        )
        .first()
    )
    cache_campaign_mode = (
        _campaign_effective_processing_mode(db, cache_campaign)
        if cache_campaign is not None
        else _CAMPAIGN_PROCESSING_STANDARD
    )
    shared_campaign_cache = cache_campaign_mode in {
        _CAMPAIGN_PROCESSING_STRICT,
        _CAMPAIGN_PROCESSING_STANDARD,
    }
    cache_user_id = (
        int(cache_campaign.owner_user_id or current_user.id)
        if shared_campaign_cache and cache_campaign is not None
        else int(current_user.id)
    )
    thresholds = get_user_embedding_thresholds(
        db,
        current_user.tenant_id,
        cache_user_id,
    )
    signature = build_results_cache_signature(
        db,
        tenant_id=current_user.tenant_id,
        user_id=cache_user_id,
        campaign_id=campaign_id,
        confirm_threshold=float(thresholds.confirm_threshold),
        verify_threshold=float(thresholds.verify_threshold),
        characteristics_threshold=float(thresholds.characteristics_threshold),
        offer_limit=offer_limit,
        conform_only=conform_only,
        useful_only=useful_only,
        min_match_score=min_match_score,
    )
    if shared_campaign_cache:
        signature.pop("user_id", None)
        signature.pop("confirm_threshold", None)
        signature.pop("verify_threshold", None)
        signature.pop("characteristics_threshold", None)
        signature["cache_scope"] = _shared_cache_profile()
    cached = get_results_cache(
        db,
        tenant_id=current_user.tenant_id,
        user_id=cache_user_id,
        campaign_id=campaign_id,
        signature=signature,
        allow_any_user=shared_campaign_cache,
    )
    if cached is not None:
        return cached

    rows = _offers_by_product(
        campaign_id,
        db,
        current_user,
        limit=offer_limit,
        conform_only=conform_only,
    )
    payload = _serialize_results_payload(
        rows,
        useful_only=useful_only,
        min_match_score=min_match_score,
    )
    # `_offers_by_product()` can materialize/update offer analyses and other
    # derived rows that are part of the cache signature. Recompute the
    # signature after building the payload so the stored cache matches the
    # post-computation database state and is immediately reusable.
    refreshed_signature = build_results_cache_signature(
        db,
        tenant_id=current_user.tenant_id,
        user_id=cache_user_id,
        campaign_id=campaign_id,
        confirm_threshold=float(thresholds.confirm_threshold),
        verify_threshold=float(thresholds.verify_threshold),
        characteristics_threshold=float(thresholds.characteristics_threshold),
        offer_limit=offer_limit,
        conform_only=conform_only,
        useful_only=useful_only,
        min_match_score=min_match_score,
    )
    if shared_campaign_cache:
        refreshed_signature.pop("user_id", None)
        refreshed_signature.pop("confirm_threshold", None)
        refreshed_signature.pop("verify_threshold", None)
        refreshed_signature.pop("characteristics_threshold", None)
        refreshed_signature["cache_scope"] = _shared_cache_profile()
    try:
        set_results_cache(
            db,
            tenant_id=current_user.tenant_id,
            user_id=cache_user_id,
            campaign_id=campaign_id,
            signature=refreshed_signature,
            payload=payload,
        )
    except Exception as exc:
        logger.warning(
            "Results cache store failed: tenant=%s user=%s campaign=%s err=%s",
            current_user.tenant_id,
            current_user.id,
            campaign_id,
            exc,
        )
        db.rollback()
    return payload


def _get_cached_results_payload_if_available(
    campaign_id: int,
    db: Session,
    current_user: User,
    *,
    conform_only: bool,
    offer_limit: int | None,
    useful_only: bool,
    min_match_score: float,
) -> list[dict[str, Any]] | None:
    def _shared_cache_profile() -> str:
        if cache_campaign_mode == _CAMPAIGN_PROCESSING_STRICT:
            return "strict_campaign_global"
        if conform_only:
            return "standard_conform_only"
        if useful_only:
            return "standard_useful_only"
        return "standard_default"

    cache_campaign = (
        db.query(Campaign)
        .filter(
            Campaign.tenant_id == current_user.tenant_id,
            Campaign.id == campaign_id,
        )
        .first()
    )
    cache_campaign_mode = (
        _campaign_effective_processing_mode(db, cache_campaign)
        if cache_campaign is not None
        else _CAMPAIGN_PROCESSING_STANDARD
    )
    shared_campaign_cache = cache_campaign_mode in {
        _CAMPAIGN_PROCESSING_STRICT,
        _CAMPAIGN_PROCESSING_STANDARD,
    }
    cache_user_id = (
        int(cache_campaign.owner_user_id or current_user.id)
        if shared_campaign_cache and cache_campaign is not None
        else int(current_user.id)
    )
    thresholds = get_user_embedding_thresholds(
        db,
        current_user.tenant_id,
        cache_user_id,
    )
    signature = build_results_cache_signature(
        db,
        tenant_id=current_user.tenant_id,
        user_id=cache_user_id,
        campaign_id=campaign_id,
        confirm_threshold=float(thresholds.confirm_threshold),
        verify_threshold=float(thresholds.verify_threshold),
        characteristics_threshold=float(thresholds.characteristics_threshold),
        offer_limit=offer_limit,
        conform_only=conform_only,
        useful_only=useful_only,
        min_match_score=min_match_score,
    )
    if shared_campaign_cache:
        signature.pop("user_id", None)
        signature.pop("confirm_threshold", None)
        signature.pop("verify_threshold", None)
        signature.pop("characteristics_threshold", None)
        signature["cache_scope"] = _shared_cache_profile()
    cached = get_results_cache(
        db,
        tenant_id=current_user.tenant_id,
        user_id=cache_user_id,
        campaign_id=campaign_id,
        signature=signature,
        allow_any_user=shared_campaign_cache,
    )
    return cached


def _simple_product_model_name() -> str:
    return str(
        getattr(config, "OPENAI_SIMPLE_PRODUCT_MODEL", "") or _SIMPLE_PRODUCT_MODEL
    ).strip() or _SIMPLE_PRODUCT_MODEL


def _ensure_simple_product_enabled() -> None:
    if not _SIMPLE_PRODUCT_ENABLED:
        raise HTTPException(status_code=404, detail="Produit simple disabled")


def _simple_product_json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _simple_product_cache_key(signature: dict[str, Any]) -> str:
    raw = _simple_product_json_dumps(signature)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _simple_product_product_cache_signature(
    *,
    campaign_id: int,
    product_id: int,
    requested_row: dict[str, Any],
    candidates: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "logic": "simple_product_product",
        "logic_version": _SIMPLE_PRODUCT_LOGIC_VERSION,
        "campaign_id": campaign_id,
        "product_id": product_id,
        "requested_row": {
            "product_name": requested_row.get("product_name"),
            "requested_qty": requested_row.get("requested_qty"),
            "requested_ocr_snippet": requested_row.get("requested_ocr_snippet"),
            "requested_product_json": requested_row.get("requested_product_json"),
        },
        "candidate_rows": [
            {
                "invoice_item_id": candidate.get("invoice_item_id"),
                "invoice_id": candidate.get("invoice_id"),
                "produit_offert": candidate.get("produit_offert"),
                "prix_unitaire_facture": candidate.get("prix_unitaire_facture"),
                "quantite_offerte": candidate.get("quantite_offerte"),
                "tva": candidate.get("tva"),
                "source_engine": candidate.get("source_engine"),
                "updated_at": str(candidate.get("updated_at") or ""),
            }
            for candidate in candidates
        ],
    }


def _simple_product_signature(
    campaign_id: int,
    db: Session,
    current_user: User,
) -> dict[str, Any]:
    thresholds = get_user_embedding_thresholds(db, current_user.tenant_id, current_user.id)
    signature = build_results_cache_signature(
        db,
        tenant_id=current_user.tenant_id,
        user_id=current_user.id,
        campaign_id=campaign_id,
        confirm_threshold=float(thresholds.confirm_threshold),
        verify_threshold=float(thresholds.verify_threshold),
        characteristics_threshold=float(thresholds.characteristics_threshold),
        offer_limit=_SIMPLE_PRODUCT_TOP_OFFERS_PER_INVOICE,
        conform_only=False,
        useful_only=False,
        min_match_score=0.0,
    )
    signature.update(
        {
            "logic": "simple_product",
            "logic_version": _SIMPLE_PRODUCT_LOGIC_VERSION,
            "simple_product_model": _simple_product_model_name(),
            "simple_product_top_candidates": _SIMPLE_PRODUCT_TOP_CANDIDATES,
            "simple_product_top_offers_per_invoice": _SIMPLE_PRODUCT_TOP_OFFERS_PER_INVOICE,
        }
    )
    return signature


def _simple_product_clip_text(value: Any, *, max_len: int = 1400) -> str | None:
    text = re.sub(r"\s+", " ", str(value or "").strip())
    if not text:
        return None
    if len(text) <= max_len:
        return text
    return text[:max_len].rstrip() + "..."


def _simple_product_tokens(value: str | None) -> list[str]:
    text = unicodedata.normalize("NFKD", str(value or "").lower()).encode(
        "ascii", "ignore"
    ).decode("ascii")
    text = _strip_name_noise_for_match(text)
    return [token for token in re.split(r"[^a-z0-9]+", text) if token and len(token) > 1]


def _simple_product_measure_tokens(value: str | None) -> set[str]:
    text = unicodedata.normalize("NFKD", str(value or "").lower()).encode(
        "ascii", "ignore"
    ).decode("ascii")
    out: set[str] = set()
    for match in re.finditer(r"(\d+(?:[.,]\d+)?)\s*(ml|cl|l|mm|cm|m|mg|g|kg|fr|ch)\b", text):
        out.add(f"{match.group(1).replace(',', '.')}{match.group(2)}")
    for match in re.finditer(r"\b(?:bt|bte|boite|boite de|box|pack|b/)\s*([0-9]+)\b", text):
        out.add(f"pack:{match.group(1)}")
    return out


def _simple_product_overlap_ratio(left: list[str], right: list[str]) -> float:
    if not left or not right:
        return 0.0
    left_set = set(left)
    right_set = set(right)
    inter = len(left_set & right_set)
    union = len(left_set | right_set)
    if union <= 0:
        return 0.0
    return inter / union


def _simple_product_requested_rows(
    campaign_id: int,
    db: Session,
    current_user: User,
) -> list[dict[str, Any]]:
    campaign = _get_campaign(db, campaign_id, current_user)
    requested_index = _build_campaign_requested_index(
        db=db,
        tenant_id=current_user.tenant_id,
        campaign_ids=[campaign_id],
    ).get(campaign_id, [])
    requested_index_by_product_id: dict[int, dict[str, Any]] = {}
    for raw in requested_index:
        try:
            product_id = int(raw.get("product_id"))
        except Exception:
            continue
        requested_index_by_product_id[product_id] = dict(raw)

    rows = (
        db.query(CampaignItem, Product)
        .join(Product, Product.id == CampaignItem.product_id)
        .filter(
            CampaignItem.tenant_id == current_user.tenant_id,
            CampaignItem.campaign_id == campaign_id,
            Product.tenant_id == current_user.tenant_id,
        )
        .order_by(CampaignItem.id.asc())
        .all()
    )
    seen_product_ids: set[int] = set()
    out: list[dict[str, Any]] = []
    for item_row, product in rows:
        product_id = int(product.id)
        if product_id in seen_product_ids:
            continue
        seen_product_ids.add(product_id)
        indexed = dict(requested_index_by_product_id.get(product_id) or {})
        requested_description = indexed.get("requested_description")
        if not isinstance(requested_description, str) or not requested_description.strip():
            requested_description = item_row.requested_description or getattr(product, "description", None)
        requested_characteristics = indexed.get("requested_characteristics")
        if not isinstance(requested_characteristics, dict):
            requested_characteristics = (
                item_row.requested_characteristics
                if isinstance(item_row.requested_characteristics, dict)
                else product.characteristics
            )
        requested_raw_product_name = _clean_display_text(
            str(indexed.get("requested_raw_product_name") or product.name or "")
        ) or None
        requested_raw_text = _simple_product_clip_text(indexed.get("requested_raw_text"), max_len=1800)
        requested_payload = _flatten_requested_components(
            requested_description=requested_description,
            requested_characteristics=requested_characteristics,
            requested_product_name=str(product.name or "").strip(),
            requested_raw_product_name=requested_raw_product_name,
            requested_raw_text=requested_raw_text,
        )
        requested_qty = _to_positive_float(item_row.requested_qty)
        requested_name = str(product.name or "").strip()
        out.append(
            {
                "product_id": product_id,
                "product_name": requested_name,
                "requested_qty": requested_qty,
                "requested_ocr_snippet": requested_raw_text
                or _simple_product_clip_text(requested_payload.get("description"), max_len=900),
                "requested_product_json": requested_payload,
                "rank_name": requested_raw_product_name or requested_name,
                "rank_tokens": _simple_product_tokens(requested_raw_product_name or requested_name),
                "rank_measure_tokens": _simple_product_measure_tokens(requested_raw_text or requested_name),
                "campaign_mode": _campaign_effective_processing_mode(db, campaign),
            }
        )
    return out


def _simple_product_offered_candidates(
    campaign_id: int,
    db: Session,
    current_user: User,
) -> list[dict[str, Any]]:
    rows = (
        db.query(InvoiceExtractionItem, InboundInvoice)
        .join(InboundInvoice, InvoiceExtractionItem.invoice_id == InboundInvoice.id)
        .filter(
            InvoiceExtractionItem.tenant_id == current_user.tenant_id,
            InvoiceExtractionItem.campaign_id == campaign_id,
            InboundInvoice.tenant_id == current_user.tenant_id,
        )
        .order_by(InvoiceExtractionItem.id.asc())
        .all()
    )
    supplier_ids = {
        int(item_row.supplier_id)
        for item_row, _invoice in rows
        if item_row.supplier_id is not None
    }
    suppliers_by_id = {
        int(row.id): row
        for row in db.query(Supplier)
        .filter(
            Supplier.tenant_id == current_user.tenant_id,
            Supplier.id.in_(supplier_ids) if supplier_ids else False,
        )
        .all()
    } if supplier_ids else {}

    out: list[dict[str, Any]] = []
    for item_row, invoice in rows:
        strict_payload = (
            dict(item_row.strict_components_payload)
            if isinstance(item_row.strict_components_payload, dict)
            else {}
        )
        offered_name = _clean_display_text(
            str(
                strict_payload.get("name")
                or item_row.corrected_product_name
                or item_row.product_name
                or ""
            )
        ) or None
        if not offered_name:
            continue
        unit_price = _to_positive_float(item_row.corrected_unit_price)
        if unit_price is None:
            unit_price = _to_positive_float(strict_payload.get("unit_price"))
        if unit_price is None:
            unit_price = _to_positive_float(item_row.unit_price)
        offered_qty = _to_positive_float(strict_payload.get("quantity"))
        if offered_qty is None:
            offered_qty = _to_positive_float(item_row.corrected_quantity)
        if offered_qty is None:
            offered_qty = _to_positive_float(item_row.quantity)
        vat_rate = _to_positive_float(strict_payload.get("vat_rate"))
        if vat_rate is None:
            vat_rate = _to_positive_float(strict_payload.get("vat_rate_calc"))
        offered_json = {
            "name": offered_name,
            "description": _clean_display_text(item_row.description),
            "raw_description": _clean_display_text(item_row.raw_description),
            "quantity": offered_qty,
            "unit": _clean_display_text(
                str(
                    strict_payload.get("unit")
                    or item_row.corrected_unit
                    or item_row.unit
                    or ""
                )
            )
            or None,
            "packaging": _clean_display_text(str(strict_payload.get("packaging") or "")) or None,
            "dimensions": _clean_display_text(str(strict_payload.get("dimensions") or "")) or None,
            "size": _clean_display_text(str(strict_payload.get("size") or "")) or None,
            "weight": _clean_display_text(str(strict_payload.get("weight") or "")) or None,
            "vat_rate": vat_rate,
        }
        offered_json = {key: value for key, value in offered_json.items() if value not in (None, "", [], {})}
        ocr_snippet = _simple_product_clip_text(
            strict_payload.get("raw_text")
            or item_row.raw_description
            or item_row.description
            or invoice.ocr_text_snapshot,
            max_len=1800,
        )
        supplier = suppliers_by_id.get(int(item_row.supplier_id or 0))
        out.append(
            {
                "supplier_id": int(item_row.supplier_id or getattr(supplier, "id", 0) or 0),
                "supplier_name": _clean_display_text(
                    str(
                        getattr(supplier, "name", None)
                        or invoice.sender_name
                        or invoice.sender_email
                        or "Inconnu"
                    )
                )
                or "Inconnu",
                "supplier_email": _clean_display_text(
                    str(getattr(supplier, "email", None) or invoice.sender_email or "")
                ),
                "supplier_phone": _clean_display_text(
                    str(getattr(supplier, "phone", None) or invoice.sender_phone or "")
                ),
                "sender_email": invoice.sender_email,
                "invoice_item_id": int(item_row.id),
                "invoice_id": int(invoice.id),
                "invoice_file_name": invoice.file_name,
                "invoice_file_path": invoice.file_path,
                "source_engine": item_row.source_engine,
                "produit_offert": offered_name,
                "produit_offert_json": offered_json,
                "produit_offert_ocr_snippet": ocr_snippet,
                "prix_unitaire_facture": unit_price,
                "tva": vat_rate,
                "quantite_offerte": offered_qty,
                "rank_tokens": _simple_product_tokens(offered_name),
                "rank_measure_tokens": _simple_product_measure_tokens(
                    ocr_snippet or offered_name
                ),
                "updated_at": getattr(item_row, "updated_at", None),
            }
        )
    return out


def _simple_product_rank_candidates(
    requested_row: dict[str, Any],
    offered_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    requested_name = str(requested_row.get("rank_name") or requested_row.get("product_name") or "")
    requested_norm = _strip_name_noise_for_match(requested_name)
    requested_tokens = list(requested_row.get("rank_tokens") or [])
    requested_measures = set(requested_row.get("rank_measure_tokens") or [])
    scored: list[tuple[float, dict[str, Any]]] = []
    for offer in offered_rows:
        offered_name = str(offer.get("produit_offert") or "")
        offered_norm = _strip_name_noise_for_match(offered_name)
        seq = difflib.SequenceMatcher(None, requested_norm, offered_norm).ratio() if requested_norm and offered_norm else 0.0
        overlap = _simple_product_overlap_ratio(requested_tokens, list(offer.get("rank_tokens") or []))
        offered_measures = set(offer.get("rank_measure_tokens") or [])
        measure_overlap = (
            len(requested_measures & offered_measures) / max(len(requested_measures | offered_measures), 1)
            if requested_measures and offered_measures
            else 0.0
        )
        local_score = (0.55 * seq) + (0.35 * overlap) + (0.10 * measure_overlap)
        candidate = dict(offer)
        candidate["local_score"] = round(float(local_score), 4)
        scored.append((local_score, candidate))
    scored.sort(
        key=lambda row: (
            -float(row[0]),
            float("inf")
            if _to_positive_float(row[1].get("prix_unitaire_facture")) is None
            else float(_to_positive_float(row[1].get("prix_unitaire_facture")) or 0.0),
            str(row[1].get("produit_offert") or ""),
        )
    )
    kept = [candidate for score, candidate in scored if score >= 0.08][: _SIMPLE_PRODUCT_TOP_CANDIDATES]
    if kept:
        return kept
    return [candidate for _score, candidate in scored[: _SIMPLE_PRODUCT_TOP_CANDIDATES]]


def _simple_product_sort_offers(
    offers: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    def _key(raw: dict[str, Any]) -> tuple[Any, ...]:
        equivalent = _to_positive_float(raw.get("prix_equivalent"))
        return (
            0 if bool(raw.get("conforme")) else 1,
            -float(raw.get("score") or 0.0),
            float("inf") if equivalent is None else equivalent,
            str(raw.get("supplier_name") or ""),
        )

    return sorted(offers, key=_key)


def _simple_product_limit_offers_per_invoice(
    offers: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    sorted_offers = _simple_product_sort_offers(offers)
    kept: list[dict[str, Any]] = []
    counts_by_invoice_id: dict[int, int] = {}
    for offer in sorted_offers:
        try:
            invoice_id = int(offer.get("invoice_id") or 0)
        except Exception:
            invoice_id = 0
        if invoice_id > 0:
            current_count = counts_by_invoice_id.get(invoice_id, 0)
            if current_count >= _SIMPLE_PRODUCT_TOP_OFFERS_PER_INVOICE:
                continue
            counts_by_invoice_id[invoice_id] = current_count + 1
        kept.append(offer)
    return kept


def _simple_product_merge_gpt_offer(
    *,
    requested_row: dict[str, Any],
    candidate: dict[str, Any],
    raw: dict[str, Any],
) -> dict[str, Any]:
    score = max(0.0, min(1.0, float(raw.get("score") or 0.0)))
    return {
        **candidate,
        "produit_demande": str(requested_row.get("product_name") or ""),
        "produit_offert": _clean_display_text(
            str(raw.get("produit_offert") or candidate.get("produit_offert") or "")
        )
        or str(candidate.get("produit_offert") or ""),
        "conforme": bool(raw.get("conforme")),
        "score": round(score, 4),
        "prix_unitaire_facture": _to_positive_float(raw.get("prix_unitaire_facture"))
        if raw.get("prix_unitaire_facture") is not None
        else _to_positive_float(candidate.get("prix_unitaire_facture")),
        "prix_equivalent": _to_positive_float(raw.get("prix_equivalent")),
        "tva": _to_positive_float(raw.get("tva"))
        if raw.get("tva") is not None
        else _to_positive_float(candidate.get("tva")),
        "quantite_demandee": _to_positive_float(raw.get("quantite_demandee"))
        if raw.get("quantite_demandee") is not None
        else _to_positive_float(requested_row.get("requested_qty")),
        "quantite_offerte": _to_positive_float(raw.get("quantite_offerte"))
        if raw.get("quantite_offerte") is not None
        else _to_positive_float(candidate.get("quantite_offerte")),
        "similarite": _simple_product_clip_text(raw.get("similarite"), max_len=300),
        "difference": [
            _simple_product_clip_text(value, max_len=220) or ""
            for value in (raw.get("difference") or [])
            if _simple_product_clip_text(value, max_len=220)
        ],
        "raison": _simple_product_clip_text(raw.get("raison"), max_len=400),
    }


def _simple_product_compare_batch_with_gpt(
    *,
    campaign_id: int,
    requested_rows: list[dict[str, Any]],
    offered_rows: list[dict[str, Any]],
) -> dict[int, list[dict[str, Any]]]:
    outputs: dict[int, list[dict[str, Any]]] = {}
    if not requested_rows:
        return outputs
    if not config.OPENAI_API_KEY:
        return {int(row.get("product_id") or 0): [] for row in requested_rows}

    model = _simple_product_model_name()
    pending: list[dict[str, Any]] = []
    total_candidates = 0
    for requested_row in requested_rows:
        product_id = int(requested_row.get("product_id") or 0)
        ranked_candidates = _simple_product_rank_candidates(requested_row, offered_rows)
        total_candidates += len(ranked_candidates)
        if not ranked_candidates:
            outputs[product_id] = []
            continue
        cache_signature = _simple_product_product_cache_signature(
            campaign_id=campaign_id,
            product_id=product_id,
            requested_row=requested_row,
            candidates=ranked_candidates,
        )
        cache_key = _simple_product_cache_key(cache_signature)
        cached = get_openai_payload_cache(
            cache_type=_SIMPLE_PRODUCT_PRODUCT_CACHE_TYPE,
            model=model,
            cache_key=cache_key,
        )
        if isinstance(cached, list):
            restored: list[dict[str, Any]] = []
            for raw in cached:
                if isinstance(raw, dict):
                    restored.append(dict(raw))
            outputs[product_id] = _simple_product_limit_offers_per_invoice(restored)
            continue
        pending.append(
            {
                "product_id": product_id,
                "requested_row": requested_row,
                "candidates": ranked_candidates,
                "cache_key": cache_key,
            }
        )
    if not pending:
        return outputs

    system_prompt = (
        "Tu compares plusieurs produits demandes avec plusieurs produits offerts. "
        "Tu travailles pour une logique metier distincte de la logique strict. "
        "Tu peux retourner plusieurs offres, mais jamais plus de 2 offres pour la meme facture. "
        "Pour chaque facture, garde seulement les 2 offres les plus conformes et les plus utiles pour ce produit demande. "
        "Le champ prix_unitaire_facture_float est toujours un float fiable et prioritaire. "
        "Ne le confonds jamais avec la TVA, le TTC, le total ligne ou le total facture. "
        "La TVA est informative seulement. "
        "Les liquides sont compensables: 1 l = 1000 ml et 1 cl = 10 ml. "
        "Si demande 10 ml et offre 20 ml, alors prix_equivalent = prix_unitaire_facture_float / 2. "
        "Les conditionnements explicites sont compensables: boite 300, bt300, b/300, 2x50, 3*100. "
        "Dans ce cas, calcule prix_equivalent sur la base de la quantite demandee. "
        "Les autres dimensions techniques ne sont pas librement compensables: longueur, diametre, gauge, taille, poids. "
        "Elles doivent etre compatibles. "
        "conforme=true seulement si la famille produit est la meme et que les differences restantes sont acceptables. "
        "score va de 0 a 1. Si le doute est important, conforme=false et score bas. "
        "Les textes similarite, difference et raison doivent etre en francais. "
        "Pour chaque produit demandé, retourne un objet avec product_id et offers. "
        "Retourne uniquement un JSON valide."
    )
    batch_size = max(
        1,
        int(getattr(config, "OPENAI_SIMPLE_PRODUCT_BATCH_SIZE", _SIMPLE_PRODUCT_GPT_BATCH_SIZE) or _SIMPLE_PRODUCT_GPT_BATCH_SIZE),
    )
    timeout_s = max(20, int(getattr(config, "OPENAI_SIMPLE_PRODUCT_TIMEOUT_SECONDS", 90) or 90))
    retry_attempts = max(1, int(getattr(config, "OPENAI_SIMPLE_PRODUCT_RETRY_ATTEMPTS", 3) or 3))
    retry_backoff_s = max(0.2, float(getattr(config, "OPENAI_SIMPLE_PRODUCT_RETRY_BACKOFF_SECONDS", 2.0) or 2.0))
    throttle_delay_s = max(
        0.0,
        float(getattr(config, "OPENAI_SIMPLE_PRODUCT_THROTTLE_SECONDS", 2.5) or 2.5),
    )

    for start in range(0, len(pending), batch_size):
        chunk = pending[start : start + batch_size]
        products_payload = []
        candidate_maps: dict[int, dict[int, dict[str, Any]]] = {}
        requested_by_product_id: dict[int, dict[str, Any]] = {}
        for entry in chunk:
            product_id = int(entry["product_id"])
            requested_row = dict(entry["requested_row"])
            candidates = list(entry["candidates"])
            requested_by_product_id[product_id] = requested_row
            candidate_maps[product_id] = {
                int(candidate["invoice_item_id"]): candidate
                for candidate in candidates
                if candidate.get("invoice_item_id") is not None
            }
            products_payload.append(
                {
                    "product_id": product_id,
                    "requested": {
                        "produit_demande": requested_row.get("product_name"),
                        "quantite_demandee": requested_row.get("requested_qty"),
                        "ocr_demande_snippet": requested_row.get("requested_ocr_snippet"),
                        "produit_demande_json": requested_row.get("requested_product_json"),
                    },
                    "offers": [
                        {
                            "invoice_item_id": candidate.get("invoice_item_id"),
                            "supplier_name": candidate.get("supplier_name"),
                            "produit_offert": candidate.get("produit_offert"),
                            "ocr_offre_snippet": candidate.get("produit_offert_ocr_snippet"),
                            "produit_offert_json": candidate.get("produit_offert_json"),
                            "prix_unitaire_facture_float": candidate.get("prix_unitaire_facture"),
                            "tva": candidate.get("tva"),
                            "quantite_offerte": candidate.get("quantite_offerte"),
                            "source_engine": candidate.get("source_engine"),
                            "local_score": candidate.get("local_score"),
                        }
                        for candidate in candidates
                    ],
                    "max_results_per_invoice": _SIMPLE_PRODUCT_TOP_OFFERS_PER_INVOICE,
                }
            )

        payload = {
            "model": model,
            "temperature": 0,
            "seed": 42,
            "messages": [
                {"role": "system", "content": system_prompt},
                {
                    "role": "user",
                    "content": _simple_product_json_dumps({"products": products_payload}),
                },
            ],
            "response_format": {
                "type": "json_schema",
                "json_schema": {
                    "name": "simple_product_compare_batch",
                    "strict": True,
                    "schema": {
                        "type": "object",
                        "additionalProperties": False,
                        "properties": {
                            "results": {
                                "type": "array",
                                "items": {
                                    "type": "object",
                                    "additionalProperties": False,
                                    "properties": {
                                        "product_id": {"type": "integer"},
                                        "offers": {
                                            "type": "array",
                                            "items": {
                                                "type": "object",
                                                "additionalProperties": False,
                                                "properties": {
                                                    "invoice_item_id": {"type": "integer"},
                                                    "produit_offert": {"type": "string"},
                                                    "conforme": {"type": "boolean"},
                                                    "score": {"type": "number", "minimum": 0, "maximum": 1},
                                                    "prix_unitaire_facture": {"type": ["number", "null"]},
                                                    "prix_equivalent": {"type": ["number", "null"]},
                                                    "tva": {"type": ["number", "null"]},
                                                    "quantite_demandee": {"type": ["number", "null"]},
                                                    "quantite_offerte": {"type": ["number", "null"]},
                                                    "similarite": {"type": ["string", "null"]},
                                                    "difference": {
                                                        "type": "array",
                                                        "items": {"type": "string"},
                                                    },
                                                    "raison": {"type": ["string", "null"]},
                                                },
                                                "required": [
                                                    "invoice_item_id",
                                                    "produit_offert",
                                                    "conforme",
                                                    "score",
                                                    "prix_unitaire_facture",
                                                    "prix_equivalent",
                                                    "tva",
                                                    "quantite_demandee",
                                                    "quantite_offerte",
                                                    "similarite",
                                                    "difference",
                                                    "raison",
                                                ],
                                            },
                                        },
                                    },
                                    "required": ["product_id", "offers"],
                                },
                            }
                        },
                        "required": ["results"],
                    },
                },
            },
        }
        parsed, _used_model = call_chat_completion_json_with_retry(
            operation="simple_product_compare_gpt_batch",
            payload=payload,
            primary_model=model,
            fallback_model="",
            timeout_s=timeout_s,
            retry_attempts=retry_attempts,
            retry_backoff_s=retry_backoff_s,
            metadata={
                "campaign_id": campaign_id,
                "product_count": len(chunk),
                "candidate_count": sum(len(entry["candidates"]) for entry in chunk),
            },
        )

        result_map: dict[int, list[dict[str, Any]]] = {}
        if isinstance(parsed, dict) and isinstance(parsed.get("results"), list):
            for row in parsed["results"]:
                if not isinstance(row, dict):
                    continue
                try:
                    product_id = int(row.get("product_id"))
                except Exception:
                    continue
                offers_raw = row.get("offers")
                if isinstance(offers_raw, list):
                    result_map[product_id] = [dict(raw) for raw in offers_raw if isinstance(raw, dict)]

        for entry in chunk:
            product_id = int(entry["product_id"])
            requested_row = requested_by_product_id.get(product_id) or dict(entry["requested_row"])
            candidate_map = candidate_maps.get(product_id) or {}
            raw_offers = result_map.get(product_id) or []
            merged_offers: list[dict[str, Any]] = []
            for raw in raw_offers:
                try:
                    invoice_item_id = int(raw.get("invoice_item_id"))
                except Exception:
                    continue
                candidate = candidate_map.get(invoice_item_id)
                if candidate is None:
                    continue
                merged_offers.append(
                    _simple_product_merge_gpt_offer(
                        requested_row=requested_row,
                        candidate=candidate,
                        raw=raw,
                    )
                )
            merged_offers = _simple_product_limit_offers_per_invoice(merged_offers)
            outputs[product_id] = merged_offers
            set_openai_payload_cache(
                cache_type=_SIMPLE_PRODUCT_PRODUCT_CACHE_TYPE,
                model=model,
                cache_key=str(entry["cache_key"]),
                payload=merged_offers,
            )
        if throttle_delay_s > 0 and (start + batch_size) < len(pending):
            time.sleep(throttle_delay_s)

    return outputs


def _simple_product_payload(
    campaign_id: int,
    db: Session,
    current_user: User,
) -> list[dict[str, Any]]:
    _ensure_simple_product_enabled()
    signature = _simple_product_signature(campaign_id, db, current_user)
    cache_key = _simple_product_cache_key(signature)
    model = _simple_product_model_name()
    cached = get_openai_payload_cache(
        cache_type=_SIMPLE_PRODUCT_CACHE_TYPE,
        model=model,
        cache_key=cache_key,
    )
    if isinstance(cached, list):
        out: list[dict[str, Any]] = []
        for raw in cached:
            if isinstance(raw, dict):
                out.append(dict(raw))
        return out

    requested_rows = _simple_product_requested_rows(campaign_id, db, current_user)
    offered_rows = _simple_product_offered_candidates(campaign_id, db, current_user)
    offers_by_product_id = _simple_product_compare_batch_with_gpt(
        campaign_id=campaign_id,
        requested_rows=requested_rows,
        offered_rows=offered_rows,
    )
    payload: list[dict[str, Any]] = []
    for requested_row in requested_rows:
        product_id = int(requested_row.get("product_id") or 0)
        offers = _simple_product_limit_offers_per_invoice(
            list(offers_by_product_id.get(product_id) or [])
        )
        row_payload = {
            "product_id": product_id,
            "product_name": str(requested_row.get("product_name") or ""),
            "requested_qty": _to_positive_float(requested_row.get("requested_qty")),
            "requested_ocr_snippet": requested_row.get("requested_ocr_snippet"),
            "requested_product_json": requested_row.get("requested_product_json"),
            "offers": offers,
            "best_offer": offers[0] if offers else None,
        }
        payload.append(row_payload)

    set_openai_payload_cache(
        cache_type=_SIMPLE_PRODUCT_CACHE_TYPE,
        model=model,
        cache_key=cache_key,
        payload=payload,
    )
    return payload


def _get_simple_product_payload_if_available(
    campaign_id: int,
    db: Session,
    current_user: User,
) -> list[dict[str, Any]] | None:
    signature = _simple_product_signature(campaign_id, db, current_user)
    cache_key = _simple_product_cache_key(signature)
    model = _simple_product_model_name()
    cached = get_openai_payload_cache(
        cache_type=_SIMPLE_PRODUCT_CACHE_TYPE,
        model=model,
        cache_key=cache_key,
    )
    if not isinstance(cached, list):
        return None
    out: list[dict[str, Any]] = []
    for raw in cached:
        if isinstance(raw, dict):
            out.append(dict(raw))
    return out


def warm_results_cache_for_campaign(
    campaign_id: int,
    db: Session,
    current_user: User,
) -> None:
    campaign = (
        db.query(Campaign)
        .filter(
            Campaign.tenant_id == current_user.tenant_id,
            Campaign.id == campaign_id,
        )
        .first()
    )
    campaign_mode = (
        _campaign_effective_processing_mode(db, campaign)
        if campaign is not None
        else _CAMPAIGN_PROCESSING_STANDARD
    )

    # Warm the exact cache signature consumed by the Strict screen
    # (cache_only=true -> useful_only=False, min_match_score=0.0, offer_limit=20).
    _cached_results_payload(
        campaign_id,
        db,
        current_user,
        conform_only=False,
        offer_limit=20,
        useful_only=False,
        min_match_score=0.0,
    )

    # campaign_results_cache stores a single row per tenant/user/campaign.
    # For strict campaigns, keep that row aligned with the Strict screen's
    # exact signature and avoid overwriting it with a different filtered view.
    if campaign_mode == _CAMPAIGN_PROCESSING_STRICT:
        return

    # Warm the lighter filtered signature used by other standard result views.
    _cached_results_payload(
        campaign_id,
        db,
        current_user,
        conform_only=False,
        offer_limit=20,
        useful_only=True,
        min_match_score=_RESULTS_MIN_GPT_MATCH,
    )
    # If no offer analysis was materialized yet but invoice items exist,
    # run a broader warm pass so Results are precomputed without manual recompute.
    try:
        offer_count = (
            db.query(func.count(OfferAnalysisResult.id))
            .filter(
                OfferAnalysisResult.tenant_id == current_user.tenant_id,
                OfferAnalysisResult.campaign_id == campaign_id,
            )
            .scalar()
            or 0
        )
        invoice_items = (
            db.query(func.count(InvoiceExtractionItem.id))
            .filter(
                InvoiceExtractionItem.tenant_id == current_user.tenant_id,
                InvoiceExtractionItem.campaign_id == campaign_id,
            )
            .scalar()
            or 0
        )
        if offer_count == 0 and invoice_items > 0:
            _cached_results_payload(
                campaign_id,
                db,
                current_user,
                conform_only=False,
                offer_limit=50,
                useful_only=False,
                min_match_score=0.0,
            )
    except Exception as exc:
        logger.warning(
            "Results cache auto-warm fallback failed: tenant=%s user=%s campaign=%s err=%s",
            current_user.tenant_id,
            current_user.id,
            campaign_id,
            exc,
        )


@router.get("", response_model=list[CampaignOut])
def list_campaigns(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignOut]:
    query = db.query(Campaign)
    if not current_user.is_admin:
        query = query.filter(
            Campaign.tenant_id == current_user.tenant_id,
            Campaign.owner_user_id == current_user.id,
        )
    campaigns = query.order_by(Campaign.created_at.desc()).all()
    return [_decorate_campaign_processing_mode(db, campaign) for campaign in campaigns]


@router.post("", response_model=CampaignOut)
def create_campaign(
    payload: CampaignCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CampaignOut:
    raw_payload = payload.dict()
    raw_payload["processing_mode"] = _normalize_campaign_processing_mode(
        raw_payload.get("processing_mode")
    )
    campaign = Campaign(
        **raw_payload,
        status=CampaignStatus.DRAFT,
        tenant_id=current_user.tenant_id,
        owner_user_id=current_user.id,
    )
    db.add(campaign)
    db.commit()
    db.refresh(campaign)
    normalized_title = _normalized_campaign_title(campaign.id, campaign.title)
    if campaign.title != normalized_title:
        campaign.title = normalized_title
        db.commit()
        db.refresh(campaign)
    return _decorate_campaign_processing_mode(db, campaign)


def _invalidate_results_cache_for_campaign_ids(
    db: Session,
    *,
    tenant_id: int,
    campaign_ids: set[int],
) -> int:
    removed = 0
    for campaign_id in sorted({int(raw) for raw in campaign_ids if int(raw) > 0}):
        removed += int(
            invalidate_results_cache(
                db,
                tenant_id=tenant_id,
                campaign_id=campaign_id,
            )
            or 0
        )
    return removed


@router.get("/imported-files", response_model=list[ImportedEmailFileOut])
def list_imported_files(
    include_tech_sheets: bool = True,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[ImportedEmailFileOut]:
    tenant_id = current_user.tenant_id
    campaigns_query = db.query(Campaign).filter(Campaign.tenant_id == tenant_id)
    if not current_user.is_admin:
        campaigns_query = campaigns_query.filter(Campaign.owner_user_id == current_user.id)
    campaigns = campaigns_query.all()
    campaign_map = {int(row.id): row for row in campaigns}
    visible_campaign_ids = sorted(campaign_map.keys())
    if not visible_campaign_ids:
        return []

    invoice_item_counts: dict[int, int] = {}
    for invoice_id, count in (
        db.query(
            InvoiceExtractionItem.invoice_id,
            func.count(InvoiceExtractionItem.id),
        )
        .filter(
            InvoiceExtractionItem.tenant_id == tenant_id,
            InvoiceExtractionItem.campaign_id.in_(visible_campaign_ids),
        )
        .group_by(InvoiceExtractionItem.invoice_id)
        .all()
    ):
        try:
            invoice_item_counts[int(invoice_id)] = int(count or 0)
        except Exception:
            continue

    tech_sheet_link_counts: dict[int, int] = {}
    if include_tech_sheets:
        for tech_sheet_id, count in (
            db.query(
                TechSheetLink.tech_sheet_id,
                func.count(TechSheetLink.id),
            )
            .filter(
                TechSheetLink.tenant_id == tenant_id,
                TechSheetLink.campaign_id.in_(visible_campaign_ids),
            )
            .group_by(TechSheetLink.tech_sheet_id)
            .all()
        ):
            try:
                tech_sheet_link_counts[int(tech_sheet_id)] = int(count or 0)
            except Exception:
                continue

    rows: list[ImportedEmailFileOut] = []
    for invoice in (
        db.query(InboundInvoice)
        .filter(
            InboundInvoice.tenant_id == tenant_id,
            InboundInvoice.campaign_id.in_(visible_campaign_ids),
        )
        .order_by(InboundInvoice.received_at.desc(), InboundInvoice.id.desc())
        .all()
    ):
        campaign = campaign_map.get(int(invoice.campaign_id))
        rows.append(
            ImportedEmailFileOut(
                file_type="invoice",
                file_id=int(invoice.id),
                file_name=str(invoice.file_name or "").strip() or f"invoice-{invoice.id}",
                sender_email=str(invoice.sender_email or "").strip() or None,
                sender_name=str(invoice.sender_name or "").strip() or None,
                subject=str(invoice.subject or "").strip() or None,
                campaign_id=int(invoice.campaign_id),
                campaign_title=str(campaign.title or "").strip() if campaign else None,
                processing_mode=_campaign_effective_processing_mode(db, campaign) if campaign else None,
                received_at=invoice.received_at or invoice.created_at or datetime.utcnow(),
                linked_items_count=invoice_item_counts.get(int(invoice.id), 0),
                parse_engine=str(invoice.parse_engine or "").strip() or None,
                parse_score=float(invoice.parse_score) if invoice.parse_score is not None else None,
            )
        )

    if include_tech_sheets:
        for sheet in (
            db.query(TechSheet)
            .filter(
                TechSheet.tenant_id == tenant_id,
                TechSheet.campaign_id.in_(visible_campaign_ids),
            )
            .order_by(TechSheet.created_at.desc(), TechSheet.id.desc())
            .all()
        ):
            campaign = campaign_map.get(int(sheet.campaign_id))
            rows.append(
                ImportedEmailFileOut(
                    file_type="tech_sheet",
                    file_id=int(sheet.id),
                    file_name=str(sheet.file_name or "").strip() or f"tech-sheet-{sheet.id}",
                    sender_email=str(sheet.sender_email or "").strip() or None,
                    sender_name=None,
                    subject=None,
                    campaign_id=int(sheet.campaign_id),
                    campaign_title=str(campaign.title or "").strip() if campaign else None,
                    processing_mode=_campaign_effective_processing_mode(db, campaign) if campaign else None,
                    received_at=sheet.created_at or datetime.utcnow(),
                    linked_items_count=tech_sheet_link_counts.get(int(sheet.id), 0),
                    parse_engine=str(sheet.parse_engine or "").strip() or None,
                    parse_score=float(sheet.parse_score) if sheet.parse_score is not None else None,
                )
            )

    rows.sort(key=lambda row: (row.received_at, row.file_id), reverse=True)
    return rows


@router.post("/imported-files/reassign", response_model=ImportedEmailFileReassignResult)
def reassign_imported_files(
    payload: ImportedEmailFileReassignPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> ImportedEmailFileReassignResult:
    tenant_id = current_user.tenant_id
    target_campaign = _get_campaign(db, int(payload.campaign_id), current_user)
    target_campaign_id = int(target_campaign.id)
    visible_campaign_query = db.query(Campaign.id).filter(Campaign.tenant_id == tenant_id)
    if not current_user.is_admin:
        visible_campaign_query = visible_campaign_query.filter(
            Campaign.owner_user_id == current_user.id
        )
    visible_campaign_ids = {int(row.id) for row in visible_campaign_query.all()}
    visible_campaign_ids.add(target_campaign_id)

    invoice_ids = sorted({int(raw) for raw in (payload.invoice_ids or []) if int(raw) > 0})
    tech_sheet_ids = sorted({int(raw) for raw in (payload.tech_sheet_ids or []) if int(raw) > 0})
    if not invoice_ids and not tech_sheet_ids:
        raise HTTPException(status_code=400, detail="No imported files selected")

    campaign_ids_to_invalidate: set[int] = {target_campaign_id}
    cleared_links = 0
    cleared_offer_analysis_results = 0
    cleared_offer_analysis_cache = 0
    moved_invoices = 0
    moved_tech_sheets = 0
    moved_item_ids: list[int] = []
    moved_tech_sheet_ids: list[int] = []

    if invoice_ids:
        invoices = (
            db.query(InboundInvoice)
            .filter(
                InboundInvoice.tenant_id == tenant_id,
                InboundInvoice.id.in_(invoice_ids),
                InboundInvoice.campaign_id.in_(visible_campaign_ids),
            )
            .all()
        )
        found_invoice_ids = {int(row.id) for row in invoices}
        missing_invoice_ids = [row for row in invoice_ids if row not in found_invoice_ids]
        if missing_invoice_ids:
            raise HTTPException(
                status_code=404,
                detail=f"Imported invoice not found: {missing_invoice_ids[0]}",
            )
        for invoice in invoices:
            old_campaign_id = int(invoice.campaign_id or 0)
            campaign_ids_to_invalidate.add(old_campaign_id)
            if old_campaign_id != target_campaign_id:
                invoice.campaign_id = target_campaign_id
                moved_invoices += 1

        invoice_item_rows = (
            db.query(InvoiceExtractionItem)
            .filter(
                InvoiceExtractionItem.tenant_id == tenant_id,
                InvoiceExtractionItem.invoice_id.in_(invoice_ids),
            )
            .all()
        )
        moved_item_ids = [int(row.id) for row in invoice_item_rows]
        for item in invoice_item_rows:
            item.campaign_id = target_campaign_id

    if tech_sheet_ids:
        tech_sheet_rows = (
            db.query(TechSheet)
            .filter(
                TechSheet.tenant_id == tenant_id,
                TechSheet.id.in_(tech_sheet_ids),
                TechSheet.campaign_id.in_(visible_campaign_ids),
            )
            .all()
        )
        found_tech_sheet_ids = {int(row.id) for row in tech_sheet_rows}
        missing_tech_sheet_ids = [row for row in tech_sheet_ids if row not in found_tech_sheet_ids]
        if missing_tech_sheet_ids:
            raise HTTPException(
                status_code=404,
                detail=f"Imported tech sheet not found: {missing_tech_sheet_ids[0]}",
            )
        for sheet in tech_sheet_rows:
            old_campaign_id = int(sheet.campaign_id or 0)
            campaign_ids_to_invalidate.add(old_campaign_id)
            if old_campaign_id != target_campaign_id:
                sheet.campaign_id = target_campaign_id
                moved_tech_sheets += 1
        moved_tech_sheet_ids = [int(row.id) for row in tech_sheet_rows]

    if moved_item_ids:
        cleared_offer_analysis_results += int(
            db.query(OfferAnalysisResult)
            .filter(
                OfferAnalysisResult.tenant_id == tenant_id,
                OfferAnalysisResult.invoice_item_id.in_(moved_item_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        cleared_offer_analysis_cache += int(
            db.query(InvoiceItemOfferAnalysisCache)
            .filter(
                InvoiceItemOfferAnalysisCache.tenant_id == tenant_id,
                InvoiceItemOfferAnalysisCache.invoice_item_id.in_(moved_item_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        cleared_links += int(
            db.query(TechSheetLink)
            .filter(
                TechSheetLink.tenant_id == tenant_id,
                TechSheetLink.invoice_item_id.in_(moved_item_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )

    if moved_tech_sheet_ids:
        cleared_offer_analysis_results += int(
            db.query(OfferAnalysisResult)
            .filter(
                OfferAnalysisResult.tenant_id == tenant_id,
                OfferAnalysisResult.tech_sheet_id.in_(moved_tech_sheet_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        cleared_offer_analysis_cache += int(
            db.query(InvoiceItemOfferAnalysisCache)
            .filter(
                InvoiceItemOfferAnalysisCache.tenant_id == tenant_id,
                InvoiceItemOfferAnalysisCache.tech_sheet_id.in_(moved_tech_sheet_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        cleared_links += int(
            db.query(TechSheetLink)
            .filter(
                TechSheetLink.tenant_id == tenant_id,
                TechSheetLink.tech_sheet_id.in_(moved_tech_sheet_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )

    invalidated_results_cache = _invalidate_results_cache_for_campaign_ids(
        db,
        tenant_id=tenant_id,
        campaign_ids=campaign_ids_to_invalidate,
    )
    db.commit()

    return ImportedEmailFileReassignResult(
        moved_invoices=moved_invoices,
        moved_tech_sheets=moved_tech_sheets,
        cleared_links=cleared_links,
        cleared_offer_analysis_results=cleared_offer_analysis_results,
        cleared_offer_analysis_cache=cleared_offer_analysis_cache,
        invalidated_results_cache=invalidated_results_cache,
    )


@router.post("/imported-files/extract", response_model=ImportedEmailFileExtractResult)
def extract_imported_files(
    payload: ImportedEmailFileExtractPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> ImportedEmailFileExtractResult:
    tenant_id = current_user.tenant_id
    visible_campaign_query = db.query(Campaign.id).filter(Campaign.tenant_id == tenant_id)
    if not current_user.is_admin:
        visible_campaign_query = visible_campaign_query.filter(
            Campaign.owner_user_id == current_user.id
        )
    visible_campaign_ids = {int(row.id) for row in visible_campaign_query.all()}

    invoice_ids = sorted({int(raw) for raw in (payload.invoice_ids or []) if int(raw) > 0})
    tech_sheet_ids = sorted({int(raw) for raw in (payload.tech_sheet_ids or []) if int(raw) > 0})
    if not invoice_ids and not tech_sheet_ids:
        raise HTTPException(status_code=400, detail="No imported files selected")

    campaigns_to_invalidate: set[int] = set()
    extracted_invoices = 0
    extracted_invoice_items = 0
    extracted_tech_sheets = 0
    relinked_tech_sheet_links = 0
    cleared_links = 0
    cleared_offer_analysis_results = 0
    cleared_offer_analysis_cache = 0

    if invoice_ids:
        invoices = (
            db.query(InboundInvoice)
            .filter(
                InboundInvoice.tenant_id == tenant_id,
                InboundInvoice.id.in_(invoice_ids),
                InboundInvoice.campaign_id.in_(visible_campaign_ids),
            )
            .order_by(InboundInvoice.id.asc())
            .all()
        )
        found_invoice_ids = {int(row.id) for row in invoices}
        missing_invoice_ids = [row for row in invoice_ids if row not in found_invoice_ids]
        if missing_invoice_ids:
            raise HTTPException(
                status_code=404,
                detail=f"Imported invoice not found: {missing_invoice_ids[0]}",
            )
        for invoice in invoices:
            items = extract_imported_invoice(
                invoice_id=int(invoice.id),
                force=bool(payload.force),
                db=db,
                current_user=current_user,
            )
            extracted_invoices += 1
            extracted_invoice_items += len(items or [])
            campaigns_to_invalidate.add(int(invoice.campaign_id))

    tech_sheet_rows: list[TechSheet] = []
    if tech_sheet_ids:
        tech_sheet_rows = (
            db.query(TechSheet)
            .filter(
                TechSheet.tenant_id == tenant_id,
                TechSheet.id.in_(tech_sheet_ids),
                TechSheet.campaign_id.in_(visible_campaign_ids),
            )
            .order_by(TechSheet.id.asc())
            .all()
        )
        found_tech_sheet_ids = {int(row.id) for row in tech_sheet_rows}
        missing_tech_sheet_ids = [row for row in tech_sheet_ids if row not in found_tech_sheet_ids]
        if missing_tech_sheet_ids:
            raise HTTPException(
                status_code=404,
                detail=f"Imported tech sheet not found: {missing_tech_sheet_ids[0]}",
            )

        cleared_offer_analysis_results += int(
            db.query(OfferAnalysisResult)
            .filter(
                OfferAnalysisResult.tenant_id == tenant_id,
                OfferAnalysisResult.tech_sheet_id.in_(tech_sheet_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        cleared_offer_analysis_cache += int(
            db.query(InvoiceItemOfferAnalysisCache)
            .filter(
                InvoiceItemOfferAnalysisCache.tenant_id == tenant_id,
                InvoiceItemOfferAnalysisCache.tech_sheet_id.in_(tech_sheet_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        cleared_links += int(
            db.query(TechSheetLink)
            .filter(
                TechSheetLink.tenant_id == tenant_id,
                TechSheetLink.tech_sheet_id.in_(tech_sheet_ids),
            )
            .delete(synchronize_session=False)
            or 0
        )
        db.flush()

        relink_requests: set[tuple[int, str | None]] = set()
        for sheet in tech_sheet_rows:
            raw_text, characteristics, score, engine = extract_tech_sheet(str(sheet.file_path))
            sheet.raw_text = raw_text
            sheet.characteristics = characteristics
            sheet.parse_score = score
            sheet.parse_engine = engine
            extracted_tech_sheets += 1
            campaigns_to_invalidate.add(int(sheet.campaign_id))
            relink_requests.add(
                (int(sheet.campaign_id), str(sheet.sender_email or "").strip().lower() or None)
            )
        db.flush()
        for campaign_id, sender_email in sorted(relink_requests):
            try:
                links = link_tech_sheets(
                    TechSheetLinkPayload(
                        campaign_id=campaign_id,
                        sender_email=sender_email,
                        min_score=0.60,
                        max_text_chars=2000,
                    ),
                    db=db,
                    current_user=current_user,
                )
                relinked_tech_sheet_links += len(links or [])
            except Exception as exc:
                logger.warning(
                    "Imported tech-sheet relink failed: tenant=%s campaign=%s sender=%s error=%s",
                    tenant_id,
                    campaign_id,
                    sender_email,
                    exc,
                )

    invalidated_results_cache = _invalidate_results_cache_for_campaign_ids(
        db,
        tenant_id=tenant_id,
        campaign_ids=campaigns_to_invalidate,
    )
    db.commit()

    return ImportedEmailFileExtractResult(
        extracted_invoices=extracted_invoices,
        extracted_invoice_items=extracted_invoice_items,
        extracted_tech_sheets=extracted_tech_sheets,
        relinked_tech_sheet_links=relinked_tech_sheet_links,
        cleared_links=cleared_links,
        cleared_offer_analysis_results=cleared_offer_analysis_results,
        cleared_offer_analysis_cache=cleared_offer_analysis_cache,
        invalidated_results_cache=invalidated_results_cache,
    )


@router.get("/invoices/{invoice_id}/file")
def download_invoice_file(
    invoice_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FileResponse:
    invoice = (
        db.query(InboundInvoice)
        .filter(
            InboundInvoice.id == invoice_id,
            InboundInvoice.tenant_id == current_user.tenant_id,
        )
        .first()
    )
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    _get_campaign(db, invoice.campaign_id, current_user)
    return _build_file_response(Path(invoice.file_path), invoice.file_name)


@router.get("/tech-sheets/{tech_sheet_id}/file")
def download_tech_sheet_file(
    tech_sheet_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FileResponse:
    sheet = (
        db.query(TechSheet)
        .filter(
            TechSheet.id == tech_sheet_id,
            TechSheet.tenant_id == current_user.tenant_id,
        )
        .first()
    )
    if not sheet:
        raise HTTPException(status_code=404, detail="Tech sheet not found")
    _get_campaign(db, sheet.campaign_id, current_user)
    return _build_file_response(Path(sheet.file_path), sheet.file_name)


@router.get("/{campaign_id}", response_model=CampaignOut)
def get_campaign(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CampaignOut:
    return _get_campaign(db, campaign_id, current_user)


@router.get("/{campaign_id}/attachments", response_model=list[CampaignAttachmentOut])
def list_attachments(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignAttachmentOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    campaign_processing_mode = _campaign_effective_processing_mode(db, campaign)
    return (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.campaign_id == campaign_id,
            CampaignAttachment.tenant_id == campaign.tenant_id,
        )
        .order_by(CampaignAttachment.created_at.desc())
        .all()
    )


@router.get("/{campaign_id}/attachments/{attachment_id}/file")
def download_attachment_file(
    campaign_id: int,
    attachment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FileResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    attachment = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.id == attachment_id,
            CampaignAttachment.campaign_id == campaign_id,
            CampaignAttachment.tenant_id == campaign.tenant_id,
        )
        .first()
    )
    if not attachment:
        raise HTTPException(status_code=404, detail="Attachment not found")
    return _build_file_response(Path(attachment.file_path), attachment.file_name)


@router.post("/{campaign_id}/attachments/upload", response_model=CampaignAttachmentOut)
async def upload_attachment(
    campaign_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CampaignAttachmentOut:
    campaign = _get_campaign(db, campaign_id, current_user)
    file_name = (file.filename or "").strip()
    if not file_name:
        raise HTTPException(status_code=400, detail="Missing filename")
    if not file_name.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Only PDF files are allowed")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Empty file")

    folder = (
        config.CAMPAIGN_ATTACHMENTS_DIR
        / f"tenant_{campaign.tenant_id}"
        / f"campaign_{campaign.id}"
    )
    folder.mkdir(parents=True, exist_ok=True)

    safe_name = _safe_attachment_name(file_name)
    stored_name = f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{uuid4().hex[:8]}_{safe_name}"
    stored_path = folder / stored_name
    stored_path.write_bytes(content)

    row = CampaignAttachment(
        tenant_id=campaign.tenant_id,
        campaign_id=campaign.id,
        file_name=file_name,
        file_path=str(stored_path),
        file_size=len(content),
        parse_engine=None,
        parse_score=None,
        extraction_preview=None,
        strict_extraction_preview=None,
        extracted_at=None,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


@router.post("/{campaign_id}/attachments/reextract", response_model=list[CampaignAttachmentOut])
async def reextract_attachments(
    campaign_id: int,
    force: bool = False,
    payload: dict = Body(default={}),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignAttachmentOut]:
    def _parse_expected_count(raw: object) -> int | None:
        try:
            value = int(raw)  # type: ignore[arg-type]
        except Exception:
            return None
        return value if value > 0 else None

    def _parse_bool(raw: object, default: bool = False) -> bool:
        if isinstance(raw, bool):
            return raw
        if raw is None:
            return default
        text = str(raw).strip().lower()
        if text in {"1", "true", "yes", "on"}:
            return True
        if text in {"0", "false", "no", "off"}:
            return False
        return default

    def _parse_doc_type_hint(raw: object) -> str | None:
        value = str(raw or "").strip().lower()
        if value in {
            "product_list",
            "requirement_sheet",
            "catalog_list",
            "accessories_appendix",
            "strict_produit",
            "mixed",
        }:
            return value
        return None

    def _parse_layout_hint(raw: object) -> str | None:
        value = str(raw or "").strip().lower()
        if value in {"table", "paragraph", "mixed"}:
            return value
        return None

    def _parse_strict_count_mode(raw: object, *, default: str) -> str:
        value = str(raw or "").strip().lower()
        if value in {"auto", "strict", "soft", "off"}:
            return value
        return default

    def _preview_products_count(preview: dict[str, Any] | None) -> int:
        if not isinstance(preview, dict):
            return 0
        products = preview.get("products")
        if isinstance(products, list):
            return len([p for p in products if isinstance(p, dict)])
        items = preview.get("items")
        if isinstance(items, list):
            return len(items)
        return 0

    def _preview_unassigned_count(preview: dict[str, Any] | None) -> int:
        if not isinstance(preview, dict):
            return 0
        rows = preview.get("unassigned")
        return len(rows) if isinstance(rows, list) else 0

    def _preview_placeholder_count(preview: dict[str, Any] | None) -> int:
        if not isinstance(preview, dict):
            return 0
        try:
            return int(preview.get("expected_count_placeholders") or 0)
        except Exception:
            return 0

    def _preview_avg_quality(preview: dict[str, Any] | None) -> float:
        if not isinstance(preview, dict):
            return 0.0
        products = preview.get("products")
        if not isinstance(products, list):
            return 0.0
        scores: list[float] = []
        for row in products:
            if not isinstance(row, dict):
                continue
            try:
                score = float(row.get("quality_score") or 0.0)
            except Exception:
                score = 0.0
            if score > 0:
                scores.append(score)
        if not scores:
            return 0.0
        return max(0.0, min(1.0, sum(scores) / len(scores)))

    def _preview_rank(preview: dict[str, Any] | None, expected: int | None) -> float:
        count = _preview_products_count(preview)
        placeholders = _preview_placeholder_count(preview)
        real_count = max(0, count - placeholders)
        if expected and expected > 0:
            coverage = min(1.0, real_count / float(expected))
        else:
            coverage = min(1.0, real_count / 20.0)
        quality = _preview_avg_quality(preview)
        unassigned = _preview_unassigned_count(preview)
        penalty = min(1.0, unassigned / float(max(1, real_count + 1)))
        return (coverage * 0.65) + (quality * 0.25) + ((1.0 - penalty) * 0.10)

    def _needs_fallback_preview(preview: dict[str, Any] | None, expected: int | None) -> bool:
        count = _preview_products_count(preview)
        placeholders = _preview_placeholder_count(preview)
        real_count = max(0, count - placeholders)
        if real_count <= 0:
            return True
        if placeholders > 0:
            return True
        if expected and expected > 0:
            min_target = max(1, int(round(expected * 0.60)))
            if real_count < min_target:
                return True
        unassigned = _preview_unassigned_count(preview)
        if unassigned > max(12, real_count * 2):
            return True
        return False

    campaign = _get_campaign(db, campaign_id, current_user)
    explicit_processing_mode = _normalize_campaign_processing_mode(campaign.processing_mode)
    expected_default: int | None = None
    expected_map: dict = {}
    strict_expected_count = True
    strict_count_mode = "strict"
    document_type_hint_default: str | None = None
    layout_hint_default: str | None = None
    document_type_hints_map: dict = {}
    layout_hints_map: dict = {}
    require_specs_for_acceptance = False
    requested_attachment_ids: set[int] = set()
    async_mode = False
    if isinstance(payload, dict):
        strict_expected_count = _parse_bool(payload.get("strict_expected_count", True), True)
        strict_count_mode = _parse_strict_count_mode(
            payload.get("strict_count_mode"),
            default=("strict" if strict_expected_count else "soft"),
        )
        require_specs_for_acceptance = _parse_bool(
            payload.get("require_specs_for_acceptance", False),
            False,
        )
        expected_default = _parse_expected_count(payload.get("expected_product_count"))
        raw_map = payload.get("expected_product_counts")
        if isinstance(raw_map, dict):
            expected_map = raw_map
        document_type_hint_default = _parse_doc_type_hint(payload.get("document_type_hint"))
        layout_hint_default = _parse_layout_hint(payload.get("layout_hint"))
        raw_doc_hints = payload.get("document_type_hints")
        if isinstance(raw_doc_hints, dict):
            document_type_hints_map = raw_doc_hints
        raw_layout_hints = payload.get("layout_hints")
        if isinstance(raw_layout_hints, dict):
            layout_hints_map = raw_layout_hints
        raw_attachment_ids = payload.get("attachment_ids")
        if isinstance(raw_attachment_ids, list):
            for raw_id in raw_attachment_ids:
                try:
                    parsed_id = int(raw_id)
                except Exception:
                    continue
                if parsed_id > 0:
                    requested_attachment_ids.add(parsed_id)
        async_mode = _parse_bool(payload.get("async_mode", False), False)

    strict_hint_requested = bool(document_type_hint_default == "strict_produit") or any(
        _parse_doc_type_hint(v) == "strict_produit" for v in document_type_hints_map.values()
    )
    campaign_processing_mode = explicit_processing_mode
    if campaign_processing_mode is None and strict_hint_requested:
        campaign_processing_mode = _CAMPAIGN_PROCESSING_STRICT
        campaign.processing_mode = _CAMPAIGN_PROCESSING_STRICT
        db.flush()
    if campaign_processing_mode is None:
        campaign_processing_mode = _campaign_effective_processing_mode(db, campaign)

    logger.info(
        "Attachment OCR requested: tenant=%s campaign=%s force=%s",
        campaign.tenant_id,
        campaign.id,
        force,
    )
    rows = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.campaign_id == campaign.id,
            CampaignAttachment.tenant_id == campaign.tenant_id,
        )
        .order_by(CampaignAttachment.created_at.desc())
        .all()
    )
    if not rows:
        raise HTTPException(status_code=400, detail="No attachments found for this campaign")

    pending_rows = [
        row
        for row in rows
        if row.extracted_at is None or row.extraction_preview is None
    ]
    target_rows = rows if force else pending_rows

    if requested_attachment_ids:
        target_rows = [row for row in target_rows if int(row.id) in requested_attachment_ids]

    if not target_rows:
        raise HTTPException(
            status_code=409,
            detail=(
                "No matching attachments to analyze"
                if requested_attachment_ids
                else "OCR already up to date for current attachments"
            ),
        )

    if async_mode:
        for row in target_rows:
            row.parse_engine = "processing"
            row.parse_score = 0.0
            row.extraction_preview = None
            row.strict_extraction_preview = None
            row.extracted_at = None
        db.commit()
        thread = threading.Thread(
            target=_run_attachment_reextract_job,
            kwargs={
                "campaign_id": int(campaign.id),
                "user_id": int(current_user.id),
                "force": force,
                "expected_default": expected_default,
                "expected_map": dict(expected_map),
                "strict_expected_count": bool(strict_expected_count),
                "strict_count_mode": str(strict_count_mode),
                "document_type_hint_default": document_type_hint_default,
                "layout_hint_default": layout_hint_default,
                "document_type_hints_map": dict(document_type_hints_map),
                "layout_hints_map": dict(layout_hints_map),
                "require_specs_for_acceptance": bool(require_specs_for_acceptance),
                "requested_attachment_ids": set(requested_attachment_ids),
            },
            daemon=True,
        )
        thread.start()
        for row in rows:
            db.refresh(row)
        return rows

    logger.info(
        "Attachment OCR run: tenant=%s campaign=%s pending=%s total=%s",
        campaign.tenant_id,
        campaign.id,
        len(target_rows),
        len(rows),
    )

    thresholds = get_user_embedding_thresholds(db, current_user.tenant_id, current_user.id)
    success_count = 0
    error_count = 0
    first_error: str | None = None
    for row in target_rows:
        try:
            expected_for_row = expected_default
            document_type_hint_for_row = document_type_hint_default
            layout_hint_for_row = layout_hint_default
            for key in (
                row.id,
                str(row.id),
                row.file_name,
                Path(row.file_name).name if row.file_name else "",
                Path(row.file_path).name if row.file_path else "",
            ):
                if not key:
                    continue
                if key in expected_map:
                    parsed_expected = _parse_expected_count(expected_map.get(key))
                    if parsed_expected is not None:
                        expected_for_row = parsed_expected
                        break
            for key in (
                row.id,
                str(row.id),
                row.file_name,
                Path(row.file_name).name if row.file_name else "",
                Path(row.file_path).name if row.file_path else "",
            ):
                if not key:
                    continue
                if key in document_type_hints_map:
                    parsed_doc_hint = _parse_doc_type_hint(document_type_hints_map.get(key))
                    if parsed_doc_hint is not None:
                        document_type_hint_for_row = parsed_doc_hint
                        break
            for key in (
                row.id,
                str(row.id),
                row.file_name,
                Path(row.file_name).name if row.file_name else "",
                Path(row.file_path).name if row.file_path else "",
            ):
                if not key:
                    continue
                if key in layout_hints_map:
                    parsed_layout_hint = _parse_layout_hint(layout_hints_map.get(key))
                    if parsed_layout_hint is not None:
                        layout_hint_for_row = parsed_layout_hint
                        break

            preview = await asyncio.to_thread(
                extract_attachment_products_preview,
                row.file_path,
                thresholds.extraction_threshold,
                expected_for_row,
                strict_expected_count if expected_for_row is not None else False,
                document_type_hint_for_row,
                layout_hint_for_row,
                strict_count_mode,
                require_specs_for_acceptance,
            )
            strict_preview = None
            if isinstance(preview, dict) and "strict_payload" in preview:
                strict_preview = preview.pop("strict_payload")

            # Robust fallback: if the first extraction is weak, retry with
            # alternative profiles and keep the best one.
            if _needs_fallback_preview(preview, expected_for_row):
                candidates: list[tuple[str | None, str | None]] = [
                    ("mixed", "mixed"),
                    ("catalog_list", "mixed"),
                    ("catalog_list", "table"),
                    ("requirement_sheet", "mixed"),
                    ("accessories_appendix", "mixed"),
                ]
                seen_profiles: set[tuple[str | None, str | None]] = {
                    (document_type_hint_for_row, layout_hint_for_row)
                }
                best_preview = preview
                best_rank = _preview_rank(preview, expected_for_row)
                for cand_doc_hint, cand_layout_hint in candidates:
                    profile = (cand_doc_hint, cand_layout_hint)
                    if profile in seen_profiles:
                        continue
                    seen_profiles.add(profile)
                    try:
                        candidate_preview = await asyncio.to_thread(
                            extract_attachment_products_preview,
                            row.file_path,
                            thresholds.extraction_threshold,
                            expected_for_row,
                            strict_expected_count if expected_for_row is not None else False,
                            cand_doc_hint,
                            cand_layout_hint,
                            strict_count_mode,
                            require_specs_for_acceptance,
                        )
                    except Exception:
                        continue
                    candidate_rank = _preview_rank(candidate_preview, expected_for_row)
                    if candidate_rank > best_rank + 0.03:
                        best_preview = candidate_preview
                        best_rank = candidate_rank
                preview = best_preview

            row.parse_engine = str(preview.get("engine") or "document_ai_ocr")
            row.parse_score = float(preview.get("avg_regex_score") or 0.0)
            if campaign_processing_mode == _CAMPAIGN_PROCESSING_STRICT and strict_preview is not None:
                row.strict_extraction_preview = strict_preview
            elif campaign_processing_mode != _CAMPAIGN_PROCESSING_STRICT:
                row.strict_extraction_preview = None
            row.extraction_preview = preview
            row.extracted_at = datetime.utcnow()
            success_count += 1
        except Exception as exc:
            error_count += 1
            if first_error is None:
                first_error = str(exc)
            logger.warning("Attachment re-extraction failed for %s: %s", row.file_path, exc)

    if success_count == 0:
        detail = "OCR failed for all selected attachments. Check Document AI credentials and server logs."
        if first_error:
            detail = f"{detail} First error: {first_error[:300]}"
        raise HTTPException(
            status_code=502,
            detail=detail,
        )

    if error_count > 0:
        logger.warning(
            "Attachment OCR completed with partial failures: tenant=%s campaign=%s success=%s errors=%s",
            campaign.tenant_id,
            campaign.id,
            success_count,
            error_count,
        )

    # Auto-fill missing requested_qty from freshly extracted attachment products.
    # No UI action required; never overwrite quantities already set by user.
    db.flush()
    auto_qty_updates = _autofill_campaign_item_quantities_from_attachments(
        db,
        tenant_id=campaign.tenant_id,
        campaign_id=campaign.id,
    )
    if auto_qty_updates > 0:
        logger.info(
            "Auto-filled requested_qty from attachment extraction: tenant=%s campaign=%s updated=%s",
            campaign.tenant_id,
            campaign.id,
            auto_qty_updates,
        )
    db.commit()
    for row in rows:
        db.refresh(row)
    return rows


@router.post("/{campaign_id}/attachments/{attachment_id}/learn-linking")
def learn_attachment_linking(
    campaign_id: int,
    attachment_id: int,
    payload: dict = Body(default={}),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    row = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.id == attachment_id,
            CampaignAttachment.campaign_id == campaign.id,
            CampaignAttachment.tenant_id == campaign.tenant_id,
        )
        .first()
    )
    if not row:
        raise HTTPException(status_code=404, detail="Attachment not found")

    products: list[dict] = []
    source = "manual_correction"
    if isinstance(payload, dict):
        source = str(payload.get("source") or "manual_correction").strip()[:80] or "manual_correction"
        raw_products = payload.get("products")
        if isinstance(raw_products, list):
            products = [item for item in raw_products if isinstance(item, dict)]

        manual_preview = payload.get("extraction_preview")
        if isinstance(manual_preview, dict):
            row.extraction_preview = manual_preview
            row.extracted_at = datetime.utcnow()
            if not products:
                raw_preview_products = manual_preview.get("products")
                if isinstance(raw_preview_products, list):
                    products = [item for item in raw_preview_products if isinstance(item, dict)]

    if not products:
        preview = row.extraction_preview if isinstance(row.extraction_preview, dict) else {}
        raw_preview_products = preview.get("products")
        if isinstance(raw_preview_products, list):
            products = [item for item in raw_preview_products if isinstance(item, dict)]

    if not products:
        raise HTTPException(
            status_code=400,
            detail="No products provided for learning. Send payload.products or extraction_preview.products.",
        )

    result = learn_anchor_linking_from_products(
        products,
        campaign_id=campaign.id,
        attachment_id=row.id,
        source=source,
    )
    db.commit()
    return {"status": "ok", **result}


@router.delete("/{campaign_id}/attachments/{attachment_id}")
def delete_attachment(
    campaign_id: int,
    attachment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    row = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.id == attachment_id,
            CampaignAttachment.campaign_id == campaign.id,
            CampaignAttachment.tenant_id == campaign.tenant_id,
        )
        .first()
    )
    if not row:
        raise HTTPException(status_code=404, detail="Attachment not found")

    path = Path(row.file_path)
    db.delete(row)
    db.commit()
    try:
        if path.exists():
            path.unlink()
    except Exception:
        pass
    return {"status": "ok"}


@router.put("/{campaign_id}", response_model=CampaignOut)
def update_campaign(
    campaign_id: int,
    payload: CampaignUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CampaignOut:
    campaign = _get_campaign(db, campaign_id, current_user)
    old_priority = campaign.priority
    updates = payload.dict(exclude_unset=True)
    if "processing_mode" in updates:
        updates["processing_mode"] = _normalize_campaign_processing_mode(updates.get("processing_mode"))
    for key, value in updates.items():
        setattr(campaign, key, value)
    if "title" in updates:
        campaign.title = _normalized_campaign_title(campaign.id, campaign.title)
    db.commit()
    if payload.priority is not None and payload.priority != old_priority:
        db.query(EmailJob).filter(
            EmailJob.campaign_id == campaign_id,
            EmailJob.tenant_id == campaign.tenant_id,
            EmailJob.status == EmailJobStatus.QUEUED,
        ).update({EmailJob.priority: campaign.priority})
        db.commit()
    db.refresh(campaign)
    return _decorate_campaign_processing_mode(db, campaign)


@router.post("/{campaign_id}/close", response_model=CampaignOut)
def close_campaign(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CampaignOut:
    campaign = _get_campaign(db, campaign_id, current_user)
    if campaign.status == CampaignStatus.SENDING:
        close_marker = f"Campaign closed manually at {datetime.utcnow().isoformat(timespec='seconds')}"
        (
            db.query(EmailJob)
            .filter(
                EmailJob.campaign_id == campaign_id,
                EmailJob.tenant_id == campaign.tenant_id,
                EmailJob.status.in_([EmailJobStatus.QUEUED, EmailJobStatus.SENDING]),
            )
            .update(
                {
                    EmailJob.status: EmailJobStatus.FAILED,
                    EmailJob.last_error: close_marker,
                },
                synchronize_session=False,
            )
        )
        (
            db.query(CampaignRecipient)
            .filter(
                CampaignRecipient.campaign_id == campaign_id,
                CampaignRecipient.tenant_id == campaign.tenant_id,
                CampaignRecipient.delivery_status == DeliveryStatus.PENDING,
            )
            .update(
                {
                    CampaignRecipient.delivery_status: DeliveryStatus.FAILED,
                    CampaignRecipient.last_error: close_marker,
                },
                synchronize_session=False,
            )
        )
    if campaign.status != CampaignStatus.CLOSED:
        campaign.status = CampaignStatus.CLOSED
        if not campaign.sent_at:
            campaign.sent_at = datetime.utcnow()
        db.commit()
        db.refresh(campaign)
    return campaign


@router.delete("/{campaign_id}")
def delete_campaign(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    db.delete(campaign)
    db.commit()
    return {"status": "ok"}


@router.post(
    "/characteristics/parse-text",
    response_model=ParseCharacteristicsFromTextOut,
)
def parse_characteristics_from_text(
    payload: ParseCharacteristicsFromTextIn,
    current_user: User = Depends(get_current_user),
) -> ParseCharacteristicsFromTextOut:
    text = (payload.text or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="text is required")

    parsed = parse_characteristics(text)
    kv = parsed.get("kv") if isinstance(parsed, dict) else {}
    normalized = _normalize_requested_characteristics_payload(
        kv if isinstance(kv, dict) else {}
    )
    pairs = parsed.get("pairs") if isinstance(parsed, dict) else []
    return ParseCharacteristicsFromTextOut(
        requested_characteristics=normalized,
        score=round(float(parsed.get("score") or 0.0), 4)
        if isinstance(parsed, dict)
        else 0.0,
        score_regex_only=round(float(parsed.get("score_regex_only") or 0.0), 4)
        if isinstance(parsed, dict)
        else 0.0,
        pairs_count=len(pairs) if isinstance(pairs, list) else 0,
    )


@router.post("/{campaign_id}/items")
def replace_items(
    campaign_id: int,
    payload: CampaignItemsPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    product_ids = {item.product_id for item in payload.items}
    _ensure_product_ids(db, product_ids, tenant_id)

    existing_items = (
        db.query(CampaignItem)
        .filter(
            CampaignItem.campaign_id == campaign_id,
            CampaignItem.tenant_id == tenant_id,
        )
        .all()
    )
    existing_qty_by_product_id: dict[int, float] = {}
    existing_qty_source_by_product_id: dict[int, str | None] = {}
    existing_qty_confidence_by_product_id: dict[int, float | None] = {}
    for row in existing_items:
        qty = _to_positive_float(row.requested_qty)
        source = _normalize_qty_source(getattr(row, "requested_qty_source", None))
        confidence = _to_positive_float(getattr(row, "requested_qty_confidence", None))
        if source is not None:
            existing_qty_source_by_product_id[int(row.product_id)] = source
        if confidence is not None:
            existing_qty_confidence_by_product_id[int(row.product_id)] = confidence
        if qty is not None:
            existing_qty_by_product_id[int(row.product_id)] = qty

    product_rows = (
        db.query(Product)
        .filter(
            Product.id.in_(product_ids),
            Product.tenant_id == tenant_id,
        )
        .all()
    )
    product_name_by_id = {int(p.id): str(p.name or "") for p in product_rows}
    qty_candidates = _collect_attachment_qty_candidates(
        db,
        tenant_id=tenant_id,
        campaign_id=campaign_id,
    )
    gpt_budget: dict[str, int] = {"calls": 0}

    db.query(CampaignItem).filter(
        CampaignItem.campaign_id == campaign_id,
        CampaignItem.tenant_id == tenant_id,
    ).delete()
    db.flush()

    for item in payload.items:
        product_id = int(item.product_id)
        old_qty = existing_qty_by_product_id.get(product_id)
        old_source = existing_qty_source_by_product_id.get(product_id)
        old_confidence = existing_qty_confidence_by_product_id.get(product_id)
        requested_qty, requested_qty_source, requested_qty_confidence = _resolve_requested_qty_for_item(
            item.requested_qty,
            existing_qty=old_qty,
            existing_source=old_source,
            existing_confidence=old_confidence,
            product_name=product_name_by_id.get(product_id),
            qty_candidates=qty_candidates,
            gpt_budget=gpt_budget,
        )
        if _to_positive_float(item.requested_qty) is not None:
            requested_qty_source, requested_qty_confidence = _refresh_qty_source_on_manual_edit(
                old_source,
                old_qty,
                requested_qty,
            )
        db.add(
            CampaignItem(
                campaign_id=campaign_id,
                product_id=item.product_id,
                requested_qty=requested_qty,
                requested_qty_source=requested_qty_source,
                requested_qty_confidence=requested_qty_confidence,
                requested_description=item.requested_description,
                requested_characteristics=item.requested_characteristics,
                requested_embedding_score=item.requested_embedding_score,
                tenant_id=tenant_id,
            )
        )

    db.commit()
    try:
        if _campaign_effective_processing_mode(db, campaign) == _CAMPAIGN_PROCESSING_STRICT:
            warm_results_cache_for_campaign(
                campaign_id=campaign_id,
                db=db,
                current_user=current_user,
            )
    except Exception as exc:
        logger.warning(
            "Strict results warm after replace_items failed: tenant=%s user=%s campaign=%s err=%s",
            current_user.tenant_id,
            current_user.id,
            campaign_id,
            exc,
        )
    return {"status": "ok"}


@router.get("/{campaign_id}/items", response_model=list[CampaignItemOut])
def list_items(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignItemOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    items = (
        db.query(CampaignItem, Product)
        .join(Product, CampaignItem.product_id == Product.id)
        .filter(
            CampaignItem.campaign_id == campaign_id,
            CampaignItem.tenant_id == tenant_id,
            Product.tenant_id == tenant_id,
        )
        .all()
    )
    return [
        {
            "id": item.id,
            "product_id": item.product_id,
            "requested_qty": item.requested_qty,
            "requested_qty_source": item.requested_qty_source,
            "requested_qty_confidence": item.requested_qty_confidence,
            "product_name": product.name,
            "requested_description": item.requested_description,
            "requested_characteristics": item.requested_characteristics,
            "requested_embedding_score": item.requested_embedding_score,
            "product_description": product.description,
            "product_characteristics": product.characteristics,
            "product_embedding_score": product.embedding_score,
        }
        for item, product in items
    ]


@router.post("/{campaign_id}/recipients", response_model=list[CampaignRecipientOut])
def create_recipients(
    campaign_id: int,
    payload: CampaignRecipientPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignRecipientOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id

    supplier_ids = _ensure_supplier_ids(db, payload.supplier_ids, tenant_id)
    group_ids = _ensure_group_ids(db, payload.group_ids, tenant_id)
    if group_ids:
        group_members = (
            db.query(GroupMember)
            .filter(GroupMember.group_id.in_(group_ids), GroupMember.tenant_id == tenant_id)
            .all()
        )
        supplier_ids.update(member.supplier_id for member in group_members)
    supplier_ids = _dedupe_supplier_ids_by_email(db, supplier_ids, tenant_id)

    created = []
    for supplier_id in supplier_ids:
        existing = (
            db.query(CampaignRecipient)
            .filter(
                CampaignRecipient.campaign_id == campaign_id,
                CampaignRecipient.supplier_id == supplier_id,
                CampaignRecipient.tenant_id == tenant_id,
            )
            .first()
        )
        if existing:
            created.append(existing)
            continue
        recipient = CampaignRecipient(
            campaign_id=campaign_id,
            supplier_id=supplier_id,
            token=create_token(),
            token_expires_at=token_expiry(),
            delivery_status=DeliveryStatus.PENDING,
            tenant_id=tenant_id,
        )
        db.add(recipient)
        created.append(recipient)

    db.commit()
    return created


@router.put("/{campaign_id}/recipients", response_model=list[CampaignRecipientOut])
def replace_recipients(
    campaign_id: int,
    payload: CampaignRecipientPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignRecipientOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id

    supplier_ids = _ensure_supplier_ids(db, payload.supplier_ids, tenant_id)
    group_ids = _ensure_group_ids(db, payload.group_ids, tenant_id)
    if group_ids:
        group_members = (
            db.query(GroupMember)
            .filter(GroupMember.group_id.in_(group_ids), GroupMember.tenant_id == tenant_id)
            .all()
        )
        supplier_ids.update(member.supplier_id for member in group_members)
    supplier_ids = _dedupe_supplier_ids_by_email(db, supplier_ids, tenant_id)

    existing = (
        db.query(CampaignRecipient)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
        )
        .all()
    )
    existing_by_supplier = {recipient.supplier_id: recipient for recipient in existing}

    for recipient in existing:
        if recipient.supplier_id not in supplier_ids:
            db.delete(recipient)

    for supplier_id in supplier_ids:
        if supplier_id in existing_by_supplier:
            continue
        recipient = CampaignRecipient(
            campaign_id=campaign_id,
            supplier_id=supplier_id,
            token=create_token(),
            token_expires_at=token_expiry(),
            delivery_status=DeliveryStatus.PENDING,
            tenant_id=tenant_id,
        )
        db.add(recipient)

    db.commit()
    return (
        db.query(CampaignRecipient)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
        )
        .all()
    )


@router.get("/{campaign_id}/recipients", response_model=list[CampaignRecipientOut])
def list_recipients(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CampaignRecipientOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    return (
        db.query(CampaignRecipient)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
        )
        .all()
    )


@router.post("/{campaign_id}/send")
def send_campaign(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    if campaign.status == CampaignStatus.CLOSED:
        raise HTTPException(status_code=400, detail="Campaign is closed")
    tenant_id = campaign.tenant_id

    settings = get_settings(db, tenant_id)
    if not settings.from_email:
        raise HTTPException(status_code=400, detail="from_email is not configured")

    items = (
        db.query(CampaignItem)
        .filter(CampaignItem.campaign_id == campaign_id, CampaignItem.tenant_id == tenant_id)
        .all()
    )
    recipients = (
        db.query(CampaignRecipient)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
        )
        .all()
    )

    campaign.status = CampaignStatus.SENDING
    db.commit()

    queued = 0
    skipped = 0
    existing_jobs = {
        row[0]
        for row in db.query(EmailJob.recipient_id)
        .filter(
            EmailJob.campaign_id == campaign_id,
            EmailJob.tenant_id == tenant_id,
            EmailJob.status.in_([EmailJobStatus.QUEUED, EmailJobStatus.SENDING]),
        )
        .all()
    }
    suppressed = {
        row.email.lower()
        for row in db.query(SuppressedEmail)
        .filter(SuppressedEmail.tenant_id == tenant_id)
        .all()
    }
    for recipient in recipients:
        if recipient.delivery_status == DeliveryStatus.SENT:
            continue
        if recipient.id in existing_jobs:
            skipped += 1
            continue

        supplier = (
            db.query(Supplier)
            .filter(
                Supplier.id == recipient.supplier_id,
                Supplier.tenant_id == tenant_id,
            )
            .first()
        )
        if not supplier or not supplier.email:
            recipient.delivery_status = DeliveryStatus.FAILED
            recipient.last_error = "Missing supplier email"
            skipped += 1
            continue
        if supplier.email.lower() in suppressed:
            recipient.delivery_status = DeliveryStatus.FAILED
            recipient.last_error = "Suppressed email"
            skipped += 1
            continue

        reply_url = f"{config.APP_URL}/r/{recipient.token}"
        html_body = render_campaign_email(campaign, supplier, items, reply_url)
        subject = render_campaign_subject(campaign.subject, campaign, supplier, reply_url)

        existing_job = (
            db.query(EmailJob)
            .filter(
                EmailJob.recipient_id == recipient.id,
                EmailJob.tenant_id == tenant_id,
            )
            .order_by(EmailJob.created_at.desc())
            .first()
        )
        if existing_job and existing_job.status == EmailJobStatus.FAILED:
            existing_job.status = EmailJobStatus.QUEUED
            existing_job.attempts = 0
            existing_job.last_error = None
            existing_job.next_attempt_at = datetime.utcnow()
            existing_job.subject = subject
            existing_job.html_body = html_body
            existing_job.priority = campaign.priority
        else:
            db.add(
                EmailJob(
                    tenant_id=tenant_id,
                    campaign_id=campaign_id,
                    recipient_id=recipient.id,
                    supplier_id=supplier.id,
                    to_email=supplier.email,
                    subject=subject,
                    html_body=html_body,
                    priority=campaign.priority,
                    status=EmailJobStatus.QUEUED,
                    attempts=0,
                    max_attempts=settings.retry_max_attempts,
                    next_attempt_at=datetime.utcnow(),
                )
            )
        recipient.delivery_status = DeliveryStatus.PENDING
        recipient.last_error = None
        recipient.sent_at = None
        queued += 1

    db.commit()

    return {"queued": queued, "skipped": skipped}


@router.get("/{campaign_id}/stats", response_model=CampaignStats)
def campaign_stats(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CampaignStats:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    recipients = (
        db.query(CampaignRecipient)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
        )
        .all()
    )
    total = len(recipients)
    sent_count = len([r for r in recipients if r.delivery_status == DeliveryStatus.SENT])
    failed_count = len([r for r in recipients if r.delivery_status == DeliveryStatus.FAILED])

    recipient_supplier_ids = {int(r.supplier_id) for r in recipients if r.supplier_id is not None}
    recipients_with_email = (
        db.query(CampaignRecipient.supplier_id, Supplier.email)
        .join(Supplier, CampaignRecipient.supplier_id == Supplier.id)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
            Supplier.tenant_id == tenant_id,
        )
        .all()
    )

    def _norm_email(value: str | None) -> str:
        return str(value or "").strip().lower()

    supplier_id_by_email: dict[str, int] = {}
    for supplier_id, supplier_email in recipients_with_email:
        sid = int(supplier_id or 0)
        if sid <= 0:
            continue
        email_key = _norm_email(supplier_email)
        if email_key and email_key not in supplier_id_by_email:
            supplier_id_by_email[email_key] = sid

    web_response_supplier_ids = {
        int(row[0])
        for row in (
            db.query(ResponseModel.supplier_id)
            .filter(
                ResponseModel.campaign_id == campaign_id,
                ResponseModel.tenant_id == tenant_id,
            )
            .distinct()
            .all()
        )
        if row and row[0] is not None and int(row[0]) > 0
    }

    email_response_supplier_ids: set[int] = set()
    invoice_sender_rows = (
        db.query(InboundInvoice.sender_email)
        .filter(
            InboundInvoice.campaign_id == campaign_id,
            InboundInvoice.tenant_id == tenant_id,
        )
        .all()
    )
    for (sender_email,) in invoice_sender_rows:
        sid = supplier_id_by_email.get(_norm_email(sender_email))
        if sid:
            email_response_supplier_ids.add(sid)

    extracted_line_supplier_ids = {
        int(row[0])
        for row in (
            db.query(InvoiceExtractionItem.supplier_id)
            .filter(
                InvoiceExtractionItem.campaign_id == campaign_id,
                InvoiceExtractionItem.tenant_id == tenant_id,
                InvoiceExtractionItem.supplier_id.isnot(None),
            )
            .distinct()
            .all()
        )
        if row and row[0] is not None and int(row[0]) > 0
    }

    responded_supplier_ids = (
        web_response_supplier_ids
        | email_response_supplier_ids
        | extracted_line_supplier_ids
    ) & recipient_supplier_ids
    responses_count = len(responded_supplier_ids)
    pending_responses = max(total - responses_count, 0)
    response_rate = (responses_count / total) if total else 0.0
    costs = _campaign_openai_cost_breakdown(
        db,
        campaign=campaign,
    )

    return CampaignStats(
        total_recipients=total,
        sent_count=sent_count,
        failed_count=failed_count,
        responses_count=responses_count,
        pending_responses_count=pending_responses,
        response_rate=response_rate,
        costs=costs,
    )


@router.get("/{campaign_id}/responses", response_model=list[ResponseOut])
def campaign_responses(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[ResponseOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    return (
        db.query(ResponseModel)
        .filter(ResponseModel.campaign_id == campaign_id, ResponseModel.tenant_id == tenant_id)
        .all()
    )


@router.get("/{campaign_id}/invoices", response_model=list[InboundInvoiceOut])
def campaign_invoices(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[InboundInvoiceOut]:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    return (
        db.query(InboundInvoice)
        .filter(
            InboundInvoice.campaign_id == campaign_id,
            InboundInvoice.tenant_id == tenant_id,
        )
        .order_by(InboundInvoice.received_at.desc())
        .all()
    )


@router.get("/{campaign_id}/preview")
def preview_email(
    campaign_id: int,
    supplier_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id

    supplier = None
    if supplier_id:
        supplier = (
            db.query(Supplier)
            .filter(
                Supplier.id == supplier_id,
                Supplier.tenant_id == tenant_id,
            )
            .first()
        )
    if not supplier:
        supplier = Supplier(id=0, name="Sample", email="sample@example.com", supplier_type=SupplierType.MEDICAMENT)

    items = (
        db.query(CampaignItem)
        .filter(CampaignItem.campaign_id == campaign_id, CampaignItem.tenant_id == tenant_id)
        .all()
    )
    reply_url = f"{config.APP_URL}/r/preview"
    html_body = render_campaign_email(campaign, supplier, items, reply_url)
    return {"html": html_body}


def _market_estimate_offer_from_record(
    record: CampaignMarketEstimateOffer,
) -> dict[str, Any]:
    payload = dict(record.payload or {})
    payload["market_estimate"] = True
    payload["market_estimate_id"] = int(record.id)
    payload["market_estimate_region"] = str(
        payload.get("market_estimate_region") or "DZ"
    ).strip() or "DZ"
    payload["market_estimate_model"] = str(
        payload.get("market_estimate_model") or record.model_name or "gpt-5.1"
    ).strip() or "gpt-5.1"
    payload.setdefault("supplier_id", 0)
    payload.setdefault("supplier_name", "GPT 5.1 - Marche Algerien")
    payload.setdefault("currency", "DZD")
    payload.setdefault("strict_alerts", ["ESTIMATION_MARCHE_DZ", "PLAFOND_PRIX_DZD"])
    payload.setdefault(
        "strict_alert_colors",
        {
            "ESTIMATION_MARCHE_DZ": "purple",
            "PLAFOND_PRIX_DZD": "purple",
        },
    )
    payload.setdefault("match_badges", ["ESTIMATION_MARCHE_DZ", "PLAFOND_PRIX_DZD"])
    payload.setdefault(
        "match_badge_breakdown",
        {
            "ESTIMATION_MARCHE_DZ": {"color": "purple"},
            "PLAFOND_PRIX_DZD": {"color": "purple"},
        },
    )
    return payload


def _inject_market_estimates_into_payload(
    payload: list[dict[str, Any]],
    *,
    db: Session,
    tenant_id: int,
    campaign_id: int,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    row_by_product_id: dict[int, dict[str, Any]] = {}
    for raw_row in payload:
        if not isinstance(raw_row, dict):
            continue
        row = dict(raw_row)
        offers = row.get("offers") or []
        normalized_offers: list[dict[str, Any]] = []
        for raw_offer in offers:
            if isinstance(raw_offer, dict):
                normalized_offers.append(dict(raw_offer))
        row["offers"] = normalized_offers
        best_offer = row.get("best_offer")
        row["best_offer"] = dict(best_offer) if isinstance(best_offer, dict) else None
        rows.append(row)
        try:
            product_id = int(row.get("product_id"))
        except Exception:
            continue
        row_by_product_id[product_id] = row

    estimates = (
        db.query(CampaignMarketEstimateOffer)
        .filter(
            CampaignMarketEstimateOffer.tenant_id == tenant_id,
            CampaignMarketEstimateOffer.campaign_id == campaign_id,
        )
        .order_by(CampaignMarketEstimateOffer.product_id.asc())
        .all()
    )
    for estimate in estimates:
        row = row_by_product_id.get(int(estimate.product_id))
        if row is None:
            continue
        offer_payload = _market_estimate_offer_from_record(estimate)
        offers = row.get("offers") or []
        exists = False
        for existing in offers:
            if not isinstance(existing, dict):
                continue
            if bool(existing.get("market_estimate")) and int(
                existing.get("market_estimate_id") or 0
            ) == int(estimate.id):
                exists = True
                break
        if not exists:
            offers.append(offer_payload)
            row["offers"] = offers
        if row.get("best_offer") is None:
            row["best_offer"] = offer_payload

    return rows


@router.get("/{campaign_id}/offers/by-product", response_model=ProductOffersPage)
def offers_by_product(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    conform_only: bool = False,
    include_market_estimates: bool = False,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    cache_only: bool = False,
) -> ProductOffersPage:
    payload: list[dict[str, Any]]
    cache_ready = True
    if cache_only:
        cached_payload = _get_cached_results_payload_if_available(
            campaign_id,
            db,
            current_user,
            conform_only=conform_only,
            offer_limit=20,
            useful_only=False,
            min_match_score=0.0,
        )
        cache_ready = cached_payload is not None
        payload = cached_payload or []
    else:
        payload = _cached_results_payload(
            campaign_id,
            db,
            current_user,
            conform_only=conform_only,
            offer_limit=20,
            useful_only=False,
            min_match_score=0.0,
        )
    if include_market_estimates:
        payload = _inject_market_estimates_into_payload(
            payload,
            db=db,
            tenant_id=current_user.tenant_id,
            campaign_id=campaign_id,
        )
    total = len(payload)
    start = (page - 1) * page_size
    end = start + page_size
    page_rows = [_lighten_product_payload(row) for row in payload[start:end]]
    return ProductOffersPage(
        items=[ProductOffers(**row) for row in page_rows],
        total=total,
        page=page,
        page_size=page_size,
        has_more=end < total,
        cache_ready=cache_ready,
    )


@router.get(
    "/{campaign_id}/offers/simple-product",
    response_model=SimpleProductOffersPage,
)
def offers_simple_product(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    cache_only: bool = False,
) -> SimpleProductOffersPage:
    _ensure_simple_product_enabled()
    payload: list[dict[str, Any]]
    cache_ready = True
    if cache_only:
        cached_payload = _get_simple_product_payload_if_available(
            campaign_id,
            db,
            current_user,
        )
        cache_ready = cached_payload is not None
        payload = cached_payload or []
    else:
        payload = _simple_product_payload(
            campaign_id,
            db,
            current_user,
        )
    total = len(payload)
    start = (page - 1) * page_size
    end = start + page_size
    page_rows = payload[start:end]
    return SimpleProductOffersPage(
        items=[SimpleProductOffers(**row) for row in page_rows],
        total=total,
        page=page,
        page_size=page_size,
        has_more=end < total,
        cache_ready=cache_ready,
    )


@router.get(
    "/{campaign_id}/offers/by-product/{product_id}/offers/{invoice_item_id}",
    response_model=OfferEntry,
)
def offer_by_product_detail(
    campaign_id: int,
    product_id: int,
    invoice_item_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> OfferEntry:
    payload = _cached_results_payload(
        campaign_id,
        db,
        current_user,
        conform_only=False,
        offer_limit=20,
        useful_only=True,
        min_match_score=_RESULTS_MIN_GPT_MATCH,
    )
    for row in payload:
        if int(row.get("product_id") or 0) != product_id:
            continue
        offers_raw = row.get("offers") or []
        for offer_raw in offers_raw:
            if not isinstance(offer_raw, dict):
                continue
            if int(offer_raw.get("invoice_item_id") or 0) != invoice_item_id:
                continue
            return OfferEntry(**_summarize_offer_components(offer_raw))
    raise HTTPException(status_code=404, detail="Offer detail not found")


@router.post(
    "/{campaign_id}/offers/by-product/{product_id}/market-estimate",
    response_model=OfferEntry,
)
def generate_market_estimate_offer(
    campaign_id: int,
    product_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> OfferEntry:
    campaign = _get_campaign(db, campaign_id, current_user)

    existing = (
        db.query(CampaignMarketEstimateOffer)
        .filter(
            CampaignMarketEstimateOffer.tenant_id == current_user.tenant_id,
            CampaignMarketEstimateOffer.campaign_id == campaign_id,
            CampaignMarketEstimateOffer.product_id == product_id,
        )
        .first()
    )
    if existing is not None:
        return OfferEntry(**_market_estimate_offer_from_record(existing))

    payload = _cached_results_payload(
        campaign_id,
        db,
        current_user,
        conform_only=False,
        offer_limit=None,
        useful_only=False,
        min_match_score=0.0,
    )
    target_row = None
    for row in payload:
        try:
            row_product_id = int(row.get("product_id") or 0)
        except Exception:
            continue
        if row_product_id == product_id:
            target_row = row
            break
    if target_row is None:
        raise HTTPException(status_code=404, detail="Requested product not found")

    requested_name = str(target_row.get("product_name") or "").strip()
    requested_qty = target_row.get("requested_qty")
    requested_source_file = str(
        target_row.get("requested_source_attachment_file") or ""
    ).strip()
    if not requested_name:
        raise HTTPException(status_code=400, detail="Requested product name missing")

    candidate_names: list[str] = []
    for raw_offer in (target_row.get("offers") or []):
        if not isinstance(raw_offer, dict):
            continue
        name = str(raw_offer.get("extracted_name") or "").strip()
        if name:
            candidate_names.append(name)

    system_prompt = (
        "Tu es un acheteur biomedical specialise en consultation du marche algerien. "
        "Tu dois utiliser la recherche web fournie par l'outil pour verifier le produit et le fournisseur suggere. "
        "Retourne strictement un JSON. "
        "Le prix doit etre une estimation MAXIMALE prudente en DZD pour une unite comparable, derivee des resultats web trouves. "
        "Si plusieurs produits existent, choisis le plus proche du besoin demande. "
        "Si plusieurs fournisseurs existent, suggere le plus credible a partir des resultats trouves. "
        "N'invente pas de source: base-toi sur les pages trouvees."
    )
    user_payload = {
        "country_market": "Algeria",
        "requested_product_name": requested_name,
        "requested_qty": requested_qty,
        "campaign_context": {
            "establishment_name": getattr(campaign, "establishment_name", None),
            "establishment_type": getattr(campaign, "establishment_type", None),
            "pharmacy_service": getattr(campaign, "pharmacy_service", None),
            "market_office_service": getattr(campaign, "market_office_service", None),
            "lot_reference": getattr(campaign, "lot_reference", None),
        },
        "requested_source_attachment_file": requested_source_file or None,
        "existing_offer_names": candidate_names[:6],
    }
    model = "gpt-5.1"
    timeout_s = max(
        45,
        int(getattr(config, "OPENAI_PRODUCT_MATCH_TIMEOUT_SECONDS", 60) or 60),
    )
    retry_attempts = max(
        1,
        int(getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_ATTEMPTS", 3) or 3),
    )
    retry_backoff_s = max(
        0.2,
        float(getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_BACKOFF_SECONDS", 1.5) or 1.5),
    )
    responses_payload = {
        "model": model,
        "input": json.dumps(user_payload, ensure_ascii=False),
        "instructions": system_prompt,
        "tools": [
            {
                "type": "web_search",
                "user_location": {
                    "type": "approximate",
                    "country": "DZ",
                },
                "search_context_size": "high",
            }
        ],
        "include": ["web_search_call.action.sources"],
        "text": {
            "format": {
                "type": "json_schema",
                "name": "market_estimate_offer",
                "schema": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "closest_product_name": {"type": "string"},
                        "commercial_description": {"type": "string"},
                        "estimated_max_unit_price_dzd": {"type": "number"},
                        "match_confidence": {
                            "type": "number",
                            "minimum": 0,
                            "maximum": 1,
                        },
                        "reason": {"type": "string"},
                        "price_basis": {"type": "string"},
                        "supplier_name_hint": {
                            "anyOf": [{"type": "string"}, {"type": "null"}]
                        },
                    },
                    "required": [
                        "closest_product_name",
                        "commercial_description",
                        "estimated_max_unit_price_dzd",
                        "match_confidence",
                        "reason",
                        "price_basis",
                        "supplier_name_hint",
                    ],
                },
                "strict": True,
            }
        },
    }

    headers = {
        "Authorization": f"Bearer {config.OPENAI_API_KEY}",
        "Content-Type": "application/json",
        "X-Client-Request-Id": f"market-estimate-{campaign_id}-{product_id}-{uuid4()}",
    }

    def _extract_responses_output_text(data: dict[str, Any] | None) -> str:
        if not isinstance(data, dict):
            return ""
        direct = data.get("output_text")
        if isinstance(direct, str) and direct.strip():
            return direct
        output = data.get("output")
        if isinstance(output, list):
            for item in output:
                if not isinstance(item, dict):
                    continue
                content = item.get("content")
                if not isinstance(content, list):
                    continue
                for part in content:
                    if not isinstance(part, dict):
                        continue
                    if str(part.get("type") or "").strip().lower() in {"output_text", "text"}:
                        text_value = part.get("text")
                        if isinstance(text_value, str) and text_value.strip():
                            return text_value
        return ""

    def _extract_web_search_sources(data: dict[str, Any] | None) -> list[str]:
        urls: list[str] = []
        if not isinstance(data, dict):
            return urls
        output = data.get("output")
        if not isinstance(output, list):
            return urls
        for item in output:
            if not isinstance(item, dict):
                continue
            if str(item.get("type") or "").strip().lower() != "web_search_call":
                continue
            action = item.get("action")
            sources = action.get("sources") if isinstance(action, dict) else None
            if not isinstance(sources, list):
                continue
            for source in sources:
                if not isinstance(source, dict):
                    continue
                url = str(source.get("url") or "").strip()
                if url and url not in urls:
                    urls.append(url)
        return urls

    responses_data: dict[str, Any] | None = None
    parsed: dict[str, Any] | None = None
    model_used = model
    last_error: Exception | None = None
    for attempt in range(retry_attempts):
        started = time.perf_counter()
        try:
            req = urllib_request.Request(
                "https://api.openai.com/v1/responses",
                data=json.dumps(responses_payload).encode("utf-8"),
                headers=headers,
                method="POST",
            )
            with urllib_request.urlopen(req, timeout=timeout_s) as resp:
                responses_data = json.loads(resp.read().decode("utf-8", errors="ignore"))
            record_openai_api_call(
                operation="campaign_market_estimate_offer_web",
                endpoint="/v1/responses",
                model=model,
                response_json=responses_data if isinstance(responses_data, dict) else None,
                latency_ms=int((time.perf_counter() - started) * 1000),
                success=True,
                metadata={
                    "campaign_id": campaign_id,
                    "product_id": product_id,
                    "mode": "market_estimate_dz_web",
                    "attempt": attempt + 1,
                },
            )
            content = _extract_responses_output_text(responses_data)
            if not content.strip():
                raise ValueError("missing_output_text")
            parsed_obj = json.loads(content)
            if not isinstance(parsed_obj, dict):
                raise ValueError("json_not_object")
            parsed = parsed_obj
            break
        except Exception as exc:
            last_error = exc if isinstance(exc, Exception) else Exception(str(exc))
            status_code = (
                getattr(exc, "code", None)
                if isinstance(exc, urllib_error.HTTPError)
                else None
            )
            record_openai_api_call(
                operation="campaign_market_estimate_offer_web",
                endpoint="/v1/responses",
                model=model,
                response_json=None,
                latency_ms=int((time.perf_counter() - started) * 1000),
                success=False,
                error_type=type(exc).__name__,
                metadata={
                    "campaign_id": campaign_id,
                    "product_id": product_id,
                    "mode": "market_estimate_dz_web",
                    "attempt": attempt + 1,
                    "http_status": status_code,
                },
            )
            if attempt < (retry_attempts - 1):
                delay = (retry_backoff_s * (2**attempt))
                time.sleep(delay)
                continue
            break

    if not isinstance(parsed, dict):
        raise HTTPException(
            status_code=502,
            detail=f"GPT web market estimate failed: {type(last_error).__name__ if last_error else 'unknown'}",
        )

    def _to_float(value: Any) -> float | None:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            try:
                return float(value.strip().replace(",", "."))
            except Exception:
                return None
        return None

    confidence = _to_float(parsed.get("match_confidence")) or 0.70
    if confidence < 0:
        confidence = 0.70
    if confidence > 1:
        confidence = 1.0
    estimated_max_unit_price_dzd = _to_float(parsed.get("estimated_max_unit_price_dzd"))
    if estimated_max_unit_price_dzd is None or estimated_max_unit_price_dzd <= 0:
        raise HTTPException(status_code=502, detail="GPT market estimate price missing")

    closest_product_name = str(parsed.get("closest_product_name") or "").strip() or requested_name
    commercial_description = str(parsed.get("commercial_description") or "").strip()
    reason = str(parsed.get("reason") or "").strip()
    price_basis = str(parsed.get("price_basis") or "").strip()
    supplier_name_hint = str(parsed.get("supplier_name_hint") or "").strip()
    market_sources = _extract_web_search_sources(responses_data)

    offer_payload: dict[str, Any] = {
        "supplier_id": 0,
        "supplier_name": supplier_name_hint or "GPT 5.1 - Marche Algerien",
        "supplier_email": None,
        "supplier_phone": None,
        "sender_email": None,
        "invoice_sender_email": None,
        "market_estimate": True,
        "market_estimate_region": "DZ",
        "market_estimate_model": str(model_used or model).strip() or model,
        "market_estimate_sources": market_sources,
        "requested_catalog_name": requested_name,
        "requested_strict_name_used": requested_name,
        "requested_qty": _to_float(requested_qty),
        "invoice_item_id": None,
        "invoice_id": None,
        "invoice_file_name": None,
        "invoice_file_path": None,
        "tech_sheet_id": None,
        "tech_sheet_file": None,
        "tech_sheet_file_path": None,
        "offered_qty": _to_float(requested_qty) or 1.0,
        "offered_qty_raw": None,
        "offered_qty_strict": _to_float(requested_qty) or 1.0,
        "unit_price": estimated_max_unit_price_dzd,
        "vat_rate": None,
        "equivalent_unit_price": estimated_max_unit_price_dzd,
        "equivalence_factor": 1.0,
        "equivalence_basis": f"market_estimate:{price_basis or 'direct_requested_unit'}",
        "equivalence_confidence": confidence,
        "currency": "DZD",
        "price_source": "market_estimate_gpt_5_1_web",
        "extracted_name": closest_product_name,
        "extracted_description": commercial_description or reason,
        "gpt_matched_name": closest_product_name,
        "gpt_match_rate": confidence,
        "embedding_match_rate": None,
        "name_match_rate": confidence,
        "match_compensable": True,
        "match_compensation_mode": "market_estimate",
        "match_compensation_factor": 1.0,
        "match_compensation_basis": "market_estimate_direct",
        "match_compensation_confidence": confidence,
        "match_badges": ["ESTIMATION_MARCHE_DZ", "PLAFOND_PRIX_DZD"],
        "match_badge_breakdown": {
            "ESTIMATION_MARCHE_DZ": {"color": "purple", "score": confidence},
            "PLAFOND_PRIX_DZD": {"color": "purple", "score": confidence},
        },
        "strict_match_score": max(confidence, 0.70),
        "strict_match_threshold": 0.65,
        "strict_match_decision": "matched",
        "strict_match_reason": "market_estimate_gpt_5_1_web",
        "strict_conversion_allowed": True,
        "strict_conversion_mode": "direct",
        "strict_conversion_factor": 1.0,
        "strict_conversion_basis": "market_estimate_requested_unit",
        "strict_final_decision": "market_estimate",
        "strict_decision_reason": reason or "Estimation GPT 5.1 avec recherche web pour le marche algerien",
        "strict_alerts": ["ESTIMATION_MARCHE_DZ", "PLAFOND_PRIX_DZD"],
        "strict_alert_colors": {
            "ESTIMATION_MARCHE_DZ": "purple",
            "PLAFOND_PRIX_DZD": "purple",
        },
    }

    estimate = CampaignMarketEstimateOffer(
        tenant_id=current_user.tenant_id,
        campaign_id=campaign_id,
        product_id=product_id,
        created_by_user_id=current_user.id,
        model_name=str(model_used or model).strip() or model,
        payload=offer_payload,
    )
    db.add(estimate)
    db.commit()
    db.refresh(estimate)
    return OfferEntry(**_market_estimate_offer_from_record(estimate))


@router.get("/{campaign_id}/offers/summary", response_model=list[OfferSummaryEntry])
def offers_summary(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[OfferSummaryEntry]:
    offers = _offers_by_product(campaign_id, db, current_user, limit=None, conform_only=False)
    summary: list[OfferSummaryEntry] = []
    for product in offers:
        if product.best_offer:
            best_offer = product.best_offer
            best_price = best_offer.equivalent_unit_price
            summary.append(
                OfferSummaryEntry(
                    product_id=product.product_id,
                    product_name=product.product_name,
                    best_price=best_price,
                    best_supplier=best_offer.supplier_name,
                    offered_qty=best_offer.offered_qty,
                    offers_count=len(product.offers),
                )
            )
        else:
            summary.append(
                OfferSummaryEntry(
                    product_id=product.product_id,
                    product_name=product.product_name,
                    best_price=None,
                    best_supplier=None,
                    offered_qty=None,
                    offers_count=0,
                )
            )
    return summary


@router.get("/{campaign_id}/export/csv")
def export_campaign_csv(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id
    rows = (
        db.query(ResponseItem, ResponseModel, Supplier, Product)
        .join(ResponseModel, ResponseItem.response_id == ResponseModel.id)
        .join(Supplier, ResponseModel.supplier_id == Supplier.id)
        .join(Product, ResponseItem.product_id == Product.id)
        .filter(
            ResponseModel.campaign_id == campaign_id,
            ResponseModel.tenant_id == tenant_id,
            ResponseItem.tenant_id == tenant_id,
            Supplier.tenant_id == tenant_id,
            Product.tenant_id == tenant_id,
        )
        .all()
    )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "campaign_id",
            "campaign_title",
            "supplier_name",
            "supplier_email",
            "product_name",
            "offered_qty",
            "unit_price",
            "currency",
            "notes",
            "submitted_at",
            "responder_email",
        ]
    )
    for item, response, supplier, product in rows:
        writer.writerow(
            [
                campaign.id,
                campaign.title,
                supplier.name,
                supplier.email,
                product.name,
                item.offered_qty,
                item.unit_price,
                item.currency,
                item.notes or "",
                response.submitted_at.isoformat(),
                response.responder_email,
            ]
        )
    content = output.getvalue()
    filename = f"campaign_{campaign_id}_responses.csv"
    return FastAPIResponse(
        content=content,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _xlsx_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        try:
            return json.dumps(value, ensure_ascii=False)
        except Exception:
            return str(value)
    return value


def _extract_note_float(notes: str | None, key: str) -> float | None:
    if not notes:
        return None
    m = re.search(rf"\b{re.escape(key)}=([0-9]+(?:\.[0-9]+)?)", notes)
    if not m:
        return None
    try:
        return float(m.group(1))
    except Exception:
        return None


def _xlsx_safe_title(value: str) -> str:
    raw = (value or "").strip() or "campaign"
    cleaned = re.sub(r'[:\\/*?\[\]]+', "_", raw)
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._")
    if not cleaned:
        cleaned = "campaign"
    return cleaned[:80]


def _strict_primary_color_token_from_offer(offer: dict[str, Any] | None) -> str:
    if not isinstance(offer, dict):
        return "gray"
    values: set[str] = set()
    strict_colors = offer.get("strict_alert_colors")
    if isinstance(strict_colors, dict):
        for raw in strict_colors.values():
            token = str(raw or "").strip().lower()
            if token:
                values.add(token)
    badge_breakdown = offer.get("match_badge_breakdown")
    if isinstance(badge_breakdown, dict):
        for raw in badge_breakdown.values():
            if not isinstance(raw, dict):
                continue
            token = str(raw.get("color") or "").strip().lower()
            if token:
                values.add(token)
    candidates = {c for c in values if c not in {"blue", "amber"}}
    if "red" in candidates:
        return "red"
    if "orange" in candidates:
        return "orange"
    if "gray" in candidates:
        return "gray"
    if "purple" in candidates:
        return "purple"
    if "yellow" in candidates:
        return "yellow"
    if "green" in candidates:
        return "green"
    if "purple" in values:
        return "purple"
    if "green" in values:
        return "green"
    return "gray"


def _strict_excel_fill_hex(color_token: str) -> str:
    token = str(color_token or "").strip().lower()
    mapping = {
        "white": "FFFFFF",
        "green": "E8F5E9",
        "yellow": "FFF8D6",
        "gray": "ECEFF3",
        "orange": "FFE8CC",
        "red": "FDECEC",
        "blue": "E8F1FF",
        "amber": "FFF1DB",
        "purple": "F2EBFF",
    }
    return mapping.get(token, "ECEFF3")


def _strict_excel_font_hex(color_token: str) -> str:
    token = str(color_token or "").strip().lower()
    if token == "white":
        return "102A43"
    if token in {"red", "orange"}:
        return "7A271A"
    if token == "green":
        return "1B4332"
    if token == "yellow":
        return "6B5A00"
    return "334155"


def _strict_excel_sort_rank(color_token: str) -> int:
    token = str(color_token or "").strip().lower()
    order = {
        "green": 0,
        "yellow": 1,
        "gray": 2,
        "purple": 3,
        "orange": 4,
        "red": 5,
        "white": 6,
    }
    return order.get(token, 6)


def _campaign_excel_header_rows(campaign: Campaign) -> list[list[Any]]:
    processing_parts: list[str] = []
    if bool(getattr(campaign, "processing_in_progress", False)):
        processing_parts.append("En cours de traitement")
    if bool(getattr(campaign, "processing_validated", False)):
        processing_parts.append("Valide")
    processing_label = " / ".join(processing_parts) if processing_parts else "-"
    deposit_date = getattr(campaign, "deposit_date", None)
    deposit_label = (
        deposit_date.strftime("%Y-%m-%d")
        if isinstance(deposit_date, datetime)
        else (str(deposit_date or "").strip() or "-")
    )
    return [
        [
            "Nom etablissement",
            getattr(campaign, "establishment_name", None) or "-",
            "Lot",
            getattr(campaign, "lot_reference", None) or "-",
            "Type etablissement",
            getattr(campaign, "establishment_type", None) or "-",
            "Date de depot",
            deposit_label,
        ],
        [
            "Adresse",
            getattr(campaign, "establishment_address", None) or "-",
            "Telephone",
            getattr(campaign, "establishment_phone", None) or "-",
            "Traitement",
            processing_label,
            "Statut campagne",
            campaign.status.value if getattr(campaign, "status", None) else "-",
        ],
        [
            "Service pharmacie",
            getattr(campaign, "pharmacy_service", None) or "-",
            "Service bureau de marche",
            getattr(campaign, "market_office_service", None) or "-",
            "Campagne",
            getattr(campaign, "title", None) or "-",
            "ID campagne",
            getattr(campaign, "id", None),
        ],
    ]


def _write_campaign_excel_header(
    ws: Any,
    *,
    campaign: Campaign,
    title: str,
    columns: list[str],
) -> int:
    from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    total_columns = max(len(columns), 8)
    title_fill = PatternFill(
        fill_type="solid",
        start_color="1F3C88",
        end_color="1F3C88",
    )
    title_font = Font(bold=True, color="FFFFFF", size=13)
    info_fill = PatternFill(
        fill_type="solid",
        start_color="EAF2FF",
        end_color="EAF2FF",
    )
    label_fill = PatternFill(
        fill_type="solid",
        start_color="D9EAF7",
        end_color="D9EAF7",
    )
    label_font = Font(bold=True, color="102A43")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.fill = title_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columns)
    summary_cell = ws.cell(row=2, column=1)
    summary_cell.value = (
        f"Campagne: {getattr(campaign, 'title', '-') or '-'} | "
        f"ID: {getattr(campaign, 'id', '-') or '-'} | "
        f"Statut: {campaign.status.value if getattr(campaign, 'status', None) else '-'}"
    )
    summary_cell.fill = info_fill
    summary_cell.font = Font(bold=True, color="102A43")
    summary_cell.alignment = Alignment(horizontal="left", vertical="center")

    row_idx = 3
    for meta_row in _campaign_excel_header_rows(campaign):
        padded = list(meta_row) + [None] * max(0, total_columns - len(meta_row))
        ws.append(padded[:total_columns])
        for col_idx in range(1, min(len(meta_row), total_columns) + 1, 2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = label_fill
            cell.font = label_font
        row_idx += 1

    ws.append([])
    ws.append(columns + [None] * max(0, total_columns - len(columns)))
    header_row_idx = ws.max_row
    for cell in ws[header_row_idx]:
        cell.fill = label_fill
        cell.font = label_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx in range(1, total_columns + 1):
        letter = get_column_letter(idx)
        if ws.column_dimensions[letter].width is None:
            ws.column_dimensions[letter].width = 16

    return header_row_idx


def _build_strict_export_rows(
    *,
    campaign: Campaign,
    db: Session,
    current_user: User,
    include_empty: bool,
    best_offer_mode: str = "lowest_total",
    exclude_display_similarity_secondary: bool = False,
    selected_offer_by_product_id: dict[int, int] | None = None,
) -> list[dict[str, Any]]:
    payload = _cached_results_payload(
        campaign.id,
        db,
        current_user,
        conform_only=False,
        offer_limit=None,
        useful_only=False,
        min_match_score=0.0,
    )
    payload = _inject_market_estimates_into_payload(
        payload,
        db=db,
        tenant_id=current_user.tenant_id,
        campaign_id=campaign.id,
    )
    requested_index = _build_campaign_requested_index(
        db=db,
        tenant_id=campaign.tenant_id,
        campaign_ids=[campaign.id],
    )
    requested_raw_name_by_product_id: dict[int, str] = {}
    for row in requested_index.get(campaign.id, []) or []:
        try:
            product_id = int(row.get("product_id"))
        except Exception:
            continue
        raw_name = str(row.get("requested_raw_product_name") or "").strip()
        if raw_name:
            requested_raw_name_by_product_id[product_id] = raw_name

    strict_default_threshold = float(
        getattr(config, "OPENAI_STRICT_REQUESTED_MATCH_MIN_CONFIDENCE", 0.65) or 0.65
    )

    def _to_float(value: Any) -> float | None:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            try:
                return float(value.strip().replace(",", "."))
            except Exception:
                return None
        return None

    def _passes_strict_threshold(offer: dict[str, Any]) -> bool:
        score = _to_float(offer.get("strict_match_score"))
        threshold = _to_float(offer.get("strict_match_threshold"))
        if threshold is None:
            threshold = strict_default_threshold
        if score is None:
            return False
        return float(score) >= float(threshold)

    strict_rows: list[dict[str, Any]] = []
    for row in payload:
        product_id = row.get("product_id")
        requested_name = requested_raw_name_by_product_id.get(
            int(product_id) if product_id is not None else -1
        ) or row.get("product_name")
        requested_qty = _to_float(row.get("requested_qty"))
        offers = row.get("offers") or []
        eligible: list[dict[str, Any]] = []
        selected_offer = None
        selected_invoice_item_id = None
        try:
            selected_invoice_item_id = int(
                (selected_offer_by_product_id or {}).get(int(product_id))
            )
        except Exception:
            selected_invoice_item_id = None
        for raw in offers:
            if not isinstance(raw, dict):
                continue
            offer = dict(raw)
            if not _passes_strict_threshold(offer):
                continue
            if exclude_display_similarity_secondary:
                reason = str(offer.get("strict_match_reason") or "").strip().lower()
                if reason == "display_similarity_threshold":
                    continue
            eq_price = _to_float(offer.get("equivalent_unit_price"))
            total_eq_requested = (
                requested_qty * eq_price
                if requested_qty is not None and requested_qty > 0 and eq_price is not None
                else None
            )
            offer["_strict_total_eq_requested"] = total_eq_requested
            eligible.append(offer)
            if selected_invoice_item_id is not None:
                try:
                    if int(offer.get("invoice_item_id")) == selected_invoice_item_id:
                        selected_offer = offer
                except Exception:
                    pass

        def _color_best_rank(offer: dict[str, Any]) -> int:
            token = _strict_primary_color_token_from_offer(offer)
            order = {
                "green": 0,
                "yellow": 1,
                "gray": 2,
                "purple": 3,
                "orange": 4,
                "red": 5,
            }
            return order.get(token, 6)

        def _offer_sort_key(offer: dict[str, Any]) -> tuple[float, float]:
            total_eq = _to_float(offer.get("_strict_total_eq_requested"))
            score = _to_float(offer.get("strict_match_score")) or -1.0
            return (
                total_eq if total_eq is not None else float("inf"),
                -float(score),
            )

        def _offer_sort_key_by_color(offer: dict[str, Any]) -> tuple[int, float, float]:
            total_eq = _to_float(offer.get("_strict_total_eq_requested"))
            score = _to_float(offer.get("strict_match_score")) or -1.0
            return (
                _color_best_rank(offer),
                total_eq if total_eq is not None else float("inf"),
                -float(score),
            )

        def _offer_sort_key_by_grouped_color(
            offer: dict[str, Any],
        ) -> tuple[int, float, float, int]:
            token = _strict_primary_color_token_from_offer(offer)
            total_eq = _to_float(offer.get("_strict_total_eq_requested"))
            score = _to_float(offer.get("strict_match_score")) or -1.0
            primary_group = 0 if token in {"green", "yellow", "gray"} else 1
            within_group_rank = {
                "green": 0,
                "yellow": 1,
                "gray": 2,
                "orange": 0,
                "red": 1,
            }.get(token, 9)
            return (
                primary_group,
                total_eq if total_eq is not None else float("inf"),
                -float(score),
                within_group_rank,
            )

        if selected_offer is not None:
            best_offer = selected_offer
        elif eligible:
            mode = str(best_offer_mode or "").strip().lower()
            if mode == "best_color":
                best_offer = min(eligible, key=_offer_sort_key_by_color)
            elif mode == "lowest_total_preferred_colors":
                best_offer = min(eligible, key=_offer_sort_key_by_grouped_color)
            else:
                best_offer = min(eligible, key=_offer_sort_key)
        else:
            best_offer = None
        if best_offer is None and not include_empty:
            continue

        unit_price = _to_float(best_offer.get("unit_price")) if best_offer else None
        eq_price = (
            _to_float(best_offer.get("equivalent_unit_price")) if best_offer else None
        )
        total_eq_requested = (
            _to_float(best_offer.get("_strict_total_eq_requested")) if best_offer else None
        )
        vat_rate = _to_float(best_offer.get("vat_rate")) if best_offer else None
        total_eq_requested_with_vat = (
            total_eq_requested * (1.0 + (vat_rate / 100.0))
            if total_eq_requested is not None and vat_rate is not None
            else None
        )
        color_token = _strict_primary_color_token_from_offer(best_offer) if best_offer else "white"
        strict_rows.append(
            {
                "campaign_id": campaign.id,
                "requested_product_id": product_id,
                "requested_product_name": requested_name,
                "market_estimate_model": (
                    str(best_offer.get("market_estimate_model") or "").strip()
                    if best_offer
                    else None
                ),
                "market_estimate_sources": (
                    " | ".join(
                        [
                            str(v).strip()
                            for v in (best_offer.get("market_estimate_sources") or [])
                            if str(v).strip()
                        ]
                    )
                    if best_offer
                    else None
                ),
                "offered_product_name": (
                    str(best_offer.get("extracted_name") or "").strip() if best_offer else None
                ),
                "requested_qty": requested_qty,
                "offered_qty": (
                    _to_float(best_offer.get("offered_qty_strict"))
                    or _to_float(best_offer.get("offered_qty"))
                    if best_offer
                    else None
                ),
                "unit_price": unit_price,
                "equivalent_unit_price": eq_price,
                "total_equivalent_requested": total_eq_requested,
                "vat_rate": vat_rate,
                "total_equivalent_requested_with_vat": total_eq_requested_with_vat,
                "strict_score": (
                    _to_float(best_offer.get("strict_match_score")) if best_offer else None
                ),
                "strict_threshold": (
                    _to_float(best_offer.get("strict_match_threshold")) if best_offer else strict_default_threshold
                ),
                "strict_decision": (
                    str(
                        best_offer.get("strict_final_decision")
                        or best_offer.get("strict_match_decision")
                        or ""
                    ).strip()
                    if best_offer
                    else None
                ),
                "compensation_mode": (
                    str(best_offer.get("match_compensation_mode") or "").strip()
                    if best_offer
                    else None
                ),
                "compensable": best_offer.get("match_compensable") if best_offer else None,
                "produit_faible": (
                    "oui" if best_offer and best_offer.get("strict_penalty_applied") is True else "non"
                ),
                "detail_produit_faible": (
                    " | ".join(
                        [
                            str(v).strip()
                            for v in (best_offer.get("strict_penalty_labels") or [])
                            if str(v).strip()
                        ]
                    )
                    if best_offer and best_offer.get("strict_penalty_applied") is True
                    else None
                ),
                "main_color": color_token,
                "badges": ",".join(
                    [str(v) for v in (best_offer.get("match_badges") or [])]
                )
                if best_offer
                else None,
            }
        )

    return strict_rows


def _strict_export_columns() -> list[str]:
    return [
        "campaign_id",
        "requested_product_id",
        "requested_product_name",
        "market_estimate_model",
        "market_estimate_sources",
        "offered_product_name",
        "requested_qty",
        "offered_qty",
        "unit_price",
        "equivalent_unit_price",
        "total_equivalent_requested",
        "vat_rate",
        "total_equivalent_requested_with_vat",
        "strict_score",
        "strict_threshold",
        "strict_decision",
        "compensation_mode",
        "compensable",
        "produit_faible",
        "detail_produit_faible",
        "main_color",
        "badges",
    ]


def _strict_rows_to_xlsx_response(
    *,
    campaign: Campaign,
    rows: list[dict[str, Any]],
    file_name: str,
    title: str,
) -> FastAPIResponse:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Font, PatternFill  # type: ignore

    def _to_float(value: Any) -> float | None:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            try:
                return float(value.strip().replace(",", "."))
            except Exception:
                return None
        return None

    columns = _strict_export_columns()
    numeric_columns = {
        "requested_qty",
        "offered_qty",
        "unit_price",
        "equivalent_unit_price",
        "total_equivalent_requested",
        "vat_rate",
        "total_equivalent_requested_with_vat",
        "strict_score",
        "strict_threshold",
    }
    numeric_number_format = "0.00"
    wb = Workbook()
    ws = wb.active
    ws.title = "strict_best_offers"
    header_row_idx = _write_campaign_excel_header(
        ws,
        campaign=campaign,
        title=title,
        columns=columns,
    )

    header_fill = PatternFill(
        fill_type="solid",
        start_color="D9EAF7",
        end_color="D9EAF7",
    )
    header_font = Font(bold=True, color="102A43")
    for cell in ws[header_row_idx]:
        cell.fill = header_fill
        cell.font = header_font

    data_start_row = header_row_idx + 1
    for row_idx, item in enumerate(rows, start=data_start_row):
        ws.append([item.get(col) for col in columns])
        color = str(item.get("main_color") or "gray").strip().lower()
        fill_hex = _strict_excel_fill_hex(color)
        font_hex = _strict_excel_font_hex(color)
        fill = PatternFill(
            fill_type="solid",
            start_color=fill_hex,
            end_color=fill_hex,
        )
        row_font = Font(color=font_hex)
        for cell in ws[row_idx]:
            cell.fill = fill
            cell.font = row_font
        for col_idx in (3, 4, 7, 8, 9):
            ws.cell(row=row_idx, column=col_idx).font = Font(
                bold=True,
                color=font_hex,
            )
        for col_idx, column_name in enumerate(columns, start=1):
            if column_name in numeric_columns:
                ws.cell(row=row_idx, column=col_idx).number_format = numeric_number_format

    total_equivalent_requested_sum = 0.0
    total_equivalent_requested_with_vat_sum = 0.0
    has_total_equivalent_requested_sum = False
    has_total_equivalent_requested_with_vat_sum = False
    for item in rows:
        total_ht = _to_float(item.get("total_equivalent_requested"))
        total_ttc = _to_float(item.get("total_equivalent_requested_with_vat"))
        if total_ht is not None:
            total_equivalent_requested_sum += total_ht
            has_total_equivalent_requested_sum = True
        if total_ttc is not None:
            total_equivalent_requested_with_vat_sum += total_ttc
            has_total_equivalent_requested_with_vat_sum = True

    summary_row_idx = ws.max_row + 2
    summary_fill = PatternFill(
        fill_type="solid",
        start_color="E2E8F0",
        end_color="E2E8F0",
    )
    summary_font = Font(bold=True, color="102A43")
    summary_values = {col: None for col in columns}
    summary_values["requested_product_name"] = "TOTALS"
    summary_values["offered_product_name"] = f"count={len(rows)}"
    summary_values["total_equivalent_requested"] = (
        total_equivalent_requested_sum if has_total_equivalent_requested_sum else None
    )
    summary_values["total_equivalent_requested_with_vat"] = (
        total_equivalent_requested_with_vat_sum
        if has_total_equivalent_requested_with_vat_sum
        else None
    )
    ws.append([])
    ws.append([summary_values.get(col) for col in columns])
    for cell in ws[summary_row_idx]:
        cell.fill = summary_fill
        cell.font = summary_font
    for col_idx, column_name in enumerate(columns, start=1):
        if column_name in numeric_columns:
            ws.cell(row=summary_row_idx, column=col_idx).number_format = numeric_number_format

    ws.freeze_panes = f"A{header_row_idx + 1}"
    from openpyxl.utils import get_column_letter  # type: ignore

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        column_letter = get_column_letter(col_idx)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 14), 42)

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    return FastAPIResponse(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename=\"{file_name}\"'},
    )


@router.get("/{campaign_id}/export/xlsx")
def export_campaign_xlsx(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id

    from openpyxl import Workbook  # type: ignore

    wb = Workbook()

    ws_campaign = wb.active
    ws_campaign.title = "campaign"
    ws_campaign.append(["id", "title", "subject", "status", "priority", "created_at", "sent_at"])
    ws_campaign.append(
        [
            campaign.id,
            campaign.title,
            campaign.subject,
            campaign.status.value if campaign.status else "",
            campaign.priority,
            campaign.created_at.isoformat() if campaign.created_at else "",
            campaign.sent_at.isoformat() if campaign.sent_at else "",
        ]
    )

    ws_items = wb.create_sheet("requested_items")
    ws_items.append(
        [
            "id",
            "product_id",
            "product_name",
            "requested_qty",
            "requested_qty_source",
            "requested_qty_confidence",
            "requested_description",
            "requested_characteristics",
        ]
    )
    item_rows = (
        db.query(CampaignItem, Product)
        .join(Product, CampaignItem.product_id == Product.id)
        .filter(
            CampaignItem.campaign_id == campaign_id,
            CampaignItem.tenant_id == tenant_id,
            Product.tenant_id == tenant_id,
        )
        .all()
    )
    for item, product in item_rows:
        ws_items.append(
            [
                item.id,
                item.product_id,
                product.name if product else "",
                item.requested_qty,
                item.requested_qty_source,
                item.requested_qty_confidence,
                item.requested_description or "",
                _xlsx_cell_value(item.requested_characteristics),
            ]
        )

    ws_recipients = wb.create_sheet("recipients")
    ws_recipients.append(
        [
            "id",
            "supplier_id",
            "supplier_name",
            "supplier_email",
            "delivery_status",
            "sent_at",
            "last_error",
        ]
    )
    recipient_rows = (
        db.query(CampaignRecipient, Supplier)
        .join(Supplier, CampaignRecipient.supplier_id == Supplier.id)
        .filter(
            CampaignRecipient.campaign_id == campaign_id,
            CampaignRecipient.tenant_id == tenant_id,
            Supplier.tenant_id == tenant_id,
        )
        .all()
    )
    for recipient, supplier in recipient_rows:
        ws_recipients.append(
            [
                recipient.id,
                recipient.supplier_id,
                supplier.name if supplier else "",
                supplier.email if supplier else "",
                recipient.delivery_status.value if recipient.delivery_status else "",
                recipient.sent_at.isoformat() if recipient.sent_at else "",
                recipient.last_error or "",
            ]
        )

    ws_invoices = wb.create_sheet("invoices")
    ws_invoices.append(
        [
            "id",
            "sender_email",
            "sender_name",
            "subject",
            "file_name",
            "file_path",
            "parse_engine",
            "parse_score",
            "received_at",
        ]
    )
    invoices = (
        db.query(InboundInvoice)
        .filter(
            InboundInvoice.campaign_id == campaign_id,
            InboundInvoice.tenant_id == tenant_id,
        )
        .order_by(InboundInvoice.received_at.desc())
        .all()
    )
    for inv in invoices:
        ws_invoices.append(
            [
                inv.id,
                inv.sender_email,
                inv.sender_name or "",
                inv.subject or "",
                inv.file_name,
                inv.file_path,
                inv.parse_engine or "",
                inv.parse_score,
                inv.received_at.isoformat() if inv.received_at else "",
            ]
        )

    ws_lines = wb.create_sheet("invoice_lines")
    ws_lines.append(
        [
            "id",
            "invoice_id",
            "supplier_id",
            "status",
            "product_name",
            "description",
            "quantity",
            "unit",
            "unit_price",
            "currency",
            "confidence",
            "similarity_score",
            "embedding_match_rate",
            "gpt_match_rate",
            "notes",
            "created_at",
            "updated_at",
        ]
    )
    line_rows = (
        db.query(InvoiceExtractionItem)
        .filter(
            InvoiceExtractionItem.campaign_id == campaign_id,
            InvoiceExtractionItem.tenant_id == tenant_id,
        )
        .order_by(InvoiceExtractionItem.id.asc())
        .all()
    )
    for row in line_rows:
        notes = row.notes or ""
        ws_lines.append(
            [
                row.id,
                row.invoice_id,
                row.supplier_id,
                row.status.value if row.status else "",
                row.product_name,
                row.description or "",
                row.quantity,
                row.unit or "",
                row.unit_price,
                row.currency,
                row.confidence,
                row.similarity_score,
                _extract_note_float(notes, "emb_score"),
                _extract_note_float(notes, "gpt_score"),
                notes,
                row.created_at.isoformat() if row.created_at else "",
                row.updated_at.isoformat() if row.updated_at else "",
            ]
        )

    ws_tech = wb.create_sheet("tech_sheets")
    ws_tech.append(
        [
            "id",
            "supplier_id",
            "sender_email",
            "file_name",
            "file_path",
            "parse_engine",
            "parse_score",
            "created_at",
        ]
    )
    tech_rows = (
        db.query(TechSheet)
        .filter(
            TechSheet.campaign_id == campaign_id,
            TechSheet.tenant_id == tenant_id,
        )
        .order_by(TechSheet.id.asc())
        .all()
    )
    for row in tech_rows:
        ws_tech.append(
            [
                row.id,
                row.supplier_id,
                row.sender_email or "",
                row.file_name,
                row.file_path,
                row.parse_engine or "",
                row.parse_score,
                row.created_at.isoformat() if row.created_at else "",
            ]
        )

    ws_links = wb.create_sheet("tech_sheet_links")
    ws_links.append(
        [
            "id",
            "invoice_item_id",
            "tech_sheet_id",
            "supplier_id",
            "score",
            "method",
            "created_at",
        ]
    )
    link_rows = (
        db.query(TechSheetLink)
        .filter(
            TechSheetLink.campaign_id == campaign_id,
            TechSheetLink.tenant_id == tenant_id,
        )
        .order_by(TechSheetLink.id.asc())
        .all()
    )
    for row in link_rows:
        ws_links.append(
            [
                row.id,
                row.invoice_item_id,
                row.tech_sheet_id,
                row.supplier_id,
                row.score,
                row.method,
                row.created_at.isoformat() if row.created_at else "",
            ]
        )

    ws_responses = wb.create_sheet("responses")
    ws_responses.append(
        [
            "response_id",
            "supplier_name",
            "supplier_email",
            "submitted_at",
            "responder_email",
            "item_id",
            "product_id",
            "offered_qty",
            "unit_price",
            "currency",
            "notes",
        ]
    )
    response_rows = (
        db.query(ResponseItem, ResponseModel, Supplier)
        .join(ResponseModel, ResponseItem.response_id == ResponseModel.id)
        .join(Supplier, ResponseModel.supplier_id == Supplier.id)
        .filter(
            ResponseModel.campaign_id == campaign_id,
            ResponseModel.tenant_id == tenant_id,
            ResponseItem.tenant_id == tenant_id,
            Supplier.tenant_id == tenant_id,
        )
        .order_by(ResponseModel.id.asc(), ResponseItem.id.asc())
        .all()
    )
    for item, response, supplier in response_rows:
        ws_responses.append(
            [
                response.id,
                supplier.name if supplier else "",
                supplier.email if supplier else "",
                response.submitted_at.isoformat() if response.submitted_at else "",
                response.responder_email or "",
                item.id,
                item.product_id,
                item.offered_qty,
                item.unit_price,
                item.currency,
                item.notes or "",
            ]
        )

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    file_name = f"{_xlsx_safe_title(campaign.title)}_{campaign.id}.xlsx"
    return FastAPIResponse(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )


@router.get("/{campaign_id}/export/top-offers.xlsx")
def export_campaign_top_offers_xlsx(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    top: int = Query(default=3, ge=1, le=10),
    include_empty: bool = True,
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id

    from openpyxl import Workbook  # type: ignore

    payload = _cached_results_payload(
        campaign_id,
        db,
        current_user,
        conform_only=False,
        offer_limit=None,
        useful_only=False,
        min_match_score=0.0,
    )

    supplier_map = {
        supplier.id: supplier
        for supplier in db.query(Supplier)
        .filter(Supplier.tenant_id == tenant_id)
        .all()
    }
    invoice_map = {
        inv.id: inv
        for inv in db.query(InboundInvoice)
        .filter(InboundInvoice.tenant_id == tenant_id)
        .all()
    }

    def _to_float(value: Any) -> float | None:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            try:
                return float(value.strip().replace(",", "."))
            except Exception:
                return None
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "top_offers"
    top_offer_columns = [
        "campaign_id",
        "requested_product_id",
        "requested_product_name",
        "offer_rank",
        "offered_product_name",
        "unit_price",
        "requested_qty",
        "requested_total",
        "requested_total_type",
        "requested_total_reason",
        "equivalent_unit_price",
        "equivalence_basis",
        "currency",
        "score",
        "supplier_name",
        "supplier_email",
        "supplier_phone",
        "invoice_id",
        "invoice_file_name",
    ]
    header_row_idx = _write_campaign_excel_header(
        ws,
        campaign=campaign,
        title="Export Results - meilleures offres",
        columns=top_offer_columns,
    )
    total_top1_requested = 0.0
    total_top1_count = 0

    for row in payload:
        product_id = row.get("product_id")
        product_name = row.get("product_name")
        requested_qty = _to_float(row.get("requested_qty"))
        offers = row.get("offers") or []
        useful_offers: list[dict[str, Any]] = []
        for raw_offer in offers:
            if not isinstance(raw_offer, dict):
                continue
            score = _to_float(raw_offer.get("gpt_match_rate"))
            if score is None or score < _RESULTS_MIN_GPT_MATCH:
                continue
            useful_offers.append(raw_offer)

        if not useful_offers and include_empty:
            ws.append(
                [
                    campaign_id,
                    product_id,
                    product_name,
                    0,
                    None,
                    None,
                    requested_qty,
                    None,
                    "na",
                    "no_offer",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )
            continue

        def _score_key(offer: dict[str, Any]) -> float:
            score = _to_float(offer.get("tech_sheet_similarity_display_score"))
            return -(score if score is not None else -1.0)

        offers_sorted = sorted(useful_offers, key=_score_key)
        for idx, offer in enumerate(offers_sorted[:top], start=1):
            eq_price = _to_float(offer.get("equivalent_unit_price"))
            unit_price = _to_float(offer.get("unit_price"))
            eq_basis = str(offer.get("equivalence_basis") or "").strip()
            is_non_convertible = eq_basis.lower().startswith("direct_")
            can_use_direct_price = is_non_convertible and unit_price is not None

            total_comparable = None
            if requested_qty is not None and requested_qty > 0:
                if eq_price is not None:
                    total_comparable = requested_qty * eq_price
                elif can_use_direct_price:
                    total_comparable = requested_qty * unit_price

            total_estimated = None
            if requested_qty is not None and requested_qty > 0 and unit_price is not None:
                total_estimated = requested_qty * unit_price

            if total_comparable is not None:
                total = total_comparable
                total_type = "comparable"
                total_reason = None
            elif total_estimated is not None:
                total = total_estimated
                total_type = "estimated"
                total_reason = "PU brut"
            else:
                total = None
                total_type = "na"
                if requested_qty is None or requested_qty <= 0:
                    total_reason = "qte demandee manquante"
                elif eq_price is None and not can_use_direct_price:
                    total_reason = "PU compare manquant" if not is_non_convertible else "PU manquant"
                else:
                    total_reason = "non calculable"

            supplier_id = offer.get("supplier_id")
            supplier = supplier_map.get(int(supplier_id)) if supplier_id else None
            invoice_id = offer.get("invoice_id")
            invoice = invoice_map.get(int(invoice_id)) if invoice_id else None
            supplier_email = (
                supplier.email
                if supplier and supplier.email
                else (invoice.sender_email if invoice else None)
            )

            offered_name = (
                str(
                    offer.get("extracted_name")
                    or offer.get("gpt_matched_name")
                    or offer.get("offered_name")
                    or ""
                )
                .strip()
                or "UNKNOWN"
            )

            ws.append(
                [
                    campaign_id,
                    product_id,
                    product_name,
                    idx,
                    offered_name,
                    unit_price,
                    requested_qty,
                    total,
                    total_type,
                    total_reason,
                    eq_price,
                    eq_basis,
                    offer.get("currency"),
                    offer.get("tech_sheet_similarity_display_score"),
                    offer.get("supplier_name"),
                    supplier_email,
                    offer.get("supplier_phone"),
                    invoice_id,
                    offer.get("invoice_file_name"),
                ]
            )
            if idx == 1 and total is not None:
                total_top1_requested += float(total)
                total_top1_count += 1

    ws.append(
        [
            campaign_id,
            None,
            "__TOTAL_TOP1__",
            None,
            None,
            None,
            None,
            total_top1_requested,
            "summary_top1",
            f"count={total_top1_count}",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )

    ws.freeze_panes = f"A{header_row_idx + 1}"
    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    file_name = f"{_xlsx_safe_title(campaign.title)}_{campaign.id}_top_offers.xlsx"
    return FastAPIResponse(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename=\"{file_name}\"'},
    )


def _gpt_export_fill_hex(score: float | None) -> str:
    if score is None:
        return "F3F4F6"
    if score >= 0.8:
        return "E8F5E9"
    if score >= 0.5:
        return "FFF8D6"
    if score >= 0.3:
        return "FFE8CC"
    return "FDECEC"


def _gpt_export_font_hex(score: float | None) -> str:
    if score is None:
        return "475467"
    if score >= 0.8:
        return "1B4332"
    if score >= 0.5:
        return "6B5A00"
    if score >= 0.3:
        return "9A3412"
    return "7A271A"


def _standard_export_to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", "."))
        except Exception:
            return None
    return None


def _standard_export_gpt_score(offer: dict[str, Any]) -> float | None:
    return _standard_export_to_float(
        offer.get("tech_sheet_similarity_display_score")
        if offer.get("tech_sheet_similarity_display_score") is not None
        else offer.get("tech_sheet_similarity_score")
    )


def _pick_standard_offer_for_export(
    offers: list[dict[str, Any]],
    *,
    mode: str,
) -> dict[str, Any] | None:
    if not offers:
        return None

    def _price(offer: dict[str, Any]) -> float:
        value = _standard_export_to_float(offer.get("unit_price"))
        return value if value is not None else float("inf")

    if mode == "lowest_price":
        offers_sorted = sorted(
            offers,
            key=lambda offer: (
                _price(offer),
                -(_standard_export_gpt_score(offer) or -1.0),
                str(offer.get("extracted_name") or "").strip().lower(),
            ),
        )
        return offers_sorted[0] if offers_sorted else None

    offers_sorted = sorted(
        offers,
        key=lambda offer: (
            -(_standard_export_gpt_score(offer) or -1.0),
            _price(offer),
            str(offer.get("extracted_name") or "").strip().lower(),
        ),
    )
    return offers_sorted[0] if offers_sorted else None


def _export_campaign_standard_gpt_selection_xlsx(
    *,
    campaign: Campaign,
    db: Session,
    current_user: User,
    mode: str,
    title: str,
    file_suffix: str,
) -> FastAPIResponse:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Font, PatternFill  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    payload = _cached_results_payload(
        campaign.id,
        db,
        current_user,
        conform_only=False,
        offer_limit=None,
        useful_only=False,
        min_match_score=0.0,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "analyse_gpt"
    columns = [
        "campaign_id",
        "requested_product_id",
        "requested_product_name",
        "offered_product_name",
        "requested_qty",
        "unit_price",
        "requested_total",
        "currency",
        "gpt_score",
        "gpt_status",
        "supplier_name",
        "supplier_phone",
        "invoice_file_name",
    ]
    header_row_idx = _write_campaign_excel_header(
        ws,
        campaign=campaign,
        title=title,
        columns=columns,
    )

    header_fill = PatternFill(
        fill_type="solid",
        start_color="D9EAF7",
        end_color="D9EAF7",
    )
    header_font = Font(bold=True, color="102A43")
    for cell in ws[header_row_idx]:
        cell.fill = header_fill
        cell.font = header_font

    for row in payload:
        product_id = row.get("product_id")
        product_name = row.get("product_name")
        requested_qty = _standard_export_to_float(row.get("requested_qty"))
        offers = [
            dict(offer)
            for offer in (row.get("offers") or [])
            if isinstance(offer, dict)
        ]
        selected = _pick_standard_offer_for_export(offers, mode=mode)
        score = _standard_export_gpt_score(selected) if selected else None
        unit_price = (
            _standard_export_to_float(selected.get("unit_price"))
            if isinstance(selected, dict)
            else None
        )
        requested_total = (
            unit_price * requested_qty
            if unit_price is not None and requested_qty is not None
            else None
        )
        ws.append(
            [
                campaign.id,
                product_id,
                product_name,
                (
                    selected.get("extracted_name")
                    or selected.get("gpt_matched_name")
                    or ""
                )
                if isinstance(selected, dict)
                else "",
                requested_qty,
                unit_price,
                requested_total,
                (selected.get("currency") or "DZD") if isinstance(selected, dict) else "",
                score,
                (
                    selected.get("tech_sheet_similarity_display_status")
                    or selected.get("tech_sheet_similarity_status")
                    or ""
                )
                if isinstance(selected, dict)
                else "",
                selected.get("supplier_name") if isinstance(selected, dict) else "",
                selected.get("supplier_phone") if isinstance(selected, dict) else "",
                selected.get("invoice_file_name") if isinstance(selected, dict) else "",
            ]
        )
        row_idx = ws.max_row
        fill = PatternFill(
            fill_type="solid",
            start_color=_gpt_export_fill_hex(score),
            end_color=_gpt_export_fill_hex(score),
        )
        row_font = Font(color=_gpt_export_font_hex(score))
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.font = row_font
        for col_name in {"requested_qty", "unit_price", "requested_total"}:
            col_idx = columns.index(col_name) + 1
            ws.cell(row=row_idx, column=col_idx).number_format = "#,##0.00"
        ws.cell(
            row=row_idx,
            column=columns.index("gpt_score") + 1,
        ).number_format = "0%"

    ws.freeze_panes = f"A{header_row_idx + 1}"
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            max(max_len + 2, 14),
            40,
        )

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    file_name = (
        f"{_xlsx_safe_title(campaign.title)}_{campaign.id}_{file_suffix}.xlsx"
    )
    return FastAPIResponse(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename=\"{file_name}\"'},
    )


@router.get("/{campaign_id}/export/gpt-best-compliance.xlsx")
def export_campaign_gpt_best_compliance_xlsx(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    if _campaign_effective_processing_mode(db, campaign) == _CAMPAIGN_PROCESSING_STRICT:
        raise HTTPException(
            status_code=400,
            detail="Cet export est reserve aux campagnes standard",
        )
    return _export_campaign_standard_gpt_selection_xlsx(
        campaign=campaign,
        db=db,
        current_user=current_user,
        mode="best_compliance",
        title="Export Analyse GPT - meilleure conformite",
        file_suffix="gpt_meilleure_conformite",
    )


@router.get("/{campaign_id}/export/gpt-lowest-price.xlsx")
def export_campaign_gpt_lowest_price_xlsx(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    if _campaign_effective_processing_mode(db, campaign) == _CAMPAIGN_PROCESSING_STRICT:
        raise HTTPException(
            status_code=400,
            detail="Cet export est reserve aux campagnes standard",
        )
    return _export_campaign_standard_gpt_selection_xlsx(
        campaign=campaign,
        db=db,
        current_user=current_user,
        mode="lowest_price",
        title="Export Analyse GPT - offre la moins chere",
        file_suffix="gpt_moins_cher",
    )


@router.get("/{campaign_id}/export/strict-top-offers.xlsx")
def export_campaign_strict_top_offers_xlsx(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    include_empty: bool = True,
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)

    strict_rows = _build_strict_export_rows(
        campaign=campaign,
        db=db,
        current_user=current_user,
        include_empty=include_empty,
        best_offer_mode="lowest_total_preferred_colors",
    )
    strict_rows.sort(
        key=lambda row: (
            _strict_excel_sort_rank(str(row.get("main_color") or "gray")),
            row["total_equivalent_requested"]
            if row["total_equivalent_requested"] is not None
            else float("inf"),
            str(row.get("requested_product_name") or "").lower(),
        )
    )
    file_name = f"{_xlsx_safe_title(campaign.title)}_{campaign.id}_strict_prix_bas.xlsx"
    return _strict_rows_to_xlsx_response(
        campaign=campaign,
        rows=strict_rows,
        file_name=file_name,
        title="Export Strict - prix bas",
    )


@router.get("/{campaign_id}/export/strict-top-offers-by-color.xlsx")
def export_campaign_strict_top_offers_by_color_xlsx(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    include_empty: bool = True,
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    strict_rows = _build_strict_export_rows(
        campaign=campaign,
        db=db,
        current_user=current_user,
        include_empty=include_empty,
        best_offer_mode="best_color",
        exclude_display_similarity_secondary=True,
    )
    strict_rows = sorted(
        strict_rows,
        key=lambda item: (
            int(item.get("requested_product_id") or 0),
            str(item.get("requested_product_name") or "").strip().lower(),
        ),
    )
    file_name = f"{_xlsx_safe_title(campaign.title)}_{campaign.id}_strict_conformite.xlsx"
    return _strict_rows_to_xlsx_response(
        campaign=campaign,
        rows=strict_rows,
        file_name=file_name,
        title="Export Strict - conformite",
    )


@router.post("/{campaign_id}/export/strict-selected-offers.xlsx")
def export_campaign_strict_selected_offers_xlsx(
    campaign_id: int,
    payload: dict[str, Any] = Body(default_factory=dict),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FastAPIResponse:
    campaign = _get_campaign(db, campaign_id, current_user)
    include_empty = bool(payload.get("include_empty", True))
    raw_selected = payload.get("selected_offer_by_product_id") or {}
    selected_offer_by_product_id: dict[int, int] = {}
    if isinstance(raw_selected, dict):
        for raw_product_id, raw_invoice_item_id in raw_selected.items():
            try:
                selected_offer_by_product_id[int(raw_product_id)] = int(raw_invoice_item_id)
            except Exception:
                continue

    strict_rows = _build_strict_export_rows(
        campaign=campaign,
        db=db,
        current_user=current_user,
        include_empty=include_empty,
        best_offer_mode="best_color",
        exclude_display_similarity_secondary=True,
        selected_offer_by_product_id=selected_offer_by_product_id,
    )
    strict_rows = sorted(
        strict_rows,
        key=lambda item: (
            int(item.get("requested_product_id") or 0),
            str(item.get("requested_product_name") or "").strip().lower(),
        ),
    )
    file_name = f"{_xlsx_safe_title(campaign.title)}_{campaign.id}_strict_selected_offers.xlsx"
    return _strict_rows_to_xlsx_response(
        campaign=campaign,
        rows=strict_rows,
        file_name=file_name,
        title="Export Strict - selection manuelle",
    )


@router.post("/{campaign_id}/purge")
def purge_campaign(
    campaign_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    campaign = _get_campaign(db, campaign_id, current_user)
    tenant_id = campaign.tenant_id

    if campaign.status != CampaignStatus.CLOSED:
        raise HTTPException(
            status_code=400,
            detail="Campaign must be CLOSED before purge",
        )

    attachment_rows = (
        db.query(CampaignAttachment.file_path)
        .filter(
            CampaignAttachment.campaign_id == campaign_id,
            CampaignAttachment.tenant_id == tenant_id,
        )
        .all()
    )
    invoice_rows = (
        db.query(InboundInvoice.file_path)
        .filter(
            InboundInvoice.campaign_id == campaign_id,
            InboundInvoice.tenant_id == tenant_id,
        )
        .all()
    )
    tech_rows = (
        db.query(TechSheet.file_path)
        .filter(
            TechSheet.campaign_id == campaign_id,
            TechSheet.tenant_id == tenant_id,
        )
        .all()
    )

    file_paths: set[str] = set()
    for (file_path,) in attachment_rows + invoice_rows + tech_rows:
        text = str(file_path or "").strip()
        if text:
            file_paths.add(text)

    counts = {
        "campaign_items": db.query(CampaignItem)
        .filter(CampaignItem.campaign_id == campaign_id, CampaignItem.tenant_id == tenant_id)
        .count(),
        "campaign_recipients": db.query(CampaignRecipient)
        .filter(CampaignRecipient.campaign_id == campaign_id, CampaignRecipient.tenant_id == tenant_id)
        .count(),
        "responses": db.query(ResponseModel)
        .filter(ResponseModel.campaign_id == campaign_id, ResponseModel.tenant_id == tenant_id)
        .count(),
        "response_items": db.query(ResponseItem)
        .join(ResponseModel, ResponseItem.response_id == ResponseModel.id)
        .filter(ResponseModel.campaign_id == campaign_id, ResponseItem.tenant_id == tenant_id)
        .count(),
        "attachments": db.query(CampaignAttachment)
        .filter(CampaignAttachment.campaign_id == campaign_id, CampaignAttachment.tenant_id == tenant_id)
        .count(),
        "invoices": db.query(InboundInvoice)
        .filter(InboundInvoice.campaign_id == campaign_id, InboundInvoice.tenant_id == tenant_id)
        .count(),
        "invoice_items": db.query(InvoiceExtractionItem)
        .filter(InvoiceExtractionItem.campaign_id == campaign_id, InvoiceExtractionItem.tenant_id == tenant_id)
        .count(),
        "tech_sheets": db.query(TechSheet)
        .filter(TechSheet.campaign_id == campaign_id, TechSheet.tenant_id == tenant_id)
        .count(),
        "tech_sheet_links": db.query(TechSheetLink)
        .filter(TechSheetLink.campaign_id == campaign_id, TechSheetLink.tenant_id == tenant_id)
        .count(),
        "offer_analysis_results": db.query(OfferAnalysisResult)
        .filter(OfferAnalysisResult.campaign_id == campaign_id, OfferAnalysisResult.tenant_id == tenant_id)
        .count(),
        "results_cache": db.query(CampaignResultsCache)
        .filter(CampaignResultsCache.campaign_id == campaign_id, CampaignResultsCache.tenant_id == tenant_id)
        .count(),
        "email_jobs": db.query(EmailJob)
        .filter(EmailJob.campaign_id == campaign_id, EmailJob.tenant_id == tenant_id)
        .count(),
    }

    db.delete(campaign)
    db.commit()

    files_removed = 0
    files_skipped = 0
    files_failed = 0
    for raw_path in file_paths:
        try:
            still_referenced = (
                db.query(CampaignAttachment)
                .filter(
                    CampaignAttachment.tenant_id == tenant_id,
                    CampaignAttachment.file_path == raw_path,
                )
                .first()
                is not None
                or db.query(InboundInvoice)
                .filter(
                    InboundInvoice.tenant_id == tenant_id,
                    InboundInvoice.file_path == raw_path,
                )
                .first()
                is not None
                or db.query(TechSheet)
                .filter(
                    TechSheet.tenant_id == tenant_id,
                    TechSheet.file_path == raw_path,
                )
                .first()
                is not None
            )
            if still_referenced:
                files_skipped += 1
                continue
            path = Path(raw_path)
            if path.exists() and path.is_file():
                path.unlink()
                files_removed += 1
        except Exception:
            files_failed += 1

    return {
        "status": "ok",
        "campaign_id": campaign_id,
        "deleted": counts,
        "files": {
            "total_linked": len(file_paths),
            "removed": files_removed,
            "skipped_referenced_elsewhere": files_skipped,
            "failed": files_failed,
        },
        # Parser learning table is intentionally preserved:
        # supplier_invoice_formats is tenant-level, not campaign-level.
        "parser_learning_preserved": True,
    }
