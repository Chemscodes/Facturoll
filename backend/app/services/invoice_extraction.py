from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import random
import shutil
import subprocess
from pathlib import Path
from typing import Any, Iterable
import difflib
import unicodedata
import base64
import time
from urllib import error as urllib_error
from urllib import request as urllib_request

from app import config
from app.models import (
    CampaignItem,
    CampaignAttachment,
    ExtractionStatus,
    InboundInvoice,
    InvoiceExtractionItem,
    InvoiceItemCharacteristicMatch,
    Product,
    SupplierInvoiceFormat,
    TechSheetLink,
)
from app.services.document_ai import process_invoice
from app.services.openai_embeddings import (
    cosine_similarity,
    enqueue_embeddings,
    fetch_embeddings,
    get_cached_embeddings,
)
from app.services.openai_payload_cache import (
    get_openai_payload_cache,
    set_openai_payload_cache,
)
from app.services.openai_observability import record_openai_api_call
from app.services.tech_sheet_extraction import _canonical_key, parse_characteristics
from sqlalchemy.orm import Session

logger = logging.getLogger("app")


_SPREADSHEET_INVOICE_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
_XLSX_PRODUCT_HEADER_HINTS = {
    "produit",
    "produits",
    "medicament",
    "medicaments",
    "designation",
    "designations",
    "désignation",
    "désignations",
    "description",
    "descriptif",
    "article",
    "libelle",
    "libellé",
    "libelles",
    "libellés",
    "nom produit",
    "nom du produit",
    "reference produit",
    "référence produit",
    "item",
    "product",
}
_XLSX_REF_HEADER_HINTS = {
    "ref",
    "reference",
    "référence",
    "ref produit",
    "reference produit",
    "référence produit",
    "code",
    "code produit",
    "code article",
    "sku",
}
_XLSX_QTY_HEADER_HINTS = {
    "qte",
    "qté",
    "qty",
    "quantity",
    "quantite",
    "quantité",
    "qte demandee",
    "qté demandée",
}
_XLSX_UNIT_HEADER_HINTS = {
    "unit",
    "unite",
    "unité",
    "u",
    "um",
    "unite de vente",
    "unité de vente",
}
_XLSX_PRICE_HEADER_HINTS = {
    "prix",
    "price",
    "pu",
    "p u",
    "p.u",
    "p.u.",
    "prix unitaire",
    "montant unitaire",
    "unit price",
    "prix ht",
    "prix ttc",
    "tarif",
    "3 m",
    "3m",
    "net",
}
_XLSX_AMOUNT_HEADER_HINTS = {"montant", "total", "amount", "net a payer"}

_UNKNOWN_NAME_TOKENS = {"", "unknown", "unk", "n/a", "na", "null", "none", "-"}
_ADMIN_NOISE_TOKENS = {
    "rib",
    "iban",
    "swift",
    "banque",
    "bank",
    "adresse",
    "address",
    "telephone",
    "tel",
    "mobile",
    "fax",
    "email",
    "site web",
    "rc",
    "nif",
    "nis",
    "ai",
    "art",
    "matricule",
    "total",
    "tva",
    "ttc",
    "ht",
    "net a payer",
    "montant global",
    "cachet",
    "signature",
}


def _deepcopy_payload(value: Any) -> Any:
    try:
        return json.loads(json.dumps(value, ensure_ascii=False))
    except Exception:
        return value


def _entity_value(entity) -> str | None:
    if entity is None:
        return None
    value = entity.mention_text or ""
    value = value.strip()
    return value or None


def _get_prop(entity, prop_type: str):
    for prop in entity.properties:
        if prop.type_ == prop_type or prop.type_ == f"line_item/{prop_type}":
            return prop
    return None


def _iter_line_items(document) -> Iterable:
    for entity in document.entities:
        if entity.type_ == "line_item":
            yield entity


def _extract_supplier_info(document) -> tuple[str | None, str | None]:
    supplier_name = None
    supplier_phone = None
    for entity in document.entities:
        etype = (entity.type_ or "").lower()
        if etype in {"supplier_name", "supplier", "vendor_name"}:
            supplier_name = _entity_value(entity) or supplier_name
        elif etype in {"supplier_phone", "supplier_phone_number", "supplier_tel", "vendor_phone", "phone"}:
            supplier_phone = _entity_value(entity) or supplier_phone
    return supplier_name, supplier_phone


def _collapse_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


_PRODUCT_WORD_HINTS = {
    "AUTOCLAVE",
    "LITRE",
    "LITRES",
    "PORTE",
    "PORTES",
    "MACHINE",
    "CHARIOT",
    "SOUDEUSE",
    "STERILISATION",
    "STERILISATEUR",
    "NEBULISEUR",
    "OTOSCOPE",
    "LARYNGOSCOPE",
    "GUERIDON",
    "PARAVENT",
}
_META_LINE_PREFIXES = (
    "MARQUE",
    "TYPE",
    "ORIGINE",
    "PAYS",
    "FABRICANT",
    "MODELE",
    "MODEL",
    "REF",
    "REFERENCE",
    "SUR COMMANDE",
    "DELAI",
    "DÃ‰LAI",
    "LIVRAISON",
    "GARANTIE",
)
_PRIMARY_META_PREFIXES = ("MARQUE", "TYPE", "ORIGINE")


def _line_starts_with_any(line: str, prefixes: tuple[str, ...]) -> bool:
    upper = _collapse_spaces(line).upper()
    return any(upper.startswith(prefix) for prefix in prefixes)


def _is_code_like(text: str) -> bool:
    token = _collapse_spaces(text).upper()
    if not token:
        return False
    if len(token) > 32:
        return False
    if not re.search(r"\d", token):
        return False
    words = token.split()
    if len(words) > 3:
        return False
    if any(word in _PRODUCT_WORD_HINTS for word in words):
        return False
    if re.search(r"[A-Z]{4,}\s+[A-Z]{4,}", token):
        return False

    if re.fullmatch(r"[A-Z]{1,5}\d{2,}[A-Z0-9\-/_.]*", token):
        return True
    if re.search(r"[-_/]", token) and re.fullmatch(r"[A-Z0-9][A-Z0-9\-/_. ]*", token):
        return True
    if len(words) == 1 and len(token) <= 16 and re.fullmatch(r"[A-Z0-9]{6,}", token):
        return True
    return False


def _prepare_clean_lines(
    description_text: str | None,
    product_code_text: str | None = None,
) -> list[str]:
    raw = (description_text or "").strip()
    if not raw:
        return []
    lines = [_collapse_spaces(line) for line in raw.splitlines()]
    lines = [line for line in lines if line]
    code = _collapse_spaces(product_code_text or "").upper()
    if code and lines:
        filtered: list[str] = []
        for idx, line in enumerate(lines):
            upper_line = line.upper()
            if upper_line == code and len(lines) > 1:
                continue
            if idx == 0 and upper_line.startswith(code + " "):
                line = line[len(code):].strip(" -:/")
            filtered.append(_collapse_spaces(line))
        lines = [line for line in filtered if line]
    if len(lines) > 1 and _is_code_like(lines[0]):
        lines = lines[1:]
    return lines


def _is_meta_line(line: str) -> bool:
    value = _collapse_spaces(line).upper()
    if not value:
        return False
    if _line_starts_with_any(value, _META_LINE_PREFIXES):
        return True
    if value.startswith("SUR COMMANDE"):
        return True
    if "DELAI" in value and "LIVRAISON" in value:
        return True
    if ":" in value:
        left = _collapse_spaces(value.split(":", 1)[0])
        if _line_starts_with_any(left, _META_LINE_PREFIXES):
            return True
    return False


def _split_name_and_meta_lines(lines: list[str]) -> tuple[list[str], list[str]]:
    if not lines:
        return [], []
    name_lines: list[str] = []
    meta_lines: list[str] = []
    for line in lines:
        if _is_meta_line(line):
            meta_lines.append(line)
        else:
            name_lines.append(line)
    if not name_lines and lines:
        name_lines = [lines[0]]
    return name_lines, meta_lines


def _derive_product_name(description_text: str | None, product_code_text: str | None = None) -> str:
    lines = _prepare_clean_lines(description_text, product_code_text)
    if not lines:
        return "UNKNOWN"
    name_lines, _ = _split_name_and_meta_lines(lines)
    candidate = _collapse_spaces(" ".join(name_lines[:3]))
    return candidate or "UNKNOWN"


def _clean_description_text(
    description_text: str | None, product_code_text: str | None = None
) -> str | None:
    lines = _prepare_clean_lines(description_text, product_code_text)
    if not lines:
        return None
    name_lines, meta_lines = _split_name_and_meta_lines(lines)
    if meta_lines:
        primary_meta = [line for line in meta_lines if _line_starts_with_any(line, _PRIMARY_META_PREFIXES)]
        selected = primary_meta if primary_meta else meta_lines
        return " | ".join(selected)
    if len(name_lines) > 1:
        tail = _collapse_spaces(" ".join(name_lines[1:]))
        return tail or None
    return None


def _merge_text_parts(base: str | None, extra: str | None) -> str | None:
    parts: list[str] = []
    for value in (base, extra):
        if not value:
            continue
        cleaned = _collapse_spaces(value)
        if cleaned:
            parts.append(cleaned)
    if not parts:
        return None
    return _collapse_spaces(" ".join(parts)) or None


def _doc_text_lines(document) -> list[str]:
    raw = (getattr(document, "text", None) or "").strip()
    if not raw:
        return []
    lines = [_collapse_spaces(line) for line in raw.splitlines()]
    return [line for line in lines if line]


def _is_product_like_ocr_line(line: str) -> bool:
    value = _collapse_spaces(line).upper()
    if len(value) < 8:
        return False
    if re.fullmatch(r"[0-9\s.,/%\-]+", value):
        return False
    words = value.split()
    if not words:
        return False
    if len(words) == 1 and _is_code_like(words[0]):
        return False
    if any(hint in value for hint in _PRODUCT_WORD_HINTS):
        return True
    if len(words) >= 3 and not _is_meta_line(value):
        return True
    return False


def _strip_leading_code_token(line: str) -> str:
    parts = _collapse_spaces(line).split()
    if len(parts) > 1 and _is_code_like(parts[0]):
        return _collapse_spaces(" ".join(parts[1:]))
    return _collapse_spaces(line)


def _fix_common_ocr_noise(text: str) -> str:
    value = _collapse_spaces(text)
    # Common OCR confusion: "38 U LITRES" where "U" is noise for the unit prefix.
    value = re.sub(
        r"\b(\d+(?:[.,]\d+)?)\s+U\s+(LITRES?|L)\b",
        r"\1 \2",
        value,
        flags=re.IGNORECASE,
    )
    return _collapse_spaces(value)


def _expand_orphan_description(orphan_text: str | None, ocr_lines: list[str]) -> str | None:
    base = _collapse_spaces(orphan_text or "")
    if not base:
        return None
    if not ocr_lines:
        return _fix_common_ocr_noise(base)

    norm_base = base.upper()
    for idx, line in enumerate(ocr_lines):
        norm_line = line.upper()
        if norm_base in norm_line or norm_line in norm_base:
            merged = line
            if idx > 0 and _is_product_like_ocr_line(ocr_lines[idx - 1]):
                prev = _strip_leading_code_token(ocr_lines[idx - 1])
                merged = _collapse_spaces(f"{prev} {line}")
            return _fix_common_ocr_noise(merged)
    return _fix_common_ocr_noise(base)


def _extract_pdf_text(path: str) -> str | None:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return None
    try:
        reader = PdfReader(path)
    except Exception:
        return None
    texts = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            continue
    content = "\n".join(texts).strip()
    return content or None


_TABLE_HEADER_NAME_PATTERNS = (
    re.compile(r"\bdesignation\b"),
    re.compile(r"\bdescription\b"),
    re.compile(r"\blibelle\b"),
    re.compile(r"\bintitule\b"),
)
_TABLE_HEADER_QTY_PATTERNS = (
    re.compile(r"\bqte\b"),
    re.compile(r"\bqty\b"),
    re.compile(r"\bqt\b"),
    re.compile(r"\bquantite\b"),
    re.compile(r"\bquantity\b"),
)
_TABLE_HEADER_UNIT_PRICE_PATTERNS = (
    re.compile(r"\bp\.?\s*u\b"),
    re.compile(r"\bpuht\b"),
    re.compile(r"\bprix\s+unitaire\b"),
    re.compile(r"\bunit\s+price\b"),
)
_TABLE_HEADER_AMOUNT_PATTERNS = (
    re.compile(r"\bmontant\b"),
    re.compile(r"\bamount\b"),
    re.compile(r"\btotal\b"),
)

_PERSONAL_EMAIL_DOMAINS = {
    "gmail.com",
    "yahoo.com",
    "hotmail.com",
    "outlook.com",
    "live.com",
    "icloud.com",
    "proton.me",
    "protonmail.com",
}

_LOCAL_ROW_INLINE_PRICES_RE = re.compile(r"\d[\d\s.,]{2,}\d\s+\d[\d\s.,]{2,}\d")
_LOCAL_ROW_INLINE_QTY_MARKER_RE = re.compile(r"\bU\b\s*\d+(?:[.,]\d+)?", re.IGNORECASE)


def _strip_accents(text: str) -> str:
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", text or "") if not unicodedata.combining(ch)
    )


def _normalize_table_header_line(line: str) -> str:
    value = _strip_accents(_collapse_spaces(line).lower())
    return value


def _line_matches_any(line: str, patterns: tuple[re.Pattern[str], ...]) -> bool:
    return any(pattern.search(line) for pattern in patterns)


def _detect_table_header_flags(line: str) -> dict[str, bool]:
    normalized = _normalize_table_header_line(line)
    return {
        "name": _line_matches_any(normalized, _TABLE_HEADER_NAME_PATTERNS),
        "qty": _line_matches_any(normalized, _TABLE_HEADER_QTY_PATTERNS),
        "unit_price": _line_matches_any(normalized, _TABLE_HEADER_UNIT_PRICE_PATTERNS),
        "amount": _line_matches_any(normalized, _TABLE_HEADER_AMOUNT_PATTERNS),
    }


def _is_table_header_line(line: str) -> bool:
    flags = _detect_table_header_flags(line)
    return bool(flags["name"] and flags["qty"])


def _first_match_position(line: str, patterns: tuple[re.Pattern[str], ...]) -> int | None:
    starts: list[int] = []
    for pattern in patterns:
        m = pattern.search(line)
        if m is not None:
            starts.append(m.start())
    if not starts:
        return None
    return min(starts)


def _infer_header_order(header_line: str) -> list[str]:
    normalized = _normalize_table_header_line(header_line)
    positions: list[tuple[int, str]] = []
    pos_name = _first_match_position(normalized, _TABLE_HEADER_NAME_PATTERNS)
    pos_qty = _first_match_position(normalized, _TABLE_HEADER_QTY_PATTERNS)
    pos_pu = _first_match_position(normalized, _TABLE_HEADER_UNIT_PRICE_PATTERNS)
    pos_amount = _first_match_position(normalized, _TABLE_HEADER_AMOUNT_PATTERNS)
    if pos_name is not None:
        positions.append((pos_name, "designation"))
    if pos_qty is not None:
        positions.append((pos_qty, "qte"))
    if pos_pu is not None:
        positions.append((pos_pu, "pu"))
    if pos_amount is not None:
        positions.append((pos_amount, "montant"))
    positions.sort(key=lambda x: x[0])
    return [label for _pos, label in positions]


def _detect_table_end_rule(lines: list[str]) -> str:
    for line in lines:
        upper = _collapse_spaces(line).upper()
        if upper.startswith("TOTAL"):
            return "total"
        if upper.startswith("TVA"):
            return "tva"
        if upper.startswith("NET"):
            return "net"
    return "eof"


def _detect_table_number_style(lines: list[str]) -> str:
    comma_decimal = 0
    dot_decimal = 0
    spaced_thousands = 0
    for line in lines:
        comma_decimal += len(re.findall(r"\d+,\d{1,3}\b", line))
        dot_decimal += len(re.findall(r"\d+\.\d{1,3}\b", line))
        spaced_thousands += len(re.findall(r"\b\d{1,3}\s\d{3}\b", line))
    if comma_decimal and dot_decimal:
        base = "mixed_decimal"
    elif comma_decimal:
        base = "comma_decimal"
    elif dot_decimal:
        base = "dot_decimal"
    else:
        base = "integer_only"
    if spaced_thousands:
        base = f"{base}_spaced_thousands"
    return base


_LOCAL_ROW_START_RE = re.compile(r"^\s*(\d{1,3})[.)-]?\s+(.*)$")
_LOCAL_NUMBER_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:\d{1,3}(?:[ \u00A0]\d{3})+|\d+)(?:[.,]\d+)?(?![A-Za-z0-9])"
)
_LOCAL_QTY_CAPTURE_RE = re.compile(
    r"\b(?:U|QTE|QTY|QT)\b\s*([0-9]+(?:[.,][0-9]+)?)",
    flags=re.IGNORECASE,
)


def _normalize_local_row_text(value: str) -> str:
    text = (value or "").replace("\u00A0", " ")
    # OCR often glues "...9 915 966.39AUTOCLAVE..."; split at numeric -> all-caps boundary.
    text = re.sub(r"(?<=\d)(?=[A-Z]{2,})", " ", text)
    return _collapse_spaces(text)


def _remove_text_spans(text: str, spans: list[tuple[int, int]]) -> str:
    if not spans:
        return _collapse_spaces(text)
    merged: list[tuple[int, int]] = []
    for start, end in sorted(spans, key=lambda x: x[0]):
        if end <= start:
            continue
        if not merged or start > merged[-1][1]:
            merged.append((start, end))
        else:
            prev_start, prev_end = merged[-1]
            merged[-1] = (prev_start, max(prev_end, end))
    parts: list[str] = []
    cursor = 0
    for start, end in merged:
        parts.append(text[cursor:start])
        cursor = end
    parts.append(text[cursor:])
    return _collapse_spaces(" ".join(parts))


def _parse_local_inline_row(line: str) -> dict[str, float | str | None] | None:
    m = _LOCAL_ROW_START_RE.match(line)
    if not m:
        return None
    body = _normalize_local_row_text(m.group(2))
    if not body or not re.search(r"[A-Za-z]", body):
        return None

    number_matches = list(_LOCAL_NUMBER_RE.finditer(body))
    if len(number_matches) < 2:
        return None

    amount_match = number_matches[-1]
    unit_price_match = number_matches[-2]
    amount_raw = amount_match.group(0)
    unit_price_raw = unit_price_match.group(0)
    amount_value = _to_float(amount_raw)
    unit_price_value = _to_float(unit_price_raw)

    qty_value: float | None = None
    qty_raw: str | None = None
    qty_span: tuple[int, int] | None = None
    marker_match = _LOCAL_QTY_CAPTURE_RE.search(body)
    if marker_match is not None:
        qty_candidate_raw = marker_match.group(1)
        candidate_qty = _to_float(qty_candidate_raw)
        span = marker_match.span(1)
        overlaps_amount = not (span[1] <= amount_match.start() or span[0] >= amount_match.end())
        overlaps_unit = not (
            span[1] <= unit_price_match.start() or span[0] >= unit_price_match.end()
        )
        if candidate_qty is not None and not overlaps_amount and not overlaps_unit:
            qty_value = candidate_qty
            qty_raw = qty_candidate_raw
            qty_span = span

    if qty_value is None and len(number_matches) >= 3:
        fallback_qty_match = number_matches[-3]
        fallback_qty_raw = fallback_qty_match.group(0)
        fallback_qty = _to_float(fallback_qty_raw)
        if fallback_qty is not None and fallback_qty > 0:
            qty_value = fallback_qty
            qty_raw = fallback_qty_raw
            qty_span = fallback_qty_match.span()

    resolved_unit_price = _resolve_unit_price(
        qty_value=qty_value,
        unit_price_value=unit_price_value,
        amount_value=amount_value,
    )

    spans_to_remove = [amount_match.span(), unit_price_match.span()]
    if qty_span is not None:
        spans_to_remove.append(qty_span)
    description = _remove_text_spans(body, spans_to_remove)
    description = re.sub(r"\b(?:U|QTE|QTY|QT)\b", " ", description, flags=re.IGNORECASE)
    description = _collapse_spaces(description.strip(" -:|/"))

    if not description or not re.search(r"[A-Za-z]", description):
        return None

    return {
        "description": description,
        "quantity": qty_value,
        "quantity_raw": qty_raw,
        "unit_price": resolved_unit_price,
        "unit_price_raw": unit_price_raw,
        "amount": amount_value,
        "amount_raw": amount_raw,
        "row_mode": "inline",
    }


def _detect_row_marker(lines: list[str]) -> str:
    marker_u = 0
    marker_qte = 0
    for line in lines:
        if re.search(r"\bU\b\s*[0-9]+(?:[.,][0-9]+)?", line, flags=re.IGNORECASE):
            marker_u += 1
        if re.search(r"\bQTE?\b\s*[0-9]+(?:[.,][0-9]+)?", line, flags=re.IGNORECASE):
            marker_qte += 1
    if marker_u == 0 and marker_qte == 0:
        return "none"
    if marker_u >= marker_qte:
        return "u"
    return "qte"


def _detect_row_layout(lines: list[str]) -> str:
    inline_hints = 0
    block_hints = 0
    for idx, line in enumerate(lines):
        if not _LOCAL_ROW_START_RE.match(line):
            continue
        body = _LOCAL_ROW_START_RE.match(line).group(2) if _LOCAL_ROW_START_RE.match(line) else ""
        if len(list(_LOCAL_NUMBER_RE.finditer(_normalize_local_row_text(body)))) >= 2:
            inline_hints += 1
        if idx + 3 < len(lines):
            qv = _to_float(lines[idx + 1])
            uv = _to_float(lines[idx + 2])
            av = _to_float(lines[idx + 3])
            if qv is not None and uv is not None and av is not None:
                block_hints += 1
    if inline_hints > block_hints and inline_hints > 0:
        return "inline"
    if block_hints > 0:
        return "block"
    return "unknown"


def _detect_signature(text: str) -> str | None:
    if not text:
        return None
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for idx, line in enumerate(lines):
        flags = _detect_table_header_flags(line)
        if not (flags["name"] and flags["qty"]):
            continue
        table_window: list[str] = []
        for tline in lines[idx + 1 :]:
            upper = _collapse_spaces(tline).upper()
            if upper.startswith(("TOTAL", "TVA", "NET")):
                table_window.append(tline)
                break
            table_window.append(tline)
        header_order = _infer_header_order(line) or ["designation", "qte"]
        row_layout = _detect_row_layout(table_window[:40])
        marker = _detect_row_marker(table_window[:40])
        number_style = _detect_table_number_style(table_window[:80])
        end_rule = _detect_table_end_rule(table_window)
        return (
            "table_v2"
            f"|cols={','.join(header_order)}"
            f"|row={row_layout}"
            f"|marker={marker}"
            f"|num={number_style}"
            f"|end={end_rule}"
        )
    return None


def _parse_local_table(text: str) -> tuple[list[dict[str, float | str | None]], float]:
    if not text:
        return [], 0.0
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    items: list[dict[str, float | str | None]] = []
    i = 0
    in_table = False
    candidate_rows = 0
    while i < len(lines):
        line = lines[i]
        if not in_table and _is_table_header_line(line):
            in_table = True
            i += 1
            continue
        if not in_table:
            i += 1
            continue
        if line.upper().startswith(("TOTAL", "TVA", "NET")):
            break
        m = _LOCAL_ROW_START_RE.match(line)
        if m:
            candidate_rows += 1
            inline_row = _parse_local_inline_row(line)
            if inline_row is not None:
                items.append(inline_row)
                i += 1
                continue

            desc = _collapse_spaces(m.group(2))
            if re.search(r"[A-Za-z]", desc):
                qty_raw = lines[i + 1] if i + 1 < len(lines) else None
                unit_price_raw = lines[i + 2] if i + 2 < len(lines) else None
                amount_raw = lines[i + 3] if i + 3 < len(lines) else None
                qty = _to_float(qty_raw) if qty_raw is not None else None
                unit_price = _to_float(unit_price_raw) if unit_price_raw is not None else None
                amount = _to_float(amount_raw) if amount_raw is not None else None
                resolved_unit_price = _resolve_unit_price(
                    qty_value=qty,
                    unit_price_value=unit_price,
                    amount_value=amount,
                )
                items.append(
                    {
                        "description": desc,
                        "quantity": qty,
                        "quantity_raw": qty_raw,
                        "unit_price": resolved_unit_price,
                        "unit_price_raw": unit_price_raw,
                        "amount": amount,
                        "amount_raw": amount_raw,
                        "row_mode": "block",
                    }
                )
                i += 4
                continue
        i += 1
    if not in_table:
        return items, 0.0
    score = len(items) / max(candidate_rows, 1)
    return items, score


def _is_personal_sender_email(sender_email: str | None) -> bool:
    if not sender_email or "@" not in sender_email:
        return False
    domain = sender_email.rsplit("@", 1)[1].strip().lower()
    return domain in _PERSONAL_EMAIL_DOMAINS


def _looks_like_inline_table_noise(desc: str | None) -> bool:
    text = _collapse_spaces(desc or "")
    if not text:
        return True
    if _LOCAL_ROW_INLINE_PRICES_RE.search(text):
        return True
    if _LOCAL_ROW_INLINE_QTY_MARKER_RE.search(text):
        return True
    return False


def _amount_consistent(qty: float | None, unit_price: float | None, amount: float | None) -> bool:
    if qty is None or unit_price is None or amount is None:
        return False
    if qty <= 0 or unit_price <= 0 or amount <= 0:
        return False
    expected = qty * unit_price
    rel_err = abs(expected - amount) / max(abs(amount), 1.0)
    return rel_err <= 0.08


def _is_local_parse_reliable(
    parsed: list[dict[str, float | str | None]],
) -> tuple[bool, dict[str, float | int]]:
    if not parsed:
        return False, {
            "rows": 0,
            "qty_rate": 0.0,
            "price_rate": 0.0,
            "amount_rate": 0.0,
            "consistency_rate": 0.0,
            "noisy_rate": 1.0,
        }
    total = len(parsed)
    with_qty = 0
    with_price = 0
    with_amount = 0
    noisy_desc = 0
    comparable = 0
    consistent = 0
    for row in parsed:
        desc = str(row.get("description") or "")
        qty = row.get("quantity")
        pu = row.get("unit_price")
        amount = row.get("amount")
        if isinstance(qty, (int, float)) and qty > 0:
            with_qty += 1
        if isinstance(pu, (int, float)) and pu > 0:
            with_price += 1
        if isinstance(amount, (int, float)) and amount > 0:
            with_amount += 1
        if _looks_like_inline_table_noise(desc):
            noisy_desc += 1
        if _amount_consistent(
            float(qty) if isinstance(qty, (int, float)) else None,
            float(pu) if isinstance(pu, (int, float)) else None,
            float(amount) if isinstance(amount, (int, float)) else None,
        ):
            consistent += 1
        if (
            isinstance(qty, (int, float))
            and isinstance(pu, (int, float))
            and isinstance(amount, (int, float))
            and qty > 0
            and pu > 0
            and amount > 0
        ):
            comparable += 1

    qty_rate = with_qty / total
    price_rate = with_price / total
    amount_rate = with_amount / total
    noisy_rate = noisy_desc / total
    consistency_rate = (consistent / comparable) if comparable else 0.0
    consistency_ok = True
    if comparable >= 2:
        consistency_ok = consistency_rate >= 0.60
    elif comparable == 1:
        consistency_ok = consistency == 1

    reliable = (
        qty_rate >= 0.60
        and price_rate >= 0.60
        and amount_rate >= 0.60
        and noisy_rate <= 0.30
        and consistency_ok
    )
    return reliable, {
        "rows": total,
        "qty_rate": round(qty_rate, 3),
        "price_rate": round(price_rate, 3),
        "amount_rate": round(amount_rate, 3),
        "consistency_rate": round(consistency_rate, 3),
        "noisy_rate": round(noisy_rate, 3),
        "comparable_rows": comparable,
    }


def _store_supplier_format(
    db: Session | None, tenant_id: int, sender_email: str | None, signature: str | None
) -> None:
    if db is None or not sender_email or not signature:
        return
    # Sender-based format persistence is unsafe for personal inboxes
    # (same mailbox receives heterogeneous supplier layouts).
    if _is_personal_sender_email(sender_email):
        return
    sender_email_norm = sender_email.strip().lower()
    signature_norm = _collapse_spaces(signature)
    if not sender_email_norm or not signature_norm:
        return

    # SessionLocal uses autoflush=False: guard against duplicate pending rows for the
    # same (tenant_id, sender_email) before we hit the DB unique constraint on commit.
    for pending in db.new:
        if (
            isinstance(pending, SupplierInvoiceFormat)
            and pending.tenant_id == tenant_id
            and (pending.sender_email or "").strip().lower() == sender_email_norm
        ):
            if pending.signature != signature_norm:
                pending.signature = signature_norm
            return
    existing = (
        db.query(SupplierInvoiceFormat)
        .filter(
            SupplierInvoiceFormat.tenant_id == tenant_id,
            SupplierInvoiceFormat.sender_email == sender_email_norm,
        )
        .first()
    )
    if existing:
        if existing.signature != signature_norm:
            existing.signature = signature_norm
        return
    fmt = SupplierInvoiceFormat(
        tenant_id=tenant_id,
        sender_email=sender_email_norm,
        signature=signature_norm,
        parser_type="table_v2",
    )
    db.add(fmt)


def _is_spreadsheet_invoice_file(file_path: str | None) -> bool:
    suffix = Path(file_path or "").suffix.lower()
    return suffix in _SPREADSHEET_INVOICE_EXTENSIONS


def _convert_spreadsheet_to_pdf(file_path: str | None) -> str | None:
    source = Path(file_path or "")
    if not source.exists() or not source.is_file():
        return None
    if source.suffix.lower() not in _SPREADSHEET_INVOICE_EXTENSIONS:
        return None

    configured_bin = str(getattr(config, "LIBREOFFICE_BIN", "") or "").strip() or str(
        os.getenv("LIBREOFFICE_BIN") or ""
    ).strip()
    soffice_bin = configured_bin or shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice_bin:
        logger.warning("Spreadsheet->PDF conversion unavailable (LibreOffice/soffice missing): %s", source)
        return None

    pdf_path = source.with_suffix(".pdf")
    try:
        if pdf_path.exists() and pdf_path.stat().st_mtime >= source.stat().st_mtime:
            return str(pdf_path)
    except Exception:
        pass

    try:
        result = subprocess.run(
            [
                soffice_bin,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(source.parent),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
    except Exception as exc:
        logger.warning("Spreadsheet->PDF conversion failed: file=%s err=%s", source, exc)
        return None

    if result.returncode != 0:
        detail = (result.stderr or "").strip() or (result.stdout or "").strip() or f"exit={result.returncode}"
        logger.warning("Spreadsheet->PDF conversion failed: file=%s err=%s", source, detail)
        return None
    if not pdf_path.exists():
        logger.warning("Spreadsheet->PDF conversion failed: file=%s missing_output_pdf", source)
        return None
    return str(pdf_path)


def resolve_invoice_processing_file_path(file_path: str | None) -> str | None:
    if _is_spreadsheet_invoice_file(file_path):
        return _convert_spreadsheet_to_pdf(file_path) or file_path
    return file_path


def _normalize_xlsx_header(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    ascii_text = (
        unicodedata.normalize("NFKD", raw)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )
    return re.sub(r"[^a-z0-9]+", " ", ascii_text).strip()


def _xlsx_header_matches(header_norm: str, hints: set[str]) -> bool:
    if not header_norm:
        return False
    if header_norm in hints:
        return True
    for hint in hints:
        if hint and hint in header_norm:
            return True
    return False


def _xlsx_to_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    if "%" in text:
        return None
    text = text.replace("\u00a0", " ").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    if text.count(".") > 1:
        left, right = text.rsplit(".", 1)
        text = left.replace(".", "") + "." + right
    try:
        return float(text)
    except ValueError:
        return None


def _detect_xlsx_header_map(row: list[Any]) -> dict[str, Any] | None:
    mapping: dict[str, Any] = {}
    price_cols: list[int] = []
    for idx, cell in enumerate(row):
        header = _normalize_xlsx_header(cell)
        if not header:
            continue
        if "product_col" not in mapping and _xlsx_header_matches(
            header, _XLSX_PRODUCT_HEADER_HINTS
        ):
            mapping["product_col"] = idx
        if "ref_col" not in mapping and _xlsx_header_matches(
            header, _XLSX_REF_HEADER_HINTS
        ):
            mapping["ref_col"] = idx
        if "qty_col" not in mapping and _xlsx_header_matches(
            header, _XLSX_QTY_HEADER_HINTS
        ):
            mapping["qty_col"] = idx
        if "unit_col" not in mapping and _xlsx_header_matches(
            header, _XLSX_UNIT_HEADER_HINTS
        ):
            mapping["unit_col"] = idx
        if "amount_col" not in mapping and _xlsx_header_matches(
            header, _XLSX_AMOUNT_HEADER_HINTS
        ):
            mapping["amount_col"] = idx
        if _xlsx_header_matches(header, _XLSX_PRICE_HEADER_HINTS):
            price_cols.append(idx)
    if "product_col" not in mapping:
        return None
    if not price_cols and "amount_col" not in mapping:
        return None
    if price_cols:
        mapping["price_cols"] = sorted(set(price_cols))
    return mapping


def _parse_xlsx_sheet_items(
    invoice: InboundInvoice,
    worksheet: Any,
) -> list[InvoiceExtractionItem]:
    header_map: dict[str, Any] | None = None
    header_row_idx = 0
    max_probe_rows = 40
    for idx, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_probe_rows, values_only=True),
        start=1,
    ):
        probe = _detect_xlsx_header_map(list(row))
        if probe is not None:
            header_map = probe
            header_row_idx = idx
            break
    if header_map is None:
        return []

    product_col = int(header_map["product_col"])
    ref_col = header_map.get("ref_col")
    qty_col = header_map.get("qty_col")
    unit_col = header_map.get("unit_col")
    amount_col = header_map.get("amount_col")
    price_cols = [int(c) for c in (header_map.get("price_cols") or [])]

    items: list[InvoiceExtractionItem] = []
    empty_streak = 0
    for row in worksheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        values = list(row)

        def _cell(col_idx: Any) -> Any:
            if not isinstance(col_idx, int):
                return None
            return values[col_idx] if 0 <= col_idx < len(values) else None

        product_raw = _cell(product_col)
        ref_raw = _cell(ref_col)
        qty_raw = _cell(qty_col)
        unit_raw = _cell(unit_col)
        amount_raw = _cell(amount_col)

        product_text = _collapse_spaces(str(product_raw or ""))
        ref_text = _collapse_spaces(str(ref_raw or ""))
        qty_value = _xlsx_to_float(qty_raw)
        amount_value = _xlsx_to_float(amount_raw)

        unit_price_value: float | None = None
        unit_price_raw: Any = None
        for pcol in price_cols:
            candidate_raw = _cell(pcol)
            candidate = _xlsx_to_float(candidate_raw)
            if candidate is not None and candidate > 0:
                # Keep the rightmost non-empty price-like cell (often discounted/final price).
                unit_price_value = candidate
                unit_price_raw = candidate_raw

        has_signal = bool(
            product_text
            or ref_text
            or (unit_price_value is not None and unit_price_value > 0)
            or (amount_value is not None and amount_value > 0)
        )
        if not has_signal:
            empty_streak += 1
            if empty_streak >= 8 and items:
                break
            continue
        empty_streak = 0

        if not product_text:
            product_text = ref_text
        if not product_text:
            continue
        # Skip section/category headers that do not carry line-level evidence.
        if (
            not ref_text
            and not (unit_price_value is not None and unit_price_value > 0)
            and not (amount_value is not None and amount_value > 0)
            and not (qty_value is not None and qty_value > 0)
        ):
            continue

        qty_norm = qty_value if (qty_value is not None and qty_value > 0) else None
        resolved_unit_price, price_source = _resolve_unit_price_with_source(
            qty_value=qty_norm,
            unit_price_value=unit_price_value,
            amount_value=amount_value,
        )

        unit_value = _collapse_spaces(str(unit_raw or "")) or None
        raw_desc = product_text
        desc_parts: list[str] = []
        if ref_text:
            desc_parts.append(f"ref: {ref_text}")
            raw_desc = f"{product_text} | REF {ref_text}"
        description_text = " | ".join(desc_parts) if desc_parts else None

        item = InvoiceExtractionItem(
            tenant_id=invoice.tenant_id,
            invoice_id=invoice.id,
            campaign_id=invoice.campaign_id,
            supplier_id=None,
            source_engine="xlsx_parser",
            status=ExtractionStatus.TO_VERIFY,
            product_name=_derive_product_name(product_text, ref_text)[:255],
            description=description_text,
            raw_description=raw_desc,
            unit=unit_value,
            quantity=qty_norm,
            unit_price=resolved_unit_price,
            currency="DZD",
            confidence=1.0,
            notes=_compose_price_notes(
                price_source=price_source,
                amount_value=amount_value,
                qty_raw=qty_raw,
                unit_price_raw=unit_price_raw,
                line_total_raw=amount_raw,
                extra_tokens=[f"sheet={_collapse_spaces(str(worksheet.title or '')).replace(' ', '_')}"],
            ),
        )
        items.append(item)
    return items


def _extract_line_items_from_spreadsheet(
    invoice: InboundInvoice,
) -> list[InvoiceExtractionItem]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        logger.warning(
            "XLSX invoice parser unavailable (openpyxl not installed): invoice=%s file=%s",
            invoice.id,
            invoice.file_path,
        )
        return []

    try:
        wb = load_workbook(invoice.file_path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning(
            "XLSX load failed: invoice=%s file=%s err=%s",
            invoice.id,
            invoice.file_path,
            exc,
        )
        return []

    items: list[InvoiceExtractionItem] = []
    try:
        for ws in wb.worksheets:
            items.extend(_parse_xlsx_sheet_items(invoice, ws))
    finally:
        try:
            wb.close()
        except Exception:
            pass

    if items:
        with_price = sum(1 for row in items if row.unit_price is not None and row.unit_price > 0)
        invoice.parse_engine = "xlsx_parser"
        invoice.parse_score = round(with_price / max(len(items), 1), 4)
    else:
        invoice.parse_engine = "xlsx_parser"
        invoice.parse_score = 0.0
    return items


def extract_line_items(invoice: InboundInvoice, db: Session | None = None) -> list[InvoiceExtractionItem]:
    processing_file_path = resolve_invoice_processing_file_path(invoice.file_path) or invoice.file_path
    if _is_spreadsheet_invoice_file(invoice.file_path):
        spreadsheet_items = _extract_line_items_from_spreadsheet(invoice)
        if spreadsheet_items:
            return spreadsheet_items
        if processing_file_path == invoice.file_path:
            return spreadsheet_items

    # Try local parser based on supplier email + signature
    text = _extract_pdf_text(processing_file_path)
    if text and not str(getattr(invoice, "ocr_text_snapshot", "") or "").strip():
        invoice.ocr_text_snapshot = str(text).strip()
    signature = _detect_signature(text) if text else None
    if (
        db is not None
        and invoice.sender_email
        and signature
        and not _is_personal_sender_email(invoice.sender_email)
    ):
        fmt = (
            db.query(SupplierInvoiceFormat)
            .filter(
                SupplierInvoiceFormat.tenant_id == invoice.tenant_id,
                SupplierInvoiceFormat.sender_email == invoice.sender_email.lower(),
            )
            .first()
        )
        if fmt and fmt.signature == signature:
            parsed, score = _parse_local_table(text or "")
            is_reliable, quality = _is_local_parse_reliable(parsed)
            if parsed and score >= config.LOCAL_PARSER_MIN_SCORE and is_reliable:
                invoice.parse_engine = "local_parser"
                invoice.parse_score = score
                items: list[InvoiceExtractionItem] = []
                for row in parsed:
                    desc = str(row.get("description") or "")
                    qty = row.get("quantity")
                    unit_price = row.get("unit_price")
                    qty_value = float(qty) if isinstance(qty, (int, float)) else None
                    unit_price_value = (
                        float(unit_price) if isinstance(unit_price, (int, float)) else None
                    )
                    name = _derive_product_name(desc)
                    description_clean = _clean_description_text(desc)
                    item = InvoiceExtractionItem(
                        tenant_id=invoice.tenant_id,
                        invoice_id=invoice.id,
                        campaign_id=invoice.campaign_id,
                        supplier_id=None,
                        source_engine="local_parser",
                        status=ExtractionStatus.TO_VERIFY,
                        product_name=name[:255],
                        description=description_clean,
                        raw_description=desc,
                        unit=None,
                        quantity=qty_value,
                        unit_price=unit_price_value,
                        currency="DZD",
                        confidence=score,
                        notes=_compose_price_notes(
                            price_source="local_table",
                            amount_value=float(row.get("amount")) if isinstance(row.get("amount"), (int, float)) else None,
                            qty_raw=row.get("quantity_raw"),
                            unit_price_raw=row.get("unit_price_raw"),
                            line_total_raw=row.get("amount_raw"),
                        ),
                    )
                    items.append(item)
                return items
            if parsed and score >= config.LOCAL_PARSER_MIN_SCORE:
                logger.warning(
                    (
                        "Local parser rejected by quality gate: "
                        "invoice=%s sender=%s rows=%s score=%.3f qty_rate=%.3f "
                        "price_rate=%.3f amount_rate=%.3f consistency_rate=%.3f noisy_rate=%.3f comparable=%s"
                    ),
                    invoice.id,
                    invoice.sender_email,
                    len(parsed),
                    score,
                    float(quality.get("qty_rate", 0.0)),
                    float(quality.get("price_rate", 0.0)),
                    float(quality.get("amount_rate", 0.0)),
                    float(quality.get("consistency_rate", 0.0)),
                    float(quality.get("noisy_rate", 0.0)),
                    int(quality.get("comparable_rows", 0)),
                )

    # Fallback: Document AI
    doc_ai_v2_enabled = bool(getattr(config, "DOC_AI_FALLBACK_V2_ENABLED", False))
    if doc_ai_v2_enabled:
        # Safe rollout: keep current extraction path until v2 logic is explicitly introduced.
        pass
    doc = process_invoice(processing_file_path)
    doc_text = str(getattr(doc, "text", None) or "").strip()
    if doc_text and not str(getattr(invoice, "ocr_text_snapshot", "") or "").strip():
        invoice.ocr_text_snapshot = doc_text
    supplier_name, supplier_phone = _extract_supplier_info(doc)
    ocr_lines = _doc_text_lines(doc)
    if supplier_name:
        invoice.sender_name = supplier_name
    if supplier_phone:
        invoice.sender_phone = supplier_phone
    items: list[InvoiceExtractionItem] = []
    pending_orphan_desc: str | None = None
    pending_orphan_code: str | None = None
    pending_orphan_conf: float | None = None

    for entity in _iter_line_items(doc):
        description = _get_prop(entity, "description")
        product_code = _get_prop(entity, "product_code")
        qty = _get_prop(entity, "quantity")
        unit_price = _get_prop(entity, "unit_price")
        amount = _get_prop(entity, "amount")
        unit = _get_prop(entity, "unit")

        qty_raw = _entity_value(qty)
        unit_price_raw = _entity_value(unit_price)
        amount_raw = _entity_value(amount)
        qty_value = _to_float(qty_raw)
        unit_price_value = _to_float(unit_price_raw)
        desc_text = _entity_value(description)
        code_text = _entity_value(product_code)
        amount_value = _to_float(amount_raw)

        # Only keep lines that have an amount (line_item/amount).
        # Some invoices split amount and description across adjacent line_items:
        # keep the description as "pending orphan" and merge it into nearest line with amount.
        if amount_value is None:
            orphan_text = _expand_orphan_description(desc_text or code_text, ocr_lines)
            if orphan_text:
                if items and (
                    items[-1].product_name == "UNKNOWN"
                    or not _collapse_spaces(items[-1].raw_description or "")
                ):
                    merged_raw = _merge_text_parts(items[-1].raw_description, orphan_text)
                    if merged_raw:
                        items[-1].raw_description = merged_raw
                        merged_name = _derive_product_name(merged_raw, code_text)
                        items[-1].product_name = (merged_name or "UNKNOWN")[:255]
                        items[-1].description = _clean_description_text(merged_raw, code_text)
                        prev_conf = items[-1].confidence or 0.0
                        items[-1].confidence = (prev_conf + (entity.confidence or 0.0)) / 2.0
                else:
                    pending_orphan_desc = orphan_text
                    pending_orphan_code = code_text
                    pending_orphan_conf = entity.confidence
            continue

        if desc_text:
            name = _derive_product_name(desc_text, code_text)
            description_text = _clean_description_text(desc_text, code_text)
            raw_description_text = desc_text
        elif code_text:
            name = _collapse_spaces(code_text)
            description_text = None
            raw_description_text = None
        else:
            name = "UNKNOWN"
            description_text = None
            raw_description_text = None
        if len(name) > 255:
            name = name[:255]

        confidence_value = entity.confidence
        if pending_orphan_desc and name == "UNKNOWN":
            merged_raw = _merge_text_parts(raw_description_text, pending_orphan_desc)
            merged_code = code_text or pending_orphan_code
            if merged_raw:
                raw_description_text = merged_raw
                name = _derive_product_name(merged_raw, merged_code)[:255]
                description_text = _clean_description_text(merged_raw, merged_code)
                confidence_value = (
                    ((entity.confidence or 0.0) + (pending_orphan_conf or 0.0)) / 2.0
                )
        pending_orphan_desc = None
        pending_orphan_code = None
        pending_orphan_conf = None

        resolved_unit_price, resolved_price_source = _resolve_unit_price_with_source(
            qty_value=qty_value,
            unit_price_value=unit_price_value,
            amount_value=amount_value,
        )

        item = InvoiceExtractionItem(
            tenant_id=invoice.tenant_id,
            invoice_id=invoice.id,
            campaign_id=invoice.campaign_id,
            supplier_id=None,
            source_engine="document_ai",
            status=ExtractionStatus.TO_VERIFY,
            product_name=name,
            description=description_text,
            raw_description=raw_description_text,
            unit=_entity_value(unit),
            quantity=qty_value,
            unit_price=resolved_unit_price,
            currency="DZD",
            confidence=confidence_value,
            notes=_compose_price_notes(
                price_source=resolved_price_source,
                amount_value=amount_value,
                qty_raw=qty_raw,
                unit_price_raw=unit_price_raw,
                line_total_raw=amount_raw,
            ),
        )
        items.append(item)

    # store format signature for this supplier
    _store_supplier_format(db, invoice.tenant_id, invoice.sender_email, signature)
    if not items and _is_spreadsheet_invoice_file(invoice.file_path):
        fallback_rows, fallback_model = _extract_invoice_rows_with_gpt_fallback(invoice)
        if fallback_rows:
            items = _build_items_from_gpt_invoice_rows(
                invoice,
                fallback_rows,
                model=fallback_model,
            )
            invoice.parse_engine = "gpt_invoice_fallback"
            invoice.parse_score = 0.85
            return items
    if items:
        avg_conf = sum((it.confidence or 0.0) for it in items) / max(len(items), 1)
        invoice.parse_engine = "document_ai"
        invoice.parse_score = avg_conf
    return items


def _normalize_text(text: str) -> str:
    return " ".join(text.strip().lower().split())


def _hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

_UNIT_ALIASES = {
    "mg": "mg",
    "g": "g",
    "kg": "kg",
    "ml": "ml",
    "l": "l",
    "cc": "ml",
    "mm": "mm",
    "cm": "cm",
    "m": "m",
    "w": "w",
    "kw": "kw",
    "v": "v",
    "Â°c": "c",
    "c": "c",
    "bar": "bar",
    "pcs": "pcs",
    "piece": "pcs",
    "piÃ¨ce": "pcs",
    "unite": "pcs",
    "unitÃ©": "pcs",
    "boite": "box",
    "boÃ®te": "box",
    "pack": "pack",
    "lot": "lot",
    "h": "h",
    "min": "min",
    "s": "s",
}

_SYNONYMS = {
    "vitamine c": ["acide ascorbique", "ascorbique"],
    "video laryngoscope": ["videolaryngoscope", "laryngoscope video"],
    "otoscope led": ["otoscope lumiÃ¨re led", "otoscope a led"],
    "nebuliseur compact": ["nÃ©buliseur compact", "nebuliseur"],
    "mano detendeur": ["manodÃ©tendeur", "mano-dÃ©tendeur", "manodetendeur"],
    "paravent": ["paravent 3 volets", "paravent trois volets"],
    "gueridon": ["guÃ©ridon", "gueridon 3 etages", "gueridon 3 Ã©tages"],
}


def _normalize_name(text: str) -> str:
    value = text.lower()
    value = re.sub(r"\\s+", " ", value).strip()
    for canonical, variants in _SYNONYMS.items():
        for v in variants:
            if v in value:
                value = value.replace(v, canonical)
    return value


def _extract_features(text: str | None) -> set[str]:
    if not text:
        return set()
    value = text.lower().replace(",", ".")
    features: set[str] = set()
    # normalize numeric + unit with simple conversions
    unit_conv = {
        # weight -> g
        "mg": ("g", 0.001),
        "g": ("g", 1.0),
        "kg": ("g", 1000.0),
        # volume -> l
        "ml": ("l", 0.001),
        "cl": ("l", 0.01),
        "l": ("l", 1.0),
        "cc": ("l", 0.001),
        # length -> m
        "mm": ("m", 0.001),
        "cm": ("m", 0.01),
        "m": ("m", 1.0),
        # time
        "s": ("s", 1.0),
        "min": ("s", 60.0),
        "h": ("s", 3600.0),
        # temperature
        "c": ("c", 1.0),
    }
    for match in re.finditer(r"(\d+(?:\.\d+)?)\s*([a-zÂ°]+)", value):
        num = match.group(1)
        unit = _UNIT_ALIASES.get(match.group(2), match.group(2))
        try:
            fnum = float(num)
        except ValueError:
            continue
        if unit in unit_conv:
            base_unit, factor = unit_conv[unit]
            fnum = fnum * factor
            features.add(f"{fnum:.6g}{base_unit}")
        else:
            features.add(f"{fnum:.6g}{unit}")
    # class patterns like "classe b"
    for match in re.finditer(r"classe\s*([a-z])", value):
        features.add(f"classe_{match.group(1)}")
    return features


def _norm_key(value: str | None) -> str:
    return _canonical_key((value or "").strip()).strip().lower()


def _norm_value(value: str | None) -> str:
    return " ".join((value or "").strip().lower().split())


def _append_unique_value(target: dict[str, list[str]], key: str, value: str | None) -> None:
    key_norm = _norm_key(key)
    val = (value or "").strip()
    if not key_norm or not val:
        return
    existing = target.setdefault(key_norm, [])
    val_norm = _norm_value(val)
    for it in existing:
        if _norm_value(it) == val_norm:
            return
    existing.append(val)


def _flatten_preview_characteristics(item: dict[str, Any]) -> dict[str, list[str]]:
    merged: dict[str, list[str]] = {}
    requested_chars = item.get("requested_characteristics")
    if isinstance(requested_chars, dict):
        for key, raw in requested_chars.items():
            if isinstance(raw, list):
                for val in raw:
                    _append_unique_value(merged, str(key), str(val))
            elif raw is not None:
                _append_unique_value(merged, str(key), str(raw))

    regex_chars = item.get("regex_characteristics")
    if isinstance(regex_chars, dict):
        for key, raw in regex_chars.items():
            if isinstance(raw, list):
                for val in raw:
                    _append_unique_value(merged, str(key), str(val))
            elif raw is not None:
                _append_unique_value(merged, str(key), str(raw))

    emb_chars = item.get("embedding_characteristics")
    if isinstance(emb_chars, dict):
        for key, raw in emb_chars.items():
            if isinstance(raw, list):
                for row in raw:
                    if isinstance(row, dict):
                        _append_unique_value(merged, str(key), str(row.get("value") or ""))
                    elif row is not None:
                        _append_unique_value(merged, str(key), str(row))
            elif isinstance(raw, dict):
                _append_unique_value(merged, str(key), str(raw.get("value") or ""))
            elif raw is not None:
                _append_unique_value(merged, str(key), str(raw))

    return merged


def _flatten_generic_characteristics_payload(payload: Any) -> dict[str, list[str]]:
    merged: dict[str, list[str]] = {}
    if payload is None:
        return merged

    def _merge_kv_map(raw_map: dict[str, Any]) -> None:
        for key, raw in raw_map.items():
            if isinstance(raw, list):
                for row in raw:
                    if isinstance(row, dict):
                        if "value" in row:
                            _append_unique_value(merged, str(key), str(row.get("value") or ""))
                        else:
                            for rv in row.values():
                                if rv is not None:
                                    _append_unique_value(merged, str(key), str(rv))
                    elif row is not None:
                        _append_unique_value(merged, str(key), str(row))
            elif isinstance(raw, dict):
                if "value" in raw:
                    _append_unique_value(merged, str(key), str(raw.get("value") or ""))
                else:
                    for rv in raw.values():
                        if rv is not None:
                            _append_unique_value(merged, str(key), str(rv))
            elif raw is not None:
                _append_unique_value(merged, str(key), str(raw))

    def _merge_key_value_list(raw_list: list[Any]) -> None:
        for row in raw_list:
            if not isinstance(row, dict):
                continue
            key = row.get("key")
            if not key:
                continue
            value = row.get("value")
            values = row.get("values")
            if isinstance(values, list):
                for v in values:
                    if v is not None:
                        _append_unique_value(merged, str(key), str(v))
                continue
            if value is not None:
                _append_unique_value(merged, str(key), str(value))

    if isinstance(payload, dict):
        # Common wrappers found in campaign_items, product.characteristics and parsed payloads.
        for wrapper in (
            "requested_characteristics",
            "regex_characteristics",
            "embedding_characteristics",
            "kv",
        ):
            raw = payload.get(wrapper)
            if isinstance(raw, dict):
                _merge_kv_map(raw)

        for wrapper in ("characteristics", "pairs"):
            raw = payload.get(wrapper)
            if isinstance(raw, list):
                _merge_key_value_list(raw)

        ignored_meta_keys = {
            "doc_type",
            "product_name",
            "score",
            "score_regex_only",
            "text_features",
            "ignored_items",
            "kv",
            "pairs",
            "characteristics",
            "requested_characteristics",
            "regex_characteristics",
            "embedding_characteristics",
        }
        plain_map = {
            k: v
            for k, v in payload.items()
            if k not in ignored_meta_keys and not k.endswith("_characteristics")
        }
        if plain_map:
            _merge_kv_map(plain_map)
    elif isinstance(payload, list):
        _merge_key_value_list(payload)

    return merged


def _collect_campaign_requested_characteristics(
    db: Session,
    tenant_id: int,
    campaign_id: int,
) -> list[dict[str, Any]]:
    rows = (
        db.query(CampaignAttachment)
        .filter(
            CampaignAttachment.tenant_id == tenant_id,
            CampaignAttachment.campaign_id == campaign_id,
            CampaignAttachment.extraction_preview.isnot(None),
        )
        .all()
    )
    aggregated: dict[str, dict[str, Any]] = {}
    for row in rows:
        preview = row.extraction_preview if isinstance(row.extraction_preview, dict) else {}
        doc_type = str(preview.get("doc_type") or "").strip().lower()
        if doc_type and doc_type not in {"requirement_sheet", "tech_sheet"}:
            continue
        items = preview.get("items") if isinstance(preview, dict) else None
        if not isinstance(items, list):
            continue
        for it in items:
            if not isinstance(it, dict):
                continue
            item_doc_type = str(it.get("doc_type") or doc_type or "").strip().lower()
            if item_doc_type and item_doc_type not in {"requirement_sheet", "tech_sheet"}:
                continue
            name = str(it.get("product_name") or "").strip()
            if not name:
                continue
            key = _norm_value(_normalize_name(name))
            if not key:
                continue
            kv = _flatten_preview_characteristics(it)
            current = aggregated.setdefault(
                key,
                {
                    "product_name": name,
                    "norm_name": key,
                    "characteristics": {},
                },
            )
            merged = current["characteristics"]
            for ck, vals in kv.items():
                for v in vals:
                    _append_unique_value(merged, ck, v)

    # Fallback: campaign items / product master data when attachment preview is absent.
    item_rows = (
        db.query(CampaignItem, Product)
        .join(Product, CampaignItem.product_id == Product.id)
        .filter(
            CampaignItem.tenant_id == tenant_id,
            CampaignItem.campaign_id == campaign_id,
            Product.tenant_id == tenant_id,
        )
        .all()
    )
    for item, product in item_rows:
        name = str(product.name or "").strip()
        if not name:
            continue
        key = _norm_value(_normalize_name(name))
        if not key:
            continue

        current = aggregated.setdefault(
            key,
            {
                "product_name": name,
                "norm_name": key,
                "characteristics": {},
            },
        )
        merged = current["characteristics"]
        requested_kv = _flatten_generic_characteristics_payload(item.requested_characteristics)
        product_kv = _flatten_generic_characteristics_payload(product.characteristics)
        for source_kv in (requested_kv, product_kv):
            for ck, vals in source_kv.items():
                for v in vals:
                    _append_unique_value(merged, ck, v)

    return list(aggregated.values())


def _select_requested_characteristics_for_product(
    product_name: str | None,
    requested_rows: list[dict[str, Any]],
) -> dict[str, list[str]]:
    name = _norm_value(_normalize_name(product_name or ""))
    if not name or not requested_rows:
        return {}

    exact = next((row for row in requested_rows if row.get("norm_name") == name), None)
    if exact is not None:
        return exact.get("characteristics") or {}

    best_row: dict[str, Any] | None = None
    best_score = -1.0
    for row in requested_rows:
        rn = str(row.get("norm_name") or "")
        if not rn:
            continue
        score = difflib.SequenceMatcher(None, name, rn).ratio()
        if score > best_score:
            best_score = score
            best_row = row
    if best_row is None or best_score < 0.55:
        return {}
    return best_row.get("characteristics") or {}


def _extract_offered_characteristics(
    db: Session,
    item: InvoiceExtractionItem,
    tenant_id: int,
    campaign_id: int,
) -> tuple[dict[str, list[str]], str]:
    # Merge offered characteristics from both sources with priority:
    # 1) linked tech sheet (primary)
    # 2) invoice line text (fallback/completion for missing keys only)
    tech_offered: dict[str, list[str]] = {}
    if item.id is not None:
        link = (
            db.query(TechSheetLink)
            .filter(
                TechSheetLink.tenant_id == tenant_id,
                TechSheetLink.campaign_id == campaign_id,
                TechSheetLink.invoice_item_id == item.id,
            )
            .order_by(TechSheetLink.score.desc())
            .first()
        )
        if link is not None and link.tech_sheet is not None:
            chars = link.tech_sheet.characteristics if isinstance(link.tech_sheet.characteristics, dict) else {}
            kv = chars.get("kv") if isinstance(chars, dict) else None
            if isinstance(kv, dict) and kv:
                for key, raw in kv.items():
                    if isinstance(raw, list):
                        for val in raw:
                            _append_unique_value(tech_offered, str(key), str(val))
                    elif raw is not None:
                        _append_unique_value(tech_offered, str(key), str(raw))
                text_features = chars.get("text_features") if isinstance(chars, dict) else None
                if isinstance(text_features, list):
                    for feat in text_features:
                        if feat:
                            _append_unique_value(tech_offered, "feature", str(feat))

    # Parse offered characteristics directly from invoice line text.
    invoice_offered: dict[str, list[str]] = {}
    text = f"{item.product_name or ''}\n{item.description or ''}".strip()
    if text:
        parsed = parse_characteristics(text)
        kv = parsed.get("kv") if isinstance(parsed, dict) else None
        if isinstance(kv, dict):
            for key, raw in kv.items():
                if isinstance(raw, list):
                    for val in raw:
                        _append_unique_value(invoice_offered, str(key), str(val))
                elif raw is not None:
                    _append_unique_value(invoice_offered, str(key), str(raw))
        text_features = parsed.get("text_features") if isinstance(parsed, dict) else None
        if isinstance(text_features, list):
            for feat in text_features:
                if feat:
                    _append_unique_value(invoice_offered, "feature", str(feat))

    if tech_offered and invoice_offered:
        merged: dict[str, list[str]] = {}
        # Primary source: tech sheet.
        for key, vals in tech_offered.items():
            for val in vals:
                _append_unique_value(merged, key, val)
        # Secondary source: invoice line (only missing keys).
        for key, vals in invoice_offered.items():
            key_norm = _norm_key(key)
            if key_norm in merged and merged[key_norm]:
                continue
            for val in vals:
                _append_unique_value(merged, key, val)
        return merged, "tech_sheet+invoice_line"

    if tech_offered:
        return tech_offered, "tech_sheet"
    if invoice_offered:
        return invoice_offered, "invoice_line"
    return {}, "invoice_line"


def _embed_texts_sync(texts: list[str]) -> dict[str, list[float]]:
    if not texts:
        return {}
    unique: list[str] = []
    seen: set[str] = set()
    for text in texts:
        t = (text or "").strip()
        if not t:
            continue
        key = _normalize_text(t)
        if key in seen:
            continue
        seen.add(key)
        unique.append(t)

    cached, missing, _ = get_cached_embeddings(unique)
    vectors: dict[str, list[float]] = dict(cached)
    if missing and config.OPENAI_API_KEY:
        unique_missing = list(dict.fromkeys(missing))
        fetched = fetch_embeddings(unique_missing)
        for text, vec in zip(unique_missing, fetched):
            vectors[_hash_text(_normalize_text(text))] = vec
    return vectors


def _compare_characteristics_with_embeddings_from_vectors(
    requested: dict[str, list[str]],
    offered: dict[str, list[str]],
    threshold: float,
    vectors: dict[str, list[float]],
) -> tuple[float | None, dict[str, Any] | None]:
    if not requested:
        return None, None
    points: list[dict[str, Any]] = []
    scores: list[float] = []
    matched = 0

    for req_key, req_vals in requested.items():
        req_values = [v for v in req_vals if (v or "").strip()]
        off_values = [v for v in offered.get(req_key, []) if (v or "").strip()]
        best_score = 0.0
        best_req = None
        best_off = None
        if req_values and off_values:
            for req_val in req_values:
                req_hash = _hash_text(_normalize_text(f"{req_key}: {req_val}"))
                req_vec = vectors.get(req_hash)
                if req_vec is None:
                    continue
                for off_val in off_values:
                    off_hash = _hash_text(_normalize_text(f"{req_key}: {off_val}"))
                    off_vec = vectors.get(off_hash)
                    if off_vec is None:
                        continue
                    score = cosine_similarity(req_vec, off_vec)
                    if score > best_score:
                        best_score = score
                        best_req = req_val
                        best_off = off_val

        is_match = best_score >= threshold if off_values else False
        if is_match:
            matched += 1
        points.append(
            {
                "key": req_key,
                "requested_values": req_values,
                "offered_values": off_values,
                "best_requested": best_req,
                "best_offered": best_off,
                "score": round(float(best_score), 4),
                "matched": is_match,
            }
        )
        scores.append(best_score)

    avg_score = sum(scores) / max(len(scores), 1) if scores else 0.0
    payload = {
        "threshold": round(float(threshold), 4),
        "matched_points": int(matched),
        "total_points": len(points),
        "points": points,
    }
    return round(float(avg_score), 4), payload


def _compare_characteristics_with_embeddings(
    requested: dict[str, list[str]],
    offered: dict[str, list[str]],
    threshold: float,
) -> tuple[float | None, dict[str, Any] | None]:
    if not requested:
        return None, None

    emb_texts: list[str] = []
    for key, vals in requested.items():
        for val in vals:
            emb_texts.append(f"{key}: {val}")
    for key, vals in offered.items():
        for val in vals:
            emb_texts.append(f"{key}: {val}")

    if not emb_texts:
        return None, None
    vectors = _embed_texts_sync(emb_texts)
    return _compare_characteristics_with_embeddings_from_vectors(
        requested=requested,
        offered=offered,
        threshold=threshold,
        vectors=vectors,
    )


def _extract_matched_name_from_notes(notes: str | None) -> str:
    text = (notes or "").strip()
    if not text:
        return ""

    patterns = [
        r"\bgpt_match=([^\n]+?)\s+gpt_score=",
        r"\bemb_match=([^\n]+?)\s+emb_score=",
        r"\bfuzzy_match=([^\n]+?)\s+fuzzy_score=",
        r"\bexact_match=([^\n]+?)(?:\s+\w+=|$)",
        r"\bmatch=([^\n]+?)\s+score=",
    ]
    for pattern in patterns:
        m = re.search(pattern, text)
        if not m:
            continue
        candidate = (m.group(1) or "").strip()
        if candidate:
            return candidate
    return ""


def _trim_for_prompt(text: str, limit: int = 600) -> str:
    value = _collapse_spaces(text or "")
    if len(value) <= limit:
        return value
    return value[:limit].rstrip() + "..."


def _clamp01(value: float) -> float:
    return max(0.0, min(1.0, float(value)))


def _resolve_product_match_score_weights() -> tuple[float, float]:
    raw_name_w = float(
        getattr(config, "OPENAI_PRODUCT_MATCH_GPT_NAME_WEIGHT", 0.60) or 0.60
    )
    raw_model_w = float(
        getattr(config, "OPENAI_PRODUCT_MATCH_GPT_MODEL_WEIGHT", 0.40) or 0.40
    )
    name_w = max(0.0, raw_name_w)
    model_w = max(0.0, raw_model_w)
    denom = name_w + model_w
    if denom <= 0:
        return 0.60, 0.40
    return name_w / denom, model_w / denom


_NAME_CORE_DROP_TOKENS = {
    "ml",
    "l",
    "litre",
    "litres",
    "mg",
    "g",
    "kg",
    "mcg",
    "ug",
    "iu",
    "ui",
    "mm",
    "cm",
    "m",
    "w",
    "kw",
    "v",
    "hz",
    "mah",
    "ah",
    "boite",
    "boites",
    "pack",
    "packs",
    "lot",
    "lots",
    "pcs",
    "piece",
    "pieces",
    "unite",
    "unites",
    "de",
    "du",
    "des",
    "la",
    "le",
    "les",
}
_NAME_CORE_NUMERIC_TOKEN_RE = re.compile(r"^\d+(?:[.,]\d+)?$")
_NAME_CORE_NUMERIC_UNIT_TOKEN_RE = re.compile(
    r"^\d+(?:[.,]\d+)?(?:ml|l|litre|litres|mg|g|kg|mcg|ug|iu|ui|mm|cm|m|w|kw|v|hz|mah|ah)$"
)
_NAME_CORE_PACK_TOKEN_RE = re.compile(r"^\d+(?:[.,]\d+)?[x\*]\d+(?:[.,]\d+)?$")


def _name_core_tokens(text: str | None) -> list[str]:
    normalized = _normalized_name_for_suspect_checks(text)
    if not normalized:
        return []
    tokens: list[str] = []
    for raw in normalized.split():
        token = str(raw or "").strip().lower()
        if not token:
            continue
        if token in _NAME_CORE_DROP_TOKENS:
            continue
        if _NAME_CORE_NUMERIC_TOKEN_RE.fullmatch(token):
            continue
        if _NAME_CORE_NUMERIC_UNIT_TOKEN_RE.fullmatch(token):
            continue
        if _NAME_CORE_PACK_TOKEN_RE.fullmatch(token):
            continue
        tokens.append(token)
    return tokens


def _name_core_similarity(left: str | None, right: str | None) -> float:
    left_tokens = _name_core_tokens(left)
    right_tokens = _name_core_tokens(right)
    if not left_tokens or not right_tokens:
        return 0.0
    left_text = " ".join(left_tokens)
    right_text = " ".join(right_tokens)
    seq_score = float(difflib.SequenceMatcher(None, left_text, right_text).ratio())
    left_set = set(left_tokens)
    right_set = set(right_tokens)
    inter = len(left_set & right_set)
    union = len(left_set | right_set)
    jaccard = (inter / union) if union else 0.0
    coverage = inter / max(1, min(len(left_set), len(right_set)))
    blended = (0.55 * seq_score) + (0.25 * coverage) + (0.20 * jaccard)
    return _clamp01(blended)


def _blend_gpt_candidate_confidence(
    *,
    item_name: str,
    candidate_name: str,
    model_confidence: float,
) -> float:
    name_w, model_w = _resolve_product_match_score_weights()
    name_score = _name_core_similarity(item_name, candidate_name)
    model_score = _clamp01(model_confidence)
    blended = (name_score * name_w) + (model_score * model_w)
    # Keep a small floor from model confidence for sparse short names.
    if name_score <= 0.01 and model_score > 0:
        blended = max(blended, model_score * model_w)
    return _clamp01(blended)


def _candidate_norm_map(candidate_names: list[str]) -> dict[str, str]:
    by_norm: dict[str, str] = {}
    for name in candidate_names:
        cleaned = str(name or "").strip()
        norm = _normalize_text(cleaned)
        if cleaned and norm and norm not in by_norm:
            by_norm[norm] = cleaned
    return by_norm


def _normalize_and_blend_gpt_rows(
    *,
    item_name: str,
    candidate_names: list[str],
    scored_rows: list[tuple[str, float]],
) -> tuple[dict[str, str], list[tuple[str, float]]]:
    by_norm = _candidate_norm_map(candidate_names)
    by_name_best: dict[str, float] = {}
    for rname, rscore in scored_rows:
        canon = by_norm.get(_normalize_text(rname))
        if not canon:
            continue
        blended_score = _blend_gpt_candidate_confidence(
            item_name=item_name,
            candidate_name=canon,
            model_confidence=rscore,
        )
        prev = by_name_best.get(canon)
        if prev is None or blended_score > prev:
            by_name_best[canon] = blended_score
    normalized_rows = sorted(by_name_best.items(), key=lambda x: x[1], reverse=True)
    return by_norm, normalized_rows


def _product_match_cache_key(
    *,
    item_text: str,
    item_name: str,
    candidate_names: list[str],
) -> str:
    normalized_candidates = sorted(
        {
            str(name or "").strip()
            for name in candidate_names
            if str(name or "").strip()
        }
    )
    name_w, model_w = _resolve_product_match_score_weights()
    payload = {
        "version": "gpt_match_name_blend_v1",
        "item_text": _collapse_spaces(item_text or ""),
        "item_name": _collapse_spaces(item_name or ""),
        "candidate_names": normalized_candidates,
        "name_weight": round(float(name_w), 6),
        "model_weight": round(float(model_w), 6),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _encode_gpt_matches_b64(rows: list[dict[str, Any]]) -> str:
    payload = json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
    return base64.urlsafe_b64encode(payload.encode("utf-8")).decode("ascii")


def _gpt_rank_requested_candidates(
    *,
    item_text: str,
    item_name: str,
    candidate_names: list[str],
) -> tuple[str | None, float | None, list[tuple[str, float]]]:
    if not config.OPENAI_API_KEY:
        return None, None, []
    if not candidate_names:
        return None, None, []

    primary_model = str(
        getattr(config, "OPENAI_PRODUCT_MATCH_MODEL", "")
        or getattr(config, "OPENAI_ATTACHMENT_MODEL", "gpt-5.1")
        or "gpt-5.1"
    ).strip()
    fallback_model = str(
        getattr(config, "OPENAI_PRODUCT_MATCH_FALLBACK_MODEL", "gpt-5.3") or ""
    ).strip()
    model_candidates: list[str] = []
    if primary_model:
        model_candidates.append(primary_model)
    if fallback_model and fallback_model != primary_model:
        model_candidates.append(fallback_model)
    if not model_candidates:
        model_candidates = ["gpt-5.1"]

    timeout_s = max(
        30,
        int(getattr(config, "OPENAI_PRODUCT_MATCH_TIMEOUT_SECONDS", 45) or 45),
    )
    retry_attempts = max(
        1,
        int(getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_ATTEMPTS", 3) or 3),
    )
    retry_backoff_s = max(
        0.2,
        float(
            getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_BACKOFF_SECONDS", 1.5)
            or 1.5
        ),
    )
    cleaned_item = _trim_for_prompt(item_text, limit=800)
    cleaned_item_name = _trim_for_prompt(item_name, limit=300)
    if not cleaned_item:
        cleaned_item = cleaned_item_name
    if not cleaned_item:
        return None, None, []

    candidate_payload = [
        {"id": idx + 1, "name": str(name or "").strip()}
        for idx, name in enumerate(candidate_names)
        if str(name or "").strip()
    ]
    if not candidate_payload:
        return None, None, []

    cache_key = _product_match_cache_key(
        item_text=cleaned_item,
        item_name=cleaned_item_name,
        candidate_names=[row["name"] for row in candidate_payload],
    )
    for model in model_candidates:
        cached_payload = get_openai_payload_cache(
            cache_type="product_match",
            model=model,
            cache_key=cache_key,
        )
        if not isinstance(cached_payload, dict):
            continue
        try:
            cached_best_name = str(cached_payload.get("best_name") or "").strip()
            cached_confidence_raw = cached_payload.get("confidence")
            cached_candidates_raw = cached_payload.get("candidates")
            cached_confidence = (
                float(cached_confidence_raw)
                if cached_confidence_raw is not None
                else None
            )
            cached_scored: list[tuple[str, float]] = []
            if isinstance(cached_candidates_raw, list):
                for row in cached_candidates_raw:
                    if not isinstance(row, dict):
                        continue
                    rname = str(row.get("name") or "").strip()
                    if not rname:
                        continue
                    try:
                        rscore = float(row.get("confidence"))
                    except Exception:
                        continue
                    cached_scored.append((rname, max(0.0, min(1.0, rscore))))
            by_norm, normalized_rows = _normalize_and_blend_gpt_rows(
                item_name=cleaned_item_name,
                candidate_names=[row["name"] for row in candidate_payload],
                scored_rows=cached_scored,
            )
            if cached_best_name.upper() == "NONE":
                return None, _clamp01(cached_confidence or 0.0), normalized_rows

            cached_best = by_norm.get(_normalize_text(cached_best_name))
            if not cached_best and normalized_rows:
                cached_best = normalized_rows[0][0]
            if not cached_best:
                return None, cached_confidence, normalized_rows

            blended_conf = None
            if cached_confidence is not None:
                blended_conf = _blend_gpt_candidate_confidence(
                    item_name=cleaned_item_name,
                    candidate_name=cached_best,
                    model_confidence=cached_confidence,
                )
            row_score = next(
                (float(score) for name, score in normalized_rows if name == cached_best),
                None,
            )
            if row_score is not None and blended_conf is not None:
                blended_conf = max(blended_conf, row_score)
            elif row_score is not None:
                blended_conf = row_score
            return cached_best, blended_conf, normalized_rows
        except Exception:
            continue

    system_prompt = (
        "Tu fais du matching produit facture -> produit demande. "
        "Priorise d'abord l'identite produit (nom exact, reference, modele, marque, capacite, classe). "
        "Les champs longs ou verboses ne doivent pas augmenter le score a eux seuls: "
        "ignore la longueur du texte et juge la precision semantique. "
        "Cas avec peu de details: base-toi surtout sur le nom principal et les references/modeles explicites; "
        "si l'evidence est faible ou ambigue, garde une confidence basse. "
        "Penalise fortement les contradictions de type/unite/capacite/classe. "
        "Evalue tous les noms presents dans la liste candidate. "
        "Retourne dans candidates tous les noms acceptables avec une confidence 0..1, "
        "tries du plus fort au plus faible. "
        "best_name doit etre le meilleur candidat de la liste retournee. "
        "Si rien n'est acceptable, retourne best_name='NONE' avec faible confidence "
        "et une liste candidates vide."
    )
    user_prompt = (
        "Produit facture a matcher:\n"
        f"{cleaned_item}\n\n"
        "Produits demandes candidats (JSON):\n"
        f"{json.dumps(candidate_payload, ensure_ascii=False)}"
    )

    payload = {
        "model": model,
        "temperature": 0,
        "seed": 42,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "response_format": {
            "type": "json_schema",
            "json_schema": {
                "name": "invoice_product_match",
                "strict": True,
                "schema": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "best_name": {"type": "string"},
                        "confidence": {
                            "type": "number",
                            "minimum": 0,
                            "maximum": 1,
                        },
                        "candidates": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "properties": {
                                    "name": {"type": "string"},
                                    "confidence": {
                                        "type": "number",
                                        "minimum": 0,
                                        "maximum": 1,
                                    },
                                },
                                "required": ["name", "confidence"],
                            },
                        },
                        "reason": {"type": "string"},
                    },
                    "required": ["best_name", "confidence", "candidates", "reason"],
                },
            },
        },
    }
    headers = {
        "Authorization": f"Bearer {config.OPENAI_API_KEY}",
        "Content-Type": "application/json",
    }

    for model in model_candidates:
        payload["model"] = model
        for attempt in range(retry_attempts):
            started = time.perf_counter()
            try:
                req = urllib_request.Request(
                    "https://api.openai.com/v1/chat/completions",
                    data=json.dumps(payload).encode("utf-8"),
                    headers=headers,
                    method="POST",
                )
                with urllib_request.urlopen(req, timeout=timeout_s) as resp:
                    data = json.loads(resp.read().decode("utf-8", errors="ignore"))
                record_openai_api_call(
                    operation="invoice_product_match_gpt",
                    endpoint="/v1/chat/completions",
                    model=model,
                    response_json=data if isinstance(data, dict) else None,
                    latency_ms=int((time.perf_counter() - started) * 1000),
                    success=True,
                    metadata={
                        "candidate_count": len(candidate_payload),
                        "attempt": attempt + 1,
                        "using_fallback_model": bool(model != primary_model),
                    },
                )
            except (urllib_error.URLError, urllib_error.HTTPError, TimeoutError, ValueError) as exc:
                status_code = getattr(exc, "code", None) if isinstance(exc, urllib_error.HTTPError) else None
                retryable = (
                    status_code in {408, 409, 429, 500, 502, 503, 504}
                    or isinstance(exc, (urllib_error.URLError, TimeoutError))
                )
                record_openai_api_call(
                    operation="invoice_product_match_gpt",
                    endpoint="/v1/chat/completions",
                    model=model,
                    response_json=None,
                    latency_ms=int((time.perf_counter() - started) * 1000),
                    success=False,
                    error_type=type(exc).__name__,
                    metadata={
                        "candidate_count": len(candidate_payload),
                        "attempt": attempt + 1,
                        "retryable": bool(retryable),
                        "http_status": status_code,
                        "using_fallback_model": bool(model != primary_model),
                    },
                )
                if retryable and attempt < (retry_attempts - 1):
                    delay = (retry_backoff_s * (2**attempt)) + random.uniform(0.0, 0.35)
                    time.sleep(delay)
                    continue
                break

            try:
                message = ((data.get("choices") or [{}])[0].get("message") or {})
                content = message.get("content")
                if isinstance(content, list):
                    parts: list[str] = []
                    for part in content:
                        if isinstance(part, dict):
                            text = part.get("text")
                            if isinstance(text, str) and text.strip():
                                parts.append(text)
                    content = "\n".join(parts)
                if not isinstance(content, str) or not content.strip():
                    raise ValueError("missing_content")
                parsed = json.loads(content)
            except Exception as parse_exc:
                if attempt < (retry_attempts - 1):
                    delay = (retry_backoff_s * (2**attempt)) + random.uniform(0.0, 0.35)
                    time.sleep(delay)
                    continue
                record_openai_api_call(
                    operation="invoice_product_match_gpt_parse",
                    endpoint="/v1/chat/completions",
                    model=model,
                    response_json=data if isinstance(data, dict) else None,
                    latency_ms=0,
                    success=False,
                    error_type=type(parse_exc).__name__,
                    metadata={
                        "candidate_count": len(candidate_payload),
                        "attempt": attempt + 1,
                        "using_fallback_model": bool(model != primary_model),
                    },
                )
                break

            raw_name = str(parsed.get("best_name") or "").strip()
            scored_rows: list[tuple[str, float]] = []
            raw_candidates = parsed.get("candidates")
            if isinstance(raw_candidates, list):
                for row in raw_candidates:
                    if not isinstance(row, dict):
                        continue
                    rname = str(row.get("name") or "").strip()
                    if not rname:
                        continue
                    try:
                        rscore = float(row.get("confidence"))
                    except Exception:
                        continue
                    rscore = max(0.0, min(1.0, rscore))
                    scored_rows.append((rname, rscore))

            by_norm, normalized_rows = _normalize_and_blend_gpt_rows(
                item_name=cleaned_item_name,
                candidate_names=[row["name"] for row in candidate_payload],
                scored_rows=scored_rows,
            )

            if not raw_name or raw_name.upper() == "NONE":
                try:
                    conf = float(parsed.get("confidence") or 0.0)
                except Exception:
                    conf = 0.0
                set_openai_payload_cache(
                    cache_type="product_match",
                    model=model,
                    cache_key=cache_key,
                    payload={
                        "best_name": "NONE",
                        "confidence": _clamp01(conf),
                        "candidates": [
                            {"name": cname, "confidence": cscore}
                            for cname, cscore in normalized_rows
                        ],
                        "scoring_version": "name_blend_v1",
                    },
                )
                return None, _clamp01(conf), normalized_rows

            matched = by_norm.get(_normalize_text(raw_name))
            if not matched and normalized_rows:
                matched = normalized_rows[0][0]
            try:
                raw_confidence = float(parsed.get("confidence"))
            except Exception:
                raw_confidence = 0.0
            confidence = _clamp01(raw_confidence)
            if matched:
                confidence = _blend_gpt_candidate_confidence(
                    item_name=cleaned_item_name,
                    candidate_name=matched,
                    model_confidence=confidence,
                )
                row_score = next(
                    (float(score) for name, score in normalized_rows if name == matched),
                    None,
                )
                if row_score is not None:
                    confidence = max(confidence, row_score)
            set_openai_payload_cache(
                cache_type="product_match",
                model=model,
                cache_key=cache_key,
                payload={
                    "best_name": matched or raw_name,
                    "confidence": confidence,
                    "candidates": [
                        {"name": cname, "confidence": cscore}
                        for cname, cscore in normalized_rows
                    ],
                    "scoring_version": "name_blend_v1",
                },
            )
            if not matched:
                return None, confidence, normalized_rows
            return matched, confidence, normalized_rows

    return None, None, []


def _compose_price_notes(
    *,
    price_source: str | None,
    amount_value: float | None = None,
    qty_raw: Any = None,
    unit_price_raw: Any = None,
    line_total_raw: Any = None,
    extra_tokens: list[str] | None = None,
) -> str:
    def _raw_token(value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        if not text:
            return None
        text = text.replace("\u00a0", " ")
        text = re.sub(r"\s+", "", text)
        return text or None

    def _num_type(raw: str | None) -> str | None:
        if not raw:
            return None
        if re.search(r"\d[.,]\d", raw):
            return "decimal"
        if re.search(r"\d", raw):
            return "integer"
        return None

    tokens: list[str] = []
    source = _collapse_spaces(str(price_source or "")).replace(" ", "_")
    if source:
        tokens.append(f"price_source={source}")
    qty_raw_token = _raw_token(qty_raw)
    if qty_raw_token:
        tokens.append(f"qty_raw={qty_raw_token}")
        qty_type = _num_type(qty_raw_token)
        if qty_type:
            tokens.append(f"qty_num_type={qty_type}")
    unit_price_raw_token = _raw_token(unit_price_raw)
    if unit_price_raw_token:
        tokens.append(f"unit_price_raw={unit_price_raw_token}")
        unit_price_type = _num_type(unit_price_raw_token)
        if unit_price_type:
            tokens.append(f"unit_price_num_type={unit_price_type}")
    line_total_raw_token = _raw_token(line_total_raw)
    if line_total_raw_token:
        tokens.append(f"line_total_raw={line_total_raw_token}")
        line_total_type = _num_type(line_total_raw_token)
        if line_total_type:
            tokens.append(f"line_total_num_type={line_total_type}")
    if isinstance(amount_value, (int, float)) and float(amount_value) > 0:
        tokens.append(f"line_amount={float(amount_value):.6f}")
    for raw in extra_tokens or []:
        token = _collapse_spaces(str(raw or "")).replace(" ", "_")
        if token:
            tokens.append(token)
    return " ".join(tokens).strip()


def _extract_note_token(notes: str | None, key: str) -> str | None:
    if not notes:
        return None
    pattern = rf"\b{re.escape(key)}=([^\s;]+)"
    match = re.search(pattern, notes)
    if not match:
        return None
    value = str(match.group(1) or "").strip()
    return value or None


def _extract_note_float(notes: str | None, key: str) -> float | None:
    raw = _extract_note_token(notes, key)
    if not raw:
        return None
    cleaned = str(raw).replace(",", ".").strip()
    try:
        return float(cleaned)
    except Exception:
        return None


def _upsert_note_token(notes: str | None, key: str, value: str | float | int | None) -> str:
    text = str(notes or "").strip()
    text = re.sub(rf"\b{re.escape(key)}=[^\s;]+", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    if value is None:
        return text
    token = _collapse_spaces(str(value)).replace(" ", "_")
    if not token:
        return text
    return f"{text} {key}={token}".strip() if text else f"{key}={token}"


def _normalized_name_for_suspect_checks(name: str | None) -> str:
    value = _collapse_spaces(str(name or "")).lower()
    if not value:
        return ""
    value = (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
    )
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return _collapse_spaces(value)


def _is_unknown_name(name: str | None) -> bool:
    norm = _normalized_name_for_suspect_checks(name)
    return norm in _UNKNOWN_NAME_TOKENS


def _looks_like_admin_noise(name: str | None) -> bool:
    norm = _normalized_name_for_suspect_checks(name)
    if not norm:
        return False
    for token in _ADMIN_NOISE_TOKENS:
        token_norm = _normalized_name_for_suspect_checks(token)
        if token_norm and token_norm in norm:
            return True
    return False


def _extract_gpt_score_from_notes(notes: str | None) -> float | None:
    if not notes:
        return None
    match = re.search(r"\bgpt_score=([0-9]+(?:\.[0-9]+)?)", notes)
    if not match:
        return None
    try:
        return float(match.group(1))
    except Exception:
        return None


def _item_line_amount(item: InvoiceExtractionItem) -> float | None:
    amount = _extract_note_float(item.notes, "line_amount")
    if isinstance(amount, (int, float)) and float(amount) > 0:
        return float(amount)
    return None


def _has_finance_inconsistency(item: InvoiceExtractionItem) -> bool:
    qty = float(item.quantity) if isinstance(item.quantity, (int, float)) else None
    unit_price = float(item.unit_price) if isinstance(item.unit_price, (int, float)) else None
    amount = _item_line_amount(item)
    if qty is None or unit_price is None or amount is None:
        return False
    if qty <= 0 or unit_price <= 0 or amount <= 0:
        return False
    expected = qty * unit_price
    rel_err = abs(expected - amount) / max(abs(amount), 1.0)
    return rel_err > 0.08


def _is_derived_price_suspicious(item: InvoiceExtractionItem) -> bool:
    notes = str(item.notes or "")
    if "price_source=derived_amount_div_qty" not in notes:
        return False
    qty = float(item.quantity) if isinstance(item.quantity, (int, float)) else None
    unit_price = float(item.unit_price) if isinstance(item.unit_price, (int, float)) else None
    if qty is not None and qty >= 100:
        return True
    if qty is not None and unit_price is not None and qty >= 20 and unit_price <= 5:
        return True
    return False


def _fallback_suspect_reasons(
    item: InvoiceExtractionItem,
    duplicate_counts: dict[str, int],
) -> list[str]:
    reasons: list[str] = []
    norm_name = _normalized_name_for_suspect_checks(item.product_name)
    is_unknown = _is_unknown_name(item.product_name)
    is_admin = _looks_like_admin_noise(item.product_name)

    if is_unknown:
        reasons.append("unknown_name")
    if is_admin:
        reasons.append("admin_noise_name")
    if duplicate_counts.get(norm_name, 0) >= 2 and (is_unknown or is_admin):
        reasons.append("duplicate_name_in_invoice")
    if _has_finance_inconsistency(item):
        reasons.append("finance_inconsistency")
    if _is_derived_price_suspicious(item):
        reasons.append("derived_price_suspicious")
    return reasons


def _invoice_fallback_cache_key(invoice: InboundInvoice, invoice_text: str) -> str:
    stat_size = 0
    stat_mtime = 0.0
    try:
        stat = Path(invoice.file_path).stat()
        stat_size = int(stat.st_size)
        stat_mtime = float(stat.st_mtime)
    except Exception:
        pass
    payload = {
        "tenant_id": int(invoice.tenant_id or 0),
        "campaign_id": int(invoice.campaign_id or 0),
        "invoice_id": int(invoice.id or 0),
        "file_name": Path(invoice.file_path or "").name,
        "file_size": stat_size,
        "file_mtime": stat_mtime,
        "text_sha256": hashlib.sha256(invoice_text.encode("utf-8")).hexdigest(),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _normalize_gpt_invoice_rows(raw_rows: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_rows, list):
        return []
    out: list[dict[str, Any]] = []
    for idx, row in enumerate(raw_rows, start=1):
        if not isinstance(row, dict):
            continue
        product_name = _collapse_spaces(str(row.get("product_name") or ""))
        if not product_name:
            continue
        try:
            line_no = int(row.get("line")) if row.get("line") is not None else idx
        except Exception:
            line_no = idx
        quantity = row.get("quantity")
        unit_price = row.get("unit_price")
        amount = row.get("amount")
        try:
            quantity = float(quantity) if quantity is not None else None
        except Exception:
            quantity = None
        try:
            unit_price = float(unit_price) if unit_price is not None else None
        except Exception:
            unit_price = None
        try:
            amount = float(amount) if amount is not None else None
        except Exception:
            amount = None
        out.append(
            {
                "line": line_no,
                "ref": _collapse_spaces(str(row.get("ref") or "")) or None,
                "product_name": product_name,
                "quantity": quantity if isinstance(quantity, (int, float)) and quantity > 0 else None,
                "unit_price": unit_price if isinstance(unit_price, (int, float)) and unit_price > 0 else None,
                "amount": amount if isinstance(amount, (int, float)) and amount > 0 else None,
            }
        )
    return out


def _build_items_from_gpt_invoice_rows(
    invoice: InboundInvoice,
    rows: list[dict[str, Any]],
    *,
    model: str,
) -> list[InvoiceExtractionItem]:
    items: list[InvoiceExtractionItem] = []
    for row in rows:
        product_name = _collapse_spaces(str(row.get("product_name") or "")) or "UNKNOWN"
        ref_text = _collapse_spaces(str(row.get("ref") or "")) or None
        qty_value = row.get("quantity")
        unit_price_value = row.get("unit_price")
        amount_value = row.get("amount")
        qty_norm = float(qty_value) if isinstance(qty_value, (int, float)) and qty_value > 0 else None
        unit_price_norm = (
            float(unit_price_value)
            if isinstance(unit_price_value, (int, float)) and unit_price_value > 0
            else None
        )
        amount_norm = (
            float(amount_value) if isinstance(amount_value, (int, float)) and amount_value > 0 else None
        )
        raw_desc = product_name
        description = f"ref: {ref_text}" if ref_text else None
        if ref_text:
            raw_desc = f"{product_name} | REF {ref_text}"
        resolved_unit_price, price_source = _resolve_unit_price_with_source(
            qty_value=qty_norm,
            unit_price_value=unit_price_norm,
            amount_value=amount_norm,
        )
        items.append(
            InvoiceExtractionItem(
                tenant_id=invoice.tenant_id,
                invoice_id=invoice.id,
                campaign_id=invoice.campaign_id,
                supplier_id=None,
                source_engine="gpt_invoice_fallback",
                status=ExtractionStatus.TO_VERIFY,
                product_name=product_name[:255],
                description=description,
                raw_description=raw_desc,
                unit=None,
                quantity=qty_norm,
                unit_price=resolved_unit_price,
                currency="DZD",
                confidence=0.85,
                notes=_compose_price_notes(
                    price_source=price_source,
                    amount_value=amount_norm,
                    qty_raw=qty_value,
                    unit_price_raw=unit_price_value,
                    line_total_raw=amount_value,
                    extra_tokens=[f"gpt_model={model}"],
                ),
            )
        )
    return items


def _extract_invoice_rows_with_gpt_fallback(invoice: InboundInvoice) -> tuple[list[dict[str, Any]], str]:
    if not config.OPENAI_API_KEY:
        return [], ""
    model = (
        str(getattr(config, "OPENAI_INVOICE_FALLBACK_MODEL", "") or "").strip()
        or str(getattr(config, "OPENAI_ATTACHMENT_MODEL", "gpt-5.1") or "gpt-5.1")
    )
    timeout_s = max(20, int(getattr(config, "OPENAI_INVOICE_FALLBACK_TIMEOUT_SECONDS", 60) or 60))
    max_chars = max(2000, int(getattr(config, "OPENAI_INVOICE_FALLBACK_MAX_INPUT_CHARS", 40000) or 40000))

    processing_file_path = resolve_invoice_processing_file_path(invoice.file_path) or invoice.file_path

    invoice_text = ""
    ocr_provider = str(getattr(config, "OPENAI_OCR_PROVIDER", "document_ai") or "document_ai").strip().lower()
    prefer_vision = _is_spreadsheet_invoice_file(invoice.file_path)
    if prefer_vision and ocr_provider in {"gpt4_vision", "openai_vision", "vision"}:
        try:
            from app.services.openai_vision_ocr import extract_text_from_file

            invoice_text = extract_text_from_file(
                file_path=processing_file_path,
                operation="invoice_spreadsheet_ocr_openai_vision",
                max_pages=getattr(config, "OPENAI_OCR_VISION_MAX_PAGES", 8),
            )
        except Exception:
            invoice_text = ""
    if not invoice_text:
        try:
            doc = process_invoice(processing_file_path)
            invoice_text = str(getattr(doc, "text", None) or "").strip()
        except Exception:
            invoice_text = ""
    if not invoice_text:
        invoice_text = _extract_pdf_text(processing_file_path) or ""
    invoice_text = invoice_text[:max_chars].strip()
    if not invoice_text:
        return [], model

    cache_key = _invoice_fallback_cache_key(invoice, invoice_text)
    cached_payload = get_openai_payload_cache(
        cache_type="invoice_fallback_extract",
        model=model,
        cache_key=cache_key,
    )
    if isinstance(cached_payload, dict):
        cached_rows = _normalize_gpt_invoice_rows(cached_payload.get("items"))
        if cached_rows:
            return cached_rows, model

    system_prompt = (
        "Extract invoice/proforma product lines from OCR text. "
        "Return only real product/service rows, exclude headers, totals, taxes, bank/admin info. "
        "Return strict JSON with key 'items'. "
        "Each item must include: line, ref, product_name, quantity, unit_price, amount. "
        "Use null when a numeric field is not available. "
        "Do not invent products."
    )
    user_prompt = json.dumps(
        {
            "invoice_file": Path(invoice.file_path or "").name,
            "ocr_text": invoice_text,
        },
        ensure_ascii=False,
    )
    payload = {
        "model": model,
        "temperature": 0,
        "seed": 42,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "response_format": {
            "type": "json_schema",
            "json_schema": {
                "name": "invoice_fallback_extract",
                "strict": True,
                "schema": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "items": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "properties": {
                                    "line": {"type": "integer"},
                                    "ref": {"type": ["string", "null"]},
                                    "product_name": {"type": "string"},
                                    "quantity": {"type": ["number", "null"]},
                                    "unit_price": {"type": ["number", "null"]},
                                    "amount": {"type": ["number", "null"]},
                                },
                                "required": [
                                    "line",
                                    "ref",
                                    "product_name",
                                    "quantity",
                                    "unit_price",
                                    "amount",
                                ],
                            },
                        }
                    },
                    "required": ["items"],
                },
            },
        },
    }
    headers = {
        "Authorization": f"Bearer {config.OPENAI_API_KEY}",
        "Content-Type": "application/json",
    }

    started = time.perf_counter()
    try:
        req = urllib_request.Request(
            "https://api.openai.com/v1/chat/completions",
            data=json.dumps(payload).encode("utf-8"),
            headers=headers,
            method="POST",
        )
        with urllib_request.urlopen(req, timeout=timeout_s) as resp:
            data = json.loads(resp.read().decode("utf-8", errors="ignore"))
    except (urllib_error.URLError, urllib_error.HTTPError, TimeoutError, ValueError) as exc:
        record_openai_api_call(
            operation="invoice_fallback_extract_gpt",
            endpoint="/v1/chat/completions",
            model=model,
            response_json=None,
            latency_ms=int((time.perf_counter() - started) * 1000),
            success=False,
            error_type=type(exc).__name__,
            metadata={"invoice_id": int(invoice.id or 0)},
        )
        return [], model
    record_openai_api_call(
        operation="invoice_fallback_extract_gpt",
        endpoint="/v1/chat/completions",
        model=model,
        response_json=data if isinstance(data, dict) else None,
        latency_ms=int((time.perf_counter() - started) * 1000),
        success=True,
        metadata={"invoice_id": int(invoice.id or 0)},
    )

    try:
        message = ((data.get("choices") or [{}])[0].get("message") or {})
        content = message.get("content")
        if isinstance(content, list):
            chunks: list[str] = []
            for part in content:
                if isinstance(part, dict):
                    text = part.get("text")
                    if isinstance(text, str) and text.strip():
                        chunks.append(text)
            content = "\n".join(chunks)
        if not isinstance(content, str) or not content.strip():
            return [], model
        parsed = json.loads(content)
    except Exception:
        return [], model

    rows = _normalize_gpt_invoice_rows(parsed.get("items"))
    if rows:
        set_openai_payload_cache(
            cache_type="invoice_fallback_extract",
            model=model,
            cache_key=cache_key,
            payload={"items": rows},
        )
    return rows, model


def _item_anchor_value(item: InvoiceExtractionItem) -> float | None:
    amount = _item_line_amount(item)
    if isinstance(amount, (int, float)) and amount > 0:
        return float(amount)
    if isinstance(item.unit_price, (int, float)) and float(item.unit_price) > 0:
        return float(item.unit_price)
    if isinstance(item.quantity, (int, float)) and float(item.quantity) > 0:
        return float(item.quantity)
    return None


def _fallback_row_anchor_value(row: dict[str, Any]) -> float | None:
    amount = row.get("amount")
    if isinstance(amount, (int, float)) and float(amount) > 0:
        return float(amount)
    unit_price = row.get("unit_price")
    if isinstance(unit_price, (int, float)) and float(unit_price) > 0:
        return float(unit_price)
    quantity = row.get("quantity")
    if isinstance(quantity, (int, float)) and float(quantity) > 0:
        return float(quantity)
    return None


def _fallback_fit_score(item: InvoiceExtractionItem, row: dict[str, Any]) -> float:
    score = 0.0
    item_anchor = _item_anchor_value(item)
    row_anchor = _fallback_row_anchor_value(row)
    if item_anchor is not None and row_anchor is not None:
        rel = abs(item_anchor - row_anchor) / max(abs(item_anchor), abs(row_anchor), 1.0)
        score += max(0.0, 1.0 - (rel * 5.0)) * 0.75
    elif row_anchor is not None:
        score += 0.10

    item_qty = float(item.quantity) if isinstance(item.quantity, (int, float)) else None
    row_qty = float(row.get("quantity")) if isinstance(row.get("quantity"), (int, float)) else None
    if item_qty is not None and row_qty is not None and item_qty > 0 and row_qty > 0:
        rel_q = abs(item_qty - row_qty) / max(abs(item_qty), abs(row_qty), 1.0)
        score += max(0.0, 1.0 - (rel_q * 3.0)) * 0.15

    current_name = str(item.product_name or "").strip()
    target_name = str(row.get("product_name") or "").strip()
    if current_name and target_name and not _is_unknown_name(current_name):
        score += _normalized_name_similarity(current_name, target_name) * 0.10
    return score


def _normalized_name_similarity(left: str | None, right: str | None) -> float:
    a = _normalized_name_for_suspect_checks(left)
    b = _normalized_name_for_suspect_checks(right)
    if not a or not b:
        return 0.0
    return float(difflib.SequenceMatcher(None, a, b).ratio())


def _row_already_covered_by_existing_item(
    row: dict[str, Any],
    existing_items: list[InvoiceExtractionItem],
) -> bool:
    row_name = str(row.get("product_name") or "").strip()
    row_anchor = _fallback_row_anchor_value(row)
    for item in existing_items:
        if _is_unknown_name(item.product_name):
            continue
        name_sim = _normalized_name_similarity(item.product_name, row_name)
        item_anchor = _item_anchor_value(item)
        rel = None
        if item_anchor is not None and row_anchor is not None:
            rel = abs(item_anchor - row_anchor) / max(abs(item_anchor), abs(row_anchor), 1.0)
        if name_sim >= 0.92:
            return True
        if rel is not None and rel <= 0.005:
            return True
        if rel is not None and rel <= 0.08 and name_sim >= 0.78:
            return True
    return False


def apply_invoice_fallback_for_low_gpt_score(
    invoice: InboundInvoice,
    items: list[InvoiceExtractionItem],
) -> dict[str, Any]:
    enabled = bool(getattr(config, "OPENAI_INVOICE_FALLBACK_ENABLED", True))
    if not enabled or not items:
        return {"enabled": enabled, "triggered": 0, "applied": 0, "model": None}

    min_gpt_score = max(
        0.0,
        min(
            1.0,
            float(getattr(config, "OPENAI_INVOICE_FALLBACK_MIN_GPT_SCORE", 0.50) or 0.50),
        ),
    )
    min_suspect_count = max(
        1,
        int(getattr(config, "OPENAI_INVOICE_FALLBACK_MIN_SUSPECT_COUNT", 1) or 1),
    )

    duplicate_counts: dict[str, int] = {}
    for item in items:
        key = _normalized_name_for_suspect_checks(item.product_name)
        if not key:
            continue
        duplicate_counts[key] = duplicate_counts.get(key, 0) + 1

    suspect_rows: list[tuple[InvoiceExtractionItem, list[str], float | None]] = []
    for item in items:
        gpt_score = _extract_gpt_score_from_notes(item.notes)
        low_or_missing = gpt_score is None or float(gpt_score) < min_gpt_score
        if not low_or_missing:
            continue
        reasons = _fallback_suspect_reasons(item, duplicate_counts)
        if not reasons:
            continue
        suspect_rows.append((item, reasons, gpt_score))

    if len(suspect_rows) < min_suspect_count:
        return {
            "enabled": enabled,
            "triggered": len(suspect_rows),
            "applied": 0,
            "model": None,
            "min_gpt_score": min_gpt_score,
            "reason": "not_enough_suspect_rows",
        }

    gpt_rows, model = _extract_invoice_rows_with_gpt_fallback(invoice)
    if not gpt_rows:
        return {
            "enabled": enabled,
            "triggered": len(suspect_rows),
            "applied": 0,
            "model": model,
            "min_gpt_score": min_gpt_score,
            "reason": "gpt_extract_empty",
        }

    suspect_items = [row[0] for row in suspect_rows]
    non_suspect_items = [it for it in items if it not in suspect_items]
    candidate_rows = [
        row
        for row in gpt_rows
        if not _row_already_covered_by_existing_item(row, non_suspect_items)
    ]
    if len(candidate_rows) < len(suspect_items):
        candidate_rows = list(gpt_rows)

    remaining_rows = list(candidate_rows)
    applied = 0
    for item, reasons, _score in suspect_rows:
        if not remaining_rows:
            break
        best_idx = -1
        best_score = -1.0
        for idx, row in enumerate(remaining_rows):
            fit = _fallback_fit_score(item, row)
            if fit > best_score:
                best_score = fit
                best_idx = idx
        if best_idx < 0:
            continue
        selected = remaining_rows.pop(best_idx)
        raw_name = _collapse_spaces(str(selected.get("product_name") or ""))
        if not raw_name:
            continue
        row_ref = _collapse_spaces(str(selected.get("ref") or "")) or None
        item.raw_description = raw_name
        item.product_name = _derive_product_name(raw_name, row_ref)[:255]
        item.description = _clean_description_text(raw_name, row_ref)

        row_qty = selected.get("quantity")
        row_price = selected.get("unit_price")
        row_amount = selected.get("amount")
        if isinstance(row_qty, (int, float)) and float(row_qty) > 0:
            item.quantity = float(row_qty)
        if isinstance(row_price, (int, float)) and float(row_price) > 0:
            item.unit_price = float(row_price)

        item.notes = _compose_price_notes(
            price_source="gpt_fallback",
            amount_value=float(row_amount) if isinstance(row_amount, (int, float)) else _item_line_amount(item),
            extra_tokens=[
                "fallback_applied=1",
                f"fallback_model={model}",
                f"fallback_reasons={','.join(reasons)}",
            ],
        )
        item.status = ExtractionStatus.TO_VERIFY
        item.confidence = max(float(item.confidence or 0.0), 0.75)
        applied += 1

    return {
        "enabled": enabled,
        "triggered": len(suspect_rows),
        "applied": int(applied),
        "model": model,
        "min_gpt_score": min_gpt_score,
        "rows_from_gpt": len(gpt_rows),
    }


def _requested_reference_signature(
    db: Session,
    tenant_id: int,
    campaign_id: int,
) -> str:
    rows = _collect_campaign_requested_characteristics(db, tenant_id, campaign_id)
    payload: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        payload.append(
            {
                "product_name": str(row.get("product_name") or "").strip(),
                "requested_characteristics": row.get("requested_characteristics"),
                "requested_description": row.get("requested_description"),
            }
        )
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _product_pool_signature(products: list[Product]) -> str:
    payload = [
        {
            "id": int(p.id or 0),
            "name": str(p.name or "").strip(),
            "description": str(p.description or "").strip(),
            "characteristics": p.characteristics,
        }
        for p in products
    ]
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _item_match_signature(
    *,
    item: InvoiceExtractionItem,
    product_pool_signature: str,
    requested_signature: str,
    threshold: float,
    embed_threshold: float,
    characteristics_match_threshold: float,
) -> str:
    payload = {
        "line": {
            "product_name": str(item.product_name or "").strip(),
            "description": str(item.description or "").strip(),
            "corrected_product_name": str(item.corrected_product_name or "").strip(),
        },
        "product_pool_signature": product_pool_signature,
        "requested_signature": requested_signature,
        "product_match_mode": str(getattr(config, "OPENAI_PRODUCT_MATCH_MODE", "embedding") or "embedding"),
        "product_match_model": str(getattr(config, "OPENAI_PRODUCT_MATCH_MODEL", "") or ""),
        "embedding_model": str(getattr(config, "OPENAI_EMBEDDING_MODEL", "") or ""),
        "threshold": round(float(threshold), 4),
        "embed_threshold": round(float(embed_threshold), 4),
        "characteristics_match_threshold": round(float(characteristics_match_threshold), 4),
        "gpt_min_conf": round(
            float(getattr(config, "OPENAI_PRODUCT_MATCH_MIN_CONFIDENCE", 0.30) or 0.30),
            4,
        ),
        "gpt_strict_min_conf": round(
            float(
                getattr(config, "OPENAI_PRODUCT_MATCH_STRICT_MIN_CONFIDENCE", 0.50)
                or 0.50
            ),
            4,
        ),
        "gpt_top_k": int(getattr(config, "OPENAI_PRODUCT_MATCH_TOP_K", 15) or 15),
        "gpt_uncertain_only": bool(
            getattr(config, "OPENAI_PRODUCT_MATCH_GPT_UNCERTAIN_ONLY", True)
        ),
        "gpt_force_to_verify_on_miss": bool(
            getattr(config, "OPENAI_PRODUCT_MATCH_FORCE_TO_VERIFY_ON_GPT_MISS", True)
        ),
        "gpt_fallback_model": str(
            getattr(config, "OPENAI_PRODUCT_MATCH_FALLBACK_MODEL", "gpt-5.3") or ""
        ),
        "gpt_retry_attempts": int(
            getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_ATTEMPTS", 3) or 3
        ),
        "gpt_retry_backoff_s": round(
            float(
                getattr(config, "OPENAI_PRODUCT_MATCH_RETRY_BACKOFF_SECONDS", 1.5)
                or 1.5
            ),
            3,
        ),
        "gpt_trigger_min": round(
            float(getattr(config, "OPENAI_PRODUCT_MATCH_GPT_TRIGGER_MIN", 0.35) or 0.35),
            4,
        ),
        "gpt_trigger_max": round(
            float(getattr(config, "OPENAI_PRODUCT_MATCH_GPT_TRIGGER_MAX", 0.88) or 0.88),
            4,
        ),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _serialize_item_match_payload(item: InvoiceExtractionItem) -> dict[str, Any]:
    return {
        "status": item.status.value if isinstance(item.status, ExtractionStatus) else str(item.status or ""),
        "notes": item.notes,
        "similarity_score": item.similarity_score,
        "similarity_threshold": item.similarity_threshold,
        "similarity_points": _deepcopy_payload(item.similarity_points),
    }


def _apply_item_match_payload(item: InvoiceExtractionItem, payload: dict[str, Any]) -> bool:
    if not isinstance(payload, dict):
        return False
    raw_status = str(payload.get("status") or "").strip().upper()
    try:
        item.status = ExtractionStatus(raw_status) if raw_status else ExtractionStatus.TO_VERIFY
    except Exception:
        item.status = ExtractionStatus.TO_VERIFY
    item.notes = payload.get("notes")
    item.similarity_score = payload.get("similarity_score")
    item.similarity_threshold = payload.get("similarity_threshold")
    item.similarity_points = _deepcopy_payload(payload.get("similarity_points"))
    return True


def _pair_cache_signature_for_item(
    *,
    item: InvoiceExtractionItem,
    offered: list[str],
    requested_specs: list[tuple[str, str, list[str]]],
    threshold: float,
    source: str,
) -> str:
    payload = {
        "line": {
            "product_name": str(item.product_name or "").strip(),
            "description": str(item.description or "").strip(),
            "corrected_product_name": str(item.corrected_product_name or "").strip(),
            "notes": str(item.notes or "").strip(),
        },
        "offered": list(offered),
        "source": source,
        "requested_specs": [
            {
                "name": requested_name,
                "norm": requested_norm,
                "characteristics": requested_chars,
            }
            for requested_name, requested_norm, requested_chars in requested_specs
        ],
        "threshold": round(float(threshold), 4),
        "embedding_model": str(getattr(config, "OPENAI_EMBEDDING_MODEL", "") or ""),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _apply_resemblance_points(
    db: Session,
    items: list[InvoiceExtractionItem],
    tenant_id: int,
    campaign_id: int,
    threshold: float,
) -> None:
    requested_rows = _collect_campaign_requested_characteristics(db, tenant_id, campaign_id)
    if not requested_rows:
        for item in items:
            item.similarity_score = None
            item.similarity_threshold = None
            item.similarity_points = None
        return

    for item in items:
        matched_name = (item.corrected_product_name or "").strip()
        matched_name_source = "corrected_product_name" if matched_name else ""
        if not matched_name:
            matched_name = _extract_matched_name_from_notes(item.notes)
            if matched_name:
                matched_name_source = "notes_match"
        if not matched_name:
            # Safe fallback: keep resemblance available even when corrected name
            # is empty, by using the extracted invoice name.
            matched_name = (item.product_name or "").strip()
            if matched_name:
                matched_name_source = "product_name"
        if not matched_name:
            item.similarity_score = None
            item.similarity_threshold = None
            item.similarity_points = None
            continue

        requested = _select_requested_characteristics_for_product(matched_name, requested_rows)
        if not requested:
            item.similarity_score = None
            item.similarity_threshold = None
            item.similarity_points = None
            continue

        offered, source = _extract_offered_characteristics(db, item, tenant_id, campaign_id)
        score, details = _compare_characteristics_with_embeddings(requested, offered, threshold)
        item.similarity_score = score
        item.similarity_threshold = threshold
        if isinstance(details, dict):
            details["source"] = source
            details["matched_product_name"] = matched_name
            details["matched_product_name_source"] = matched_name_source
        item.similarity_points = details


def compute_resemblance_points_for_items(
    db: Session,
    items: list[InvoiceExtractionItem],
    threshold: float,
) -> None:
    if not items:
        return
    sample = next((it for it in items if it is not None), None)
    if sample is None:
        return
    thr = float(threshold)
    _apply_resemblance_points(
        db=db,
        items=items,
        tenant_id=sample.tenant_id,
        campaign_id=sample.campaign_id,
        threshold=thr,
    )


def rebuild_characteristics_pair_cache(
    db: Session,
    tenant_id: int,
    campaign_id: int,
    threshold: float,
    *,
    item_ids: list[int] | None = None,
) -> int:
    query = db.query(InvoiceExtractionItem).filter(
        InvoiceExtractionItem.tenant_id == tenant_id,
        InvoiceExtractionItem.campaign_id == campaign_id,
    )
    if item_ids:
        query = query.filter(InvoiceExtractionItem.id.in_(item_ids))
    items = query.all()
    if not items:
        return 0
    target_item_ids = [item.id for item in items if item.id is not None]
    if not target_item_ids:
        return 0

    requested_rows = _collect_campaign_requested_characteristics(
        db=db,
        tenant_id=tenant_id,
        campaign_id=campaign_id,
    )
    requested_specs_by_norm: dict[str, dict[str, Any]] = {}
    for row in requested_rows:
        product_name = str(row.get("product_name") or "").strip()
        if not product_name:
            continue
        product_norm = _normalize_text(product_name)
        if not product_norm:
            continue
        characteristics = row.get("characteristics")
        if not isinstance(characteristics, dict):
            characteristics = {}
        bucket = requested_specs_by_norm.setdefault(
            product_norm,
            {
                "product_name": product_name,
                "characteristics": {},
            },
        )
        merged_chars = bucket["characteristics"]
        for key, vals in characteristics.items():
            if isinstance(vals, list):
                for val in vals:
                    _append_unique_value(merged_chars, str(key), str(val))
            elif vals is not None:
                _append_unique_value(merged_chars, str(key), str(vals))
    requested_specs: list[tuple[str, str, dict[str, list[str]]]] = []
    for product_norm, payload in requested_specs_by_norm.items():
        requested_specs.append(
            (
                str(payload.get("product_name") or product_norm),
                product_norm,
                payload.get("characteristics") or {},
            )
        )
    if not requested_specs:
        stale_query = db.query(InvoiceItemCharacteristicMatch).filter(
            InvoiceItemCharacteristicMatch.tenant_id == tenant_id,
            InvoiceItemCharacteristicMatch.campaign_id == campaign_id,
        )
        stale_query = stale_query.filter(InvoiceItemCharacteristicMatch.invoice_item_id.in_(target_item_ids))
        stale_query.delete(synchronize_session=False)
        return 0

    # Idempotent rebuild for target items: clear existing rows first, then reinsert.
    db.query(InvoiceItemCharacteristicMatch).filter(
        InvoiceItemCharacteristicMatch.tenant_id == tenant_id,
        InvoiceItemCharacteristicMatch.campaign_id == campaign_id,
        InvoiceItemCharacteristicMatch.invoice_item_id.in_(target_item_ids),
    ).delete(synchronize_session=False)
    db.flush()

    offered_by_item_id: dict[int, dict[str, list[str]]] = {}
    source_by_item_id: dict[int, str] = {}
    pair_cache_key_by_item_id: dict[int, str] = {}
    emb_texts: list[str] = []
    pair_cache_model = str(getattr(config, "OPENAI_EMBEDDING_MODEL", "") or "embedding")

    for item in items:
        if item.id is None:
            continue
        offered, source = _extract_offered_characteristics(
            db=db,
            item=item,
            tenant_id=tenant_id,
            campaign_id=campaign_id,
        )
        offered_by_item_id[item.id] = offered
        source_by_item_id[item.id] = source
        pair_cache_key_by_item_id[item.id] = _pair_cache_signature_for_item(
            item=item,
            offered=[f"{key}: {val}" for key, vals in offered.items() for val in vals],
            requested_specs=[
                (
                    requested_name,
                    requested_norm,
                    [f"{key}: {val}" for key, vals in requested_chars.items() for val in vals],
                )
                for requested_name, requested_norm, requested_chars in requested_specs
            ],
            threshold=threshold,
            source=source,
        )
        cached_payload = get_openai_payload_cache(
            cache_type="invoice_pair_cache",
            model=pair_cache_model,
            cache_key=pair_cache_key_by_item_id[item.id],
        )
        if isinstance(cached_payload, dict) and isinstance(cached_payload.get("rows"), list):
            continue
        for key, vals in offered.items():
            for val in vals:
                emb_texts.append(f"{key}: {val}")

    for _name, _norm, req_chars in requested_specs:
        for key, vals in req_chars.items():
            for val in vals:
                emb_texts.append(f"{key}: {val}")

    vectors = _embed_texts_sync(emb_texts) if emb_texts else {}

    upserted = 0

    for item in items:
        if item.id is None:
            continue
        offered = offered_by_item_id.get(item.id, {})
        source = source_by_item_id.get(item.id, "invoice_line")
        matched_name = (
            (item.corrected_product_name or "").strip()
            or _extract_matched_name_from_notes(item.notes)
            or (item.product_name or "").strip()
        )
        pair_cache_key = pair_cache_key_by_item_id.get(item.id)
        if pair_cache_key:
            cached_payload = get_openai_payload_cache(
                cache_type="invoice_pair_cache",
                model=pair_cache_model,
                cache_key=pair_cache_key,
            )
            if isinstance(cached_payload, dict) and isinstance(cached_payload.get("rows"), list):
                for cached_row in cached_payload.get("rows") or []:
                    if not isinstance(cached_row, dict):
                        continue
                    requested_name = str(cached_row.get("requested_product_name") or "").strip()
                    requested_norm = str(cached_row.get("requested_product_norm") or "").strip()
                    if not requested_name or not requested_norm:
                        continue
                    row = InvoiceItemCharacteristicMatch(
                        tenant_id=tenant_id,
                        campaign_id=campaign_id,
                        invoice_item_id=item.id,
                        requested_product_name=requested_name[:255],
                        requested_product_norm=requested_norm[:255],
                    )
                    db.add(row)
                    row.requested_product_name = requested_name[:255]
                    row.requested_product_norm = requested_norm[:255]
                    cached_matched_name = str(cached_row.get("matched_product_name") or "").strip()
                    row.matched_product_name = cached_matched_name[:255] if cached_matched_name else None
                    try:
                        row.threshold = float(cached_row.get("threshold") or threshold)
                    except Exception:
                        row.threshold = float(threshold)
                    try:
                        row.score = (
                            float(cached_row.get("score"))
                            if cached_row.get("score") is not None
                            else None
                        )
                    except Exception:
                        row.score = None
                    try:
                        row.matched_points = int(cached_row.get("matched_points") or 0)
                    except Exception:
                        row.matched_points = 0
                    try:
                        row.total_points = int(cached_row.get("total_points") or 0)
                    except Exception:
                        row.total_points = 0
                    points = cached_row.get("points")
                    row.points = points if isinstance(points, list) else []
                    row.source = str(cached_row.get("source") or source or "invoice_line")[:50]
                    upserted += 1
                continue

        cache_rows_payload: list[dict[str, Any]] = []
        for requested_name, requested_norm, requested_chars in requested_specs:
            score, details = _compare_characteristics_with_embeddings_from_vectors(
                requested=requested_chars,
                offered=offered,
                threshold=threshold,
                vectors=vectors,
            )
            if not isinstance(details, dict):
                details = {
                    "threshold": round(float(threshold), 4),
                    "matched_points": 0,
                    "total_points": 0,
                    "points": [],
                }
            details["source"] = source
            details["matched_product_name"] = matched_name

            row = InvoiceItemCharacteristicMatch(
                tenant_id=tenant_id,
                campaign_id=campaign_id,
                invoice_item_id=item.id,
                requested_product_name=requested_name[:255],
                requested_product_norm=requested_norm[:255],
            )
            db.add(row)
            row.requested_product_name = requested_name[:255]
            row.requested_product_norm = requested_norm[:255]
            row.matched_product_name = matched_name[:255] if matched_name else None
            row.threshold = float(details.get("threshold") or threshold)
            row.score = float(score) if score is not None else None
            try:
                row.matched_points = int(details.get("matched_points")) if details.get("matched_points") is not None else 0
            except Exception:
                row.matched_points = 0
            try:
                row.total_points = int(details.get("total_points")) if details.get("total_points") is not None else 0
            except Exception:
                row.total_points = 0
            points = details.get("points")
            row.points = points if isinstance(points, list) else []
            row.source = str(details.get("source") or source or "invoice_line")[:50]
            upserted += 1
            cache_rows_payload.append(
                {
                    "requested_product_name": requested_name[:255],
                    "requested_product_norm": requested_norm[:255],
                    "matched_product_name": matched_name[:255] if matched_name else None,
                    "threshold": float(details.get("threshold") or threshold),
                    "score": float(score) if score is not None else None,
                    "matched_points": int(row.matched_points or 0),
                    "total_points": int(row.total_points or 0),
                    "points": row.points if isinstance(row.points, list) else [],
                    "source": row.source,
                }
            )
        if pair_cache_key and cache_rows_payload:
            set_openai_payload_cache(
                cache_type="invoice_pair_cache",
                model=pair_cache_model,
                cache_key=pair_cache_key,
                payload={"rows": cache_rows_payload},
            )

    return upserted


def match_products_by_characteristics(
    items: list[InvoiceExtractionItem],
    products: list[Product],
    threshold: float,
) -> None:
    if not items or not products:
        return

    product_features: list[tuple[str, set[str]]] = []
    for p in products:
        base = f"{p.name or ''}".strip()
        feats = _extract_features(base)
        product_features.append((p.name, feats))

    for item in items:
        text = f"{item.product_name or ''} {item.description or ''}".strip()
        item_feats = _extract_features(text)
        if not item_feats:
            continue

        best_score = -1.0
        best_name = None
        for name, feats in product_features:
            if not feats:
                continue
            common = item_feats.intersection(feats)
            score = len(common) / max(len(feats), 1)
            if score > best_score:
                best_score = score
                best_name = name

        if best_name is None:
            continue

        prev = item.notes or ""
        item.notes = f"{prev} char_match={best_name} score={best_score:.2f}".strip()
        # Caracteristiques = information uniquement (ne confirme pas)


def match_products_hybrid(
    items: list[InvoiceExtractionItem],
    products: list[Product],
    threshold: float,
    *,
    embed_threshold: float | None = None,
    embed_min_threshold: float | None = None,
    characteristics_match_threshold: float | None = None,
    preserve_corrected_on_reject: bool = False,
    db: Session | None = None,
) -> None:
    if not items or not products:
        return

    match_items = items
    match_cache_keys: dict[int, tuple[str, str]] = {}
    if db is not None:
        sample = next((it for it in items if it is not None), None)
        if sample is not None:
            effective_embed_threshold = (
                config.OPENAI_EMBEDDING_THRESHOLD if embed_threshold is None else float(embed_threshold)
            )
            effective_char_threshold = (
                config.OPENAI_CHARACTERISTICS_MATCH_THRESHOLD
                if characteristics_match_threshold is None
                else float(characteristics_match_threshold)
            )
            requested_signature = _requested_reference_signature(db, sample.tenant_id, sample.campaign_id)
            pool_signature = _product_pool_signature(products)
            misses: list[InvoiceExtractionItem] = []
            for item in items:
                cache_key = _item_match_signature(
                    item=item,
                    product_pool_signature=pool_signature,
                    requested_signature=requested_signature,
                    threshold=float(threshold),
                    embed_threshold=effective_embed_threshold,
                    characteristics_match_threshold=effective_char_threshold,
                )
                cache_model = str(getattr(config, "OPENAI_PRODUCT_MATCH_MODEL", "") or "hybrid")
                cached = get_openai_payload_cache(
                    cache_type="invoice_product_match",
                    model=cache_model,
                    cache_key=cache_key,
                )
                if isinstance(cached, dict) and _apply_item_match_payload(item, cached):
                    continue
                match_cache_keys[id(item)] = (cache_key, cache_model)
                misses.append(item)
            match_items = misses

    if not match_items:
        return

    # characteristics score
    match_products_by_characteristics(match_items, products, threshold)

    # fuzzy match on normalized names
    product_names = [_normalize_name(p.name or "") for p in products]
    item_names = []
    for i in match_items:
        # If user corrected the extracted name, use it as the primary text
        # for fuzzy/embedding rematch.
        base_name = (i.corrected_product_name or i.product_name or "")
        combined = f"{base_name} {i.description or ''}".strip()
        item_names.append(_normalize_name(combined))
    fuzzy_scores: list[float] = []
    fuzzy_best_indices: list[int] = []
    for item, item_name in zip(match_items, item_names):
        if not item_name:
            fuzzy_scores.append(0.0)
            fuzzy_best_indices.append(-1)
            continue
        best_score = -1.0
        best_index = -1
        for idx, prod_name in enumerate(product_names):
            if not prod_name:
                continue
            score = difflib.SequenceMatcher(None, item_name, prod_name).ratio()
            if score > best_score:
                best_score = score
                best_index = idx
        if best_index >= 0:
            matched_name = products[best_index].name
            item.notes = (item.notes or "") + f" fuzzy_match={matched_name} fuzzy_score={best_score:.2f}"
        fuzzy_scores.append(max(best_score, 0.0))
        fuzzy_best_indices.append(best_index)

    # embeddings score on names (with synonyms normalization)
    if not config.OPENAI_API_KEY:
        # Still compute characteristics resemblance points when DB context exists.
        if db is not None:
            sample = next((it for it in match_items if it is not None), None)
            if sample is not None:
                char_threshold = (
                    config.OPENAI_CHARACTERISTICS_MATCH_THRESHOLD
                    if characteristics_match_threshold is None
                    else float(characteristics_match_threshold)
                )
                _apply_resemblance_points(
                    db=db,
                    items=match_items,
                    tenant_id=sample.tenant_id,
                    campaign_id=sample.campaign_id,
                    threshold=char_threshold,
                )
                for item in match_items:
                    cache_meta = match_cache_keys.get(id(item))
                    if cache_meta is None:
                        continue
                    cache_key, cache_model = cache_meta
                    set_openai_payload_cache(
                        cache_type="invoice_product_match",
                        model=cache_model,
                        cache_key=cache_key,
                        payload=_serialize_item_match_payload(item),
                    )
        return

    combined = product_names + item_names
    cached, missing_texts, _ = get_cached_embeddings(combined)
    if missing_texts:
        enqueue_embeddings(missing_texts)
        for item in match_items:
            if not item.notes:
                item.notes = "embedding_pending"
        # Do not call OpenAI here; wait for worker to fill cache.
        return

    product_embeddings = [
        cached[_hash_text(_normalize_text(name or ""))] for name in product_names
    ]
    item_embeddings = [
        cached[_hash_text(_normalize_text(name or ""))] for name in item_names
    ]

    embed_threshold = (
        config.OPENAI_EMBEDDING_THRESHOLD if embed_threshold is None else embed_threshold
    )
    # Kept in function signature for API compatibility, but status assignment
    # no longer uses a second "verify/pending" threshold.
    use_gpt_match = str(
        getattr(config, "OPENAI_PRODUCT_MATCH_MODE", "embedding") or "embedding"
    ).strip().lower() in {"gpt", "gpt_only", "gpt_hybrid", "hybrid_gpt"}
    gpt_uncertain_only = bool(
        getattr(config, "OPENAI_PRODUCT_MATCH_GPT_UNCERTAIN_ONLY", True)
    )
    gpt_trigger_min = max(
        0.0,
        min(
            1.0,
            float(getattr(config, "OPENAI_PRODUCT_MATCH_GPT_TRIGGER_MIN", 0.35) or 0.35),
        ),
    )
    gpt_trigger_max = max(
        gpt_trigger_min,
        min(
            1.0,
            float(getattr(config, "OPENAI_PRODUCT_MATCH_GPT_TRIGGER_MAX", 0.88) or 0.88),
        ),
    )
    gpt_min_conf = max(
        0.0,
        min(
            1.0,
            float(getattr(config, "OPENAI_PRODUCT_MATCH_MIN_CONFIDENCE", 0.30) or 0.30),
        ),
    )
    gpt_strict_min_conf = max(
        0.0,
        min(
            1.0,
            float(
                getattr(config, "OPENAI_PRODUCT_MATCH_STRICT_MIN_CONFIDENCE", 0.50)
                or 0.50
            ),
        ),
    )
    gpt_force_to_verify_on_miss = bool(
        getattr(config, "OPENAI_PRODUCT_MATCH_FORCE_TO_VERIFY_ON_GPT_MISS", True)
    )
    gpt_top_k = max(3, int(getattr(config, "OPENAI_PRODUCT_MATCH_TOP_K", 15) or 15))
    gpt_max_calls_per_run = max(
        1,
        int(getattr(config, "OPENAI_PRODUCT_MATCH_GPT_MAX_CALLS_PER_RUN", 120) or 120),
    )
    gpt_calls_used = 0
    product_display_names = [str(p.name or "").strip() for p in products]

    for item, item_embedding, _fuzzy_score, fuzzy_best_index in zip(
        match_items, item_embeddings, fuzzy_scores, fuzzy_best_indices
    ):
        # Keep this guard for externally pre-confirmed/manual cases.
        if item.status == ExtractionStatus.CONFIRMED:
            continue
        pair_scores: list[tuple[float, int]] = []
        best_score = -1.0
        best_index = -1
        for idx, prod_emb in enumerate(product_embeddings):
            score = cosine_similarity(item_embedding, prod_emb)
            pair_scores.append((score, idx))
            if score > best_score:
                best_score = score
                best_index = idx
        if best_index < 0:
            continue

        should_call_gpt = bool(use_gpt_match)
        if should_call_gpt and gpt_uncertain_only:
            # Skip GPT on confident easy cases and very low-score hard misses.
            # GPT is reserved for the "gray zone" only.
            if (
                best_score < gpt_trigger_min
                or best_score >= gpt_trigger_max
                or _fuzzy_score >= 0.92
            ):
                should_call_gpt = False

        if should_call_gpt:
            if gpt_calls_used >= gpt_max_calls_per_run:
                should_call_gpt = False
                prev = item.notes or ""
                item.notes = f"{prev} gpt_budget_exhausted=1".strip()
            else:
                gpt_calls_used += 1
        if should_call_gpt:
            candidate_indices = [
                idx
                for _score, idx in sorted(pair_scores, key=lambda x: x[0], reverse=True)[:gpt_top_k]
            ]
            if fuzzy_best_index >= 0 and fuzzy_best_index not in candidate_indices:
                candidate_indices.insert(0, fuzzy_best_index)
            candidate_names = [
                product_display_names[idx]
                for idx in candidate_indices
                if 0 <= idx < len(product_display_names) and product_display_names[idx]
            ]
            gpt_item_text = _collapse_spaces(
                f"{item.corrected_product_name or item.product_name or ''} {item.description or ''}"
            )
            gpt_match_name, gpt_score, gpt_scored = _gpt_rank_requested_candidates(
                item_text=gpt_item_text,
                item_name=(item.corrected_product_name or item.product_name or ""),
                candidate_names=candidate_names,
            )
            gpt_candidates_payload: list[dict[str, Any]] = []
            for cname, cscore in gpt_scored:
                gpt_candidates_payload.append(
                    {
                        "name": cname,
                        "score": round(float(cscore), 4),
                    }
                )
            gpt_blob = (
                _encode_gpt_matches_b64(gpt_candidates_payload)
                if gpt_candidates_payload
                else ""
            )
            eligible_gpt_rows = [
                (cname, cscore)
                for cname, cscore in gpt_scored
                if cscore >= gpt_min_conf
            ]
            primary_gpt_name: str | None = None
            primary_gpt_score: float | None = None
            if eligible_gpt_rows:
                primary_gpt_name, primary_gpt_score = eligible_gpt_rows[0]
            elif (
                gpt_match_name
                and gpt_score is not None
                and gpt_score >= gpt_min_conf
            ):
                primary_gpt_name = gpt_match_name
                primary_gpt_score = float(gpt_score)

            if primary_gpt_name and primary_gpt_score is not None:
                prev = item.notes or ""
                item.notes = (
                    f"{prev} gpt_match={primary_gpt_name} gpt_score={float(primary_gpt_score):.2f}"
                ).strip()
                if gpt_blob:
                    item.notes = f"{item.notes} gpt_matches_b64={gpt_blob}".strip()
                if primary_gpt_score >= gpt_strict_min_conf:
                    item.status = ExtractionStatus.CONFIRMED
                else:
                    # Strict GPT gate in GPT mode:
                    # score below threshold stays TO_VERIFY (no auto-confirm).
                    item.status = ExtractionStatus.TO_VERIFY
                continue

            elif gpt_blob:
                prev = item.notes or ""
                item.notes = f"{prev} gpt_matches_b64={gpt_blob}".strip()

        if use_gpt_match and gpt_force_to_verify_on_miss:
            # In GPT mode, if GPT score is absent we do not auto-confirm with emb.
            # Keep emb score as evidence only, but route to manual verification.
            matched_name = products[best_index].name
            prev = item.notes or ""
            item.notes = (
                f"{prev} emb_match={matched_name} emb_score={best_score:.2f} gpt_missing=1"
            ).strip()
            item.status = ExtractionStatus.TO_VERIFY
            continue

        matched_name = products[best_index].name
        prev = item.notes or ""
        item.notes = f"{prev} emb_match={matched_name} emb_score={best_score:.2f}".strip()
        # Embeddings logic (simplified):
        # - >= embed_threshold -> CONFIRMED
        # - else -> TO_VERIFY
        if best_score >= embed_threshold:
            item.status = ExtractionStatus.CONFIRMED
        else:
            item.status = ExtractionStatus.TO_VERIFY

    if db is not None:
        sample = next((it for it in match_items if it is not None), None)
        if sample is not None:
            char_threshold = (
                config.OPENAI_CHARACTERISTICS_MATCH_THRESHOLD
                if characteristics_match_threshold is None
                else float(characteristics_match_threshold)
            )
            _apply_resemblance_points(
                db=db,
                items=match_items,
                tenant_id=sample.tenant_id,
                campaign_id=sample.campaign_id,
                threshold=char_threshold,
            )
            for item in match_items:
                cache_meta = match_cache_keys.get(id(item))
                if cache_meta is None:
                    continue
                cache_key, cache_model = cache_meta
                set_openai_payload_cache(
                    cache_type="invoice_product_match",
                    model=cache_model,
                    cache_key=cache_key,
                    payload=_serialize_item_match_payload(item),
                )



def match_products_with_embeddings(
    items: list[InvoiceExtractionItem],
    product_names: list[str],
) -> None:
    if not items or not product_names:
        return
    if not config.OPENAI_API_KEY:
        return

    # Exact match shortcut (case-insensitive, normalized)
    normalized_products = {
        _normalize_text(p): p for p in product_names if p
    }
    for item in items:
        name = (item.product_name or "").strip()
        if not name:
            continue
        key = _normalize_text(name)
        if key in normalized_products:
            matched = normalized_products[key]
            item.status = ExtractionStatus.CONFIRMED
            item.notes = f"exact_match={matched}"

    remaining = [i for i in items if i.status != ExtractionStatus.CONFIRMED]
    if not remaining:
        return

    item_names = [item.product_name or "" for item in remaining]
    combined = product_names + item_names

    cached, missing_texts, _ = get_cached_embeddings(combined)
    if missing_texts:
        enqueue_embeddings(missing_texts)
        for item in remaining:
            item.status = ExtractionStatus.TO_VERIFY
            item.notes = "embedding_pending"
        return

    product_embeddings = [
        cached[_hash_text(_normalize_text(name or ""))] for name in product_names
    ]
    item_embeddings = [
        cached[_hash_text(_normalize_text(name or ""))] for name in item_names
    ]

    for item, item_embedding in zip(remaining, item_embeddings):
        name = item.product_name or ""
        if not name.strip():
            continue
        best_score = -1.0
        best_index = -1
        for idx, prod_emb in enumerate(product_embeddings):
            score = cosine_similarity(item_embedding, prod_emb)
            if score > best_score:
                best_score = score
                best_index = idx

        if best_index >= 0:
            matched_name = product_names[best_index]
            item.notes = f"match={matched_name} score={best_score:.3f}"
            if best_score >= config.OPENAI_MATCH_CONFIRM_THRESHOLD:
                item.status = ExtractionStatus.CONFIRMED
            elif best_score >= config.OPENAI_MATCH_VERIFY_THRESHOLD:
                item.status = ExtractionStatus.TO_VERIFY
            else:
                item.status = ExtractionStatus.TO_VERIFY


def _to_float(value: str | None) -> float | None:
    if value is None:
        return None
    cleaned = value.replace(" ", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return None


def _resolve_unit_price(
    qty_value: float | None,
    unit_price_value: float | None,
    amount_value: float | None,
) -> float | None:
    value, _source = _resolve_unit_price_with_source(
        qty_value=qty_value,
        unit_price_value=unit_price_value,
        amount_value=amount_value,
    )
    return value


def _resolve_unit_price_with_source(
    qty_value: float | None,
    unit_price_value: float | None,
    amount_value: float | None,
) -> tuple[float | None, str]:
    # Prefer the explicit unit_price extracted by parser/OCR when present.
    # In real invoices, amount/qty can be noisy (taxes, discounts, OCR qty errors).
    if unit_price_value is not None and unit_price_value > 0:
        return unit_price_value, "ocr_unit_price"

    # Fallback: derive from line amount and quantity.
    if amount_value is None:
        return None, "missing"
    if qty_value is not None and qty_value > 0:
        return round(amount_value / qty_value, 6), "derived_amount_div_qty"
    # Last resort when quantity is missing.
    return amount_value, "amount_only"

